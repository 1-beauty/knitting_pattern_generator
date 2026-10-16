# -*- coding: utf-8 -*-
"""eq(6) 受控双行：oMath 在 child12(第二个∀)前拆为两个 oMathPara(中心格两段)，
编号格垂直居中；并移除 9.5pt 的 w:rPr 恢复 11pt。"""
import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_ALIGN_VERTICAL
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
M = qn('m:oMath'); MP = qn('m:oMathPara'); MR = qn('m:r'); WRPR = qn('w:rPr')
WP = qn('w:p'); WPPR = qn('w:pPr')

doc = Document(MAIN)
t = doc.tables[7]
cell = t.cell(0, 1)
om = cell._tc.find('.//' + M)

# 恢复 11pt：移除仅含 sz/szCs 的 w:rPr
for r in om.iter(MR):
    rpr = r.find(WRPR)
    if rpr is not None and all(c.tag in (qn('w:sz'), qn('w:szCs')) for c in rpr):
        r.remove(rpr)

kids = list(om)
SPLIT = 12  # 第二个 ∀ 之前
new_ompara = OxmlElement('m:oMathPara')
opr = OxmlElement('m:oMathParaPr')
jc = OxmlElement('m:jc'); jc.set(qn('m:val'), 'center'); opr.append(jc)
new_ompara.append(opr)
new_om = OxmlElement('m:oMath')
new_ompara.append(new_om)
for k in kids[SPLIT:]:
    new_om.append(k)  # lxml move

first_p = cell._tc.find(WP)
new_p = OxmlElement('w:p')
ppr = first_p.find(WPPR)
if ppr is not None:
    new_p.append(copy.deepcopy(ppr))
new_p.append(new_ompara)
first_p.addnext(new_p)

t.cell(0, 2).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
doc.save(MAIN)
print('eq6 split into 2 lines; line1 children:', SPLIT, ' line2 children:', len(kids) - SPLIT)
