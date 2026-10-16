# -*- coding: utf-8 -*-
"""重建正文表 1-7：英文内容 + 正确表头(恢复 V_h/V_c/T) + 三线表 + 居中 + 英文表题在上。
下标记法：用 § 包裹，如 'V§h§=0' -> V + 下标h + =0。"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)

VH = 'V§h§'; VC = 'V§c§'; T = 'T'

TABLES = {
 0: dict(caption='Table 1 The 15 discrete action labels retained after sample-level data cleaning.',
   rows=[['ID','Action code','Semantic class'],
     ['0','f_knit','LOOP'],['1','b_knit','LOOP'],['2','xfer_f_b','TRANSFER'],['3','xfer_b_f','TRANSFER'],
     ['4','b_tuck','TUCK'],['5','miss','EMPTY'],['6','xfer_l1','TRANSFER'],['7','f_tuck','TUCK'],
     ['8','f_split','SPLIT'],['9','b_split','SPLIT'],['10','f_tuck_2','TUCK'],['11','carrier','DROP'],
     ['12','split_v','SPLIT'],['13','b_tuck_2','TUCK'],['14','xfer_r1','TRANSFER']]),
 1: dict(caption='Table 2 Mapping from visual semantics to process realization.',
   rows=[['Visual semantics','Process realization','Event pattern','Key constraints'],
     ['Wave tuck','Front-bed base + split/tuck combination','diagonal_wave','Generated in groups; whole group dropped if it does not fit the boundary'],
     ['Honeycomb','Alternating front/back tuck','honeycomb_tuck','Tuck spacing; consecutive stacking limit'],
     ['Sparse mesh','Base + float/transfer local events','lattice_miss','Hole density; neighbor isolation'],
     ['Relief / seed stitch','Front-bed base + periodic tuck','staggered_tuck / seed_tuck','Tuck density; boundary safety margin'],
     ['Closed geometry (diamond, hexagon, etc.)','Geometric mask + reverse-loop outline','— (outline expressed by boundary action back_loop_outline)','Centered, symmetric, closed'],
     ['Rectangular frame','Back-bed or rib border','— (compiled directly from frame thickness)','Continuously closed; not broken by inner structures'],
     ['Cable / displacement','Left/right transfer combination','column_xfer','Left/right load balance; non-conflicting targets']]),
 2: dict(caption='Table 3 Field-level accuracy of PatternPlan semantic parsing.',
   rows=[['Field','Correct / Total','Accuracy'],
     ['Layout','100 / 100','100.0%'],['Shape','97 / 100','97.0%'],['Texture','96 / 100','96.0%'],
     ['Position','90 / 100','90.0%'],['Parent','100 / 100','100.0%'],['Params','100 / 100','100.0%'],
     ['Field-level overall','583 / 600','97.2%'],['Fully-correct samples','83 / 100','83.0%']]),
 3: dict(caption='Table 4 Comparison of generation methods on manufacturing constraints and semantic fidelity.',
   rows=[['Method',VH,VC,T,'Type-Coverage↑'],
     ['Ours (full system)','0','0','0.080','0.692'],
     ['Plan-Grid (internal baseline)','0.3','0','0.087','0.642'],
     ['Plan-Grid+Repair (internal baseline)','0.4','0','0.087','0.642'],
     ['Plan-Grid+CSP (internal baseline)','0','0','0.085','0.642'],
     ['LLM-Direct','60.3','0.2','0.307','0.867'],
     ['LLM-Direct+Repair','57.6','0','0.313','0.867'],
     ['T2I + color quantization','565.7','0','0.680','0.050'],
     ['Random Valid Map','0','0','0.000','0.050']]),
 4: dict(caption='Table 5 Summary of module ablation results.',
   rows=[['Setting','Focus',VH+'=0 cases',VC+'=0 cases',VH+'='+VC+'=0 cases','Mean '+VH,'Mean '+VC,'Mean '+T,'Main impact'],
     ['full','Full system','8 / 8','8 / 8','8 / 8','0','0','0.131','Best action validity and structural stability'],
     ['no_repair','Conservative self-repair','8 / 8','5 / 8','5 / 8','0','0.500','0.131','Weaker repair of process-heuristic violations'],
     ['no_refine','Local compilation refinement','8 / 8','8 / 8','8 / 8','0','0','0.111','Lower switching-intensity score; less texture richness'],
     ['no_aesthetic','Aesthetic post-processing','8 / 8','8 / 8','8 / 8','0','0','0.143','Higher switching-intensity score; affects visual regularity']]),
 5: dict(caption='Table 6 Controlled comparison of internal baselines.',
   rows=[['Method','Mean '+VH,'Mean '+VC,VH+'=0 cases','Region IoU↑','Leakage↓','Grouping completeness↑'],
     ['Ours (full system)','0','0','8 / 8','0.357','0.240','1.000'],
     ['Plan-Grid','0.25','0','6 / 8','0.309','0.311','0.953'],
     ['Plan-Grid+Repair','0.13','0','7 / 8','0.309','0.311','0.953'],
     ['Plan-Grid+CSP','0','0','8 / 8','0.309','0.311','0.953']]),
 6: dict(caption='Table 7 Whole-image quantitative evaluation of the structure-aware visual preview (50 validation samples).',
   rows=[['Stage','MSE ↓','SSIM ↑','LPIPS ↓','SeamErr_diff ↓','BoundaryF1 ↑','TextureSim ↑'],
     ['Stage 1','0.0328','0.636','0.277','0.0859','0.222','0.917'],
     ['Stage 2','0.0203','0.760','0.172','0.0565','0.235','0.929'],
     ['Stage 3','0.0205','0.758','0.170','0.0573','0.234','0.930'],
     ['Stage 4','0.0648','-0.148','0.515','0.0774','0.310','0.846']]),
}

def set_cell(cell, text, bold=False):
    for p in list(cell.paragraphs):
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    p = cell.paragraphs[0]
    p.style = doc.styles['Compact']
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = text.split('§')
    for i, seg in enumerate(parts):
        if seg == '':
            continue
        r = p.add_run(seg)
        r.bold = bold
        if i % 2 == 1:
            r.font.subscript = True

def three_line(tbl):
    tblPr = tbl._tbl.tblPr
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
    borders = OxmlElement('w:tblBorders')
    for edge, sz in (('top', '12'), ('bottom', '12')):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000'); borders.append(e)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'none'); e.set(qn('w:sz'), '0')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto'); borders.append(e)
    tblPr.append(borders)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'autofit'); tblPr.append(lay)
    for cell in tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')):
            tcPr.remove(old)
        tcb = OxmlElement('w:tcBorders')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '8')
        b.set(qn('w:space'), '0'); b.set(qn('w:color'), '000000'); tcb.append(b)
        tcPr.append(tcb)

def set_caption_para(para, text):
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    m = re.match(r'(Table\s+\d+)\s*(.*)', text)
    if m:
        r1 = para.add_run(m.group(1) + ' '); r1.bold = True
        r2 = para.add_run(m.group(2)); r2.bold = False
    else:
        para.add_run(text)
    para.style = doc.styles['TableCaption']

# 定位每个主表及其上方题注段
body = doc.element.body
from docx.table import Table as DxTable
from docx.text.paragraph import Paragraph
seq = []
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        seq.append(['p', Paragraph(child, doc)])
    elif tag == 'tbl':
        seq.append(['t', DxTable(child, doc)])

tbl_pos = [i for i, (k, o) in enumerate(seq) if k == 't']
for idx in range(0, 7):
    pos = tbl_pos[idx]
    old = seq[pos][1]
    # 上方题注：必须是加粗段且以 表N 开头（避免误伤 bold=False 的正文引用句）
    cap_para = None
    for j in range(pos - 1, max(-1, pos - 5), -1):
        if seq[j][0] != 'p':
            continue
        pj = seq[j][1]
        is_bold = any(r.bold for r in pj.runs if r.bold is not None)
        if is_bold and re.match(r'^表\s*' + str(idx + 1) + r'\b', pj.text.strip()):
            cap_para = pj; break
    spec = TABLES[idx]
    nrows = len(spec['rows']); ncols = len(spec['rows'][0])
    new = doc.add_table(rows=nrows, cols=ncols)
    for ri, rowdata in enumerate(spec['rows']):
        for ci, val in enumerate(rowdata):
            set_cell(new.cell(ri, ci), val, bold=(ri == 0))
    three_line(new)
    # 移动到旧表位置并删除旧表
    old._tbl.addnext(new._tbl)
    old._tbl.getparent().remove(old._tbl)
    if cap_para is not None:
        set_caption_para(cap_para, spec['caption'])
        print(f'table[{idx}] rebuilt + caption EN')
    else:
        print(f'table[{idx}] rebuilt (caption para NOT found)')

doc.save(SRC)
print('saved')
