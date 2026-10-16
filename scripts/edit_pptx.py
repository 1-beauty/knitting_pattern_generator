"""Layout restructuring for knitting_ai_competition.pptx.
Edits are keyed by (slide_index, shape_id) so they stay stable.
Preserves all styling; only repositions shapes and splits jammed bullet runs.
"""
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

PATH = r"e:\DeskTop\knitting_ai_competition.pptx"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

prs = Presentation(PATH)


def shp_map(slide_idx):
    slide = prs.slides[slide_idx - 1]
    return {sh.shape_id: sh for sh in slide.shapes}


def set_pos(sh, l=None, t=None, w=None, h=None):
    if l is not None:
        sh.left = Inches(l)
    if t is not None:
        sh.top = Inches(t)
    if w is not None:
        sh.width = Inches(w)
    if h is not None:
        sh.height = Inches(h)


def _set_spacing(pPr, line_pct=None, aft_pts=None):
    # lnSpc / spcBef / spcAft must be the first children of pPr (in this order)
    if line_pct is not None:
        for e in pPr.findall(f"{A}lnSpc"):
            pPr.remove(e)
        ln = etree.SubElement(pPr, f"{A}lnSpc")
        pct = etree.SubElement(ln, f"{A}spcPct")
        pct.set("val", str(int(line_pct * 1000)))
        pPr.insert(0, ln)
    if aft_pts is not None:
        for e in pPr.findall(f"{A}spcAft"):
            pPr.remove(e)
        aft = etree.SubElement(pPr, f"{A}spcAft")
        pts = etree.SubElement(aft, f"{A}spcPts")
        pts.set("val", str(int(aft_pts * 100)))
        # place after lnSpc/spcBef
        idx = 0
        for i, ch in enumerate(pPr):
            if ch.tag in (f"{A}lnSpc", f"{A}spcBef"):
                idx = i + 1
        pPr.insert(idx, aft)


def split_runs_to_paragraphs(sh, line_pct=None, aft_pts=None):
    """Take a text frame whose bullets are multiple runs in one paragraph and
    rebuild it so each run becomes its own paragraph. Formatting preserved."""
    txBody = sh.text_frame._txBody
    paras = txBody.findall(f"{A}p")
    items = []  # (pPr_element_or_None, run_element)
    for p in paras:
        pPr = p.find(f"{A}pPr")
        for r in p.findall(f"{A}r"):
            items.append((pPr, r))
    for p in paras:
        txBody.remove(p)
    for pPr, r in items:
        newp = etree.SubElement(txBody, f"{A}p")
        new_pPr = deepcopy(pPr) if pPr is not None else etree.SubElement(newp, f"{A}pPr")
        if pPr is not None:
            newp.append(new_pPr)
        if line_pct is not None or aft_pts is not None:
            if new_pPr.getparent() is None:
                newp.insert(0, new_pPr)
            _set_spacing(new_pPr, line_pct, aft_pts)
        newp.append(deepcopy(r))


# ============================================================
# SLIDE 3 — split jammed bullets + rebalance the two tall boxes
# ============================================================
s3 = shp_map(3)
# split bullets (left ✦ x4, right ▪ x5) into stacked lines with breathing room
split_runs_to_paragraphs(s3[9], line_pct=115, aft_pts=9)     # left: 设计师的表达 bullets
split_runs_to_paragraphs(s3[115], line_pct=115, aft_pts=9)   # right: 横机需要的 bullets
# shrink the two 5.4in panels so content is balanced (bottom ~5.55)
set_pos(s3[6],  t=1.05, h=4.55)     # left panel
set_pos(s3[15], t=1.05, h=4.55)     # right panel
# distribute left column content inside the panel
set_pos(s3[8],  t=1.62)             # quote block
set_pos(s3[9],  t=2.95, h=2.4)      # left bullets, more room
# right column: bullets start just under the "每个格子" caption
set_pos(s3[115], t=3.5, h=2.0)
# recenter the middle "断层" group to the new panel center (~3.3)
set_pos(s3[10], t=2.7)              # 断层 box
set_pos(s3[11], t=2.75)            # ❌断层 title text
set_pos(s3[12], t=3.1)            # 高度依赖/专业打版师
set_pos(s3[13], t=3.25)            # left connector line
set_pos(s3[14], t=3.25)            # right connector line

# ============================================================
# SLIDE 5 — align the "输出" pill inside card 4
# ============================================================
s5 = shp_map(5)
set_pos(s5[31], l=10.05)   # green output pill bg  (was 9.85, card4 left=10.05)
set_pos(s5[32], l=10.1)    # output pill text

# ============================================================
# SLIDE 10 — stop "PatternPlan" x-axis labels from wrapping
# ============================================================
s10 = shp_map(10)
# widen + recenter the two "w/o PatternPlan" label boxes (0.75 -> 1.0)
set_pos(s10[11], l=1.795, w=1.0)   # left chart
set_pos(s10[27], l=7.845, w=1.0)   # right chart

prs.save(PATH)
print("saved edits: S3, S5, S10")
