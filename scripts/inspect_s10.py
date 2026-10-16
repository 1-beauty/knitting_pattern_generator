# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
from pptx.util import Emu

P = r"e:\DeskTop\knitting_ai_competition.pptx"
IDX = int(sys.argv[1]) if len(sys.argv) > 1 else 10
prs = Presentation(P)
sl = prs.slides[IDX - 1]
def IN(v): return round(Emu(v).inches, 2) if v is not None else None
for sh in sl.shapes:
    txt = ""
    if sh.has_text_frame:
        txt = " | ".join(p.text for p in sh.text_frame.paragraphs if p.text)[:60]
    print(f"#{sh.shape_id:>3} {sh.shape_type} name={sh.name[:16]:<16} "
          f"L{IN(sh.left)} T{IN(sh.top)} W{IN(sh.width)} H{IN(sh.height)} :: {txt}")
