# -*- coding: utf-8 -*-
"""P0-1 抽离补充材料：
- 主文件 docs/论文to_cii.docx 删除 '补充材料' H2 及之后全部 body 子元素(保留末尾 sectPr)；
- 副本 docs/Supplementary_Material.docx 删除 '补充材料' 之前全部 body 子元素，标题改英文；
- 主文件在 参考文献 前插入 Supplementary data online 句。
"""
import sys, io, shutil
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
shutil.copyfile(MAIN, SUPP)

def body_children(doc):
    return list(doc.element.body.iterchildren())

def find_supp_heading(doc):
    for child in doc.element.body.iterchildren():
        if child.tag.endswith('}p'):
            from docx.text.paragraph import Paragraph
            p = Paragraph(child, doc)
            if p.style.name == 'Heading 2' and p.text.strip().startswith('补充材料'):
                return child
    return None

# ---- 主文件：删除补充材料及之后 ----
doc = Document(MAIN)
h = find_supp_heading(doc)
assert h is not None, 'supp heading not found in main'
body = doc.element.body
removing = False
removed = 0
for child in list(body.iterchildren()):
    if child is h:
        removing = True
    if removing and child.tag != qn('w:sectPr'):
        body.remove(child); removed += 1
print('main removed elements:', removed)

# 在 参考文献 前插入 supplementary online 句
ref_h = None
for child in body.iterchildren():
    if child.tag.endswith('}p'):
        from docx.text.paragraph import Paragraph
        p = Paragraph(child, doc)
        if p.style.name == 'Heading 2' and p.text.strip() == '参考文献':
            ref_h = child; break
from docx.oxml import OxmlElement
newp = OxmlElement('w:p')
ppr = OxmlElement('w:pPr'); pstyle = OxmlElement('w:pStyle'); pstyle.set(qn('w:val'), 'BodyText'); ppr.append(pstyle); newp.append(ppr)
r = OxmlElement('w:r'); t = OxmlElement('w:t')
t.text = ('Supplementary data to this article can be found online at https://doi.org/<DOI-to-be-assigned> '
          '(Supplementary Material, Tables S1-S19 and Figures S1-S3).')
t.set(qn('xml:space'), 'preserve'); r.append(t); newp.append(r)
ref_h.addprevious(newp)
print('main: supplementary-online sentence inserted before 参考文献')
doc.save(MAIN)

# ---- 副本：删除补充材料之前 ----
s = Document(SUPP)
hs = find_supp_heading(s)
assert hs is not None, 'supp heading not found in supp'
sbody = s.element.body
for child in list(sbody.iterchildren()):
    if child is hs:
        break
    if child.tag != qn('w:sectPr'):
        sbody.remove(child)
# 标题改英文
from docx.text.paragraph import Paragraph
hp = Paragraph(hs, s)
for rr in list(hp.runs):
    rr._r.getparent().remove(rr._r)
hp.add_run('Supplementary Material')
s.save(SUPP)
print('supp saved; first heading ->', Paragraph(hs, s).text)

# 计数核验
d2 = Document(MAIN); s2 = Document(SUPP)
print('main tables:', len(d2.tables), ' main paras:', len(d2.paragraphs))
print('supp tables:', len(s2.tables), ' supp paras:', len(s2.paragraphs))
