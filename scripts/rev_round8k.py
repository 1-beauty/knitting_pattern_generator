# -*- coding: utf-8 -*-
"""第八轮 k（修 round8j 引发的 Supp→正文表号失同步 + 既有 Supp 交叉引用缺陷）：
1) Supp 中指向正文内容表的"表 N"(无 S, N∈2-5) 同步 +1 → 3-6（正文表号已重排）。
2) 畸形前缀"正文 4.正文 1.2 节"→"正文 4.3.2 节"(s079, 内容=Region IoU 形态学, 对应正文4.3.2)；
   "正文 4.正文 3.1 节"→"正文 4.3.1 节"(s142/s161, 内容=TAP口径/协议A, 对应正文4.3.1)。
3) 悬空 4.1.x（正文已无 4.1.1/4.1.2/4.1.3 子节, 内容迁至 Supp）：
   正文 4.1.1/4.1.2 → 补充材料 Section M（Dataset, Metrics）；正文 4.1.3 → 补充材料 Section J（Patch, Resolution）。
全部在 w:t 级替换, 不重建段落, 不动 oMath。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import re
from docx import Document
from docx.oxml.ns import qn
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
s = Document(SUPP)

SHIFT = {'2': '3', '3': '4', '4': '5', '5': '6'}
tab_rx = re.compile(r'(表)(\s*)([2-5])(?!\d)')
MAL = {'正文 4.正文 1.2 节': '正文 4.3.2 节', '正文 4.正文 3.1 节': '正文 4.3.1 节'}
dang_rx = re.compile(r'正文\s*4\.1\.([123])\s*节?')
DANG = {'1': '补充材料 Section M', '2': '补充材料 Section M', '3': '补充材料 Section J'}

n_tab = n_mal = n_dang = 0
for p in s.paragraphs:
    for wt in p._p.findall('.//' + qn('w:t')):
        t = wt.text
        if not t:
            continue
        orig = t
        t, c1 = tab_rx.subn(lambda mo: mo.group(1) + mo.group(2) + SHIFT[mo.group(3)], t)
        n_tab += c1
        for a, b in MAL.items():
            if a in t:
                n_mal += t.count(a); t = t.replace(a, b)
        t, c3 = dang_rx.subn(lambda mo: DANG[mo.group(1)], t)
        n_dang += c3
        if t != orig:
            wt.text = t
s.save(SUPP)
print('table-ref shifts=%d (expect 12) | malformed fixes=%d (expect 3) | dangling 4.1.x fixes=%d (expect 5)'
      % (n_tab, n_mal, n_dang))

# 复扫残留
s2 = Document(SUPP)
rem_tab = sorted(int(mo.group(1)) for i, p in enumerate(s2.paragraphs)
                 for mo in re.finditer(r'(?<![S\d])表\s*([1-9])(?!\d)', p.text))
rem_mal = [i for i, p in enumerate(s2.paragraphs) if re.search(r'正文\s*\d\.\s*正文', p.text)]
rem_dang = [i for i, p in enumerate(s2.paragraphs) if re.search(r'4\.1\.\d', p.text)]
print('all bare 表N in Supp (should be subset of 3-6):', rem_tab)
print('remaining malformed 正文N.正文:', rem_mal)
print('remaining dangling 4.1.x:', rem_dang)
assert n_tab == 12 and n_mal == 3 and n_dang == 5
assert set(rem_tab) <= {3, 4, 5, 6} and not rem_mal and not rem_dang
print('OK round8k')
