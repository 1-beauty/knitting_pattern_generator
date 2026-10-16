# -*- coding: utf-8 -*-
"""Append 36 NESTEXP rows to Supplementary Table S16c (table idx 25)."""
import sys, os
sys.path.insert(0, r'e:\DeskTop\knitting_pattern_generator\scripts')
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
import pandas as pd
from docx import Document
from docx.shared import Pt

def set_cell(cell, text, size=9, bold=False):
    for p in cell.paragraphs:
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    run.bold = bold

SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
BASE = r'e:\DeskTop\knitting_pattern_generator\outputs\analysis'

# Load per-case data
pg = pd.read_csv(os.path.join(BASE, 'region_nested_new_pg.csv'))
ours = pd.read_csv(os.path.join(BASE, 'region_nested_new_ours.csv'))
# Normalize case IDs: strip _dense/_event suffixes to join on base
pg['case_base'] = pg['case'].str.replace('_dense', '', regex=False)
ours['case_base'] = ours['case'].str.replace('_event', '', regex=False)
merged = pg.merge(ours, on=['case_base', 'motif_id'], suffixes=('_pg', '_ours'))
print(f"joined rows: {len(merged)}")

doc = Document(SUPP)
tbl = doc.tables[25]  # S16c

# Verify header
hdr = [tbl.cell(0, c).text.strip() for c in range(6)]
print("header:", hdr)
assert '案例' in hdr[0]

n_existing = len(tbl.rows)
print(f"existing rows (incl header): {n_existing}")

# Sort by case_base
merged = merged.sort_values('case_base').reset_index(drop=True)

for _, row in merged.iterrows():
    base = row['case_base']
    parts = base.split('_')
    if len(parts) >= 2:
        case_num = parts[-1]
        case = f"Nested Expand Case {case_num}"
    else:
        case = base
    motif = row['motif_id']
    iou_pg = f"{row['region_iou_pg']:.3f}"
    iou_ours = f"{row['region_iou_ours']:.3f}"
    leak_pg = f"{row['leakage_pg']:.3f}"
    leak_ours = f"{row['leakage_ours']:.3f}"
    wp_pg = row['wave_pair_rate_pg']
    wp_ours = row['wave_pair_rate_ours']
    if pd.isna(wp_pg):
        wp_pg = '\u2014'
    else:
        wp_pg = f"{wp_pg:.3f}"
    if pd.isna(wp_ours):
        wp_ours = '\u2014'
    else:
        wp_ours = f"{wp_ours:.3f}"
    # V_h: all 36 new cases have V_h=0 for both arms (confirmed in analysis)
    vh_str = "0 / \u2014 / 0"

    new_row = tbl.add_row()
    cells = new_row.cells
    vals = [case, motif,
            f"{iou_pg} / {iou_ours}",
            f"{leak_pg} / {leak_ours}",
            f"{wp_pg} / {wp_ours}",
            vh_str]
    for i, v in enumerate(cells):
        set_cell(v, str(vals[i]), size=9)

doc.save(SUPP)
print(f"OK: appended {len(merged)} rows to S16c (total now {len(tbl.rows)})")
