# -*- coding: utf-8 -*-
"""rev_round10a: CiI style alignment + insertion of 16 new Computers in Industry citations.

Design notes
------------
* Style markers are derived from 21 REAL CiI abstracts (see scripts/_cii_abstracts.txt):
  open on the industrial problem; "this paper presents"; name the industrial actors
  (pattern engineers / designers / practitioners / operators); frame validation as a real
  industrial case study; use business-value vocabulary (lead time, iteration, reuse);
  close with guidance for practitioners and researchers.
* New citations are inserted as NON-NUMERIC placeholders [N1]..[N16] so that the later
  numeric renumbering pass cannot touch them and they cannot be confused with real refs.
* Every target paragraph is asserted oMath==0 before whole-paragraph rewrite, and carries a
  must-keep token list (all pre-existing citation tokens, numbers, terms) asserted before write.
"""
from docx import Document
from docx.oxml.ns import qn
import re

PATH = r"e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx"
d = Document(PATH)
M_OMATH = qn('m:oMath')
WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'\u2019\.\-]*")

DASH = "\u2014"   # em dash
EN = "\u2013"     # en dash
APO = "\u2019"    # typographic apostrophe

def find(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None

def omath(p):
    return p._p.findall('.//' + M_OMATH)

SEG = []

# ============================ ABSTRACT ============================
SEG.append(("abstract", "Flat-knitting action graphs link design semantics",
"Flat knitting is a core technology of flexible manufacturing and mass personalization, yet "
"converting design intent into machine instructions remains manual and experience-dependent: "
"pattern engineers must specify every needle cell" + APO + "s action class, lengthening "
"development lead time and excluding designers without a knitting background. This paper addresses "
"that need by formulating the task as semantics-to-action compilation and presenting a "
"hierarchical semantic-constraint compilation framework with editable generation. A "
"knitting-specific manufacturing intermediate representation, PatternPlan, captures geometry, "
"texture and parent-child hierarchy and is compiled into a two-dimensional action graph through "
"sparse event planning, with hard-constraint checking and conservative self-repair, region-level "
"structural editing, and a conditional flow-matching visual preview for design-stage verification. "
"On a 15-class action space cleaned from the Neural Inverse Knitting dataset, PatternPlan reaches "
"97.2% field-level parsing accuracy over 100 inputs within rule-library coverage. Baseline and "
"controlled experiments show that direct action-matrix generation and visual-image quantization "
"fail on action validity or design semantics, and that replacing event-level compilation with "
"dense filling degrades wave-pair grouping integrity (p<0.001, n=44); against a CP-SAT global "
"constraint-projection baseline, action validity can be closed a posteriori by constraint solving, "
"whereas grouping integrity is injected explicitly by compilation. The action color-code map is a "
"machine-specific-code-agnostic abstract representation; five designs were converted through the "
"partner factory" + APO + "s standard patterning workflow and knitted into samples consistent with "
"the designs (Fig. 7), confirming accuracy in a real manufacturing setting. "
"Within rule-library coverage, semantics-to-action compilation thus ensures action legality and "
"region-level editability, which plain text-to-image generation does not provide.",
["97.2%", "100 inputs", "p<0.001, n=44", "CP-SAT", "PatternPlan", "15-class",
 "Neural Inverse Knitting", "Fig. 7", "grouping integrity"]))

# ============================ 1.1 industry background ============================
SEG.append(("p008", "Flat knitting, a core technology of flexible manufacturing",
"Flat knitting is a core technology of flexible manufacturing and mass personalization, and its "
"industrial practice depends on transforming design intent precisely into machine instructions. "
"In industry the action color map is the dominant pattern representation and machine-input "
"format: pattern engineers must specify the action class of every needle cell by cell in CAD "
"systems such as Stoll M1 Plus and Shima Seiki SDS-ONE APEX [1, 2]. This tabular representation "
"is also the mainstream paradigm in knitting-machine programming research [3, 4], yet its "
"automated generation remains an unmet industrial need. The transformation is still carried out "
"manually by professional pattern engineers. As Industry 4.0 programmes push manufacturing "
"execution systems toward tighter integration and faster response to changing customer demand "
"[N3], this experience-dependent mode has become a production bottleneck, and the availability of "
"skilled personnel is itself reported as a barrier to digital adoption in the textile and "
"clothing industry [N4]: the professional compilation barrier between design intent and "
"manufacturable actions prevents industrial designers without a knitting background from "
"producing production-ready patterns. Removing this barrier so that non-expert designs can enter "
"production directly is an instance of the design-for-X family of approaches, in which "
"manufacturing requirements are addressed at the design stage rather than after it [N1], and it "
"matches the wider objective of reducing or even eliminating the need for expert knowledge in "
"data-driven industrial frameworks [N2]. This paper refers to that goal as democratized "
"design-for-manufacturing and treats it as a core problem in knitting industrial informatics. "
"Designers describe textures and layouts in unstructured natural language, whereas computerized "
"flat-knitting machines require stitch-level actions defined precisely on a discrete grid; this "
"cross-modal, cross-level transformation exhibits a pronounced representation gap" + DASH + "the "
"ambiguity of high-level semantics and the strict constraints of the low-level action space lack "
"an effective mapping mechanism, leading to long and uncontrollable idea-to-product cycles and to "
"a high development lead time.",
["[1, 2]", "[3, 4]", "Stoll M1 Plus", "SDS-ONE APEX", "democratized design-for-manufacturing",
 "representation gap", "[N1]", "[N2]", "[N3]", "[N4]"]))

# ============================ 1.2 challenge opener ============================
SEG.append(("p010", "Although generative AI has progressed substantially",
"Although generative AI has progressed substantially in textile visual design (e.g., ColorwAI [5] "
"and FabricDiffusion [6]; see [7] for a review), applying it to the generation of flat-knitting "
"action color maps remains difficult. Engineering design is commonly characterized as the process "
"of solving technical problems within requirements and constraints in order to create new "
"artifacts [N5]; in flat knitting the constraint side of that characterization is unusually "
"strict, because the artifact is a discrete action structure whose legality is decided by "
"physical process rules rather than by appearance. Integrating data-driven generation into such "
"a setting requires the constraints to be carried explicitly through the pipeline rather than "
"checked afterwards. Three fundamental challenges follow.",
["[5]", "[6]", "[7]", "ColorwAI", "FabricDiffusion", "[N5]"]))

# ============================ 1.3 present work ============================
SEG.append(("p015", "Although color-map generation is the mainstream demand",
"Although color-map generation is the mainstream demand of the flat-knitting industry, existing "
"generative methods mostly remain at the level of visual texture images [5, 6, 13] or linear "
"machine-instruction sequences [15], and no unified framework takes the stitch-level action grid "
"as the direct generation target. To address the difficulty of mapping natural-language design "
"semantics directly onto two-dimensional stitch-action structures, this paper presents a "
"framework that formulates flat-knitting action-map generation as a semantics-to-action "
"compilation problem oriented to discrete manufacturing constraints, centered on hierarchical "
"semantic constraint compilation with editable generation as its goal. The core proposition is "
"to map natural-language design semantics within the rule-library coverage domain into a "
"hierarchical structural plan with parent" + EN + "child region relations together with a sparse "
"event representation, and to generate action grids that satisfy the discrete process rules "
"through event-level constraint compilation rather than pixel-level post-processing. The main "
"contributions are as follows:",
["[5, 6, 13]", "[15]", "semantics-to-action compilation"]))

# ============================ 1.3 factory verification ============================
SEG.append(("p020", "Finally, the proposed system has completed factory verification",
"Finally, the proposed system was validated in a real industrial setting through a factory "
"production test carried out in cooperation with Ningbo Cixing Co., Ltd. (Section 4.5): the "
"action color maps and structure-aware visual preview images were delivered to the factory, "
"converted through its pattern-making workflow, and knitted into finished samples on a real "
"flat-knitting machine; the region geometry and texture structure of the physical samples are "
"consistent with the designs, and the finished products are shown in Fig. 7, confirming the "
"accuracy of the proposed approach under real manufacturing conditions.",
["Ningbo Cixing Co., Ltd.", "Section 4.5", "Fig. 7"]))

# ============================ 2.1 NL-driven design generation ============================
SEG.append(("p023", "In recent years, natural-language-driven generation",
"In recent years, natural-language-driven generation has progressed substantially in image "
"synthesis and instruction-driven editing [16], parametric layout, and 3D model generation "
"[8, 17]; among these, LayoutGPT [18] verified the feasibility of using LLMs to generate spatial "
"layouts directly in stylesheet languages. LLMs have demonstrated the ability to convert design "
"intent into structured scripts: LLM4CAD [8] uses multimodal LLMs to generate three-dimensional "
"CAD models, and Text2CAD [19] supports generating sequential CAD modeling sequences from text "
"prompts at different refinement levels; the two verify the feasibility of parametric modeling "
"under multimodal and pure-text conditions, respectively. Moreover, recent studies have further "
"explored introducing LLMs into CAD-to-manufacturing design flows to enhance generative capability "
"in engineering design [17], and Design2Code [20] provides a contrast from the perspective of "
"automatic conversion from visual designs to front-end code. Within industrial informatics "
"specifically, LLMs have been assessed as digital intelligent assistants in assembly "
"manufacturing, where they improved operators" + APO + " experience, reduced cognitive load and "
"raised the quality of process outputs [N6], and natural language processing has been used to "
"construct design-requirements knowledge bases from unstructured design guidelines [N7]; these "
"studies confirm that unstructured language can be turned into structured, industrially usable "
"artifacts, but neither addresses a discrete manufacturing action space. Existing methods mainly "
"target visual content or geometric model generation; they lack process-semantic modeling for the "
"knitting manufacturing process and cannot generate action structures directly usable for flat "
"knitting. Unlike general visual objects, knit structures require the generated result to align "
"strictly with physical action semantics on a discrete grid, which makes existing vision-driven "
"methods difficult to transfer directly to knitting manufacturing scenarios.",
["[16]", "[8, 17]", "[18]", "[8]", "[19]", "[17]", "[20]", "LayoutGPT", "LLM4CAD", "Text2CAD",
 "Design2Code", "[N6]", "[N7]"]))

# ============================ 2.2 knitting pattern generation ============================
SEG.append(("p025", "Existing research on digital knitting design falls",
"Existing research on digital knitting design falls into three categories. (1) Visual texture "
"generation: ColorwAI [5] and FabricDiffusion [6] use GANs or diffusion models [21, 22] for "
"high-fidelity texture transfer and color matching, with conditional control [23] and "
"Transformer-based architectures [24] improving controllability and scale; Knit-FLUX [13] "
"enhances photorealism via low-rank adaptation but outputs only RGB pixels without action "
"semantics. (2) Instruction inverse recovery and synthesis: Neural Inverse Knitting [15] derives "
"machine instructions from sample images, while deepKnit [25], AutoKnit [26], and Knit Sketching "
"[27] map 3D meshes, cut-and-sew patterns, or text to instruction sequences; these focus on "
"linear instruction streams and struggle with the topological relations of two-dimensional action "
"grids. (3) New representations and compilers: knitting skeletons [28], TopoKnit [29], visual "
"flat-knitting programming [4], KNITIT [30], 3D machine-knitting compilers [31] with their formal "
"semantics and scheduling framework [32], skCAD [33], and KnitGIST [34] simplify the process "
"through high-level primitives but lack flexibility for open-ended natural-language semantics. In "
"the adjacent garment-engineering literature, CAD methods for three-dimensional garment design "
"have been surveyed systematically [N8], and associative-feature approaches have been used to "
"define garment patterns so that high-level design features propagate into manufacturable pattern "
"geometry [N9]; the same design-to-manufacture propagation problem arises in flat knitting, except "
"that the target there is a continuous pattern outline whereas here it is a discrete stitch-action "
"grid. Overall, these methods target visual texture, inverse recovery, or constraint optimization, "
"and none unifies generation from natural-language intent to two-dimensional action structures. "
"From the perspective of representation and constraint-injection timing, existing pipelines split "
"into linear instruction streams (Neural Inverse Knitting [15], knitout [48], 3D knitting "
"compilers [31, 32]), which treat one-dimensional sequences as objects and handle constraints at "
"late-compilation or post-hoc stages, and program-synthesis texture generation (KnitGIST [34]); "
"this paper instead takes the two-dimensional stitch action grid as the direct generation target, "
"injects region attribution and grouped-pairing semantics at the event-planning stage, and "
"supports field-level region editing.",
["[5]", "[6]", "[13]", "[15]", "[25]", "[26]", "[27]", "[28]", "[29]", "[4]", "[30]", "[31]",
 "[32]", "[33]", "[34]", "[48]", "[21, 22]", "[23]", "[24]", "[N8]", "[N9]"]))

# ============================ 2.3 constraint-aware design ============================
SEG.append(("p027", "Physics-constraint-driven generation (constraint-aware generation)",
"Physics-constraint-driven generation (constraint-aware generation) is becoming a research hotspot "
"in manufacturing [35]. Related representational efforts in industrial informatics include "
"formalized data models and taxonomies for collaborative manufacturing environments [N10] and "
"ontology-based knowledge representation that bridges design and process information in additive "
"manufacturing [N11]; both indicate that an explicit manufacturing-oriented representation is a "
"prerequisite for automating downstream decisions, which is precisely the role PatternPlan plays "
"in this paper. Traditional knitting CAD systems (e.g., Stoll M1 Plus and SDS-ONE APEX [1, 2]) "
"support fine-grained stitch editing but are essentially \u201cdrawing + compilation\u201d tools "
"lacking automatic detection of process compliance and real-time repair capability. Academia has "
"recently attempted to introduce process rules into diffusion processes, among which KnitDiff [14] "
"is the closest work to this paper: it likewise targets constraint-aware generation of knitting "
"patterns, but its constraints intervene through post-hoc verification and local optimization, its "
"generation target is not the two-dimensional stitch action grid, and it reports no physical "
"knitting verification; in contrast, this paper places constraints upfront at the event-level "
"compilation stage, takes the two-dimensional action grid as the direct generation target, and "
"completes factory knitting verification. Reinforcement-learning optimization has also been "
"applied to knitting scheduling [36]. In addition, some studies track loop dependencies through "
"automated loop modeling for three-dimensional machine knitting [26]. Nevertheless, existing work "
"mostly focuses on post-hoc verification or local optimization and struggles to systematically "
"coordinate strongly coupled constraints such as loop support, transfer offset, and tension "
"balance during the generation stage. This paper constructs a knowledge-guided action compiler "
"that places constraint checking upfront in the generation pipeline, achieving a closed loop of "
"constraint checking and conservative repair at the generation stage.",
["[35]", "[1, 2]", "[14]", "[36]", "[26]", "Stoll M1 Plus", "SDS-ONE APEX", "KnitDiff",
 "[N10]", "[N11]", "\u201cdrawing + compilation\u201d"]))

# ============================ 2.4 constraint solving ============================
SEG.append(("p029", "Two research lines relate directly",
"Two research lines relate directly to the strong-constraint generation paradigms compared in "
"Section 4. The first is constrained decoding: PICARD [11] checks formal grammar constraints "
"during autoregressive decoding, Grammar Prompting [10] guides LLMs to generate legal programs "
"using a domain-specific-language grammar, and automaton-based decoding [37] generalizes this to "
"regular and context-free constraints; these guarantee formal legality at the symbol level but do "
"not address structural semantics. The second is constraint-solving-driven synthesis: constraint "
"programming [38] declares and solves constraint satisfaction problems over discrete domains, with "
"solvers such as OR-Tools CP-SAT [39] widely verified in scheduling and combinatorial "
"optimization. Within industrial informatics the same technique has been used to generate "
"alternative material handling system configurations [N12], and constraint management in "
"parametric design has been addressed through redundancy elimination to improve design reuse "
"[N13]. In knitting, KnitGIST [34] synthesizes textures and the formal semantics and scheduling "
"framework of knitting compilers [32] incorporates constraint solving into compilation. Both "
"guarantee local constraint satisfiability at the symbolic level, but structural semantics such as "
"region attribution and grouped pairing cannot be derived from the constraints themselves and must "
"be carried explicitly by the generation process. Accordingly, this paper encodes the H1" + EN +
"H8 hard constraints as a CP-SAT model, constructs a global constraint-projection baseline "
"(definitions in Supplementary Section M), and quantitatively examines in Section 4.3.2 the "
"boundary of replacing event-level compilation with constraint solving.",
["[11]", "[10]", "[37]", "[38]", "[39]", "[34]", "[32]", "CP-SAT", "PICARD",
 "Grammar Prompting", "KnitGIST", "Supplementary Section M", "Section 4.3.2", "[N12]", "[N13]",
 "H1" + EN + "H8"]))

# ============================ 3.4 constraint-aware refinement ============================
SEG.append(("p070", "This section integrates action-legality constraints",
"This section integrates action-legality constraints, process heuristic constraints, and "
"grouped-structure integrity checks into the constraint-aware refinement stage; the checking and "
"repair loop is shown in Fig. 5. Their decidability follows from discrete definitions: every rule "
"is checked by a deterministic program on a finite grid, without physical simulation or human "
"judgment. In purpose this corresponds to the semi-automated compliance checking frameworks used "
"elsewhere in industrial informatics, where engineering designs are verified against codified "
"rules by combining ontologies with natural language processing [N14]; the difference is that the "
"rule set here is discrete and the checking is fully deterministic, so no interpretation step is "
"required. Taking Generation Case 01 as an example, Fig. 5 shows four columns: (a) the candidate "
"ActionGrid (this case's real 40\u00d740 color map, 4 action classes), which passes all checks in "
"the first round (0 hard-constraint and 0 process violations, 0 repair rounds, 0 guard deletions), "
"so the loop idles; (b) four deterministic check groups" + DASH + "write-time structural guards "
"(ascending risk, outlines before textures), hard constraints H1" + EN + "H8, process heuristics "
"C1" + EN + "C6, and grouped-structure integrity" + DASH + "each violation v localized to a small "
"support domain support(v) and checked in fixed order; (c) the four operator contracts of "
"conservative repair R(A,v): roll back only stable base-fabric actions, cancel incomplete event "
"groups, delete or downgrade high-risk actions, and never create new events, with modifications "
"confined to support(v); and (d) re-check and release, where the dashed loop returns along an "
"orthogonal path, the violation count is monotonically non-increasing (V_hard converges to 0), and "
"region outlines and the base fabric are not reshaped; on passing, the action color map and machine "
"export are output. The archived statistics give the repair path's true triggering frequency: "
"among 263 runs, 33 found process violations at first check, 17 entered the repair loop, 100 "
"produced conservative guard deletions, and 0 hard violations occurred during detection; because "
"pre-repair grids are not exported, the repair path is shown only through the four operator "
"contracts (aligned with the four chips in columns (b) and (d)).",
["Generation Case 01", "40\u00d740", "H1" + EN + "H8", "C1" + EN + "C6", "support(v)",
 "R(A,v)", "263 runs", "V_hard", "Fig. 5", "[N14]"]))

# ============================ 3.5 visual preview ============================
SEG.append(("p084", "As shown in Fig. 6, the model adopts",
"As shown in Fig. 6, the model adopts a UNet-based [41] conditional backbone with a "
"conditional-flow-matching framework [42] (the same continuous-time paradigm as Rectified Flow "
"[43]), fusing the center and 8-neighborhood action conditions through a neighbor-aware encoder "
"and introducing action-semantics-based process-attribute encoding (encoding, attribute "
"definitions, ODE inference, and sampling in Supplementary Section D; the pipeline is further "
"encapsulated as a desktop prototype supporting interactive incremental editing, Supplementary "
"Section L). Generative models are increasingly applied to engineering design tasks in industrial "
"informatics, for example diffusion models for the intelligent design of component dimensions in "
"reinforced-concrete frames [N15]; the present module follows the same paradigm but conditions on "
"a discrete action map rather than on geometric parameters, and its output is an appearance proxy "
"for design review rather than a design artifact. Taking Generation Case 01 as an example, Fig. 6 "
"shows four columns: (a) the conditional inputs are this case's real action color map A "
"(40\u00d740, classes f_knit/b_knit/f_tuck/f_split) and the neighborhood window N8 (center + 8 "
"neighbors); (b) the neighbor-aware encoder fuses per-class physical priors (tension, elasticity, "
"stability) into a condition vector c injected into the UNet; (c) conditional flow matching learns "
"the velocity field v(x,t,c), and ODE midpoint integration (300 steps, globally correlated noise, "
"seed 42, patch 16) generates the preview x1 from noise x0 on a 20\u00d720 patch grid; and (d) the "
"appearance preview embeds this case's real Stage 4 output (320\u00d7320, the same verified asset "
"as Fig. 1), giving feedback on region boundaries, texture, and continuity. The module reads only "
"the action map A and introduces no yarn material or machine state; the preview is an appearance "
"rendering rather than a yarn-level physical simulation and does not constitute a proof of "
"knittability.",
["[41]", "[42]", "[43]", "Supplementary Section D", "Supplementary Section L",
 "Generation Case 01", "40\u00d740", "f_knit/b_knit/f_tuck/f_split", "N8", "v(x,t,c)",
 "300 steps", "seed 42", "patch 16", "20\u00d720", "320\u00d7320", "Fig. 1", "[N15]"]))

# ============================ 4.5 factory production test ============================
SEG.append(("p130", "The factory production test was carried out",
"The factory production test was carried out on a Cixing CX2-52EC-U computerized flat-knitting "
"machine of Ningbo Cixing Co., Ltd. (Fig. 8; principal specifications in Table 8), a 52-inch "
"single-carriage dual-system machine whose motor-controlled direct-selection-needle cam system "
"executes knit, tuck, and transfer in one-to-one correspondence with the action primitives defined "
"in Fig. 2; the test used the machine's standard production configuration. As a "
"machine-specific-code-agnostic abstract action representation, the action color map was converted "
"by the factory through its pattern-making workflow into the machine program of this model. The "
"five samples were knitted on this machine model; Fig. 8 is a manufacturer product photograph, and "
"the samples serve to verify the feasibility of the proposed approach rather than to characterize "
"the machine's physical output. This validation pattern" + DASH + "confirming generative output "
"physically on the production platform" + DASH + "follows established practice in industrial "
"informatics, where data-driven inverse design has been confirmed by physical validation on the "
"machine platform within specified tolerance bands [N16]; here the verification target is the "
"region geometry and texture of the knitted sample rather than a dimensional tolerance.",
["Cixing CX2-52EC-U", "Ningbo Cixing Co., Ltd.", "Fig. 8", "Table 8", "Fig. 2", "52-inch",
 "[N16]"]))

# ============================ 5 conclusion: practitioner/researcher guidance ============================
SEG.append(("p142", "Overall, the core value of this paper lies in",
"Overall, the core value of this paper lies in connecting design semantics with manufacturing "
"actions through a hierarchical manufacturing IR, and in achieving hard-constraint zero violations "
"of the action representation on the tested inputs (exceptions in Section 4.3.3), region structure, "
"and editability through event-level constraint compilation; the constraint results should be "
"understood as verification under predefined discrete rules. From the stability of semantic parsing "
"(std 2.6 pp), the zero compilation hard violations and grouping integrity (p<0.001), and the "
"quantified distribution" + DASH + "boundary trade-off of the visual preview to the verification "
"on 5 factory-knitted samples, the four-level evidence chain closes the complete link from design "
"intent to physical product. On this basis, the system of this paper has passed the real landing "
"test of a cooperating factory, all 5 representative samples were knitted successfully with region "
"structures and textures consistent with the designs, and the chain from digital generation to "
"physical knitting has been confirmed accurate by the real landing test (Section 4.5; finished "
"products in Fig. 7). At the level of industrial informatics, this framework demonstrates an "
"end-to-end compilation path from natural-language design semantics to manufacturable actions, "
"providing a reusable methodological paradigm for democratized design-for-manufacturing in flat "
"knitting. For practitioners, it offers a route by which designers without a knitting background "
"can obtain production-ready action maps within rule-library coverage, shortening the "
"idea-to-product cycle that currently depends on professional pattern engineers; for researchers, "
"it identifies the representation and constraint-injection points at which further work is "
"worthwhile. Future work will focus on extending knitting-structure rules, region-level structural "
"reconstruction, structure-aware visual evaluation, controlled user studies, and factory landing "
"tests oriented to richer structures and complete patterns.",
["Section 4.3.3", "std 2.6 pp", "p<0.001", "5 factory-knitted samples", "Section 4.5",
 "Fig. 7", "democratized design-for-manufacturing"]))

# ============================ APPLY ============================
print("=== APPLY round10a: style alignment + new CiI citations ===")
report = []
for key, prefix, new, tokens in SEG:
    p = find(prefix)
    assert p is not None, "%s: locate fail (%s)" % (key, prefix[:40])
    assert len(omath(p)) == 0, "%s: oMath!=0 -> refuse whole-paragraph rewrite" % key
    missing = [t for t in tokens if t not in new]
    assert not missing, "%s: new text missing tokens %s" % (key, missing)
    w_before = len(WORD.findall(p.text))
    runs = p.runs
    assert runs, "%s: no runs" % key
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""
    w_after = len(WORD.findall(new))
    report.append((key, w_before, w_after))
    print("  %-8s %4d -> %4d  (%+d)" % (key, w_before, w_after, w_after - w_before))

# abstract word-count guard
abs_p = find("Flat knitting is a core technology")
awc = len(abs_p.text.split())
print("\nabstract whitespace words:", awc)
assert awc <= 250, "abstract exceeds 250 words: %d" % awc

# placeholder inventory
allph = sorted(set(re.findall(r"\[N\d+\]", "\n".join(p.text for p in d.paragraphs))))
print("placeholders present:", allph)
assert len(allph) == 16, "expected 16 placeholders, got %d" % len(allph)

# ---- save ----
try:
    d.save(PATH); saved = PATH
except PermissionError:
    saved = PATH.replace(".docx", "_fix10a.docx")
    d.save(saved); print("WARNING locked, bypass:", saved)
print("saved:", saved)
print("net word change:", sum(a - b for _, b, a in report))
