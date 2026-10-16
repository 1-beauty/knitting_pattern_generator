# -*- coding: utf-8 -*-
"""v2：移除 run 包裹的 brk(无效)，改插 bare <m:brk/> 为 oMath 直接子元素。
位置同 v1：(1)before7 (3)before3 (6)before12+30 (11)before6。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
M = qn('m:oMath'); MR = qn('m:r'); MB = qn('m:brk')
BREAKS = {0: [7], 3: [3], 7: [12, 30], 12: [6]}

doc = Document(MAIN)
for ti, idxs in BREAKS.items():
    om = doc.tables[ti].cell(0, 1)._tc.find('.//' + M)
    # 移除 v1 的 run 包裹 brk（run 仅含一个 brk）
    for r in list(om.findall(MR)):
        if len(r) == 1 and r[0].tag == MB:
            om.remove(r)
    children = list(om)
    for i in idxs:
        brk = OxmlElement('m:brk')
        children[i].addprevious(brk)
    print(f'table[{ti}] {doc.tables[ti].cell(0,2).text.strip()}: bare brk inserted, children now {len(list(om))}')
doc.save(MAIN)
print('saved')
