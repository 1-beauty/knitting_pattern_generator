# -*- coding: utf-8 -*-
"""第七轮修订（Supp）：S1 列名张力等级→切换风险等级；S16c 注补统计披露；
Section I 补三类映射实现边界注；Section F 补工厂测试条件注（慈星/常规横机/5送5成）；
正文文献重编号同步（>=31 减 1）。全部断言通过才落盘。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import re, json
from docx import Document
from docx.oxml.ns import qn

SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
STAT_SUPP = json.load(open(r'scripts\_r7_stat.json', encoding='utf-8'))['STAT_SUPP']
s = Document(SUPP)
ps = s.paragraphs

def find(pre): return next(p for p in ps if p.text.startswith(pre))
def append(p, text):
    import copy
    from docx.oxml import OxmlElement
    nr = copy.deepcopy(p.runs[-1]._r)
    for t in nr.findall(qn('w:t')): nr.remove(t)
    for om in nr.findall('.//' + qn('m:oMath')): om.getparent().remove(om)
    nt = OxmlElement('w:t'); nt.set(qn('xml:space'), 'preserve'); nt.text = text
    nr.append(nt); p._p.append(nr)

# 1) 表 S1 列名 + 全文统一改名
n = 0
for tbl in s.tables:
    for row in tbl.rows:
        for c in row.cells:
            for t in c._tc.findall('.//' + qn('w:t')):
                if t.text and '张力等级' in t.text:
                    t.text = t.text.replace('张力等级', '切换风险等级'); n += 1
for p in ps:
    for t in p._p.findall('.//' + qn('w:t')):
        if t.text and '张力等级' in t.text:
            n += t.text.count('张力等级'); t.text = t.text.replace('张力等级', '切换风险等级')
assert n >= 2, n

# 2) S16c 注补统计
append(find('注：Plan-Grid 与 Plan-Grid+Repair 的结构级指标逐行相同'), STAT_SUPP)

# 3) Section I 末补映射边界注
j_head = next(p for p in ps if p.style.name == 'Heading 3' and p.text.startswith('J. '))
body_style = next(p.style for p in ps if p.style.name == 'Body Text')
np1 = s.add_paragraph('三类映射的实现边界（对应正文 3.3 节说明）：column_xfer 类绞花在 ±1 针动作空间下为位移纹理近似，跨针交叉所需的横移针床（racking）与暂存针逻辑不在当前动作空间；sparse mesh 以 miss 浮线与 transfer 组合呈现镂空视觉效应，需空针留孔的真眼孔（lace）组织不在当前规则库；wave_tuck 为本文以 split+tuck 组合定义的复合动作纹理，不声称对应标准组织学名称。', style=body_style)
j_head._p.addprevious(np1._p)

# 4) Section F 补工厂条件注
f_head = next(p for p in ps if p.style.name == 'Heading 3' and p.text.startswith('Supplementary Section F'))
np2 = s.add_paragraph('工厂测试条件（对应正文 4.5 节）：验证与宁波慈星股份有限公司合作完成；动作色码图由该公司按其标准制版流程转换为机台控制程序，并在该公司常规工业横机上织出；机号与纱线规格随工厂对应织物的常规生产配置，未为本测试单独定制；送厂 5 片、成功织出并寄回 5 片，无返工记录。', style=body_style)
f_head._p.addnext(np2._p)

# 5) 正文文献重编号同步
def remap(tok):
    nums = re.findall(r'\d+', tok)
    out, pos = [], 0
    for nstr in nums:
        i = tok.index(nstr, pos); out.append(tok[pos:i]); pos = i + len(nstr)
        n = int(nstr)
        out.append(str(n - 1 if n >= 31 else n))
    out.append(tok[pos:])
    return ''.join(out)
for t in s.element.body.findall('.//' + qn('w:t')):
    if t.text and re.search(r'\[(\d+(?:,\s*\d+)*)\]', t.text):
        t.text = re.sub(r'\[(\d+(?:,\s*\d+)*)\]', lambda mm: '[' + remap(mm.group(1)) + ']', t.text)

txt = '\n'.join(p.text for p in s.paragraphs)  # 重新取段落列表，包含步骤3/4新插入段
assert '切换风险等级' in txt and '张力等级' not in txt
assert STAT_SUPP[:20] in txt and '宁波慈星股份有限公司合作完成' in txt
assert '[31, 32]' in txt and '[40]' in txt and '[32, 33]' not in txt
s.save(SUPP)
print('OK round7 supp | S1 rename=%d | stat+scope+factory notes inserted | cites remapped' % n)
