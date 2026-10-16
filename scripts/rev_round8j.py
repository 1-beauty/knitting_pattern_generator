# -*- coding: utf-8 -*-
"""第八轮 j（响应评审 P0-1/P0-2 降级口径 + P0-3 正文自足 + P1-5 预览透明度）：
1) 重排正文内容表号：旧 Table/表 2-5 → 3-6（单遍 w:t 级替换，不动 oMath）。
2) 3.4 插入 H1–H8 硬约束简表（新 Table 2，源=Supp T02 前4列，零改数值）+ p066 加"正文表 2"指针。
3) 3.3 插入规则库结构与覆盖度段（源=Supp s151/表S3a，114 条映射，保 [12][26][34]）。
4) 4.4 插入 Stage1–4 全图级指标表（新 Table 7，源=Supp T32，零改数值）+ p117 加"正文表 7"指针与真实配对/OOD 说明。
5) 降级口径：摘要 p002 末句去"better served…than T2I"；结论 p129"系统性下降"→"当前样本下呈下降趋势（方向性证据）"。
6) 修 Supp s147 畸形前缀"正文 3.正文 3.2 节"/"正文 4.正文 3.2 节"。
"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
m = Document(MAIN); s = Document(SUPP)
ps = m.paragraphs


def para(pre):
    return next(p for p in ps if p.text.startswith(pre))


def rewrite(p, new):
    assert not p._p.findall('.//' + qn('m:oMath')), ('oMath', p.text[:20])
    old = set(re.findall(r'\[\d+[^\]]*\]', p.text))
    miss = [c for c in old if c not in new]
    assert not miss, (p.text[:20], miss)
    rr = p.runs; rr[0].text = new
    for r in rr[1:]: r._element.getparent().remove(r._element)


def par_after(ref_p, text, style):
    np = OxmlElement('w:p'); ref_p._p.addnext(np)
    p = Paragraph(np, ref_p._parent); p.style = m.styles[style]
    if text: p.add_run(text)
    return p


def table_after(ref_p, data, style='Table'):
    t = m.add_table(rows=len(data), cols=len(data[0]))
    try: t.style = m.styles[style]
    except Exception: pass
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci); cell.text = str(val)
            if ri == 0:
                for pp in cell.paragraphs:
                    for rr in pp.runs: rr.bold = True
    ref_p._p.addnext(t._tbl)
    return t


# ---------- 1) 重排表号 2-5 -> 3-6（单遍，w:t 级，保 oMath） ----------
SHIFT = {'2': '3', '3': '4', '4': '5', '5': '6'}
rx = re.compile(r'(表|Table)(\s*)([2-5])(?!\d)')
cnt = 0
for p in ps:
    for wt in p._p.findall('.//' + qn('w:t')):
        if wt.text and rx.search(wt.text):
            wt.text, n = rx.subn(lambda mo: mo.group(1) + mo.group(2) + SHIFT[mo.group(3)], wt.text)
            cnt += n
assert cnt == 14, ('renumber count', cnt)  # p090,91,97x2,98,99,102,106,107x3,108x2,109
print('renumber done, replacements =', cnt)

# ---------- 2) H1–H8 简表（新 Table 2）+ p066 指针 ----------
p066 = para('动作图中的四类离散结构错误')
rewrite(p066, p066.text.replace('（逐条定义与修复策略见补充材料表 S3）',
                                '（正文表 2 给出 H1–H8 简表，逐条修复策略见补充材料表 S3）'))
H = [['编号', '约束名称', '检查对象', '违反含义'],
     ['H1', '非法动作索引', '动作编号', '编号不在 15 类动作标签空间内'],
     ['H2', '下方支撑', '成圈/集圈/移针/劈针', '需要支撑的动作下方为空'],
     ['H3', '移针偏移上限', '移针', '横向偏移超过允许范围'],
     ['H4', '移针目标越界', '移针', '目标列超出画布'],
     ['H5', '同行目标冲突', '移针', '多个来源落到同一目标'],
     ['H6', '集圈连续堆叠', '集圈', '同列连续集圈超过上限'],
     ['H7', '辅助控制动作上方清空', '辅助控制动作', '动作上方仍存在不兼容支撑结构'],
     ['H8', '同行床间方向冲突', '床间移针', '同行同时存在前后相反方向床间移针']]
cap2 = par_after(p066, 'Table 2 Action-legality hard constraints H1–H8 (check targets and '
                       'violation semantics; per-constraint repair strategies in Sec. 3.4 and '
                       'Supplementary Table S3).', 'Table Caption')
table_after(cap2, H)

# ---------- 3) 规则库结构与覆盖度段（3.3，插到 p048 之后） ----------
p048 = para('该分解是对两类失败模式的直接回应')
rule_txt = ('本文的知识引导体现在一个可枚举、可审计的语义—工艺规则库上：规则库按（视觉语义 effect，形状 shape）'
            '二元组组织，当前包含 114 条显式映射项，覆盖 10 种视觉语义（plain、reverse、wave_tuck、cable、rib、'
            'links、lace、relief、honeycomb、seed）与 18 种形状取值（11 种闭合几何与 7 种布局型），其依据为针织工艺'
            '手册 [12] 记载的组织结构—动作对应关系，并参考自动机器编织 [26] 与针织纹理程序合成 [34] 研究中的动作语义'
            '整理而成。规则库外的输入经形状别名归一化与优先级回退后仍无法匹配的，统一降级为素色底布，因此语义覆盖范围'
            '是事先可界定的——这也是 4.2 节解析测试集能够限定为“规则库覆盖范围内输入”的前提。需说明三类映射的实现'
            '边界：column_xfer 类绞花在 ±1 针动作空间下为位移纹理近似（跨针交叉所需的横移针床与暂存针逻辑不在当前'
            '动作空间）；sparse mesh 以 miss 浮线与 transfer 组合呈现镂空视觉效应，需空针留孔的真眼孔（lace）组织'
            '不在当前规则库；wave_tuck 为本文以 split+tuck 组合定义的复合动作纹理，不声称对应标准组织学名称。完整'
            '规模统计与典型映射见补充材料表 S3a、表 S20。')
par_after(p048, rule_txt, 'Body Text')

# ---------- 4) 4.4 预览指标表（新 Table 7）+ p117 指针/OOD 说明 ----------
p117 = para('本节评估结构感知视觉预览对第一阶段动作结构信息的呈现能力')
rewrite(p117, '本节评估结构感知视觉预览对第一阶段动作结构信息的呈现能力；鉴于其设计阶段辅助定位与验证边界'
              '（3.5 节、4.6 节），同时报告像素级重建指标（SSIM [44]、LPIPS [45] 等）与结构感知指标。正文表 7 '
              '给出 Stage 1–4 在数据集验证集 50 张真实配对样片上的主要全图级指标（MSE、SSIM、LPIPS、SeamErr_diff、'
              'BoundaryF1、TextureSim）；评估在真实数据配对上进行，系统生成动作图对预览模型属分布外输入，完整全图级'
              '与窗口级指标见补充材料表 S10、表 S9。其权衡为：Stage 4 以更小 patch 与边界感知加权牺牲像素级全局一致性、'
              '换取边界与局部纹理可读性提升，与其设计阶段外观代理定位一致。')
PV = [['Stage', 'MSE ↓', 'SSIM ↑', 'LPIPS ↓', 'SeamErr_diff ↓', 'BoundaryF1 ↑', 'TextureSim ↑'],
      ['Stage 1', '0.0328', '0.636', '0.277', '0.0859', '0.222', '0.917'],
      ['Stage 2', '0.0203', '0.760', '0.172', '0.0565', '0.235', '0.929'],
      ['Stage 3', '0.0205', '0.758', '0.170', '0.0573', '0.234', '0.930'],
      ['Stage 4', '0.0648', '-0.148', '0.515', '0.0774', '0.310', '0.846']]
cap7 = par_after(p117, 'Table 7 Whole-image quantitative evaluation of the structure-aware visual '
                       'preview across Stage 1–4 on 50 real dataset validation pairs (full whole-image '
                       'and window-level metrics in Supplementary Tables S10 and S9).', 'Table Caption')
table_after(cap7, PV)

# ---------- 5) 降级口径 ----------
p002 = para('Flat-knitting action graphs link design semantics')
old_last = ('Within rule-library coverage, flat-knitting action-graph generation is therefore better '
            'served by manufacturing-constraint-oriented semantics-to-action compilation than by plain '
            'text-to-image generation.')
new_last = ('Within rule-library coverage, semantics-to-action compilation thus guarantees action legality '
            'and region-level editability, which plain text-to-image generation does not provide.')
assert old_last in p002.text
rewrite(p002, p002.text.replace(old_last, new_last))

p129 = para('在第二项核心方法上，事件级约束编译')
rewrite(p129, p129.text.replace(
    '以区域内密集填充替换事件级编译会使纹理泄漏率与成组完整率等结构指标系统性下降（4.3.2 节）',
    '以区域内密集填充替换事件级编译会使纹理泄漏率与成组完整率等结构指标在当前样本下呈下降趋势'
    '（方向性证据而非显著性结论，4.3.2 节）'))

m.save(MAIN)

# ---------- 6) 修 Supp s147 畸形前缀 ----------
fixed = 0
for p in s.paragraphs:
    for wt in p._p.findall('.//' + qn('w:t')):
        if wt.text and '正文 3.正文 3.2 节' in wt.text:
            wt.text = wt.text.replace('正文 3.正文 3.2 节', '正文 3.2 节'); fixed += 1
        if wt.text and '正文 4.正文 3.2 节' in wt.text:
            wt.text = wt.text.replace('正文 4.正文 3.2 节', '正文 4.3.2 节'); fixed += 1
s.save(SUPP)
print('supp malformed-prefix fixes =', fixed)

# ---------- 校验 ----------
txt = '\n'.join(p.text for p in m.paragraphs)
assert 'better served' not in txt and '系统性下降' not in txt
assert '正文表 2 给出 H1–H8' in txt and '正文表 7 给出 Stage 1–4' in txt
assert '114 条显式映射项' in txt
aw = len(ps[2].text.split())
print('OK round8j | abstract split-words=%d (<=250: %s)' % (aw, aw <= 250))
assert aw <= 250
