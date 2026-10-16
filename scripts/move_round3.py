# -*- coding: utf-8 -*-
"""第三层精简 · 第三轮（C+D+E）：
C) 4.5 para173 工厂样片逐行 walkthrough 凝缩（图 10 与工厂口径句保留，para172 不动）；
D) para109 的 S1/S3（Type-Coverage 告诫 + TAP 口径限制）与 para124 的撰写/标注协议细节
   元素级切分移入补充材料新 Section H；正文留定义 + 指针；
E) 4.6（178–184）与 5 结论（188–191）深凝缩；局限六 para185 与 177 不动。
纪律：工厂口径不弱化；数字/引用/交叉引用保留；含 oMath 段落只做元素级切分不整段重建。"""
import copy
import re
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MAIN = r'docs\论文to_cii.docx'
SUPP = r'docs\Supplementary_Material.docx'
W, M = qn('w:t'), qn('m:t')

def split_runs(p, positions):
    """在 p.text 字符边界处劈开跨边界的直接子 run（oMath 内 m:r 不动），使 segs 可按元素分组。"""
    cuts = sorted(set(positions))
    pos = 0
    for ch in list(p._p):
        if ch.tag != qn('w:r'):
            continue
        txt = ''.join(x.text or '' for x in ch.iter(W))
        s, e = pos, pos + len(txt)
        pos = e
        inner = [c for c in cuts if s < c < e]
        if not inner:
            continue
        frags, last = [], s
        for c in inner:
            frags.append(txt[last - s:c - s]); last = c
        frags.append(txt[last - s:])
        tnode = ch.find(W)
        tnode.text = frags[0]
        prev = ch
        for fr in frags[1:]:
            nr = copy.deepcopy(ch)
            nr.find(W).text = fr
            prev.addnext(nr)
            prev = nr

def segs(p, bounds):
    """按 p.text 的字符边界把段落子元素（run/oMath）切分为 len(bounds)+1 组；oMath 零宽归当前组。"""
    groups = [[] for _ in range(len(bounds) + 1)]
    pos, gi = 0, 0
    for ch in list(p._p):
        if ch.tag == qn('w:pPr'):
            continue
        txt = ''.join(t.text or '' for t in ch.iter(W))
        while gi < len(bounds) and pos >= bounds[gi]:
            gi += 1
        groups[gi].append(ch)
        pos += len(txt)
    return groups

def new_para_from(sup, elems, body_style, fixes=()):
    np = OxmlElement('w:p')
    for e in elems:
        np.append(copy.deepcopy(e))
    sup.element.body.find(qn('w:sectPr')).addprevious(np)
    para = next(q for q in sup.paragraphs if q._p is np)
    para.style = body_style
    for r in para.runs:
        for old, new in fixes:
            if old in r.text:
                r.text = r.text.replace(old, new)
    return para

doc = docx.Document(MAIN)
ps = doc.paragraphs
ts = [p.text for p in ps]
i109 = next(i for i, t in enumerate(ts) if t.startswith('由式（10）可见'))
i124 = next(i for i, t in enumerate(ts) if t.startswith('为构建语义解析测试集'))
i173 = next(i for i, t in enumerate(ts) if t.startswith('图 10 将上述证据整合'))
p109, p124, p173 = ps[i109], ps[i124], ps[i173]

sup = docx.Document(SUPP)
head_style = next(p.style for p in sup.paragraphs if p.style.name == 'Heading 3')
body_style = next(p.style for p in sup.paragraphs if p.style.name == 'Body Text')

# ---------- D1: para109 切分 ----------
t = p109.text
b = [t.index('目标动作纯度（TAP）'), t.index('TAP 的口径'), t.index('动作多样性定义为')]
split_runs(p109, b)
G = segs(p109, b)
assert len(G) == 4 and all(G[k] for k in range(4))
move109 = G[0] + G[2]
# 正文：删除移动组；在含 F1 的 run 末尾加指针
for e in move109:
    p109._p.remove(e)
hit = [r for r in p109.runs if 'F1' in r.text]
assert len(hit) == 1
hit[0].text = hit[0].text.rstrip() + '其口径限制与使用约束详见补充材料 Section H。'

# ---------- D2: para124 切分 ----------
t = p124.text
b2 = [t.index('撰写前仅告知'), t.index('汇总经去重'), t.index('（标注者间一致性'), t.index('标注者间一致性'), t.index('Parent 按 motif'), t.index('字段级准确率衡量')]
# G1 keep [0,撰写前) | G2 move [撰写前,汇总经去重) | G3 keep [汇总,（标注者间) | G4 keep [（标注者间,Parent) | G5 move [Parent,字段级) | G6 keep
G2 = segs(p124, [b2[0], b2[1], b2[3], b2[4], b2[5]]) if False else None
split_runs(p124, [b2[0], b2[1], b2[3], b2[4], b2[5]])
G2 = segs(p124, [b2[0], b2[1], b2[3], b2[4], b2[5]])
move124 = G2[1] + G2[4]
for e in move124:
    p124._p.remove(e)
ptr = OxmlElement('w:r')
tt = OxmlElement('w:t')
tt.set(qn('xml:space'), 'preserve')
tt.text = '撰写协议与标注口径详见补充材料 Section H）。'
ptr.append(tt)
G2[3][-1].addnext(ptr)

# ---------- 补充材料 Section H ----------
sup.add_paragraph('H. 语义解析测试集撰写协议与评价指标口径（Test-set Writing Protocol and Metric Calibers）', style=head_style)
new_para_from(sup, move109, body_style, fixes=(
    ('见 4.3.1 节', '见正文 4.3.1 节'),
    ('以下述结构级指标为准', '以正文的结构级指标为准'),
))
new_para_from(sup, move124, body_style, fixes=(
    ('不参与计分）。', '不参与计分。'),
))
sup.save(SUPP)
print('supp H created')

# ---------- C: para173 凝缩 ----------
COND173 = ('图 10 将上述证据整合为单张对照表，各行依次给出 PatternPlan 摘要、40×40 动作色码图、结构感知视觉预览图与工厂实织样片。'
    '五个工厂样片均成功织出，区域轮廓、纹理类型与周期走向均与设计对应（如 Sample 3 的连续罗纹纵条、Sample 1/2 圆形与菱形区域内'
    '与设计点阵对应的吊目纹理、Sample 4 的蜂窝状点阵、Sample 5 的 X 形斜向纹理），各样片边框与主体区域的边界位置均与设计网格对应。'
    '上述结果表明，本文生成的动作色码图经工厂打版流程转换后可在真实横机上织出区域结构与纹理同设计一致的织物：'
    '语义到动作编译思路已通过真实制造链路的落地测试，思路的准确性得到实物证实。')
assert not p173._p.findall('.//' + qn('m:oMath'))
runs = p173.runs
runs[0].text = COND173
for r in runs[1:]:
    r._element.getparent().remove(r._element)

# ---------- E: 4.6 与结论深凝缩 ----------
E = {}
E['从方法设计角度看'] = ('从方法设计角度看，PatternPlan 的作用不只是将自然语言转化为中间 JSON 表示，而是为后续动作编译提供显式的几何作用域、'
    '纹理归属和父子层级约束；可编辑生成实验中单字段编辑的局部性（4.3.3 节）正是这一表示层作用的直接体现：语义解析与结构规划应被视为'
    '连接设计意图与制造动作的关键表示层，而非可有可无的前处理步骤。')
E['在评估层面'] = ('在评估层面，本文已将 H1–H8 硬约束形式化为 CP-SAT 约束模型并纳入对比（Plan-Grid+CSP，见 4.1.4 与 4.3.2 节）；'
    '后续可进一步将 C1–C6 工艺启发式约束与成组配对、作用域交集等结构语义编码进约束模型，并探索不依赖 PatternPlan 引导的端到端约束合成。')
E['当前系统仍存在六方面局限'] = ('当前系统仍存在六方面局限。首先，语义到工艺实现的映射依赖预定义规则集合：现有规则覆盖波浪吊目、蜂巢、网眼、罗纹和'
    '部分绞花等典型组织，但尚不能表达任意开放式针织结构；当输入涉及规则库之外的复杂组织时，系统采用保守降级策略，优先保留几何轮廓和动作合法性，'
    '但可能牺牲纹理丰富度和设计细节。')
E['其次，当前自修复机制'] = ('其次，当前自修复机制主要基于局部动作替换，适合处理孤立的支撑缺失、移针目标冲突和成组组织残缺，但对区域级结构冲突'
    '（如大面积密集镂空、连续前后床切换、高密度局部纹理）仍然不足；压力测试中的失败案例表明，后续需要引入区域级重构机制，例如自动降低镂空密度、'
    '插入支撑列、生成加固边框或重新规划局部纹理分布。')
E['第三，本文中的动作切换强度评分'] = ('第三，动作切换强度评分和工艺启发式约束属于代理指标，不能替代真实织造中的纱线张力、断纱风险、脱圈概率或织物平整度测量；'
    '该指标适合在同一规则体系下比较不同动作图的相对风险，后续需结合具体纱线材料、机台参数和实织样片反馈进行校准和验证。')
E['第四，视觉预览模块'] = ('第四，视觉预览模块是基于公开数据集学习得到的条件生成模型，定位为设计阶段的结构预览，而不是物理织造仿真模型：'
    'Stage 1–3 更适合像素级重建，Stage 4 更强调边界响应和局部纹理可读性，但二者都不能替代真实织造结果，也不能作为实际可织造性的最终证明；'
    '未来需要引入更贴近针织结构的评价指标（如面向 slip/tuck 配色组织的纱线级仿真与计算设计方法 [48]）并结合真实样片验证。此外，本文全部评估基于'
    '客观指标与代理指标，尚未包含设计者参与的用户研究；视觉预览与可编辑接口对设计迭代效率的实际增益、以及受控领域之外开放输入的泛化表现，'
    '需在后续受控用户研究与外部设计者独立撰写的开放输入集上验证。')
E['第五，本文在样本级数据清理'] = ('第五，本文在样本级数据清理中删除了包含两类低频动作的完整配对样本，当前模型和约束检查器仅支持过滤后保留的 15 类动作，'
    '对删除的两类动作及其相关组织结构不具备生成或检查能力；覆盖完整动作空间需重新引入相应样本并扩展约束规则。')
E['在第一项核心方法上'] = ('在第一项核心方法上，PatternPlan 显式表达形状、纹理、空间位置、工艺参数和父子层级，将开放式自然语言转化为可编辑、可校验和'
    '可编译的结构计划。100 条受控领域输入上 97.2% 的字段级解析准确率与 83.0% 的样本完全正确率（4.2 节），以及局部编辑实验中两个编辑通道各有'
    '四个变体实现非目标区域动作零改动（4.3.3 节），表明该表示能够较稳定地承接设计意图，并可作为交互式设计的结构控制接口，已封装为桌面原型系统的'
    '增量编辑与版本回溯功能（3.6 节）。')
E['在第二项核心方法上'] = ('在第二项核心方法上，事件级约束编译通过逐层语义降低将结构计划与纹理事件转换为规则化动作网格。基线对比表明，直接从自然语言或'
    '图像生成完整动作矩阵难以同时保证针迹级动作合法性与设计语义（4.3.1 节）；同约束体系下的控变量消融进一步显示，以区域内密集填充替换事件级编译会使'
    '纹理泄漏率与成组完整率等结构指标系统性下降（4.3.2 节）；与 CP-SAT 全局约束投影强基线的对比进一步区分了两类性质不同的差距——动作合法性可由约束'
    '求解后验闭合，而区域归属与成组配对必须由编译过程显式注入。在该链路中，保守自修复与结构守卫作为兜底安全网提升高风险输入下的工艺稳定性；'
    '在此保障下，方法在非模板化开放输入上生成多类动作组合的动作图而未退化为单一平针填充。')
E['作为辅助机制'] = ('作为辅助机制，结构感知视觉预览模块以条件流匹配模型将动作色码图转换为针织外观预览，为设计阶段提供边界、纹理与连续性的视觉反馈；'
    'Stage 1–3 与 Stage 4 之间的重建—边界权衡同时表明，单一图像重建指标不足以完整评价面向设计辅助的针织预览效果（4.4 节）。')
E['总体而言'] = ('总体而言，本文的核心价值在于以层次化制造中间表示连接设计语义与制造动作，以事件级约束编译保证动作表示的合法性、区域结构与可编辑性；'
    '约束结果应理解为预定义离散规则下的验证。在此基础上，本文系统已通过合作工厂的真实落地检测，5 片代表性样片均成功织出且区域结构与纹理均与设计一致，'
    '从数字生成到物理织造的完整链路证实了思路的准确性（4.5 节，成品实物见图 10）。未来工作将围绕针织组织规则扩展、区域级结构重构、结构感知视觉评价、'
    '受控用户研究以及面向更丰富组织与完整花型的工厂落地测试展开。')

nE = 0
for p in ps:
    for pre, new in E.items():
        if p.text.startswith(pre):
            assert not p._p.findall('.//' + qn('m:oMath')), pre
            old_cites = set(re.findall(r'\[\d+[^\]]*\]', p.text))
            miss = [c for c in old_cites if c not in new]
            assert not miss, (pre, miss)
            rr = p.runs
            rr[0].text = new
            for r in rr[1:]:
                r._element.getparent().remove(r._element)
            nE += 1
            break
assert nE == len(E) == 11, nE
doc.save(MAIN)

# ---------- 校验 ----------
chk = docx.Document(MAIN)
cts = [p.text for p in chk.paragraphs]
ref = next(i for i, t in enumerate(cts) if t.strip() == 'References')
tot = sum(len(re.findall(r'[\u4e00-\u9fff]', t)) for t in cts[:ref])
body = '\n'.join(cts[:ref])
for probe in ['补充材料 Section H', '详见补充材料 Section H', '思路的准确性得到实物证实', '[48]', 'Cohen']:
    assert probe in body, probe
print('E paras rewritten:', nE, '| main CN chars now:', tot)
s2 = docx.Document(SUPP)
st = [p.text for p in s2.paragraphs]
assert any(t.startswith('H. 语义解析测试集撰写协议') for t in st)
assert any('TAP 的口径' in t for t in st) and any('撰写前仅告知' in t for t in st)
print('supp H ok, paragraphs:', len(s2.paragraphs))
