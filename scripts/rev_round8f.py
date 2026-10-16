# -*- coding: utf-8 -*-
"""第八轮 f（辅助/次要内容结构性下放 + 辅助段凝缩）：
不触碰任何编号公式（式1–式8 全留正文），只处理无 oMath 的辅助/讨论/代理指标段。
1) 3.5 视觉预览：p078 架构细节下放 Supp Section D，正文保留紧凑桩（保 [41][42][43]）；p077 凝缩（保 [40]）。
2) 4.4 视觉预览评估：p117 凝缩（保 [44][45]）。
3) 4.6 讨论：p125/p126（方法层解读与后续约束建模方向）整段下放 Supp Section N，p124 追加指针；p127 六局限概述凝缩（保 [47][48]）。
4) 3.4：p066 四类结构错误/H1–H8 散文凝缩（逐条定义已在 Supp 表 S3，保 [29][12][26][31][32]）。
5) 4.3.1：p102 Plan-Grid+CSP 凝缩为结果句（机制细节已在 Supp Section K/C）。
全部断言通过才落盘。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import copy, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
m, s = Document(MAIN), Document(SUPP)
ps = m.paragraphs
sectPr_s = s.element.body.findall(qn('w:sectPr'))[-1]


def find(pre): return next(p for p in ps if p.text.startswith(pre))


def rewrite(pre, new):
    p = find(pre)
    assert not p._p.findall('.//' + qn('m:oMath')), ('oMath', pre)
    old_cites = set(re.findall(r'\[\d+[^\]]*\]', p.text))
    miss = [c for c in old_cites if c not in new]
    assert not miss, (pre[:20], miss)
    rr = p.runs; rr[0].text = new
    for r in rr[1:]: r._element.getparent().remove(r._element)
    return p


def append(p, text):
    nr = copy.deepcopy(p.runs[-1]._r)
    for t in nr.findall(qn('w:t')): nr.remove(t)
    for om in nr.findall('.//' + qn('m:oMath')): om.getparent().remove(om)
    nt = OxmlElement('w:t'); nt.set(qn('xml:space'), 'preserve'); nt.text = text
    nr.append(nt); p._p.append(nr)


def supp_anchor(prefix):
    sp = s.paragraphs
    i = next(k for k, p in enumerate(sp) if p.style.name == 'Heading 3' and p.text.startswith(prefix))
    for k in range(i + 1, len(sp)):
        if sp[k].style.name in ('Heading 1', 'Heading 2', 'Heading 3'):
            return sp[k]._p
    return sectPr_s


sbt = next(p.style for p in s.paragraphs if p.style.name == 'Body Text')
rx = re.compile(r'(?<!正文 )(\d\.\d+(?:\.\d+)?) 节')


def move_to_supp(pre, anchor_prefix):
    """整段 deepcopy 到 Supp 指定节末，迁段内指向正文的节号加'正文'前缀，删原段。"""
    p = find(pre)
    el = copy.deepcopy(p._p)
    for t in el.findall('.//' + qn('w:t')):
        if t.text: t.text = rx.sub(r'正文 \1 节', t.text)
    supp_anchor(anchor_prefix).addprevious(el)
    p._p.getparent().remove(p._p)


def supp_add(anchor_prefix, text):
    np = s.add_paragraph(text, style=sbt)
    el = np._p; el.getparent().remove(el)
    supp_anchor(anchor_prefix).addprevious(el)


# ---------- 1) 3.5 视觉预览 ----------
supp_add('D. ', '结构感知视觉预览的完整架构：模型以基于 UNet [41] 的条件生成骨干，经邻居感知条件编码器融合中心动作与 8 邻域动作条件，'
                '并在后期阶段引入基于动作语义定义的工艺属性编码（相对切换风险等级、弹性等级与稳定性等级，见表 S1）；'
                '生成框架采用条件流匹配（Conditional Flow Matching）[42]，与 Rectified Flow [43] 同属连续时间生成范式，'
                '学习噪声 patch 到目标针织 patch 的速度场。')
rewrite('如图 5 所示，模型采用基于 UNet',
        '如图 5 所示，模型采用基于 UNet [41] 的条件生成骨干与条件流匹配框架 [42]（与 Rectified Flow [43] 同属连续时间生成范式），经邻居感知条件编码器融合中心与 8 邻域动作条件，'
        '并引入基于动作语义的工艺属性编码；邻居感知编码、工艺属性定义、ODE 推理与采样细节见补充材料 Section D，'
        '完整管线另封装为支持交互增量编辑的桌面原型系统（补充材料 Section L）。')
rewrite('动作网格是面向机器与规则的离散表示',
        '动作网格是面向机器与规则的离散表示，设计者难以直接从 15 类色码检查区域边界与纹理连续性——这是引入视觉预览模块的直接动因：'
        '将内部动作图转换为针织外观预览图，使第一阶段动作结构可视化。该阶段定义为条件图像生成任务，基于公开数据集配对样本学习动作图到织物外观的映射，'
        '角色边界明确：既辅助非专业设计者理解动作网格、暴露边界断裂与纹理不连续等风险，又不显式模拟纱线、机台参数与织造过程，'
        '既非线圈级物理仿真（如 Yarn-Level Simulation [40]），其输出也不能作为可织造性证明。')

# ---------- 2) 4.4 视觉预览评估 ----------
rewrite('本节评估结构感知视觉预览对第一阶段',
        '本节评估结构感知视觉预览对第一阶段动作结构信息的呈现能力；鉴于其设计阶段辅助定位与验证边界（3.5 节、4.6 节），'
        '同时报告像素级重建指标（SSIM [44]、LPIPS [45] 等）与结构感知指标。Stage 1–4 全图级与窗口级定量结果见补充材料 Section C；'
        '其权衡为：Stage 4 以更小 patch 与边界感知加权牺牲像素级全局一致性、换取边界与局部纹理可读性提升，与其设计阶段外观代理定位一致。')

# ---------- 3) 4.6 讨论：p125/p126 下放 Supp N + p124 指针 + p127 凝缩 ----------
move_to_supp('从方法设计角度看，PatternPlan 的作用不只是', 'N. ')
move_to_supp('在评估层面，本文已将 H1–H8 硬约束形式化为 CP-SAT', 'N. ')
append(find('上述实验表明，本文方法的主要优势在于'),
       '方法层面的详细解读、PatternPlan 表示层作用与后续约束建模方向见补充材料 Section N。')
rewrite('当前系统仍存在六方面局限，此处概述',
        '当前系统仍存在六方面局限（详细展开见补充材料 Section N）：其一，语义到工艺实现的映射依赖预定义规则集合，规则库外复杂组织采用保守降级、'
        '可能牺牲纹理丰富度；其二，自修复以局部动作替换为主，区域级结构冲突（大面积密集镂空、连续前后床切换、高密度局部纹理）仍需区域级重构机制；'
        '其三，动作切换强度评分与工艺启发式约束属代理指标，不能替代真实织造中的纱线张力、断纱风险与织物平整度测量；其四，结构感知视觉预览定位为设计阶段结构预览'
        '而非物理织造仿真，未来需引入更贴近针织结构的纱线级仿真与计算设计方法 [47] 并结合真实样片验证；其五，样本级清理删除了含两类低频动作的配对样本，'
        '暂不覆盖完整动作空间；其六，本文动作色码图是机器专有色码无关的抽象动作表示 [48]，颜色仅用于区分内部动作类别，4.5 节 5 片实物样片的区域几何与纹理均与设计一致，'
        '为思路在真实制造条件下提供了可行性证据（proof-of-concept），批量泛化留待后续工作。工业部署语境中的系统定位、领域可迁移性与人力角色边界亦见补充材料 Section N。')

# ---------- 4) 3.4 四类结构错误/H1–H8 凝缩 ----------
rewrite('动作图中的四类离散结构错误',
        '动作图中的四类离散结构错误——动作索引非法、局部支撑缺失、移针目标冲突、成组结构残缺——在像素级图像生成中没有对应物：'
        '颜色仅是内部动作类别编码，任意相邻颜色组合并不必然合法，须满足线圈支撑、目标针位与床间方向等局部关系。这些关系源于横机物理机制：'
        '线圈下方支撑与移针偏移上限（通常不超过 3 针 [29]）是基本工艺约束 [12]，移针目标冲突与床间方向约束已在自动机器编织 [26] 与针织编译器 [31, 32] 中被形式化处理。'
        '本文将此类不可协商的结构要求定义为 H1–H8 八类硬约束（逐条定义与修复策略见补充材料表 S3）：硬约束不评价外观优劣，'
        '而给出动作网格“具有明确结构含义”的可检查下界，违反任一条即视为结构错误而非风格差异。')

# ---------- 5) 4.3.1 Plan-Grid+CSP 凝缩 ----------
rewrite('在贪心修复之上引入全局约束求解',
        '在贪心修复之上引入全局约束求解（Plan-Grid+CSP）：以 Plan-Grid 密集目标为投影目标，在满足 H1–H8 的可行空间内求解汉明距离最小的网格。'
        '全局投影在 10 条输入上全部取得硬约束零违例（9 条 Plan-Grid 输出本已可行、投影为恒等映射；唯一含残余违例的 Baseline Case 07 上，'
        '确定性规则修复无法消除的 3 处硬违例被以 6 格改动最优消除，平均求解耗时 0.171 s/例，与密集编译本身 0.093 s/例同量级）；'
        '但其结构级指标与 Plan-Grid 完全相同（表 5），合法性闭合并不带来区域结构语义改善。投影机制、外部基线对称修复对比与逐方法 TAP 口径分析见补充材料 Section K 与 Section C。')

# ---------- 校验 ----------
ref = next(i for i, p in enumerate(m.paragraphs) if p.text.strip() == 'References')
txt = '\n'.join(p.text for p in m.paragraphs)
body = '\n'.join(p.text for p in m.paragraphs[:ref])
CN = len(re.findall(r'[\u4e00-\u9fff]', txt))
cited = set(int(x) for t in re.findall(r'\[(\d+(?:[,，]\s*\d+)*)\]', body) for x in re.findall(r'\d+', t))
orphan = [c for c in range(1, 49) if c not in cited]
assert not orphan, ('ORPHAN', orphan)
figs = [int(re.match(r'^Fig\. (\d+)\.', p.text).group(1)) for p in m.paragraphs if re.match(r'^Fig\. \d+\.', p.text)]
assert figs == [1, 2, 3, 4, 5, 6], figs
# 移动段不再出现在正文
for bad in ['从方法设计角度看，PatternPlan 的作用不只是', '在评估层面，本文已将 H1–H8 硬约束形式化为 CP-SAT']:
    assert bad not in txt, bad
stxt = '\n'.join(p.text for p in s.paragraphs)
assert '正文 4.3.3 节' in stxt and '结构感知视觉预览的完整架构' in stxt
m.save(MAIN); s.save(SUPP)
print('OK round8f | main para CN=%d | orphan=%s | figs=%s' % (CN, orphan, figs))
