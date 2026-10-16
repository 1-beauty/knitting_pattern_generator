# -*- coding: utf-8 -*-
"""Round 8o: Add KID + CLIPScore columns to Table 7; update p119 text."""
import sys, copy
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)

# ============================================================
# Helper: three-line borders + set_cell
# ============================================================
def set_cell(cell, text, size=9, bold=False):
    for p in cell.paragraphs:
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    run.bold = bold

def three_line(tbl):
    """Apply three-line (booktabs) border: top, header-bottom, bottom."""
    for row_idx, row in enumerate(tbl.rows):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            # Remove old borders element if any
            old = tcPr.find(qn('w:tcBorders'))
            if old is not None:
                tcPr.remove(old)
            # Top of first row
            if row_idx == 0:
                top = OxmlElement('w:top')
                top.set(qn('w:val'), 'single')
                top.set(qn('w:sz'), '12')
                tcPr.append(borders)
                borders.append(top)
            # Bottom of header row
            if row_idx == 0:
                bot = OxmlElement('w:bottom')
                bot.set(qn('w:val'), 'single')
                bot.set(qn('w:sz'), '6')
                borders.append(bot)
            # Bottom of last row
            if row_idx == len(tbl.rows) - 1:
                bot2 = OxmlElement('w:bottom')
                bot2.set(qn('w:val'), 'single')
                bot2.set(qn('w:sz'), '12')
                tcPr.append(borders)
                borders.append(bot2)
            elif borders.getparent() is None:
                tcPr.append(borders)

# ============================================================
# Rebuild Table 7 (index 16): 7 -> 9 columns
# ============================================================
old_tbl_idx = 16
old_tbl = doc.tables[old_tbl_idx]
old_tbl_el = old_tbl._tbl

# New data
data = [
    ['Stage', 'MSE \u2193', 'SSIM \u2191', 'LPIPS \u2193', 'SeamErr_diff \u2193',
     'BoundaryF1 \u2191', 'TextureSim \u2191', 'KID \u2193', 'CLIPScore \u2191'],
    ['Stage 1', '0.0328', '0.636', '0.277', '0.0859', '0.222', '0.917', '0.110', '0.283'],
    ['Stage 2', '0.0203', '0.760', '0.172', '0.0565', '0.235', '0.929', '0.042', '0.283'],
    ['Stage 3', '0.0205', '0.758', '0.170', '0.0573', '0.234', '0.930', '0.044', '0.283'],
    ['Stage 4', '0.0648', '-0.148', '0.515', '0.0774', '0.310', '0.846', '0.107', '0.282'],
]

n_rows, n_cols = len(data), len(data[0])
new_tbl = doc.add_table(rows=n_rows, cols=n_cols)
# No style assignment — use three_line borders only

for r in range(n_rows):
    for c in range(n_cols):
        bold = (r == 0)
        set_cell(new_tbl.cell(r, c), data[r][c], size=9, bold=bold)

three_line(new_tbl)

# Move new table element to replace old
new_tbl_el = new_tbl._tbl
parent = old_tbl_el.getparent()
idx = list(parent).index(old_tbl_el)
# Remove from wherever add_table put it
parent.remove(new_tbl_el)
parent.insert(idx, new_tbl_el)
# Remove old
parent.remove(old_tbl_el)

print("Table 7 rebuilt (9 cols)")

# ============================================================
# Update p119: add KID, CLIPScore to metric list
# ============================================================
p119 = doc.paragraphs[119]
old_text = p119.text
# Replace: "MSE、SSIM、LPIPS、SeamErr_diff、BoundaryF1、TextureSim）"
# New: "MSE、SSIM、LPIPS、SeamErr_diff、BoundaryF1、TextureSim、KID、CLIPScore）"
target = "BoundaryF1\u3001TextureSim\uff09"
replacement = "BoundaryF1\u3001TextureSim\u3001KID\u3001CLIPScore\uff09"
found = False
for t in p119._p.findall('.//' + qn('w:t')):
    if t.text and target in t.text:
        t.text = t.text.replace(target, replacement)
        found = True
        break
if not found:
    # Try across runs
    full = p119.text
    if target in full:
        # rewrite entire paragraph (safe if no oMath)
        assert not p119._p.findall('.//' + qn('m:oMath')), "p119 has oMath"
        new_full = full.replace(target, replacement)
        # clear runs after first
        runs = p119.runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        runs[0].text = new_full
        found = True
assert found, "p119 metric list not found for update"
print("p119 updated with KID/CLIPScore")

# ============================================================
# Add KID/CLIPScore note to p119 (distribution + alignment metrics)
# Insert explanation after the metrics list
# ============================================================
# Add a sentence explaining KID and CLIPScore role
p119_text = p119.text
# Find insertion point: after "主要全图级指标" section mentioning the metric list
ins_marker = "\u8865\u5145\u6750\u6599\u8868"
# We'll append after the existing text. Add one explanatory sentence at end.
append_note = ("\u5176\u4e2d\uff0cKID\uff08Kernel Inception Distance\uff09\u4f5c\u4e3a\u5c0f\u6837\u672c\u65e0\u504f\u5206\u5e03\u6307\u6807\u8861\u91cf\u751f\u6210\u6837\u672c\u4e0e\u771f\u5b9e\u6837\u672c\u7684\u5206\u5e03\u5dee\u8ddd\uff1b"
               "CLIPScore\uff08\u57fa\u4e8e ViT-B/32 \u56fe\u6587\u4f59\u5f26\u76f8\u4f3c\u5ea6\uff09\u8861\u91cf\u9884\u89c8\u56fe\u4e0e\u8bbe\u8ba1\u63cf\u8ff0\u7684\u8bed\u4e49\u5bf9\u9f50\u7a0b\u5ea6\u3002")
# Append to end of paragraph
p119.add_run(append_note)
print("p119 KID/CLIP explanation appended")

# ============================================================
# Save
# ============================================================
doc.save(SRC)
print("\nOK round8o")
