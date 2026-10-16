# -*- coding: utf-8 -*-
"""
scripts/make_factory_delivery_docx.py
生成送厂试织对照表 Word 文档：开头说明色块含义与预览图呈现效果，
随后 5 行「色码图 | 视觉预览图」并排对照，供工厂对照改错。

色码图采用论文定义的色块颜色（instructions.json 色表，内部色码图 *_output.bmp），
不使用已废弃的工厂机器色表。

运行（本机 base 环境即可，无需 GPU）：
    E:\\ProgramData\\miniconda3\\python.exe scripts/make_factory_delivery_docx.py

输出：outputs/factory_delivery/送厂试织对照表_色码图与预览图.docx
"""
import os
import tempfile

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DELIVERY = os.path.join(ROOT, 'outputs', 'factory_delivery')
RUNS = os.path.join(ROOT, 'outputs', 'runs')
OUT_DOCX = os.path.join(DELIVERY, '送厂试织对照表_色码图与预览图.docx')

# 色块含义（论文色表 instructions.json）：色值 / 针法含义 / 织物表面效果
COLOR_ROWS = [
    ('FF0010', '#FF0010', '前床成圈（正针）', '正面平针，V 形针圈纵列'),
    ('2BCE48', '#2BCE48', '后床成圈（反针）', '正面呈反面针效果（反平针纹理）'),
    ('00A4FF', '#00A4FF', '前床集圈（吊目）', '正面横向凸纹 / 蜂窝状凹点孔眼'),
    ('FFE606', '#FFE606', '前床劈针（成圈+翻针复合，仅 F1/F2 少量格）',
     '波纹纹理的过渡针步，试织时过渡行按打版师惯例处理'),
]

# 5 行对照：编号与图案 / 内部色码图（论文色块） / 预览图 PNG
PAIR_ROWS = [
    ('F1', '方框 + 圆形 + 圆内波纹吊目点阵',
     'patternplan/pp_01/color_map/pp_01_output.bmp',
     'F1_pp01_circle_wave_preview.png'),
    ('F2', '方框 + 菱形 + 菱内波纹吊目点阵',
     'edit/edt_05/color_map/edt_05_output.bmp',
     'F2_edt05_diamond_wave_preview.png'),
    ('F3', '方框 + 框内竖罗纹条',
     'patternplan/pp_04/color_map/pp_04_output.bmp',
     'F3_pp04_frame_rib_preview.png'),
    ('F4', '方框 + 圆环 + 圆内蜂窝吊目点阵',
     'edit_native/edtp_01/color_map/edtp_01_output.bmp',
     'F4_edtp01_circle_honeycomb_preview.png'),
    ('F5', '方框 + X 形斜向交叉条',
     'random/rnd_03/color_map/rnd_03_output.bmp',
     'F5_rnd03_frame_cross_preview.png'),
]

IMG_W = Cm(5.6)          # 并排图宽（两图+编号列合计须 < A4 可用宽 15.9cm）
LABEL_W = Cm(3.6)        # 编号列宽
CELL_PAD = Cm(0.3)       # 图片单元格留白


def shade_cell(cell, hex_fill):
    """给表格单元格填纯色，用作色块示意。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def set_cell_width(cell, width):
    cell.width = width


def upscale_color_map(bmp_path, png_path, size=400):
    """40x40 色码图最近邻放大，保证 Word 里色块边缘锐利。"""
    Image.open(bmp_path).convert('RGB').resize(
        (size, size), Image.Resampling.NEAREST).save(png_path)


def add_pair_table(doc, tmp_dir):
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ('编号与图案', '色码图（论文色块定义，40×40 针）',
                                '视觉预览图（模拟织出效果）')):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
    set_cell_width(hdr[0], LABEL_W)
    set_cell_width(hdr[1], IMG_W + CELL_PAD)
    set_cell_width(hdr[2], IMG_W + CELL_PAD)

    for fid, desc, src_rel, png_name in PAIR_ROWS:
        row = table.add_row().cells
        set_cell_width(row[0], LABEL_W)
        set_cell_width(row[1], IMG_W + CELL_PAD)
        set_cell_width(row[2], IMG_W + CELL_PAD)

        p0 = row[0].paragraphs[0]
        r0 = p0.add_run(fid)
        r0.bold = True
        r0.font.size = Pt(12)
        p0b = row[0].add_paragraph()
        r0b = p0b.add_run(desc)
        r0b.font.size = Pt(9)

        upscaled = os.path.join(tmp_dir, fid + '_colormap_big.png')
        upscale_color_map(os.path.join(RUNS, src_rel), upscaled)
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[1].paragraphs[0].add_run().add_picture(upscaled, width=IMG_W)

        row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[2].paragraphs[0].add_run().add_picture(
            os.path.join(DELIVERY, png_name), width=IMG_W)


def main():
    doc = Document()

    # 中文字体兜底
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    title = doc.add_heading('送厂试织对照表：色码图与视觉预览图', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        '本表供工厂试织时对照改错：左列为色码图（40×40 针，色块含义见第一节），'
        '右列为模型由该色码图推理出的视觉预览图（模拟织出效果，320×320 像素）。'
        '两图逐格对应：色码图每一格对应预览图中 8×8 像素的纹理块，位置严格对齐。')

    doc.add_heading('一、色块含义', level=1)
    ct = doc.add_table(rows=1, cols=4)
    ct.style = 'Table Grid'
    hdr = ct.rows[0].cells
    for cell, text in zip(hdr, ('色块', '色值', '针法含义', '织物表面效果')):
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    for fill, hexval, action, effect in COLOR_ROWS:
        cells = ct.add_row().cells
        shade_cell(cells[0], fill)
        cells[0].paragraphs[0].add_run('')
        cells[1].paragraphs[0].add_run(hexval)
        cells[2].paragraphs[0].add_run(action)
        cells[3].paragraphs[0].add_run(effect)

    doc.add_heading('二、预览图呈现效果与对照方法', level=1)
    doc.add_paragraph(
        '预览图中：前床编织格呈现正面平针（V 形针圈），后床编织格呈现反平针纹理，'
        '前吊目格呈现横向凸纹或蜂窝状凹点孔眼；罗纹条、波纹点阵、蜂窝点阵等纹理'
        '在预览图中与色码图区域一一对应。')
    doc.add_paragraph(
        '对照方法：将试织实物与预览图比对，若某区域纹理不符，按对齐关系定位到色码图'
        '对应格，核对该格色块的针法含义与实际织法是否一致。F1/F2 含少量黄色前床劈针格'
        '（见上表末行），属波纹纹理的过渡针步。')

    doc.add_heading('三、色码图与预览图逐行对照（5 片）', level=1)
    with tempfile.TemporaryDirectory() as tmp_dir:
        add_pair_table(doc, tmp_dir)

    doc.save(OUT_DOCX)
    print('saved:', OUT_DOCX)


if __name__ == '__main__':
    main()
