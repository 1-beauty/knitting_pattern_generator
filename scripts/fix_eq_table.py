# -*- coding: utf-8 -*-
"""把已加制表位的显示公式段改为 1x3 无边框表：[空 | 公式居中 | 编号右对齐]。
保证超宽公式换行时编号仍贴右页边。"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Twips

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)
M = qn('m:oMath')
W = [900, 7226, 900]  # twips, 合计 9026 = 版心宽

def no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'none'); e.set(qn('w:sz'), '0')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto'); borders.append(e)
    tblPr.append(borders)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), '9026')
    tblPr.append(tw)

done = 0
for p in list(doc.paragraphs):
    ppr = p._p.find(qn('w:pPr'))
    if ppr is None or ppr.find(qn('w:tabs')) is None:
        continue
    om = p._p.find('.//' + M)
    if om is None:
        continue
    # 取编号 run（最后一个含 (n) 的 w:r）
    num = None; num_run = None
    for r in p._p.findall(qn('w:r')):
        t = ''.join(x.text or '' for x in r.findall(qn('w:t')))
        if re.match(r'^\(\s*\d+\s*\)$', t.strip()):
            num = t.strip(); num_run = r
    tbl = doc.add_table(rows=1, cols=3)
    no_borders(tbl)
    for ci, w in enumerate(W):
        cell = tbl.cell(0, ci)
        tcPr = cell._tc.get_or_add_tcPr()
        tcw = OxmlElement('w:tcW'); tcw.set(qn('w:type'), 'dxa'); tcw.set(qn('w:w'), str(w))
        tcPr.append(tcw)
    # 公式入中列
    c1 = tbl.cell(0, 1).paragraphs[0]
    for r in list(c1.runs):
        r._r.getparent().remove(r._r)
    om.getparent().remove(om)
    c1._p.append(om)
    c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 编号入右列
    c2 = tbl.cell(0, 2).paragraphs[0]
    for r in list(c2.runs):
        r._r.getparent().remove(r._r)
    if num_run is not None:
        num_run.getparent().remove(num_run)
        c2._p.append(num_run)
    c2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # 替换原段
    p._p.addprevious(tbl._tbl)
    p._p.getparent().remove(p._p)
    done += 1
    print('eq -> table:', num)

doc.save(SRC)
print('converted:', done)
