"""仅另存图7/图9的展示层副本；不重跑GUI，不调用PowerPoint COM。"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
from zipfile import ZipFile

sys.dont_write_bytecode = True
from PIL import Image, ImageChops, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt
import paper_case_labels

ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / 'docs/figures'
FIG9 = FIGURES / 'fig_sim_stages_composite.png'
FIG7 = FIGURES / 'gui/fig_gui_edit_history_composite.png'
OUT9 = FIGURES / 'fig_sim_stages_composite_CaseLabels_EN.pptx'
OUT7 = FIGURES / 'fig_gui_edit_history_composite_CaseLabels_EN.pptx'
CASES9 = ('sem_01', 'sem_03', 'sem_06')
STAGES = tuple(FIGURES / f'fig_sim_{c.replace("_", "")}_stages.png' for c in CASES9)
ROW_Y = (0, 372, 744)
EMU_PER_PX = 9525
FONT = 'Times New Roman'
# v2 原生版式（px @96dpi）：源图烧入的 Arial 16px 列标题条弃用且不嵌入，
# 列标题以原生 Times New Roman 粗体置顶一行；竖排案例标签同字号（视觉一致性决策）。
DATA_X = 68
COL_PITCH = 328
PREVIEW_PX = 320
GUTTER_X, GUTTER_W = 12, 44
HEADER_Y, HEADER_H = 14, 46
SLIDE_ROW_Y = (76, 424, 772)
SLIDE_W, SLIDE_H = 1388, 1108
TITLE_PX = 34
GUI_EVIDENCE = {
    'edtp_00': 'History row 3, column 3: run-ID line.',
    'edtp_01': 'History row 2, column 6: run-ID line.',
    'edtp_02': 'History row 3, column 1: run-ID line.',
    'edtp_03': 'History row 3, column 2: run-ID line.',
    'edtp_04': 'History row 2, column 5: run-ID line.',
    'edtp_05': 'History row 2, column 4: run-ID line.',
    'edt_00': 'History row 3, column 3: baseline description, not the run-ID line.',
    'NEST-04': 'History row 3, columns 4 and 5: visible title prefix; one key entry.',
}
EXPECTED_HASHES = {
    'docs/figures/fig_sim_stages_composite.png': '794f60bd21773fa4a75da2dad64a91aee2ef4da07691e42c45940f652965de83',
    'docs/figures/gui/fig_gui_edit_history_composite.png': '18dc4a01006c1f657fa8a19033dbbb14aaf984decfc8721ff45a97ea2439c859',
    'docs/figures/fig_sim_sem01_stages.png': '397bd2de713336fa4bd3c5d658af459e52dd30bd84d40f1c14620b8f74982ac4',
    'docs/figures/fig_sim_sem03_stages.png': '2746dbaf1d747d76ba61558868027c62fa33e16aa0d7afc9937434f87fbc0947',
    'docs/figures/fig_sim_sem06_stages.png': '3fc4789ba6e5749c3b43e7bf689e742d2cabbc3041b6455a5cc0b022bacf0957',
    'docs/figures/gui/shot_history.png': '69f1f6b0a1ba48e2e586cbeeef0c37e273fb4cc7451ec094c74d9949b80b85f3',
    'docs/figures/gui/shot_workspace_edit.png': '3dc23ffe2583db2f7c73acb2483020e9a2396f4d077a6114c35d4544a0d9c255',
}


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def sha_bytes(blob):
    return hashlib.sha256(blob).hexdigest()


def source_records():
    paths = [FIG9, FIG7, *STAGES, FIGURES / 'gui/shot_history.png',
             FIGURES / 'gui/shot_workspace_edit.png', Path(paper_case_labels.__file__)]
    records = {}
    for path in paths:
        blob = path.read_bytes()
        record = {'sha256': sha_bytes(blob), 'bytes': len(blob)}
        if path.suffix == '.png':
            with Image.open(BytesIO(blob)) as image:
                record['size_px'] = list(image.size)
        records[path.relative_to(ROOT).as_posix()] = record
    return records


def rgb(path_or_blob):
    source = BytesIO(path_or_blob) if isinstance(path_or_blob, bytes) else path_or_blob
    with Image.open(source) as image:
        return image.convert('RGB')


def equal_pixels(left, right):
    return left.size == right.size and left.tobytes() == right.tobytes()


def inspect_sources():
    records = source_records()
    for relpath, expected_hash in EXPECTED_HASHES.items():
        require(records[relpath]['sha256'] == expected_hash,
                f'Visually reviewed source changed: {relpath}')
    image = rgb(FIG9)
    require(image.size == (1360, 1104), 'Unexpected Figure 9 dimensions.')
    core = Image.new('RGB', (1304, 1104), 'white')
    rows = []
    for case, path, y in zip(CASES9, STAGES, ROW_Y):
        row = rgb(path)
        require(row.size == (1304, 360), 'Unexpected stage-row dimensions.')
        for stage in range(4):
            x = stage * COL_PITCH
            box = [56 + x, y + 40, 56 + x + PREVIEW_PX, y + 360]
            require(equal_pixels(image.crop(box), row.crop((x, 40, x + PREVIEW_PX, 360))),
                    f'Stage preview differs: {case} stage {stage + 1}')
            core.paste(row.crop((x, 40, x + PREVIEW_PX, 360)), (x, y + 40))
        rows.append({'run_id': case, 'display_label': paper_case_labels.case_label(case),
                     'source': str(path.relative_to(ROOT)),
                     'preview_crops_px': [[56 + s * COL_PITCH, y + 40, 56 + s * COL_PITCH + PREVIEW_PX, y + 360]
                                          for s in range(4)],
                     'pixel_identical_to_original_stage_previews': True})
    blanked = image.crop((56, 0, 1360, 1104)).copy()
    draw = ImageDraw.Draw(blanked)
    for y in ROW_Y:
        draw.rectangle([0, y, 1304, y + 39], fill='white')
    require(equal_pixels(core, blanked),
            'The retained area contains content outside the twelve stage previews and header strips.')
    removed = image.crop((0, 0, 56, 1104))
    ink_box = ImageChops.difference(removed, Image.new('RGB', removed.size, 'white')).getbbox()
    header_ink = [ImageChops.difference(rgb(p).crop((0, 0, 1304, 40)),
                                        Image.new('RGB', (1304, 40), 'white')).getbbox() for p in STAGES]
    return {'sources': records, 'coordinate_convention': 'Pixels; [left, top, right, bottom); origin at top-left.',
            'figure9': {'removed_external_label_strip_px': [0, 0, 56, 1104],
                        'removed_strip_nonwhite_bbox_px': ink_box,
                        'discarded_burned_in_header_strips_px': [[56, y, 1360, y + 40] for y in ROW_Y],
                        'discarded_header_ink_bbox_px': header_ink,
                        'discarded_header_reason': 'Arial 16px burned-in column titles; replaced by native '
                                                   'Times New Roman text and never embedded.',
                        'retained_full_content_px': [56, 0, 1360, 1104],
                        'rows': rows, 'retained_area_exactly_reconstructs': True},
            'figure7': {'image_size_px': list(rgb(FIG7).size),
                        'crop': 'None; original composite PNG embedded byte-for-byte.',
                        'visible_id_evidence': GUI_EVIDENCE,
                        'evidence_method': 'Direct visual inspection of the original composite and history capture.',
                        'not_inferred': 'No hidden cards, truncated path suffixes or timestamp-to-case associations inferred.'}}


def native_texts(prs):
    return [s.text for s in prs.slides[0].shapes if s.has_text_frame]


def audit_presentations(p9, p7):
    require(Counter(native_texts(p9)) == Counter(
        [f'Stage {s + 1}' for s in range(4)] + [paper_case_labels.case_label(c) for c in CASES9]),
        'Figure 9 must contain exactly four native column titles and three native English case labels.')
    pictures = [s for s in p9.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    require(len(pictures) == 12, 'Figure 9 requires exactly 12 independent stage previews.')
    require({rgb(s.image.blob).size for s in pictures} == {(PREVIEW_PX, PREVIEW_PX)},
            'Burned-in header strips or resampled previews are forbidden.')
    pictures7 = [s for s in p7.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    require(len(pictures7) == 1, 'Figure 7 must retain one independent original composite image.')
    require(pictures7[0].image.blob == FIG7.read_bytes(), 'GUI screenshot bytes changed.')
    require('Original GUI capture; run IDs retained' in native_texts(p7), 'Missing capture provenance statement.')
    all_text7 = '\n'.join(native_texts(p7))
    for run_id in GUI_EVIDENCE:
        require(all_text7.count(paper_case_labels.case_label(run_id)) == 1, f'Duplicate or missing key: {run_id}')
    expected_key = [key_text(c) for c in GUI_EVIDENCE]
    require(Counter(native_texts(p7)) == Counter([
        'English case key', 'Original GUI capture; run IDs retained', *expected_key]),
        'Unexpected native GUI text or duplicate key.')
    require((p9.slide_width, p9.slide_height) == (px(SLIDE_W), px(SLIDE_H)), 'Figure 9 slide size changed.')
    full_gui = rgb(FIG7)
    require((p7.slide_width, p7.slide_height) == (px(full_gui.width), px(full_gui.height + 280)),
            'Figure 7 slide size changed.')
    require(shape_box(pictures7[0]) == [0, 0, full_gui.width, full_gui.height],
            'The original GUI capture was cropped, translated or scaled.')
    geometry9, geometry7 = audit_geometry(p9), audit_geometry(p7)
    core = Image.new('RGB', (1304, 1104), 'white')
    crops = []
    for row_number, (case, path, y) in enumerate(zip(CASES9, STAGES, SLIDE_ROW_Y), 1):
        stage = rgb(path)
        for s in range(4):
            name = f'row{row_number}_stage{s + 1}_preview'
            shape = next(sh for sh in pictures if sh.name == name)
            actual = rgb(shape.image.blob)
            require(equal_pixels(actual, stage.crop((s * COL_PITCH, 40, s * COL_PITCH + PREVIEW_PX, 360))),
                    f'Preview pixels changed: {name}')
            require(shape_box(shape) == [DATA_X + s * COL_PITCH, y, PREVIEW_PX, PREVIEW_PX],
                    f'Unexpected preview placement: {name}')
            core.paste(actual, (s * COL_PITCH, ROW_Y[row_number - 1] + 40))
            crops.append({'picture': name, 'source': path.relative_to(ROOT).as_posix(),
                          'source_crop_px': [s * COL_PITCH, 40, s * COL_PITCH + PREVIEW_PX, 360],
                          'original_composite_crop_px': [56 + s * COL_PITCH, ROW_Y[row_number - 1] + 40,
                                                         56 + s * COL_PITCH + PREVIEW_PX,
                                                         ROW_Y[row_number - 1] + 360],
                          'decoded_rgb_sha256': sha_bytes(actual.tobytes())})
        label = next(s for s in p9.slides[0].shapes if s.name == f'row{row_number}_english_case_label')
        require(label.text == paper_case_labels.case_label(case), 'Wrong row-to-label mapping.')
        require(shape_box(label) == [GUTTER_X, y, GUTTER_W, PREVIEW_PX], 'Case label is not in its external gutter.')
        require(label.text_frame._txBody.bodyPr.get('vert') == 'vert270', 'Case label must be native vertical text.')
    for s in range(4):
        header = next(sh for sh in p9.slides[0].shapes if sh.name == f'column_header_stage{s + 1}')
        require(header.text == f'Stage {s + 1}', 'Wrong column-title mapping.')
        require(shape_box(header) == [DATA_X + s * COL_PITCH, HEADER_Y, PREVIEW_PX, HEADER_H],
                'Column title is not centered over its preview column.')
    retained = rgb(FIG9).crop((56, 0, 1360, 1104))
    blanked = retained.copy()
    draw = ImageDraw.Draw(blanked)
    for y in ROW_Y:
        draw.rectangle([0, y, 1304, y + 39], fill='white')
    require(equal_pixels(core, blanked), 'Reassembled PPTX pictures do not preserve the full original data area.')
    require(all(s.image.blob != FIG9.read_bytes() for s in pictures),
            'The old labelled composite must never be embedded underneath replacement labels.')
    return {
        'figure9': {'independent_stage_previews': 12, 'native_column_titles': 4,
                    'native_case_labels': 3, 'burned_in_header_strips_embedded': False,
                    'original_labelled_composite_embedded': False,
                    'all_retained_pixels_identical': True, 'retained_rgb_sha256': sha_bytes(core.tobytes()),
                    'image_crop_manifest': crops, 'geometry': geometry9},
        'figure7': {'independent_original_images': 1, 'original_png_bytes_identical': True,
                    'embedded_original_sha256': sha_bytes(pictures7[0].image.blob),
                    'crop_px': [0, 0, full_gui.width, full_gui.height],
                    'unique_visible_id_keys': 8, 'geometry': geometry7},
    }


def px(value):
    return round(value * EMU_PER_PX)


def shape_box(shape):
    return [round(v / EMU_PER_PX, 6) for v in (shape.left, shape.top, shape.width, shape.height)]


def key_text(run_id):
    suffix = ' (baseline text)' if run_id == 'edt_00' else ''
    return f'{paper_case_labels.case_label(run_id)} \u2194 {run_id}{suffix}'


def audit_geometry(prs):
    require(len(prs.slides) == 1, 'Only one figure slide is allowed.')
    slide = prs.slides[0]
    names = [s.name for s in slide.shapes]
    require(len(names) == len(set(names)), 'Shape names are not unique.')
    require(not slide.shapes.title, 'An unexpected title placeholder was added.')
    width, height = prs.slide_width / EMU_PER_PX, prs.slide_height / EMU_PER_PX
    bounds, text_metrics = [], []
    for shape in slide.shapes:
        x, y, w, h = shape_box(shape)
        require(0 <= x and 0 <= y and x + w <= width and y + h <= height,
                f'Object outside slide: {shape.name}')
        require(shape.rotation == 0, 'Unexpected object rotation.')
        for other_name, (ox, oy, ow, oh) in bounds:
            overlap = min(x + w, ox + ow) > max(x, ox) and min(y + h, oy + oh) > max(y, oy)
            require(not overlap, f'Overlapping objects: {other_name} / {shape.name}')
        bounds.append((shape.name, [x, y, w, h]))
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            require(all(v == 0 for v in (shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom)),
                    f'Non-destructive picture crop is forbidden: {shape.name}')
            require((w, h) == shape.image.size, f'Raster picture was scaled: {shape.name}')
        else:
            require(shape.has_text_frame, 'Unexpected non-image, non-text object.')
            require(all(ord(c) < 128 or c == '\u2194' for c in shape.text), 'New text must be English.')
            require(len(shape.text_frame.paragraphs) == 1, 'Unexpected multiline native label.')
            p = shape.text_frame.paragraphs[0]
            require(len(p.runs) == 1, 'Unexpected text runs.')
            run = p.runs[0]
            require(run.font.name == FONT, 'Unexpected label font.')
            font_px = round(run.font.size.pt * 96 / 72)
            font_file = 'timesbd.ttf' if run.font.bold else 'times.ttf'
            font = ImageFont.truetype(font_file, font_px)
            x0, y0, x1, y1 = font.getbbox(shape.text)
            vertical = shape.text_frame._txBody.bodyPr.get('vert') == 'vert270'
            avail_w, avail_h = (h, w) if vertical else (w, h)
            require(x1 - x0 + 4 <= avail_w and y1 - y0 + 4 <= avail_h,
                    f'Label text does not fit: {shape.name}')
            text_metrics.append({'name': shape.name, 'text': shape.text, 'box_px': [x, y, w, h],
                                 'font_pt': run.font.size.pt, 'text_fits': True,
                                 'text_orientation': 'vert270' if vertical else 'horizontal'})
    return {'slide_size_px_at_96dpi': [width, height], 'shape_count': len(bounds),
            'overlapping_pairs': [], 'off_slide_shapes': [], 'picture_cropping': False,
            'picture_resampling': False, 'native_text': text_metrics,
            'render_status': 'PPTX structural and font-metric audit only; final rendering delegated to the main agent.'}


def new_figure(width, height, title):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(width), px(height)
    prs.core_properties.title = title
    prs.core_properties.subject = 'English display-name layer; original experimental run IDs retained'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    return prs, slide


def add_picture(slide, name, image_or_blob, x, y):
    if isinstance(image_or_blob, bytes):
        blob = image_or_blob
    else:
        stream = BytesIO()
        image_or_blob.save(stream, format='PNG')
        blob = stream.getvalue()
    size = rgb(blob).size
    shape = slide.shapes.add_picture(BytesIO(blob), px(x), px(y), width=px(size[0]), height=px(size[1]))
    shape.name = name
    return shape


def add_text(slide, name, text, x, y, w, h, size_px=24, bold=False, vertical=False, align='left'):
    shape = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    shape.name = name
    shape.fill.background()
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if vertical:
        frame._txBody.bodyPr.set('vert', 'vert270')
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if (vertical or align == 'center') else PP_ALIGN.LEFT
    p.space_before = p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size_px * 72 / 96)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string('172033')
    return shape


def build_figure9():
    prs, slide = new_figure(SLIDE_W, SLIDE_H, 'Visual preview stages: English semantic case labels')
    for s in range(4):
        add_text(slide, f'column_header_stage{s + 1}', f'Stage {s + 1}',
                 DATA_X + s * COL_PITCH, HEADER_Y, PREVIEW_PX, HEADER_H,
                 size_px=TITLE_PX, bold=True, align='center')
    for i, (case, path, y) in enumerate(zip(CASES9, STAGES, SLIDE_ROW_Y), 1):
        row = rgb(path)
        # 只从无案例标签的原始行素材真实裁切预览区；烧入的Arial表头条弃用不嵌入。
        for stage in range(4):
            x = stage * COL_PITCH
            add_picture(slide, f'row{i}_stage{stage + 1}_preview',
                        row.crop((x, 40, x + PREVIEW_PX, 360)), DATA_X + x, y)
        add_text(slide, f'row{i}_english_case_label', paper_case_labels.case_label(case),
                 GUTTER_X, y, GUTTER_W, PREVIEW_PX, size_px=TITLE_PX, bold=True, vertical=True)
    return prs


def build_figure7():
    blob = FIG7.read_bytes()
    width, height = rgb(blob).size
    prs, slide = new_figure(width, height + 280, 'Original GUI capture with an external English case key')
    # 直接嵌入原PNG字节，不转码、不重画、不裁边、不添加界面内文字。
    add_picture(slide, 'original_gui_capture_unmodified', blob, 0, 0)
    add_text(slide, 'external_key_title', 'English case key', 40, height + 20, width - 80, 36, bold=True)
    add_text(slide, 'capture_provenance', 'Original GUI capture; run IDs retained',
             40, height + 62, width - 80, 34, size_px=22)
    for i, run_id in enumerate(GUI_EVIDENCE):
        row, col = divmod(i, 2)
        col_width = (width - 120) // 2
        add_text(slide, f'external_case_key_{i + 1}', key_text(run_id),
                 40 + col * (col_width + 40), height + 112 + row * 40,
                 col_width, 32, size_px=22)
    return prs


def checked_presentation_bytes(prs):
    stream = BytesIO()
    prs.save(stream)
    blob = stream.getvalue()
    with ZipFile(BytesIO(blob)) as archive:
        require(archive.testzip() is None, 'Corrupt PPTX archive.')
        require(not any('vbaProject' in name for name in archive.namelist()), 'Unexpected macro payload.')
    return blob


def render_spec():
    width7, height7 = rgb(FIG7).size
    specs = []
    for path, size in ((OUT9, (SLIDE_W, SLIDE_H)), (OUT7, (width7, height7 + 280))):
        specs.append({'pptx': str(path), 'pdf': str(path.with_suffix('.pdf')),
                      'png': str(path.with_suffix('.png')), 'slide': 1,
                      'slide_inches': [v / 96 for v in size], 'png_size_px_at_192dpi': [v * 2 for v in size]})
    return {'outputs': specs, 'powerpoint_com_called': False,
            'pdf_rendering': 'Main agent exports the single slide to the named PDF, preserving slide size.',
            'png_command_after_pdf_export': f'& "{sys.executable}" -B "{Path(__file__).resolve()}" --rasterize-pdf'}


def audit_delivery():
    evidence = inspect_sources()
    for output in (OUT9, OUT7):
        require(output.exists(), f'Missing approved output: {output.name}')
    p9, p7 = Presentation(OUT9), Presentation(OUT7)
    evidence['pptx_audit'] = audit_presentations(p9, p7)
    for prs in (p9, p7):
        notes = json.loads(prs.slides[0].notes_slide.notes_text_frame.text)
        require(notes['sources'] == evidence['sources'], 'Embedded source-hash audit does not match the files.')
        require(notes['pptx_audit'] == evidence['pptx_audit'], 'Embedded geometry audit no longer matches the slide.')
    require(source_records() == evidence['sources'], 'An input changed during auditing.')
    evidence['source_hashes_unchanged'] = True
    evidence['render_spec'] = render_spec()
    return evidence


def build_outputs():
    evidence = inspect_sources()
    p9, p7 = build_figure9(), build_figure7()
    evidence['pptx_audit'] = audit_presentations(p9, p7)
    evidence['render_spec'] = render_spec()
    # 审计存入PPTX备注和标准输出，不创建额外清单或文档文件。
    for prs in (p9, p7):
        prs.slides[0].notes_slide.notes_text_frame.text = json.dumps(evidence, ensure_ascii=True, indent=2)
    blob9, blob7 = checked_presentation_bytes(p9), checked_presentation_bytes(p7)
    audit_presentations(Presentation(BytesIO(blob9)), Presentation(BytesIO(blob7)))
    require(source_records() == evidence['sources'], 'An input changed during generation.')
    OUT9.write_bytes(blob9)
    OUT7.write_bytes(blob7)
    return audit_delivery()


def rasterize_pdfs():
    # 只处理主代理已经导出的PDF；不启动Office、GUI、浏览器或远端任务。
    import pymupdf as fitz
    audit_delivery()
    specs = render_spec()['outputs']
    for spec in specs:
        require(Path(spec['pdf']).is_file(), f'Main-agent PDF export required: {spec["pdf"]}')
    result = []
    for spec in specs:
        with fitz.open(spec['pdf']) as doc:
            require(len(doc) == 1, 'Expected a single-slide PDF.')
            expected = [v * 72 for v in spec['slide_inches']]
            require(abs(doc[0].rect.width - expected[0]) < 0.1 and abs(doc[0].rect.height - expected[1]) < 0.1,
                    'PDF dimensions do not match the approved figure slide.')
            pixmap = doc[0].get_pixmap(dpi=192, alpha=False)
            pixmap.save(spec['png'])
            result.append({'png': spec['png'], 'size_px': [pixmap.width, pixmap.height]})
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument('--inspect', action='store_true', help='Read-only source integrity and geometry inspection.')
    modes.add_argument('--audit', action='store_true', help='Read-only audit of both saved PPTX files.')
    modes.add_argument('--render-spec', action='store_true', help='Print exact rendering paths, slide sizes and PNG command.')
    modes.add_argument('--rasterize-pdf', action='store_true', help='Rasterize only the PDFs already exported by the main agent.')
    args = parser.parse_args()
    if args.inspect:
        result = inspect_sources()
    elif args.audit:
        result = audit_delivery()
    elif args.render_spec:
        result = render_spec()
    elif args.rasterize_pdf:
        result = rasterize_pdfs()
    else:
        result = build_outputs()
    print(json.dumps(result, ensure_ascii=True, indent=2))


if __name__ == '__main__':
    main()
