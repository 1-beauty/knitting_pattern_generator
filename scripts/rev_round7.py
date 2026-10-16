# -*- coding: utf-8 -*-
"""第七轮修订（模拟审稿修复·正文）：
R1 摘要中心命题收敛+工厂口径精确化(<=250词)；R2 "开放式"限定；R3 "保证/证实"措辞校准；
R4 97.2%口径披露；R5 弱基线定位；R6 配对Wilcoxon+bootstrap CI披露(数据取自表S16c)；
R7 4.5 慈星合作方/常规横机/送厂5片成5片；R8 预览价值重述；R9 术语/H3/carrier澄清；
R10 映射实现边界；R11 KnitDiff专段+能力定位段；R12 4.6新增部署/迁移/人因三段；
R13 Data availability边界；R14 图注排版工单文字清除；R15 重复文献[4]/[30]合并+全局重编号。
铁律：不编造数字；统计量由表S16c逐案例配对值实算。全部断言通过才落盘。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import copy, re, itertools, random
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
m = Document(MAIN)
ps = m.paragraphs

def find(pre): return next(p for p in ps if p.text.startswith(pre))
def wt(p): return p._p.findall('.//' + qn('w:t'))

# ---------- 统计量（表 S16c 配对逐 motif 值，实算） ----------
IOU_PG = [0.406, 0.750, 0.279, 0.038, 0.698, 0.003, 0.126, 0.307, 0.173]
IOU_OU = [0.452, 0.679, 0.460, 0.167, 0.726, 0.106, 0.048, 0.349, 0.228]
GRP_PG = [1.000, 0.982, 0.973, 0.857, 0.952]
GRP_OU = [1.000, 1.000, 1.000, 1.000, 1.000]

def wilcoxon_exact(d):
    d = [x for x in d if x != 0]
    n = len(d)
    order = sorted(range(n), key=lambda i: abs(d[i]))
    rank = [0] * n
    for r, i in enumerate(order, 1): rank[i] = r
    Wobs = sum(r for i, r in enumerate(rank) if d[i] > 0)
    le = ge = 0
    for signs in itertools.product((1, -1), repeat=n):
        W = sum(r for r, s in zip(rank, signs) if s == 1)
        le += (W <= Wobs); ge += (W >= Wobs)
    return 2 * min(le, ge) / (2 ** n)

d_iou = [round(o - p, 3) for o, p in zip(IOU_OU, IOU_PG)]
p_iou = wilcoxon_exact(d_iou)
mean_d = sum(d_iou) / len(d_iou)
rnd = random.Random(42); B = 10000
means = sorted(sum(rnd.choices(d_iou, k=len(d_iou))) / len(d_iou) for _ in range(B))
ci_lo, ci_hi = means[int(0.025 * B)], means[int(0.975 * B) - 1]
d_grp = [round(o - p, 3) for o, p in zip(GRP_OU, GRP_PG)]
p_grp = wilcoxon_exact(d_grp)
assert 0.15 < p_iou < 0.25 and ci_lo < 0 < ci_hi, (p_iou, ci_lo, ci_hi)
STAT_MAIN = ('对 9 个配对 motif 的配对 Wilcoxon 符号秩检验（表 S16c）显示，Region IoU 差值'
             '（Ours − Plan-Grid，均值 +%.3f）双侧 p = %.3f，配对 bootstrap 95%% 置信区间'
             '（%.3f, %.3f）含零，故该结构差异在当前样本量下为方向性证据而非显著性结论；其方向在四档'
             '形态学设置下保持一致（表 S16e），扩大 motif 样本以检验显著性留待后续工作。'
             % (mean_d, p_iou, ci_lo, ci_hi))
STAT_SUPP = ('统计补充：Region IoU 差值（Ours − Plan-Grid）的精确配对 Wilcoxon 符号秩检验双侧 '
             'p = %.3f（n = 9 motif）；配对 bootstrap（10000 次重采样，种子 42）均值差 95%% 置信区间为'
             '（%.3f, %.3f），含零；成组完整率在 5 条定义输入上 4 正 1 零差，双侧 p = %.3f；泄漏率方向不一致'
             '（4 更低 / 3 更高 / 2 持平），不做显著性检验。上述表明表 5 的结构差异在当前样本量下为方向性证据。'
             % (p_iou, ci_lo, ci_hi, p_grp))

# ---------- 工具 ----------
def span(p, old, new, cnt=1):
    n = 0
    for t in wt(p):
        if t.text and old in t.text:
            n += t.text.count(old); t.text = t.text.replace(old, new)
    assert n == cnt, ('span', old, n)

def append(p, text):
    nr = copy.deepcopy(p.runs[-1]._r)
    for t in nr.findall(qn('w:t')): nr.remove(t)
    for om in nr.findall('.//' + qn('m:oMath')): om.getparent().remove(om)
    nt = OxmlElement('w:t'); nt.set(qn('xml:space'), 'preserve'); nt.text = text
    nr.append(nt); p._p.append(nr)

def rewrite(p, new):
    assert not p._p.findall('.//' + qn('m:oMath')), 'oMath in rewrite target'
    old_cites = set(re.findall(r'\[\d+[^\]]*\]', p.text))
    miss = [c for c in old_cites if c not in new]
    assert not miss, (p.text[:30], miss)
    rr = p.runs; rr[0].text = new
    for r in rr[1:]: r._element.getparent().remove(r._element)

# ---------- R1 摘要 ----------
ABS = ('Flat-knitting action graphs link design semantics to stitch-level action structures, whose '
       'generation must satisfy discrete manufacturing constraints such as loop support, transfer targets '
       'and action combinations. We formulate this task as semantics-to-action compilation and propose a '
       'hierarchical semantic-constraint compilation and editable generation framework. A knitting-specific '
       'manufacturing intermediate representation, PatternPlan, captures geometry, texture and parent-child '
       'hierarchy; it is compiled into a two-dimensional action graph via sparse event planning, with '
       'hard-constraint checking and conservative self-repair suppressing structural conflicts, region-level '
       'structural editing, and a conditional flow-matching visual preview for verification. On a 15-class '
       'action space cleaned from the Neural Inverse Knitting dataset, PatternPlan achieves 97.2% field-level '
       'parsing accuracy over 100 inputs within rule-library coverage. Baseline and controlled experiments show '
       'that direct action-matrix generation and visual-image quantization fail on action validity or design '
       'semantics, and that replacing event-level compilation with dense filling degrades region-structure '
       'metrics; comparison with a CP-SAT global constraint-projection baseline further shows that action '
       'validity can be closed a posteriori by constraint solving, whereas region attribution and grouping must '
       'be injected explicitly by compilation, an advantage that is directional under the current motif sample. '
       'The action color-code map is a machine-specific, color-agnostic abstract action representation; five '
       'representative designs were converted through the partner factory\u2019s standard patterning workflow and '
       'knitted on its routine industrial flat-knitting machines into physical samples consistent with the '
       'designs (Fig. 10), a proof-of-concept of the digital-to-physical pipeline. Within rule-library coverage, '
       'flat-knitting action-graph generation is therefore better served by manufacturing-constraint-oriented '
       'semantics-to-action compilation than by plain text-to-image generation.')
assert len(ABS.split()) <= 250, len(ABS.split())
rewrite(find('Flat-knitting action graphs link design semantics'), ABS)

# ---------- R2/R3/R5/R7 span ----------
span(find('色码图生成虽是横机针织工业的主流需求'), '将开放式自然语言语义映射为', '将规则库覆盖域内的自然语言设计语义映射为')
p014 = find('（1）提出层次化制造中间表示')
span(p014, '连接开放式设计语义与离散制造', '连接规则库覆盖域内的设计语义与离散制造')
span(p014, '将开放式自然语言描述转化为可校验', '将覆盖域内的自然语言描述转化为可校验')
p018 = find('最后，本文系统已通过合作工厂的真实落地检测')
span(p018, '本文系统已通过合作工厂的真实落地检测（4.5 节）', '本文系统已与宁波慈星股份有限公司合作完成工厂验证（4.5 节）')
span(p018, '从数字生成到物理织造的完整链路证实了思路的准确性。', '从数字生成到物理织造的链路在受控 5 例上初步验证了思路的可行性。')
span(find('条件化生成保证。'), '条件化生成保证。', '条件化生成的经验命题。')
span(find('本节是全文证据链的物理落点'), '工厂依据色码图织出样片并寄回成品。',
     '工厂在该公司常规工业横机（宁波慈星股份有限公司）上依据色码图织出样片并寄回成品；机号与纱线规格随工厂对应织物的常规生产配置，未为本测试单独定制，送厂 5 片、成功织出 5 片。')
span(find('图 10 将上述证据整合为单张对照表'), '语义到动作编译思路已通过真实制造链路的落地测试，思路的准确性得到实物证实。',
     '语义到动作编译思路在受控 5 例上通过了真实制造链路的可行性验证：经工厂标准制版流程转换后，常规工业横机可织出区域结构与纹理同设计一致的织物；该验证为可行性证据（proof-of-concept），对批量生产的泛化性留待后续工作。')
span(find('在第一项核心方法上'), '将开放式自然语言转化为可编辑', '将规则库覆盖域内的自然语言转化为可编辑')
p186 = find('总体而言，本文的核心价值在于以层次化制造中间表示连接设计语义')
span(p186, '以事件级约束编译保证动作表示的合法性', '以事件级约束编译在受测输入上实现动作表示的硬约束零违例（例外见 4.3.3 节）')
span(p186, '从数字生成到物理织造的完整链路证实了思路的准确性（4.5 节', '从数字生成到物理织造的链路在受控 5 例上初步验证了思路的可行性（4.5 节')
span(find('学术界近期尝试将工艺规则引入扩散过程') if any(p.text.startswith('学术界近期尝试') for p in ps) else find('受物理约束驱动的生成'),
     '学术界近期尝试将工艺规则引入扩散过程（如 KnitDiff [14]），或通过强化学习优化编织调度',
     '学术界近期尝试将工艺规则引入扩散过程，其中 KnitDiff [14] 是与本文最接近的工作：其同样面向针织图案的约束感知生成，但约束以后置校验与局部优化方式介入、生成目标非二维针迹动作网格、且未报告物理织造验证；本文则将约束前置于事件级编译阶段、以二维动作网格为直接生成目标并完成工厂实织验证。针织调度亦有过强化学习优化')

# ---------- R4/R5/R6/R8/R9/R10/R11 append ----------
append(find('结果如表 2 所示'), '需要强调，上述总体准确率按“未显式给出即默认填充且不计分”的口径计算，显式参数字段仅占 7.7%（54/705，表 S7a），故 97.2% 不应外推为对任意参数组合的解析能力；Position（90.0%）是当前语义解析的主要瓶颈字段。')
append(find('三类外部基线并非与本文方法同等能力的强基线'), '上述三类弱基线的定位是任务困难性诊断；建模立场的主证据为 4.3.2 节的控变量结论——动作合法性可由约束求解后验闭合，而区域归属与成组配对必须由编译过程显式注入。')
append(find('在相同 PatternPlan 下，区域内密集'), STAT_MAIN)
append(find('Stage 4 呈现出不同的指标权衡'), '作为设计阶段的学习型外观代理（learned surrogate），预览的价值在于以秒级代价为设计者提供边界与纹理的视觉反馈、减少物理打样轮次，而非替代物理验证或作为可织造性证明。')
append(find('本文使用的动作标签来源于'), '术语约定：tuck 在本文统一译为“集圈”，“吊目”仅作为“波浪吊目”等纹理名中的行业别名出现；carrier 表示该格的导纱器（纱嘴）控制动作，因不参与成圈纹理形成而归入 DROP 语义类并计入背景集合。此外，xfer_l2 的删除使当前动作空间的最大同床横移为 ±1 针，故 H3（移针偏移上限）在该空间下结构性不激活，本文将其保留为面向完整动作空间的预留断言；其余硬约束与各级守卫的实际触发情况见 4.3 节与补充材料 Section K。')
append(find('规则库按（视觉语义，形状）二元组组织'), '需要说明三类映射的实现边界：column_xfer 类绞花在当前 ±1 针动作空间下为位移纹理近似（跨针交叉所需的横移针床与暂存针逻辑不在当前动作空间）；sparse mesh 以 miss 浮线与 transfer 组合呈现镂空视觉效应，需空针留孔的真眼孔（lace）组织不在当前规则库；wave_tuck 为本文以 split+tuck 组合定义的复合动作纹理，不声称对应标准组织学名称。三者的工艺依据与对照见补充材料 Section I。')
append(find('针对针织领域的数字化设计'), '从表示与约束注入时机看，既有针织管线可分为线性指令流（Neural Inverse Knitting [15]、knitout [49]、三维机器编织编译器 [32, 33]）与程序合成式纹理生成（KnitGIST [35]）两类：前者以一维指令序列为对象、约束多在编译后期或后验阶段处理，后者面向视觉纹理程序；本文的差异在于以二维针迹动作网格为直接生成目标、在事件规划阶段显式注入区域归属与成组配对语义，并字段级支持区域编辑。')

# ---------- R12 4.6 新增三段 ----------
h5 = next(p for p in ps if p.style.name == 'Heading 2' and p.text.strip().startswith('5 结论'))
body_style = next(p.style for p in ps if p.style.name == 'Body Text')
NEW_PARAS = [
 '在工业部署语境中，本系统定位于设计端的意图承接与合法性检查环节：其产出为经约束认证的动作色码图与外观预览，供打版师在其 CAD 流程中复核与转换；色码图到机台控制程序的最后一公里当前仍由工厂标准制版流程完成（4.5 节），本文未消除该环节，而是将其输入从自由绘制替换为合法性已认证的结构化表示。向生产尺度（数百针×上千行）的扩展需分块编译与约束复杂度控制，属后续工作。',
 '本文框架可拆分为领域无关层（语义前端—制造 IR—事件级编译—约束认证）与领域相关层（动作标签集、硬约束与启发式约束定义、语义—工艺规则库）。前者可迁移至其他离散动作制造域：刺绣的针迹序列与换色约束、机织的穿综与提花 draft 及浮长约束、增材制造路径规划的支撑与悬垂约束，均可表述为“动作网格 + 硬约束 + 区域语义”；但各域结构语义的来源不同（如增材制造的支撑约束源于几何力学而非邻域静态规则），迁移须重建领域相关层，本文不声称跨域即插即用。',
 '在人力角色上，本系统应理解为打版师的增强工具而非替代：其承担意图结构化与合法性检查，工艺判断与最后一公里转换仍由人完成；若生成表示进入生产并引发织造故障，验证责任边界应以工厂制版复核为准，本文的约束认证仅覆盖预定义离散规则层面。语义前端依赖外部大语言模型接口，设计文本将离开本地环境，存在设计 IP 与供应商锁定顾虑；架构上前端可替换为本地部署的开源权重模型以规避该风险，接口版本漂移对解析行为的影响需回归验证，本文未做跨版本钉扎实验。',
]
for t in NEW_PARAS:
    np_ = m.add_paragraph(t, style=body_style)
    h5._p.addprevious(np_._p)

# ---------- R13 Data availability ----------
rewrite(find('The semantic-parsing test inputs'),
        'The semantic-parsing test inputs (100 cases), per-case outputs and the code will be made publicly '
        'available with the publication of this article. Factory-test artifacts (color-code maps and physical '
        'sample photographs) are subject to the partner\u2019s confidentiality terms and are available from the '
        'corresponding author on reasonable request; the factory patterning conversion itself is not '
        'reproducible outside the partner\u2019s workflow.')

# ---------- R14 图注排版工单文字 ----------
gcnt = {}
for pat, exp in [(r'All figure text in English, Times New Roman, delivered as native editable PPTX vector objects \(no SVG[,;]?\s*', 3),
                 (r'no bitmap pasting\)\.', 2), (r'all text in Times New Roman; ', 1),
                 (r'the only bitmap is the true Stage 4 preview output\)\.', 1)]:
    n = 0
    for p in ps:
        for t in wt(p):
            if t.text:
                t.text, k = re.subn(pat, '', t.text); n += k
    assert n == exp, (pat, n)
    gcnt[pat[:20]] = n

# ---------- R15 重复文献合并 + 全局重编号 ----------
p023 = find('针对针织领域的数字化设计')
span(p023, '[30]', '[4]')
ent30 = next(p for p in ps if p.text.startswith('[30] '))
ent30._p.getparent().remove(ent30._p)
def remap(tok):
    nums = re.findall(r'\d+', tok)
    out, pos = [], 0
    for nstr in nums:
        i = tok.index(nstr, pos); out.append(tok[pos:i]); pos = i + len(nstr)
        n = int(nstr)
        out.append(str(n - 1 if n >= 31 else n))
    out.append(tok[pos:])
    return ''.join(out)
for t in m.element.body.findall('.//' + qn('w:t')):
    if t.text and re.search(r'\[\d+(?:,\s*\d+)*\]', t.text):
        t.text = re.sub(r'\[(\d+(?:,\s*\d+)*)\]', lambda mm: '[' + remap(mm.group(1)) + ']', t.text)

# ---------- 校验 ----------
txt = '\n'.join(p.text for p in m.paragraphs)
nums = sorted(set(int(n) for t in re.findall(r'\[(\d+(?:,\s*\d+)*)\]', txt) for n in re.findall(r'\d+', t)))
assert nums == list(range(1, 49)), nums[:5] + nums[-5:]
ents = [p.text for p in m.paragraphs if re.match(r'^\[\d+\] ', p.text)]
assert [int(re.match(r'^\[(\d+)\]', e).group(1)) for e in ents] == list(range(1, 49)), len(ents)
assert '条件化生成保证' not in txt and '证实了思路的准确性' not in txt
assert '宁波慈星' in txt and STAT_MAIN[:20] in txt
m.save(MAIN)
print('OK round7 main | wilcoxon p=%.3f CI=(%.3f, %.3f) p_grp=%.3f | abstract words=%d | cites 1-48, entries=%d'
      % (p_iou, ci_lo, ci_hi, p_grp, len(ABS.split()), len(ents)))
import json
json.dump({'STAT_SUPP': STAT_SUPP}, open(r'scripts\_r7_stat.json', 'w', encoding='utf-8'), ensure_ascii=False)
