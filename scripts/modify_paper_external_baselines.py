# -*- coding: utf-8 -*-
"""Modify paper Table 4 + 4.3.1 narrative + references per EXTERNAL_BASELINES_NOTE results."""
import re, copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)
log_lines = []
def log(s):
    log_lines.append(s)
    print(s)

def set_cell(cell, text, bold=False):
    for p in list(cell.paragraphs):
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
    p = cell.paragraphs[0]
    try:
        p.style = doc.styles['Compact']
    except KeyError:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = text.split('\xc2\xa7')
    for i, seg in enumerate(parts):
        if seg == '':
            continue
        r = p.add_run(seg)
        r.bold = bold
        if i % 2 == 1:
            r.font.subscript = True

def three_line(tbl):
    tblPr = tbl._tbl.tblPr
    for b in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(b)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
    borders = OxmlElement('w:tblBorders')
    for edge, sz in (('top', '12'), ('bottom', '12')):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000'); borders.append(e)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'none'); e.set(qn('w:sz'), '0')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto'); borders.append(e)
    tblPr.append(borders)
    for cell in tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn('w:tcBorders')):
            tcPr.remove(old)
        tcb = OxmlElement('w:tcBorders')
        b = OxmlElement('w:bottom'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '8')
        b.set(qn('w:space'), '0'); b.set(qn('w:color'), '000000'); tcb.append(b)
        tcPr.append(tcb)

def replace_para_text(para, new_text):
    style = para.style
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    parts = new_text.split('\xc2\xa7')
    for i, seg in enumerate(parts):
        if seg == '':
            continue
        run = para.add_run(seg)
        if i % 2 == 1:
            run.font.subscript = True
    para.style = style

# ══════════════════════════════════════════════════════════════════
# 1. REPLACE TABLE 4
# ══════════════════════════════════════════════════════════════════
target_tbl = None
for ti, tbl in enumerate(doc.tables):
    if len(tbl.rows) >= 8:
        h = tbl.rows[0].cells[0].text.strip()
        r1 = tbl.rows[1].cells[0].text.strip()
        if 'Method' in h and 'Ours' in r1:
            target_tbl = tbl
            log(f"[1] Found Table 4 at index {ti} ({len(tbl.rows)}x{len(tbl.columns)})")
            break

if target_tbl is None:
    raise RuntimeError("Cannot locate Table 4")

SEP = '\u00a7'
VH = 'V' + SEP + 'h' + SEP
VC = 'V' + SEP + 'c' + SEP

new_rows = [
    ['Method', 'Manuf. rate\u2191', 'Fidelity\u2191', 'Mean ' + VH + '\u2193', 'Mean ' + VC + '\u2193', 't (s)\u2193'],
    ['SD3-medium [49]', '0.000', '0.346', '128.4', '0.0', '17.63'],
    ['PixArt-\u03a3 [50]', '0.000', '0.349', '694.8', '0.0', '15.29'],
    ['VAR-d16 [51]', '0.000', '0.350', '449.8', '0.6', '0.01'],
    ['SD v1.5 [22]', '0.000', '0.313', '739.3', '0.0', '1.07'],
    ['SDXL-Turbo [52]', '0.000', '0.253', '592.6', '0.0', '0.39'],
    ['LLM-Direct [46]', '0.000', '0.738', '64.2', '0.1', '1.11'],
    ['Random Valid Map', '1.000', '0.071', '0.0', '0.0', '1.18'],
    ['Ours (full system)', '0.933', '0.628', '0.2', '0.0', '1.55'],
]

old_ncols = len(target_tbl.columns)
if old_ncols < 6:
    log(f"  Columns {old_ncols}->6, rebuilding table")
    new_tbl = doc.add_table(rows=len(new_rows), cols=6)
    for ri, rowdata in enumerate(new_rows):
        for ci, val in enumerate(rowdata):
            set_cell(new_tbl.cell(ri, ci), val, bold=(ri == 0))
    three_line(new_tbl)
    target_tbl._tbl.addnext(new_tbl._tbl)
    target_tbl._tbl.getparent().remove(target_tbl._tbl)
else:
    while len(target_tbl.rows) > len(new_rows):
        tr = target_tbl.rows[-1]._tr
        target_tbl._tbl.remove(tr)
    while len(target_tbl.rows) < len(new_rows):
        target_tbl.add_row()
    for ri, rowdata in enumerate(new_rows):
        for ci, val in enumerate(rowdata):
            set_cell(target_tbl.rows[ri].cells[ci], val, bold=(ri == 0))
    three_line(target_tbl)
log("[1] Table 4 replaced (8 methods x 30 cases)")

for p in doc.paragraphs:
    if 'Table 4' in p.text and 'Comparison of generation methods' in p.text:
        replace_para_text(p, 'Table 4 Comparison with external generative baselines on manufacturability and semantic fidelity (n\u2009=\u200930).')
        log("[1] Table 4 caption updated")
        break

# ══════════════════════════════════════════════════════════════════
# 2. REWRITE 4.3.1 PARAGRAPHS
# ══════════════════════════════════════════════════════════════════
paras = doc.paragraphs

# --- p101 area: Table 4 intro ---
for i, p in enumerate(paras):
    if '\u8868 4 \u6bd4\u8f83\u4e86' in p.text and '\u5916\u90e8\u57fa\u7ebf' in p.text:
        new_p = (
            "\u8868 4 \u5728\u8bed\u4e49\u89e3\u6790\u6d4b\u8bd5\u96c6\u524d 30 \u6761\u8f93\u5165\uff08sem_01\u2013sem_30\uff0c4 \u540d\u9488\u7ec7\u4e13\u4e1a\u4eba\u5458\u72ec\u7acb\u64b0\u5199\uff09\u4e0a\uff0c"
            "\u5c06\u5b8c\u6574\u7cfb\u7edf\u4e0e\u4e94\u7c7b\u8fd1\u5e74\u901a\u7528\u56fe\u50cf\u751f\u6210\u6a21\u578b\u3001LLM-Direct \u53ca Random Valid Map \u5728\u53ef\u5236\u9020\u7387\u4e0e\u56fe\u6848\u4fdd\u771f\u5ea6\u4e24\u4e2a\u7ef4\u5ea6\u8fdb\u884c\u5bf9\u6bd4\u3002"
            "\u6240\u6709\u751f\u6210\u7c7b\u65b9\u6cd5\u5171\u4eab\u540c\u4e00\u6761\u8bc4\u6d4b\u94fe\u8def\uff1a\u6587\u672c/\u7c7b\u522b\u6761\u4ef6 \u2192 \u751f\u6210\u56fe\u50cf \u2192 40\u00d740 \u6700\u8fd1\u90bb\u8272\u7801\u91cf\u5316 \u2192 "
            "H1\u2013H8 \u786c\u7ea6\u675f\u68c0\u67e5 + \u540c\u65cf\u5f52\u5e76\u4fdd\u771f\u5ea6\u8ba1\u7b97\u3002"
            "\u4e09\u4e2a Tier A \u5f3a\u6a21\u578b\uff08SD3-medium [49]\u3001PixArt-\u03a3 [50]\u3001VAR-d16 [51]\uff09\u4e0e\u4e24\u4e2a Tier B \u6a21\u578b\uff08SD v1.5 [22]\u3001SDXL-Turbo [52]\uff09"
            "\u5728 30 \u4f8b\u4e0a\u53ef\u5236\u9020\u7387\u5747\u4e3a 0.000\uff0c\u5e73\u5747\u786c\u7ea6\u675f\u8fdd\u4f8b " + VH + " \u4ecb\u4e8e 128.4\u2013739.3\uff1b"
            "LLM-Direct \u867d\u4fdd\u771f\u5ea6\u8fbe 0.738\uff0c\u53ef\u5236\u9020\u7387\u540c\u6837\u4e3a 0.000\uff08\u5e73\u5747 " + VH + " = 64.2\uff09\u3002"
            "Random Valid Map \u53ef\u5236\u9020\u7387 1.000 \u4f46\u4fdd\u771f\u5ea6\u4ec5 0.071\uff0c\u8bf4\u660e\u52a8\u4f5c\u5408\u6cd5\u4e0e\u8bbe\u8ba1\u8bed\u4e49\u4e00\u81f4\u4e4b\u95f4\u5b58\u5728\u6839\u672c\u6027\u9e3f\u6c9f\u3002"
            "\u5b8c\u6574\u7cfb\u7edf\u5728 30 \u4f8b\u4e2d 28 \u4f8b\u540c\u65f6\u6ee1\u8db3 " + VH + "=0 \u4e14 " + VC + "=0\uff08\u53ef\u5236\u9020\u7387 0.933\uff09\uff0c"
            "\u4fdd\u771f\u5ea6 0.628\uff0c\u5e73\u5747 " + VH + " = 0.2\uff0c"
            "\u8868\u660e\u4e8b\u4ef6\u7ea7\u7ea6\u675f\u7f16\u8bd1\u80fd\u591f\u5728\u7edf\u4e00\u79bb\u6563\u89c4\u5219\u4e0b\u5c06\u8bbe\u8ba1\u8bed\u4e49\u7a33\u5b9a\u8f6c\u5316\u4e3a\u53ef\u68c0\u67e5\u7684\u52a8\u4f5c\u7f51\u683c\u3002"
        )
        replace_para_text(p, new_p)
        log(f"[2] Rewrote p{i} (Table 4 intro)")
        break

# --- Table 4 note ---
for i, p in enumerate(paras):
    if p.text.strip().startswith('\u6ce8\uff1aPlan-Grid') and 'Section M' in p.text:
        new_note = (
            '\u6ce8\uff1a\u53ef\u5236\u9020\u7387\uff08Manuf. rate\uff09= \u540c\u65f6\u6ee1\u8db3 ' + VH + '=0 \u4e14 ' + VC + '=0 \u7684\u8f93\u51fa\u6bd4\u4f8b\uff1b'
            '\u4fdd\u771f\u5ea6\uff08Fidelity\uff09= \u221a(Texture-F \u00b7 Region-F)\uff0c'
            '\u91c7\u7528\u540c\u65cf\u5f52\u5e76\u53e3\u5f84\uff08f_tuck_2\u2192f_tuck, b_tuck_2\u2192b_tuck, split_v\u2192f_split, xfer_*\u2192xfer\uff09\uff0c'
            '\u5bf9\u6240\u6709\u65b9\u6cd5\u4e00\u81f4\u751f\u6548\uff08\u4e25\u683c\u53e3\u5f84\u5217\u4e8e\u8865\u5145\u6750\u6599\u8868 S11\uff09\u3002'
            'SD3-medium\u3001PixArt-\u03a3\u3001SD v1.5\u3001SDXL-Turbo \u63a5\u53d7\u6587\u672c\u6761\u4ef6\uff0c\u4ee5\u8bbe\u8ba1\u610f\u56fe\u4e3a prompt \u76f4\u63a5\u751f\u6210\uff1b'
            'VAR-d16 \u4ec5\u63a5\u53d7 ImageNet \u7c7b\u522b\u6761\u4ef6\uff08\u53d6 wool \u7c7b 911\uff09\u3002'
            '\u539f Plan-Grid \u7cfb\u5217\u5185\u90e8\u57fa\u7ebf\u5bf9\u6bd4\u7531\u8868 6 \u63a7\u53d8\u91cf\u5b9e\u9a8c\u627f\u8f7d\u3002'
            '\u91cf\u5316\u5668\u654f\u611f\u6027\u5206\u6790\u89c1\u8865\u5145\u6750\u6599\u9644\u5f55 S\u3002\u9010\u6848\u4f8b\u7ed3\u679c\u89c1\u8865\u5145\u6750\u6599\u8868 S11\u3002'
        )
        replace_para_text(p, new_note)
        log(f'[2] Rewrote p{i} (Table 4 note)')
        break

# --- Three-tier evidence chain ---
for i, p in enumerate(paras):
    if '\u4e0a\u8ff0\u7ed3\u679c\u8868\u660e\u542b' in p.text and '\u8868\u9762\u7eb9\u7406\u7c7b\u578b\u8986\u76d6' in p.text:
        new_ev = (
            '\u4e0a\u8ff0\u7ed3\u679c\u6784\u6210\u4e09\u6863\u8bc1\u636e\u94fe\uff1a'
            '\u7b2c\u4e00\uff0c\u5f3a\u6a21\u578b\u8bc1\u4f2a\u2014\u2014\u4e09\u4e2a Tier A \u6a21\u578b\u8986\u76d6\u6269\u6563 Transformer\uff08SD3: Rectified Flow DiT, ICML 2024\uff09\u3001'
            '\u53ef\u63a7 DiT\uff08PixArt-\u03a3, ECCV 2024\uff09\u4e0e\u89c6\u89c9\u81ea\u56de\u5f52\uff08VAR-d16, NeurIPS 2024 Best Paper\uff09\u4e09\u79cd\u5f53\u524d\u6700\u4f18\u8303\u5f0f\uff0c'
            '\u53ef\u7ec7\u7387\u5747\u4e3a 0.000\uff0c\u6392\u9664\u4e86\u4e2a\u522b\u6a21\u578b\u80fd\u529b\u4e0d\u8db3\u7684\u89e3\u91ca\u3002'
            '\u7b2c\u4e8c\uff0c\u8303\u5f0f\u7ea7\u5931\u6548\u2014\u2014\u8de8\u5e74\u4ee3\uff082022\u20132024\uff09\u3001\u8de8\u67b6\u6784\uff08LDM\u2192DiT\u2192Flow\u2192\u81ea\u56de\u5f52\uff09\u7684\u53ef\u7ec7\u7387\u4e00\u81f4\u6027\u4e3a 0 '
            '\u6307\u5411\u201c\u89c6\u89c9\u989c\u8272\u2194\u52a8\u4f5c\u8bed\u4e49\u201d\u7684\u8868\u793a\u65ad\u5c42\uff1a\u901a\u7528\u751f\u56fe\u6a21\u578b\u5728\u50cf\u7d20\u7a7a\u95f4\u4f18\u5316\u89c6\u89c9\u4f3c\u7136\uff0c'
            '\u65e0\u6cd5\u8868\u8fbe\u7ebf\u5708\u652f\u6491\u3001\u79fb\u9488\u76ee\u6807\u3001\u5806\u53e0\u4e0a\u9650\u7b49\u79bb\u6563\u65f6\u5e8f\u7ed3\u6784\u7ea6\u675f\u3002'
            '\u4ee5 PixArt-\u03a3 \u4e3a\u4f8b\uff0c30 \u4f8b\u91cf\u5316\u7ed3\u679c\u4e2d\u6210\u5708\u52a8\u4f5c\u50cf\u7d20\u5360\u6bd4\u4ec5 0.1%\uff0829/30 \u4f8b\u65e0\u4efb\u4f55\u6210\u5708\uff09\uff0c'
            '97.7% \u96c6\u4e2d\u4e8e tuck_2/miss/carrier\uff0c\u8fde\u7ec7\u7269\u57fa\u5e95\u6761\u4ef6\u90fd\u4e0d\u6210\u7acb\u3002'
            '\u7b2c\u4e09\uff0c\u5408\u6cd5\u2260\u8bed\u4e49\u2014\u2014Random Valid Map \u53ef\u7ec7\u7387 1.000 \u4f46\u4fdd\u771f\u5ea6\u4ec5 0.071\uff0c'
            'LLM-Direct \u4fdd\u771f\u5ea6 0.738 \u4f46\u53ef\u7ec7\u7387 0.000\uff1a\u4e24\u7ef4\u6307\u6807\u4e92\u8865\uff0c\u5355\u7ef4\u9ad8\u5206\u4e0d\u6784\u6210\u53ef\u5236\u9020\u82b1\u6837\uff0c'
            '\u53cd\u886c\u8bed\u4e49\u2013\u5408\u6cd5\u6027\u8054\u5408\u4f18\u5316\u7684\u5fc5\u8981\u6027\u3002'
        )
        replace_para_text(p, new_ev)
        log(f'[2] Rewrote p{i} (three-tier evidence)')
        break
