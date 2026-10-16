# -*- coding: utf-8 -*-
"""Generate Cover Letter for Computers in Industry submission."""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
# Set margins
for sec in doc.sections:
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

def add_para(text, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    return p

# Header
add_para('Cover Letter', bold=True, size=14, space_after=12)
add_para('')
add_para('Dear Editor-in-Chief,', space_after=10)
add_para('')

# Paragraph 1: Submission statement
add_para(
    'We submit the manuscript entitled "Hierarchical semantic-constraint compilation and '
    'editable generation of action graphs for flat knitting machines" for consideration by '
    'Computers in Industry.',
    space_after=10
)

# Paragraph 2: Problem + contribution
add_para(
    'This work addresses a core bottleneck in industrial informatics for textile manufacturing: '
    'the gap between natural-language design intent and machine-executable stitch-level action '
    'specifications in flat-knitting production. We formulate this gap as a semantics-to-action '
    'compilation problem and propose a hierarchical framework that (1) parses free-text design '
    'descriptions into a structured, editable intermediate representation (PatternPlan), and '
    '(2) compiles this representation into action color-code maps through sparse event-level '
    'planning with discrete manufacturing constraints, guaranteeing zero hard-violation outputs '
    'on all tested inputs.',
    space_after=10
)

# Paragraph 3: Key evidence
add_para(
    'The experimental evaluation provides a four-level evidence chain: semantic parsing achieves '
    '97.2% field-level accuracy (Cohen\'s kappa = 0.92, run-to-run std 2.6 pp); controlled ablation '
    'on 44 hard-nested inputs demonstrates statistically significant superiority in wave-pair '
    'grouping completeness (p < 0.001) while maintaining constructive V_h = 0 guarantees; '
    'distribution-level metrics (KID) and structural metrics (BoundaryF1) jointly characterize the '
    'fidelity–readability tradeoff of the visual preview; and five physical samples successfully '
    'knitted at a partner factory confirm the end-to-end digital-to-physical pipeline.',
    space_after=10
)

# Paragraph 4: Fit for CiI
add_para(
    'We believe this work aligns strongly with the scope of Computers in Industry. It contributes '
    'a reusable design-for-manufacturing compilation paradigm that democratizes access to '
    'industrial knitting: non-expert designers can input natural-language descriptions and receive '
    'immediately manufacturable action graphs, bypassing the traditional skilled pattern-programming '
    'bottleneck. The constraint-aware compilation approach, hierarchical intermediate representation, '
    'and factory-validated deployment directly address the journal\'s focus on information and '
    'communication technologies in industrial production systems.',
    space_after=10
)

# Paragraph 5: Ethical statements
add_para(
    'This manuscript is original, has not been published previously, and is not under consideration '
    'by any other journal or conference. All authors have read and approved the submission. '
    'No conflicts of interest exist. The research did not involve human subjects or animal testing.',
    space_after=10
)

# Closing
add_para('')
add_para('We look forward to your consideration.', space_after=10)
add_para('')
add_para('Sincerely,', space_after=4)
add_para('<Corresponding Author Name>', space_after=2)
add_para('<Affiliation>', space_after=2)
add_para('<Email>', space_after=2)

doc.save(r'e:\DeskTop\knitting_pattern_generator\docs\Cover_Letter.docx')
print('Cover_Letter.docx generated')
