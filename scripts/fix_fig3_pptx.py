# -*- coding: utf-8 -*-
"""Fix Fig.3 PPTX: remove Part C case example, enlarge fonts, improve spacing."""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from copy import deepcopy
import copy

PPTX_PATH = r'e:\DeskTop\knitting_pattern_generator\docs\figures\Fig3_Event_Level_Constraint_Compilation_EN.pptx'

prs = Presentation(PPTX_PATH)
slide = prs.slides[0]

# === Step 1: Remove case example shapes from Part C (IDs 373-376) ===
remove_ids = {373, 374, 375, 376}
shapes_to_remove = []
for sh in slide.shapes:
    if sh.shape_id in remove_ids:
        shapes_to_remove.append(sh)

sp_tree = slide.shapes._spTree
for sh in shapes_to_remove:
    sp_tree.remove(sh._element)
print(f"Removed {len(shapes_to_remove)} case-example shapes (IDs {remove_ids})")

# === Step 2: Enlarge all fonts by +2pt ===
FONT_MAP = {22: 24, 24: 26, 26: 28, 30: 32}
enlarged = 0
for sh in slide.shapes:
    if not sh.has_text_frame:
        continue
    for para in sh.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                old_pt = int(run.font.size / 12700)
                new_pt = FONT_MAP.get(old_pt, old_pt + 2)
                run.font.size = Pt(new_pt)
                enlarged += 1
print(f"Enlarged {enlarged} runs (22->24, 24->26, 26->28, 30->32 pt)")

# === Step 3: Redistribute Part C chips vertically ===
# After removing case text, redistribute the 3 remaining chips evenly
# Part C panel: T=2.18" to T=8.50" (H=6.32")
# 3 chips of H=1.06" each

chip_ids = [
    # (auto_shape_id, title_id, body_id)
    (364, 365, 366),
    (367, 368, 369),
    (370, 371, 372),
]

panel_top = Inches(2.18)
panel_height = Inches(6.32)
chip_height = Inches(1.06)
n_chips = 3

# Calculate even spacing
total_chip_h = n_chips * chip_height
gap = (panel_height - total_chip_h) / (n_chips + 1)

shape_map = {sh.shape_id: sh for sh in slide.shapes}

for idx, (rect_id, title_id, body_id) in enumerate(chip_ids):
    new_top = panel_top + gap * (idx + 1) + chip_height * idx
    
    # Move rect
    if rect_id in shape_map:
        shape_map[rect_id].top = int(new_top)
    # Move title (offset 0.08" from chip top)
    if title_id in shape_map:
        shape_map[title_id].top = int(new_top + Inches(0.08))
    # Move body (offset 0.51" from chip top)
    if body_id in shape_map:
        shape_map[body_id].top = int(new_top + Inches(0.51))

print(f"Redistributed Part C: 3 chips, gap={gap/914400:.2f}\"")

# === Step 4: Widen Part C chips slightly for better aesthetics ===
# Current width: 3.73", panel inner width ~3.85"
# Make chips slightly wider: 3.73 -> 3.75 (minimal, already near max)
# Actually keep as is - already good width

# === Step 5: Save ===
prs.save(PPTX_PATH)
print(f"\nSaved: {PPTX_PATH}")
print("Done. Next: COM export to PNG/PDF.")
