# -*- coding: utf-8 -*-
"""第八轮 l（第二轮评审的纯写作项，零实验零改数）：
1) 4.2 p092 口径修正：表3分母=6字段×100=600，未显式参数按默认填充并记为正确（故 Params 100/100），
   删除与表格矛盾的"不计分"表述；显式参数实例 54/705(7.7%) 为承载解析难度的子集。
2) 摘要软化强模态："must be injected" → "are instead injected"（去 necessity 断言，保 ≤250 词）。
3) 4.4 p117 把 Stage 4 取舍讲硬：补表7实数（SSIM 0.758→−0.148；BoundaryF1 0.234→0.310），
   明确设计评审场景边界/局部可读性优先于全局像素保真，故 Stage 4 为最终档位。
4) 4.1 p088 补规模边界：全部实验在 40×40 画布与清理后 15 类动作空间上进行。
"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import re
from docx import Document
from docx.oxml.ns import qn
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
m = Document(MAIN); ps = m.paragraphs


def para(pre):
    return next(p for p in ps if p.text.startswith(pre))


def rewrite(p, new):
    assert not p._p.findall('.//' + qn('m:oMath')), ('oMath', p.text[:20])
    old = set(re.findall(r'\[\d+[^\]]*\]', p.text))
    miss = [c for c in old if c not in new]
    assert not miss, (p.text[:20], miss)
    rr = p.runs; rr[0].text = new
    for r in rr[1:]: r._element.getparent().remove(r._element)


# 1) 4.2 口径
p092 = para('结果如表 3 所示：PatternPlan 字段级总体准确率')
rewrite(p092, p092.text.replace(
    '需要强调，总体准确率按“未显式给出即默认填充且不计分”口径计算，显式参数字段仅占 7.7%（54/705，表 S7a），'
    '故 97.2% 不应外推为对任意参数组合的解析能力；',
    '需要强调，总体准确率的分母为六字段×100 条输入=600，其中未显式给出工艺参数的输入按默认值填充并记为正确'
    '（故 Params 计 100/100）；真正承载解析难度的显式参数实例仅占 7.7%（54/705，表 S7a），因此 97.2% 主要由'
    '五个非参数字段贡献，不应外推为对任意参数组合的解析能力；'))

# 2) 摘要软化 must
p002 = para('Flat-knitting action graphs link design semantics')
assert 'must be injected explicitly by compilation' in p002.text
rewrite(p002, p002.text.replace('must be injected explicitly by compilation',
                                'are instead injected explicitly by compilation'))

# 3) 4.4 Stage 4 取舍讲硬
p117 = para('本节评估结构感知视觉预览对第一阶段动作结构信息的呈现能力')
rewrite(p117, p117.text.replace(
    '其权衡为：Stage 4 以更小 patch 与边界感知加权牺牲像素级全局一致性、换取边界与局部纹理可读性提升，'
    '与其设计阶段外观代理定位一致。',
    '其权衡为：Stage 4 以更小 patch 与边界感知加权牺牲像素级全局一致性（SSIM 由 0.758 降至 −0.148）、'
    '换取边界与局部纹理可读性提升（BoundaryF1 由 0.234 升至 0.310，表 7）；在设计评审场景下边界与局部纹理'
    '可读性优先于全局像素保真，故 Stage 4 作为最终预览档位，与其设计阶段外观代理定位一致。'))

# 4) 4.1 规模边界
p088 = para('实验设置摘要：第一阶段语义解析测试集')
rewrite(p088, p088.text.rstrip() + '全部生成与评估实验在 40×40 画布与清理后 15 类动作空间上进行，'
                                   '更大画布与完整动作空间下的结论超出当前实验范围。')

m.save(MAIN)

txt = '\n'.join(p.text for p in m.paragraphs)
assert '不计分' not in txt and 'must be injected' not in txt
assert '40×40 画布与清理后 15 类动作空间' in txt
assert 'SSIM 由 0.758 降至 −0.148' in txt
aw = len(ps[2].text.split())
print('OK round8l | abstract split-words=%d (<=250: %s)' % (aw, aw <= 250))
assert aw <= 250
