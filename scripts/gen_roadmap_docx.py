# -*- coding: utf-8 -*-
"""生成第二篇论文新想法 Word 文档（简洁、自然语言版）。

输出: docs/第二篇论文路线.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(BASE, "docs", "figures", "paper2_roadmap.png")
OUT = os.path.join(BASE, "docs", "第二篇论文路线.docx")

FONT = "Microsoft YaHei"
NAVY = (0x12, 0x39, 0x5F)
CENTER = WD_ALIGN_PARAGRAPH.CENTER


def zh(run, size=11, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def para(doc, text, size=11, bold=False, align=None, after=8, before=0, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    zh(r, size, bold, color)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.4
    return p


doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

para(doc, "第二篇论文的新想法", size=20, bold=True, align=CENTER, after=12)

para(doc, "第一篇已经能把一段文字变成可织造的针织图案，但它要求用户把话说得很准、流程是写死的、"
          "而且一次生成。第二篇我想让它更像“人和 AI 一起做设计”，主要是下面三个新想法。", after=14)

para(doc, "一、多模态输入", size=14, bold=True, color=NAVY, before=4, after=5)
para(doc, "不再只能用文字。用户可以画一张草图，再配几句话，系统把图和文字一起理解成一个设计意图。"
          "这样普通人也能上手——不会写精确的描述，画个大概、说一句“我想要几何花纹的围巾”就够了。"
          "技术上用视觉语言模型（VLM）把草图和文字融合，落到第一篇已有的 PatternPlan 上，"
          "后面的编译和生成可以直接复用。", after=14)

para(doc, "二、输入之后能一直改，而且不挑输入的标准", size=14, bold=True, color=NAVY, before=4, after=5)
para(doc, "这一条有两层意思。", after=6)
para(doc, "一是持续修改。用户给完草图和文字，系统不直接出最终图，而是先给一版草案和示意图；"
          "用户觉得哪里不对就继续说，系统再改，来回几轮，满意了才真正生成色码图和针织图。"
          "这样试错很便宜，不会一次生成不满意就得全部重来。", after=6)
para(doc, "二是不同标准的输入。草图可以是专业的，也可以是随手涂鸦，系统要能自适应："
          "看得懂的就精确还原，看不懂的用合理默认补上，遇到关键的歧义就主动问用户。", after=14)

para(doc, "三、让 AI 参与更多的 Agent 系统", size=14, bold=True, color=NAVY, before=4, after=5)
para(doc, "第一篇是一条写死的流水线，八个步骤按固定顺序跑。第二篇我想让一个 Agent 来主导整个过程："
          "它负责理解输入、维护对话和草案的状态，判断什么时候该补全、什么时候该问、"
          "什么时候调用生成或修复工具、什么时候可以停下来交付。AI 从“流水线里被动的一步”"
          "变成“能看情况主动决策的中枢”，这也是和第一篇最大的区别。", after=14)

para(doc, "这三个想法合起来，就是把一个专业工具，变成一个普通人能和 AI 对话、边聊边改、"
          "最后拿到可织造图案的设计系统。", bold=True, after=16)

doc.add_picture(FIG, width=Inches(6.6))
doc.paragraphs[-1].alignment = CENTER
para(doc, "图 1  三个新想法与整体流程", size=9.5, align=CENTER, color=(0x66, 0x66, 0x66), after=4)

try:
    doc.save(OUT)
    print("saved:", OUT)
except PermissionError:
    alt = OUT.replace(".docx", "_新版.docx")
    doc.save(alt)
    print("原文件被 Word 占用，已另存为:", alt)
