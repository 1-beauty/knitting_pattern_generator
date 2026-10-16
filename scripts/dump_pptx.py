import sys
from pptx import Presentation
from pptx.util import Emu

path = r"e:\DeskTop\knitting_ai_competition.pptx"
prs = Presentation(path)
print("SLIDE SIZE:", prs.slide_width, prs.slide_height,
      "=> in:", round(Emu(prs.slide_width).inches, 2), "x", round(Emu(prs.slide_height).inches, 2))
print("NUM SLIDES:", len(prs.slides))
print("=" * 80)
for i, slide in enumerate(prs.slides):
    print(f"\n##### SLIDE {i+1}  layout={slide.slide_layout.name!r}")
    for sh in slide.shapes:
        kind = sh.shape_type
        try:
            pos = f"L{Emu(sh.left).inches:.2f} T{Emu(sh.top).inches:.2f} W{Emu(sh.width).inches:.2f} H{Emu(sh.height).inches:.2f}"
        except Exception:
            pos = "pos?"
        txt = ""
        if sh.has_text_frame:
            lines = []
            for p in sh.text_frame.paragraphs:
                runs = "".join(r.text for r in p.runs)
                lines.append(runs)
            txt = " | ".join(l for l in lines if l)
        print(f"  - id={sh.shape_id} name={sh.name!r} type={kind} {pos}")
        if txt:
            print(f"      TEXT: {txt[:200]}")
