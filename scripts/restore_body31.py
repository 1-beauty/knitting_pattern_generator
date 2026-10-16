# -*- coding: utf-8 -*-
"""恢复被 captions_en 误覆盖的两个正文句(main para[31], para[175])。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
B31 = ('图 1 概括了从设计意图、PatternPlan 层次化制造 IR、事件级约束编译到可编辑动作图与工厂实物验证的整体框架。'
       '视觉预览作为动作图的辅助检视，用户可通过更新 IR 中的目标字段触发局部重新编译。其中，动作生成与辅助预览的逻辑链路为：')
B175 = ('图 10 将上述证据整合为单张对照表，各行依次给出 PatternPlan 摘要、40×40 动作色码图、结构感知视觉预览图与工厂实织样片。'
        '从结果可以看出，PatternPlan 将自然语言中的几何主体、内部纹理和区域层级转化为结构化计划，约束编译器进一步将其编译为满足硬约束的动作色码图：'
        '五个工厂样片均成功织出——Factory Sample 3 样片呈现与设计周期一致的连续罗纹纵条，Factory Sample 1、Factory Sample 2 样片分别在圆形与菱形区域内形成与设计点阵对应的吊目纹理且区域轮廓清晰可辨，'
        'Factory Sample 4 样片在圆形区域内呈现蜂窝状点阵纹理，Factory Sample 5 样片呈现 X 形斜向纹理走向，各样片边框与主体区域的边界位置均与设计网格对应。'
        '上述结果表明，本文生成的动作色码图不仅满足离散动作级约束，而且经工厂打版流程转换后可在真实横机上织出区域结构与纹理同设计一致的织物：'
        '本文方法能够在统一框架下完成从设计意图到动作表示的完整生成过程，其语义到动作编译思路已通过真实制造链路的落地测试，思路的准确性得到实物证实。')
doc = Document(MAIN)
ps = doc.paragraphs
assert ps[31].text.strip().startswith('Fig. 1.'), ps[31].text[:30]
assert ps[175].text.strip().startswith('Fig. 10.'), ps[175].text[:30]
for p, txt in ((ps[31], B31), (ps[175], B175)):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(txt)
doc.save(MAIN)
d = Document(MAIN)
n = sum(1 for p in d.paragraphs if p.text.strip().startswith('Fig.'))
print('restored; Fig. captions now:', n)
