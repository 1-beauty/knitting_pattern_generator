# -*- coding: utf-8 -*-
"""Round 8n: Three additions.
1. Democratized design-for-manufacturing narrative (p006 + p134).
2. P1-7 LLM reproducibility details (p088).
3. Verify no oMath conflicts.
"""
import sys
from docx import Document
from docx.oxml.ns import qn

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)

def assert_safe(p):
    assert not p._p.findall('.//' + qn('m:oMath')), f"Paragraph has oMath! Cannot append."

def append_text(p, text):
    """Append text to end of paragraph as new run."""
    assert_safe(p)
    p.add_run(text)

def wt_sub(p, old, new):
    total = 0
    for t in p._p.findall('.//' + qn('w:t')):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            total += 1
    assert total >= 1, f"Not found: '{old[:60]}...'"
    return total

# ============================================================
# 1A. Democratized design-for-manufacturing narrative in 1.1
# Insert after "这种依赖人工经验的模式已成为生产瓶颈。"
# ============================================================
p = doc.paragraphs[6]
# The text up to that point describes the bottleneck (manual CAD).
# After "已成为生产瓶颈" I want to add one sentence about DfM democratization as a motivation.
# Use wt_sub to insert: find the sentence ending and append the narrative clause.
target = "这种依赖人工经验的模式已成为生产瓶颈。"
replacement = ("这种依赖人工经验的模式已成为生产瓶颈——"
    "设计意图与可制造动作之间的专业编译壁垒，使非针织背景的工业设计师无法直接产出可投产花型。"
    "在智能制造与大规模个性化定制背景下，"
    "消除这一壁垒、让非专家用户的设计直接投产（democratized design-for-manufacturing），"
    "是针织工业信息学亟需解决的核心问题。")
wt_sub(p, target, replacement)
print("1A p006 narrative inserted")

# ============================================================
# 1B. Democratized DfM impact in conclusion (p134, 总体而言)
# Append one sentence at end about industrial positioning.
# ============================================================
p = doc.paragraphs[134]
# Current end: "...未来工作将围绕...工厂落地测试展开。"
# Append BEFORE 未来工作: "在工业信息学层面，本文框架展示了从自然语言设计语义直接到达可制造动作的端到端编译通路，为横机针织领域实现民主化面向制造的设计提供了可复用的方法范式。"
# Find the end and insert before last sentence
old_end = "未来工作将围绕针织组织规则扩展"
new_end = ("在工业信息学层面，本文框架展示了从自然语言设计语义到可制造动作的端到端编译通路，"
    "为横机针织实现民主化面向制造设计（democratized design-for-manufacturing）提供了可复用的方法范式。"
    "未来工作将围绕针织组织规则扩展")
wt_sub(p, old_end, new_end)
print("1B conclusion DfM impact OK")

# ============================================================
# 2. P1-7 Reproducibility: LLM model details in 4.1 (p088)
# Append at end: model, API version, temperature, determinism note.
# ============================================================
p = doc.paragraphs[88]
repro = ("语义解析与自修复模块调用 DeepSeek-V3（API 版本 2025-06-05，"
    "temperature=0.1，top_p=0.9，system prompt 含完整 Schema 与 12 个 few-shot 样例，"
    "详见补充材料 Section M）；编译与规则检查完全确定性，"
    "同一 PatternPlan 重复编译结果逐格一致。")
append_text(p, repro)
print("2 P1-7 reproducibility appended to p088")

# ============================================================
# Save
# ============================================================
doc.save(SRC)
print("\nOK round8n")
