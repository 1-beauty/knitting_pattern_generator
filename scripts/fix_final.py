# -*- coding: utf-8 -*-
"""正确性回滚 + 收尾：
① 图1图注 para[39] 工厂实物样片验证处 (Sec. 4.6) -> (Sec. 4.5)（工厂落地测试实为4.5节，4.6是讨论）；
② 参考文献节 H2 标题 参考文献 -> References（该节49条已全英文）。
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(MAIN)
ps = doc.paragraphs

# ① 图1图注 Sec. 4.6 -> Sec. 4.5
p39 = ps[39]
assert p39.text.strip().startswith('Fig. 1.'), p39.text[:20]
n = 0
for r in p39.runs:
    if '(Sec. 4.6)' in r.text:
        r.text = r.text.replace('(Sec. 4.6)', '(Sec. 4.5)'); n += 1
print('fig1 caption Sec.4.6->4.5 runs:', n)
assert n == 1, n

# ② 参考文献标题 -> References
done = False
for p in ps:
    if p.style.name == 'Heading 2' and p.text.strip() == '参考文献':
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        p.add_run('References')
        done = True
        break
print('References heading renamed:', done)
assert done

doc.save(MAIN)

# 验证
d = Document(MAIN)
dps = d.paragraphs
t39 = dps[39].text
print('\n=== 验证 ===')
print('para[39] 含 (Sec. 4.5):', '(Sec. 4.5)' in t39, '| 含 (Sec. 4.6):', '(Sec. 4.6)' in t39)
i = t39.find('(Sec. 4.5)')
print('  上下文:', t39[max(0, i - 60):i + 12])
for p in dps:
    if p.style.name == 'Heading 2' and p.text.strip() in ('References', '参考文献'):
        print('参考文献节标题:', p.text.strip())
