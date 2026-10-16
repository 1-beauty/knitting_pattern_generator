"""PPT 视觉语法工具箱。

与现有幻灯片保持同一套规范（数值均取自原稿实测，勿擅改）：
  画布 13.333 × 7.5 in，底色 #F4F6FA，顶部 0.10 in 蓝条，
  标题 22pt 粗体 #1E3A8A + 1.5pt 蓝色下划线，
  底部 0.20 in 深蓝页脚条 + 8pt #BFDBFE 页脚文字与页码。
形状一律不带 p:style（原稿如此，否则会多出主题阴影/描边）。
"""

from __future__ import annotations

import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------- 画布 ----------
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_L = 0.45
CONTENT_W = 12.45          # 0.45 → 12.90
BODY_TOP = 1.10            # 标题下划线之下的可用起点
BODY_BOTTOM = 7.20         # 页脚条之上的可用终点

# ---------- 配色 ----------
BG = "F4F6FA"
NAVY = "1E3A8A"
BLUE = "2563EB"
PURPLE = "7C3AED"
GREEN = "059669"
ORANGE = "D97706"
RED = "DC2626"
TEAL = "0891B2"

BLUE_L = "EFF6FF"
PURPLE_L = "F5F3FF"
GREEN_L = "F0FDF4"
ORANGE_L = "FFFBEB"
RED_L = "FEF2F2"
TEAL_L = "ECFEFF"

INK = "1A1A2E"
BODY = "2D3561"
MUTED = "64748B"
LINE = "CBD5E1"
WHITE = "FFFFFF"
PANEL = "F8FAFC"
FOOTTXT = "BFDBFE"

FONT = "Microsoft YaHei"
MONO = "Consolas"
FOOTER_TEXT = "第八届中国研究生人工智能创新大赛"

_PAGENO = re.compile(r"\d+\s*/\s*\d+")

_ANCHOR = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE,
           "bot": MSO_ANCHOR.BOTTOM}
# python-pptx 的自动图形模板会给首段写死 algn="ctr"，而原稿形状根本没有 algn。
# 因此 align=None 必须显式落成左对齐，否则每个文本框的第一段都会莫名居中。
_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT,
          None: PP_ALIGN.LEFT}
_SHAPE = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE,
          "arrow": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW,
          "chev": MSO_SHAPE.CHEVRON, "oval": MSO_SHAPE.OVAL,
          "pent": MSO_SHAPE.PENTAGON, "diamond": MSO_SHAPE.DIAMOND}


def C(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


def _strip_style(sh):
    """去掉 python-pptx 默认写入的主题样式，避免出现阴影和主题描边。"""
    el = sh._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)


# ============================================================
# 基础图元
# ============================================================

def box(slide, l, t, w, h, *, fill=None, line=None, lw=1.0, shape="rect",
        adj=None):
    """实心/描边矩形。fill=None 表示无填充，line=None 表示无描边。"""
    sh = slide.shapes.add_shape(_SHAPE[shape], Inches(l), Inches(t),
                                Inches(w), Inches(h))
    _strip_style(sh)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = C(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = C(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if adj is not None and len(sh.adjustments):
        sh.adjustments[0] = adj
    sh.text_frame.word_wrap = True
    return sh


def txt(slide, l, t, w, h, lines, *, size=9, bold=False, color=BODY,
        font=FONT, align=None, anchor="mid", ls=1.4, wrap=True, space=0):
    """文本框。lines 可为字符串或字符串列表（每项一段）。

    列表项也可写成 (文本, 覆盖参数字典)，用于同一框内混排不同字号/颜色。
    """
    sh = box(slide, l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = _ANCHOR[anchor]
    items = lines if isinstance(lines, (list, tuple)) else [lines]
    for i, item in enumerate(items):
        over = {}
        if isinstance(item, (list, tuple)):
            item, over = item[0], item[1]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = over.get("ls", ls)
        al = _ALIGN[over.get("align", align)]
        if al is not None:
            p.alignment = al
        sp = over.get("space", space)
        if sp:
            p.space_after = Pt(sp)
        r = p.add_run()
        r.text = str(item)
        f = r.font
        f.size = Pt(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.name = over.get("font", font)
        f.color.rgb = C(over.get("color", color))
    return sh


def hline(slide, l, t, w, *, color=LINE, lw=1.0):
    sh = box(slide, l, t, w, 0, line=color, lw=lw)
    return sh


def vline(slide, l, t, h, *, color=LINE, lw=1.0):
    return box(slide, l, t, 0, h, line=color, lw=lw)


def pic(slide, path, l, t, w=None, h=None):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(l), Inches(t), **kw)


def pic_fit(slide, path, l, t, w, h, *, frame=True, frame_fill=WHITE,
            frame_line=LINE):
    """把图片按比例塞进 (l,t,w,h) 框内并居中，可选画一个底框。"""
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    if frame:
        box(slide, l, t, w, h, fill=frame_fill, line=frame_line)
    return pic(slide, path, l + (w - dw) / 2, t + (h - dh) / 2, w=dw, h=dh)


# ============================================================
# 组合件
# ============================================================

def chip(slide, l, t, w, h, label, *, fill=BLUE, color=WHITE, size=9,
         bold=True, shape="round", adj=0.20, align="c", font=FONT):
    box(slide, l, t, w, h, fill=fill, line=fill, shape=shape, adj=adj)
    return txt(slide, l, t, w, h, label, size=size, bold=bold, color=color,
               align=align, anchor="mid", ls=1.0, font=font)


def tag(slide, l, t, w, h, label, *, accent=BLUE, light=BLUE_L, size=8.5):
    """浅底 + 同色描边 + 同色文字的小标签。"""
    box(slide, l, t, w, h, fill=light, line=accent, shape="round", adj=0.22)
    return txt(slide, l, t, w, h, label, size=size, bold=True, color=accent,
               align="c", ls=1.0)


def card(slide, l, t, w, h, header, body, *, accent=BLUE, light=BLUE_L,
         head_h=0.42, head_size=10.5, body_size=8.8, body_color=BODY,
         body_ls=1.35, body_align=None):
    """标题条 + 浅底正文的卡片，四层架构页同款。"""
    box(slide, l, t, w, head_h, fill=accent, line=accent)
    txt(slide, l + 0.08, t + 0.04, w - 0.16, head_h - 0.08, header,
        size=head_size, bold=True, color=WHITE, align="c", ls=1.0)
    box(slide, l, t + head_h, w, h - head_h, fill=light, line=accent)
    if body:
        txt(slide, l + 0.12, t + head_h + 0.08, w - 0.24, h - head_h - 0.16,
            body, size=body_size, color=body_color, anchor="top", ls=body_ls,
            align=body_align)


def kpi(slide, l, t, w, h, label, value, *, value_color=GREEN, label_size=8.5,
        value_size=10):
    """白底指标卡：上灰标签、下彩色数值。"""
    box(slide, l, t, w, h, fill=WHITE, line=LINE)
    txt(slide, l + 0.10, t + 0.03, w - 0.20, h * 0.40, label,
        size=label_size, color=MUTED, ls=1.0)
    txt(slide, l + 0.10, t + h * 0.42, w - 0.20, h * 0.52, value,
        size=value_size, bold=True, color=value_color, ls=1.0)


def band(slide, l, t, w, h, content, *, fill=NAVY, color=WHITE, size=10.5,
         align="c", bold=False):
    """整行深色强调条。"""
    box(slide, l, t, w, h, fill=fill, line=fill)
    return txt(slide, l + 0.10, t + 0.04, w - 0.20, h - 0.08, content,
               size=size, color=color, align=align, bold=bold, ls=1.25)


def arrow(slide, l, t, w, h, *, color=BLUE, shape="arrow"):
    return box(slide, l, t, w, h, fill=color, line=color, shape=shape)


def connector(slide, l, t, w, *, color=BLUE, lw=2.0):
    """卡片之间的短横线连接（四层架构页同款）。"""
    return hline(slide, l, t, w, color=color, lw=lw)


def bullets(slide, l, t, w, h, items, *, size=9, color=BODY, dot=BLUE,
            dot_size=0.07, row_h=0.30, dot_dx=0.02, text_dx=0.18):
    """圆点列表：每项一行，圆点与文字分离绘制，行距可控。"""
    for i, item in enumerate(items):
        y = t + i * row_h
        box(slide, l + dot_dx, y + row_h / 2 - dot_size / 2, dot_size,
            dot_size, fill=dot, line=dot, shape="oval")
        txt(slide, l + text_dx, y, w - text_dx, row_h, item, size=size,
            color=color, ls=1.15)


def table(slide, l, t, col_w, rows, *, row_h=0.30, head_h=None,
          head_fill=NAVY, head_color=WHITE, body_fill=WHITE,
          alt_fill=PANEL, grid=LINE, size=8.5, head_size=None,
          aligns=None, cell_color=None, cell_bold=None, col_colors=None):
    """轻量表格：用矩形 + 文本框拼，渲染比 pptx 原生表格更可控。

    rows[0] 为表头。aligns 为每列对齐（'l'/'c'/'r'）。
    cell_color / cell_bold 为 {(行, 列): 值} 的按格覆盖（行 0 为表头）。
    col_colors 为 {列: 颜色} 的整列文字色覆盖（不含表头）。
    """
    head_h = head_h or row_h
    head_size = head_size or size
    aligns = aligns or ["l"] * len(col_w)
    cell_color = cell_color or {}
    cell_bold = cell_bold or {}
    col_colors = col_colors or {}
    total_w = sum(col_w)

    y = t
    for ri, row in enumerate(rows):
        h = head_h if ri == 0 else row_h
        if ri == 0:
            box(slide, l, y, total_w, h, fill=head_fill, line=head_fill)
        else:
            fill = body_fill if ri % 2 == 1 else alt_fill
            box(slide, l, y, total_w, h, fill=fill, line=grid)
        x = l
        for ci, cell in enumerate(row):
            if ri == 0:
                col, bold, sz = head_color, True, head_size
            else:
                col = cell_color.get((ri, ci),
                                     col_colors.get(ci, BODY))
                bold = cell_bold.get((ri, ci), False)
                sz = size
            txt(slide, x + 0.07, y, col_w[ci] - 0.14, h, str(cell),
                size=sz, bold=bold, color=col, align=aligns[ci], ls=1.0)
            x += col_w[ci]
        y += h
    return y          # 返回表格底边 y，便于继续排版


# ============================================================
# 页面骨架
# ============================================================

def add_slide(prs):
    """新建一页并铺好底色、顶条、页脚条与页脚文字（页码稍后统一写）。"""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    box(slide, 0, 0, SLIDE_W, SLIDE_H, fill=BG, line=BG)
    box(slide, 0, 0, SLIDE_W, 0.10, fill=BLUE, line=BLUE)
    box(slide, 0, 7.30, SLIDE_W, 0.20, fill=NAVY, line=NAVY)
    txt(slide, 0.40, 7.31, 6.00, 0.18, FOOTER_TEXT, size=8, color=FOOTTXT,
        ls=1.0)
    return slide


def head(slide, title, sub=None):
    """标题 + 下划线；sub 为右端小字补充说明。"""
    txt(slide, MARGIN_L, 0.25, 12.30, 0.55, title, size=22, bold=True,
        color=NAVY, ls=1.0)
    hline(slide, MARGIN_L, 0.88, 12.30, color=BLUE, lw=1.5)
    if sub:
        txt(slide, 6.90, 0.34, 5.85, 0.38, sub, size=10, color=MUTED,
            align="r", ls=1.0)


def page(prs, title, sub=None):
    slide = add_slide(prs)
    head(slide, title, sub)
    return slide


# ============================================================
# 幻灯片顺序与页码
# ============================================================

def move_slide(prs, src: int, dst: int):
    """把第 src 页（0 基）移动到第 dst 页位置。"""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = ids[src]
    lst.remove(el)
    ids = list(lst)
    lst.insert(dst, el)


def reorder(prs, order: list[int]):
    """按 order（原下标序列）整体重排幻灯片。"""
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        lst.remove(el)
    for i in order:
        lst.append(ids[i])


def _slide_rids(prs):
    """建立幻灯片部件 → rId 的映射。"""
    out = {}
    for rid, rel in prs.part.rels.items():
        if rel.reltype.endswith("/slide"):
            out[rel.target_part] = rid
    return out


def set_order(prs, slides):
    """按给定的幻灯片对象序列重排；未列出的幻灯片连带关系一起删除。

    用对象而不用下标，避开“删一页后后面所有下标都移位”的错位风险。
    """
    lst = prs.slides._sldIdLst
    rid_of = _slide_rids(prs)
    els = {el.get(qn("r:id")): el for el in list(lst)}
    keep = [els[rid_of[s.part]] for s in slides]
    kept = set(id(el) for el in keep)
    for el in list(lst):
        lst.remove(el)
    for el in keep:
        lst.append(el)
    for rid, el in els.items():
        if id(el) not in kept:
            prs.part.drop_rel(rid)


def renumber(prs, *, skip_first=True):
    """把每页右下角页码统一改写为 'n / 总数'；缺页码的页补一个。

    只认页脚条上的形状（top ≥ 7.2 in），否则正文里的
    “20 / 20”之类指标会被误认为页码而遭覆盖。
    """
    total = len(prs.slides._sldIdLst)
    foot_top = Inches(7.2)
    for i, slide in enumerate(prs.slides, start=1):
        if skip_first and i == 1:
            continue
        target = None
        for sh in slide.shapes:
            if not sh.has_text_frame or sh.top is None or sh.top < foot_top:
                continue
            t = sh.text_frame.text.strip()
            if _PAGENO.fullmatch(t):
                target = sh
                break
        label = f"{i} / {total}"
        if target is None:
            txt(slide, 12.00, 7.31, 0.70, 0.18, label, size=8,
                color=FOOTTXT, font="Arial", align="r", ls=1.0)
            continue
        for p in target.text_frame.paragraphs:
            for r in p.runs:
                r.text = ""
        p0 = target.text_frame.paragraphs[0]
        if p0.runs:
            p0.runs[0].text = label
        else:
            r = p0.add_run()
            r.text = label
            r.font.size = Pt(8)
            r.font.name = "Arial"
            r.font.color.rgb = C(FOOTTXT)
