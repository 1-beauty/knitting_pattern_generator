# -*- coding: utf-8 -*-
"""移除 11 个公式表(1x3 含 oMath)的 tblStyle 引用与条件格式，消除浅灰边框线。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)
n = 0
for t in doc.tables:
    if len(t.rows) == 1 and len(t.columns) == 3 and t.cell(0, 1)._tc.findall('.//' + qn('m:oMath')):
        tblPr = t._tbl.tblPr
        for s in tblPr.findall(qn('w:tblStyle')):
            tblPr.remove(s)
        lk = tblPr.find(qn('w:tblLook'))
        if lk is None:
            lk = OxmlElement('w:tblLook'); tblPr.append(lk)
        for a, v in (('w:firstRow', '0'), ('w:firstColumn', '0'), ('w:lastRow', '0'),
                     ('w:lastColumn', '0'), ('w:noHBand', '1'), ('w:noVBand', '1'), ('w:val', '0000')):
            lk.set(qn(a), v)
        n += 1
doc.save(SRC)
print('eq tables de-styled:', n)
