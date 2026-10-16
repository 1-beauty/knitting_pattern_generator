# -*- coding: utf-8 -*-
"""Translate Chinese tables to English + fix Table 5 column widths."""
import sys, re
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Cm, Pt

DOCX = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(DOCX)

# ============================================================
# 1. Algorithm box (tbl[6]) -> English
# ============================================================
algo_translations = {
    0: (2, 'Natural-language description D'),
    1: (1, 'Internal action color-map P, statistics report Q'),
    2: (1, '\tParse D into the hierarchical structural plan PatternPlan'),
    3: (1, '\tRender PatternPlan into VisualSpec'),
    4: (1, '\tMap VisualSpec to RealizationPlan'),
    5: (1, '\tGenerate sparse EventPoints from event_pattern'),
    6: (1, '\tInitialize the 40\u00d740 ActionGrid and fill regional base fabric'),
    7: (1, '\tDraw boundaries according to border_style'),
    8: (1, '\tWrite sparse events and run local transition solving'),
    9: (1, '\tRun refinement, aesthetic pass, and hierarchical clipping'),
    10: (1, '\tRun compiler guard, canvas-boundary safeguarding, and group-integrity check'),
    11: (1, '\tCheck hard/craft constraints and switching-intensity score; run conservative self-repair'),
    12: (1, '\tExport the internal action color-map, visual-preview input, and statistics report'),
}
cell = doc.tables[6].cell(0, 0)
for para_idx, (run_idx, new_text) in algo_translations.items():
    p = cell.paragraphs[para_idx]
    if run_idx < len(p.runs):
        p.runs[run_idx].text = new_text
print("Algorithm box (tbl[6]) translated to English")

# ============================================================
# 2. Table 2 (tbl[7], H1-H8 constraints) -> English
# ============================================================
table2_data = [
    ['ID', 'Constraint', 'Check target', 'Violation semantics'],
    ['H1', 'Illegal action index', 'Action index', 'Index outside the 15-class action-label space'],
    ['H2', 'Lower support', 'Knit / tuck / transfer / split', 'Cell below a support-requiring action is empty'],
    ['H3', 'Transfer offset limit', 'Transfer', 'Lateral offset exceeds the allowed range'],
    ['H4', 'Transfer target out of bounds', 'Transfer', 'Target column lies outside the canvas'],
    ['H5', 'Same-row target conflict', 'Transfer', 'Multiple sources land on the same target'],
    ['H6', 'Consecutive tuck stacking', 'Tuck', 'Consecutive tucks in a column exceed the limit'],
    ['H7', 'Auxiliary action cleared above', 'Auxiliary control action', 'Incompatible support structure remains above the action'],
    ['H8', 'Same-row bed-direction conflict', 'Inter-bed transfer', 'Opposite-direction inter-bed transfers coexist in one row'],
]
tbl2 = doc.tables[7]
for r in range(len(tbl2.rows)):
    for c in range(len(tbl2.columns)):
        cell = tbl2.cell(r, c)
        new_text = table2_data[r][c]
        # Preserve first run formatting, clear others
        if cell.paragraphs and cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].text = new_text
            for extra_run in cell.paragraphs[0].runs[1:]:
                extra_run.text = ''
        else:
            cell.text = new_text
print("Table 2 (tbl[7]) translated to English")

# ============================================================
# 3. Table 5 (tbl[13]) column widths - widen "Main impact"
# ============================================================
tbl5 = doc.tables[13]
widths = [1.5, 2.5, 1.4, 1.4, 1.6, 1.2, 1.2, 1.1, 4.0]  # total 15.9
for i, w in enumerate(widths):
    tbl5.columns[i].width = Cm(w)
    for row in tbl5.rows:
        row.cells[i].width = Cm(w)
print(f"Table 5 widths fixed: {widths} = {sum(widths):.1f}cm")

# ============================================================
# Save
# ============================================================
doc.save(DOCX)
print(f"\nSaved: {DOCX}")
