# -*- coding: utf-8 -*-
"""第七轮补丁 d（正文）：4.5/4.6 残留强口径校准——
p167 "系统已落地"、p172 "准确织出"、p180 "证实了…准确性" 三处与已批准措辞校准对齐。
节名"工厂落地测试"（p098/p166/p189）为测试名称，保留。全部断言通过才落盘。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
m = Document(MAIN)
ps = m.paragraphs

def find(pre): return next(p for p in ps if p.text.startswith(pre))
def span(p, old, new, cnt=1):
    n = 0
    for t in p._p.findall('.//' + qn('w:t')):
        if t.text and old in t.text:
            n += t.text.count(old); t.text = t.text.replace(old, new)
    assert n == cnt, ('span', old, n)

span(find('本节是全文证据链的物理落点'),
     '本文系统已落地、成品已织出、实物已在图 10 中展示。',
     '本文系统已完成受控工厂验证、样片已织出并寄回、实物在图 10 中展示。')
span(find('上述实验表明，本文方法的主要优势在于将自然语言设计描述转化为可编辑、可检查'),
     '受约束的动作输出经工厂打版流程转换后可在真实横机上准确织出。',
     '受约束的动作输出经工厂打版流程转换后可在真实横机上以与设计一致的区域结构与纹理织出。')
span(find('第六，需要明确本文的定位'),
     '4.5 节已将技术思路送合作工厂完成落地测试，5 片实物样片的区域几何与纹理结构均与设计一致，证实了本文思路在真实制造条件下的准确性。',
     '4.5 节已与宁波慈星股份有限公司合作完成受控工厂验证，5 片实物样片的区域几何与纹理结构均与设计一致，为本文思路在真实制造条件下提供了可行性证据（proof-of-concept）。')

txt = '\n'.join(p.text for p in m.paragraphs)
assert '本文系统已落地' not in txt and '证实了本文思路' not in txt and '准确织出' not in txt
assert '可行性证据（proof-of-concept）' in txt
m.save(MAIN)
print('OK round7d main | overclaim residuals calibrated: 3')
