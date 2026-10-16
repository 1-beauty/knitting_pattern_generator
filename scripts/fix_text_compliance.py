# -*- coding: utf-8 -*-
"""文字层合规修复：
1) 摘要 320 词 -> <=250 词（CiI 官方上限），保留全部主张/数字/工厂验证口径；
2) 11 条显示公式补正文引用（式（1）–式（11）各一次）；
3) para98 碎片标签句 "截图说明与英文对照键。" 改写为连贯 prose。"""
import re
import docx
from docx.oxml.ns import qn

SRC = r'docs\论文to_cii.docx'
doc = docx.Document(SRC)
ps = doc.paragraphs

NEW_ABSTRACT = (
    'Flat-knitting action graphs link design semantics to stitch-level action structures, whose '
    'generation must satisfy discrete manufacturing constraints such as loop support, transfer '
    'targets and action combinations. We formulate this task as semantics-to-action compilation '
    'and propose a hierarchical semantic-constraint compilation and editable generation framework. '
    'A knitting-specific manufacturing intermediate representation, PatternPlan, captures geometry, '
    'texture and parent-child hierarchy; it is compiled into a two-dimensional action graph via '
    'sparse event planning, with hard-constraint checking and conservative self-repair suppressing '
    'structural conflicts, region-level structural editing, and a conditional flow-matching visual '
    'preview for verification. On a 15-class action space cleaned from the Neural Inverse Knitting '
    'dataset, PatternPlan achieves 97.2% field-level parsing accuracy over 100 inputs within '
    'rule-library coverage. Baseline and controlled experiments show that direct action-matrix '
    'generation and visual-image quantization fail on action validity or design semantics, and that '
    'replacing event-level compilation with dense filling degrades region-structure metrics; '
    'comparison with a CP-SAT global constraint-projection baseline further shows that action '
    'validity can be closed a posteriori by constraint solving, whereas region attribution and '
    'grouping must be injected explicitly by compilation. The action color-code map is a '
    'machine-specific, color-agnostic abstract action representation; five representative designs '
    'were validated in a real factory production test: their color-code maps were converted through '
    'the factory patterning workflow and knitted into physical samples consistent with the designs '
    '(Fig. 10). These results indicate that flat-knitting action-graph generation is better modeled '
    'as manufacturing-constraint-oriented semantics-to-action compilation than as plain '
    'text-to-image generation.'
)
assert len(NEW_ABSTRACT.split()) <= 250, len(NEW_ABSTRACT.split())

def set_para_text(p, new_text):
    """整段换文：保留首 run 的 rPr 格式，其余 run 删除。"""
    runs = p.runs
    assert runs, 'empty paragraph'
    runs[0].text = new_text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)

def replace_in_para(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return True
    # 跨 run 且无内联公式：整段重建安全
    if old in p.text and not p._p.findall('.//' + qn('m:oMath')):
        set_para_text(p, p.text.replace(old, new, 1))
        return True
    return False

# --- 1) 摘要 ---
ab_idx = next(i for i, p in enumerate(ps) if p.text.startswith('Flat-knitting action graphs'))
old_words = len(ps[ab_idx].text.split())
set_para_text(ps[ab_idx], NEW_ABSTRACT)
print('abstract:', old_words, '->', len(NEW_ABSTRACT.split()), 'words')

# --- 2) 公式引用：(para_idx, old, new) ---
EQ_CITES = [
    (32,  '其中，', '式（1）中，'),
    (44,  '其中，', '式（2）中，'),
    (56,  '其中，', '式（3）中，'),
    (63,  '基于上述两类关系，', '基于上述两类关系及式（5）定义的有效作用域，'),
    (64,  '其中', '式（6）中，'),
    (80,  '其中，', '式（7）中，'),
    (83,  '若 split 或 tuck 任一侧', '按式（8）的检查规则，若 split 或 tuck 任一侧'),
    (88,  '推理通过常微分方程', '推理通过对式（9）所学速度场进行常微分方程'),
    (111, 'Type-Coverage 仅衡量', '由式（10）可见，Type-Coverage 仅衡量'),
    (113, '此外，父子包含率为', '除式（11）的两项指标外，父子包含率为'),
]
for idx, old, new in EQ_CITES:
    assert replace_in_para(ps[idx], old, new), (idx, old)
    print('eq cite ok: para', idx)
# 式（4）：para61 末尾 "标注动作：" -> "标注动作（式（4））："
p61 = ps[61]
assert p61.text.rstrip().endswith('：')
done = False
for r in reversed(p61.runs):
    if r.text.rstrip().endswith('：'):
        r.text = r.text.rstrip()[:-1] + '（式（4））：'
        done = True
        break
assert done
print('eq cite ok: para 61 (tail)')

# --- 3) para98 碎片句（容游离空格：整段重建，保留首 run 格式） ---
p98 = ps[98]
pat98 = re.compile(r'截图说明与英文对照键。\s*GUI\s*截图均为未经修改的原始采集，保留原始运行\s*ID、时间\s*戳与文件名。')
NEW98 = ('图 7 的 GUI 截图均为未经修改的原始采集，保留原始运行 ID、时间戳与文件名，'
         '界面字段与英文对照键的对应关系如下。')
assert pat98.search(p98.text), p98.text[:60]
set_para_text(p98, pat98.sub(NEW98, p98.text, count=1))
print('para98 rewritten')

doc.save(SRC)

# --- 校验 ---
chk = docx.Document(SRC)
ts = [p.text for p in chk.paragraphs]
ref = next(i for i, t in enumerate(ts) if t.strip() == 'References')
body = ts[:ref]
print('abstract words now:', len(body[ab_idx].split()))
cites = sorted(int(m.group(1)) for t in body for m in re.finditer(r'式（(\d+)）', t))
print('eq citations present:', cites)
print('para98 head:', ts[98][:52])
