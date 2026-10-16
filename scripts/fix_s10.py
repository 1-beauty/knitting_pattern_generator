# -*- coding: utf-8 -*-
"""S10: remove w/o PatternPlan bars from both charts, re-space remaining 4 groups."""
from pptx import Presentation
from pptx.util import Inches as IN

P = r"e:\DeskTop\knitting_ai_competition.pptx"
prs = Presentation(P)
sl = prs.slides[9]  # S10
m = {sh.shape_id: sh for sh in sl.shapes}

# 1) delete the w/o PatternPlan group (left: value+label; right: bar+value+label)
for sid in (10, 11, 25, 26, 27):
    el = m[sid]._element
    el.getparent().remove(el)

# 2) re-space remaining 4 groups; bars sit at value_x + 0.05
moves = {
    # left chart (panel 0.45..6.25): 1.10 / 2.35 / 3.60 / 4.85
    8: 1.10, 9: 1.10,            # Ours value+label
    12: 2.40, 13: 2.35, 14: 2.35,  # LLM bar+value+label
    15: 3.65, 16: 3.60, 17: 3.60,  # T2I bar+value+label
    18: 4.85, 19: 4.85,            # Random value+label
    # right chart (panel 6.5..12.3): 7.15 / 8.40 / 9.65 / 10.90
    22: 7.20, 23: 7.15, 24: 7.15,  # Ours bar+value+label
    28: 8.45, 29: 8.40, 30: 8.40,  # LLM bar+value+label
    31: 9.65, 32: 9.65,            # T2I value+label
    33: 10.95, 34: 10.90, 35: 10.90,  # Random bar+value+label
}
for sid, left in moves.items():
    m[sid].left = IN(left)

prs.save(P)
print("S10 fixed OK")

# 3) scan whole deck for leftover "w/o"/"PatternPlan消融" texts
prs2 = Presentation(P)
for i, s in enumerate(prs2.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and ("w/o" in sh.text_frame.text or "消融" in sh.text_frame.text):
            print(f"LEFTOVER slide {i} #{sh.shape_id}: {sh.text_frame.text[:80]!r}")
print("scan done")
