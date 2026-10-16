# -*- coding: utf-8 -*-
"""回滚损坏的 bare brk + 加宽公式中间列消除丑换行。
① 移除 oMath 直接子元素 m:brk(损坏源)；② 11个公式表列宽 [900,7226,900] -> [200,8226,600]，
   使 (2)(3)(6)(11) 单行容纳；(1) 链式断点在→处可接受。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Twips
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
M = qn('m:oMath'); MB = qn('m:brk')
W = [Twips(200), Twips(8226), Twips(600)]

doc = Document(MAIN)
removed = 0
for t in doc.tables:
    if len(t.rows) == 1 and len(t.columns) == 3:
        om = t.cell(0, 1)._tc.find('.//' + M)
        if om is None:
            continue
        for b in list(om.findall(MB)):  # 直接子元素 brk
            om.remove(b); removed += 1
print('bare brk removed:', removed)

n = 0
for t in doc.tables:
    if len(t.rows) == 1 and len(t.columns) == 3 and t.cell(0, 1)._tc.find('.//' + M) is not None:
        for j, w in enumerate(W):
            t.columns[j].width = w
            for row in t.rows:
                row.cells[j].width = w
        n += 1
print('eq tables widened:', n)
doc.save(MAIN)
print('saved')
