# -*- coding: utf-8 -*-
"""Compact Fig.3: reduce slide height, fill panels, match other figures' density."""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from pptx import Presentation
from pptx.util import Pt, Emu, Inches

PPTX = r'e:\DeskTop\knitting_pattern_generator\docs\figures\Fig3_Event_Level_Constraint_Compilation_EN.pptx'
prs = Presentation(PPTX)
slide = prs.slides[0]

# === Step 1: Reduce slide height 9.0" -> 8.0" ===
prs.slide_height = Inches(8.0)
print("Slide height: 9.0 -> 8.0 in")

shape_map = {sh.shape_id: sh for sh in slide.shapes}

def setpos(sid, top=None, height=None, left=None, width=None):
    sh = shape_map.get(sid)
    if sh is None:
        print(f"  WARN: shape {sid} not found")
        return
    if top is not None: sh.top = Inches(top)
    if height is not None: sh.height = Inches(height)
    if left is not None: sh.left = Inches(left)
    if width is not None: sh.width = Inches(width)

# === Step 2: Outer panel boxes H=8.20 -> 7.20 ===
for sid in [2, 8, 14, 20]:
    setpos(sid, top=0.40, height=7.20)
print("Outer panels: H 8.20 -> 7.20")

# === Step 3: Container boxes H=6.32 -> 5.10 ===
for sid in [7, 13, 19, 25]:
    setpos(sid, top=2.18, height=5.10)
print("Containers: H 6.32 -> 5.10 (T=2.18 to 7.28)")

# === Step 4: Connector lines between panels -> T=4.70 ===
for sid in [26, 27, 28]:
    setpos(sid, top=4.70)

# === Panel (a): enlarge masks, reposition labels ===
# Masks: 347/201 (left), 348/346 (right) - enlarge 1.79 -> 1.88, move up
setpos(347, top=2.60, height=1.88, left=0.30, width=1.88)
setpos(201, top=2.61, height=1.86, left=0.31, width=1.86)
setpos(348, top=2.60, height=1.88, left=2.28, width=1.88)
setpos(346, top=2.61, height=1.86, left=2.29, width=1.86)
# Labels: spread from 4.90
setpos(349, top=4.90)
setpos(350, top=5.60)
setpos(351, top=6.30)
print("Panel (a): masks enlarged+up, labels spread")

# === Panel (b): 4 chips, tops at 2.53/3.67/4.81/5.95 ===
chip_b = [
    (352, 353, 354, 2.53),
    (355, 356, 357, 3.67),
    (358, 359, 360, 4.81),
    (361, 362, 363, 5.95),
]
for rect, title, body_txt, top in chip_b:
    setpos(rect, top=top)
    setpos(title, top=top + 0.08)
    setpos(body_txt, top=top + 0.51)
print("Panel (b): 4 chips redistributed")

# === Panel (c): 3 chips, tops at 2.58/4.20/5.82 ===
chip_c = [
    (364, 365, 366, 2.58),
    (367, 368, 369, 4.20),
    (370, 371, 372, 5.82),
]
for rect, title, body_txt, top in chip_c:
    setpos(rect, top=top)
    setpos(title, top=top + 0.08)
    setpos(body_txt, top=top + 0.51)
print("Panel (c): 3 chips redistributed")

# === Panel (d): 4 layers + 3 arrows + footer ===
setpos(377, top=2.50); setpos(378, top=2.61)   # L0
setpos(379, top=3.16)                            # arrow
setpos(380, top=3.50); setpos(381, top=3.61)   # L1
setpos(382, top=4.16)                            # arrow
setpos(383, top=4.50); setpos(384, top=4.61)   # L2
setpos(385, top=5.16)                            # arrow
setpos(386, top=5.50); setpos(387, top=5.61)   # L3
setpos(388, top=6.55)                            # footer
print("Panel (d): layers redistributed")

# === Save ===
prs.save(PPTX)
print(f"\nSaved: {PPTX}")
