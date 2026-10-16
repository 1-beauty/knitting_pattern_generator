"""改版 汪瑞knitting_ai_competition.pptx。

在原稿视觉规范内做三件事：
  1. 新增 9 页内容页（系统框架、编译四步、动作空间、H1-H8、C1-C6、修复回路、
     分步走查、GUI 总览、增量编辑与部署）并重建目录页；
  2. 就地修正原第 5 页的措辞与零违例口径；
  3. 重排为 21 页并统一页码。

所有数字取自真实运行 outputs/runs/gui/20260731_155816 与对落盘网格的重跑检查，
规则条目逐条对照 generator_one/rules 下的源码。
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

import ppt_kit as K
from ppt_kit import (BLUE, BLUE_L, GREEN, GREEN_L, INK, MUTED, NAVY, ORANGE,
                     ORANGE_L, PURPLE, PURPLE_L, TEAL, TEAL_L, band, card,
                     chip, kpi, page, pic_fit, renumber, set_order, table,
                     txt, vline)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "scripts" / "_ppt_in.pptx"
OUT = ROOT / "scripts" / "_ppt_out.pptx"
FIG = ROOT / "docs" / "figures" / "ppt"
GUI = ROOT / "docs" / "figures" / "gui"


def retext(slide, old, new):
    """把某页中文本恰为 old 的形状改写为 new，保留原有字体设置。"""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.text_frame.text.strip() != old:
            continue
        p = sh.text_frame.paragraphs[0]
        if p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
        return True
    return False


def numbered(slide, x, y, n, title, body, *, w_title=3.40, w_body=6.55,
             accent=BLUE, cw=0.40, size=11, body_size=8.6):
    """编号方块 + 标题 + 说明的一条走查项。"""
    chip(slide, x, y + 0.03, cw, cw, f"{n:02d}", fill=accent, size=10.5,
         shape="rect", font="Arial")
    tx = x + cw + 0.16
    txt(slide, tx, y, w_title, 0.28, title, size=size, bold=True, color=INK,
        ls=1.0)
    txt(slide, tx, y + 0.28, w_body, 0.70, body, size=body_size, color=K.BODY,
        anchor="top", ls=1.3)


# ============================================================
# 目录（重建）
# ============================================================

TOC = [
    ("问题背景", "设计与制造之间的断层"),
    ("现有方法局限", "为什么现有方案不够用？"),
    ("系统总览与框架", "四层架构 · 端到端数据流"),
    ("创新点与确定性编译", "PatternPlan · 四步变换"),
    ("动作空间", "15 类横机动作与工艺分级"),
    ("规则体系", "硬约束 H1–H8 · 工艺启发式 C1–C6"),
    ("检查与修复", "保守降级 · 失败边界"),
    ("流程分步走查", "一句话指令到双图产物"),
    ("实验结果", "语义解析 · 基线对比 · 可编辑生成"),
    ("桌面工作台", "本地界面 + 远程 GPU 推理"),
]


def make_toc(prs):
    slide = page(prs, "目录")
    ys = [1.75, 2.83, 3.91, 4.99, 6.07]
    for i, (title, sub) in enumerate(TOC):
        right = i >= 5
        y = ys[i % 5]
        cx = 7.85 if right else 1.35
        tx = 8.40 if right else 1.90
        sw = 4.60 if right else 5.00
        chip(slide, cx, y, 0.42, 0.42, f"{i + 1:02d}", fill=BLUE, size=13,
             shape="rect", font="Arial")
        txt(slide, tx, y, 3.80, 0.24, title, size=13, bold=True, color=INK,
            ls=1.0)
        txt(slide, tx, y + 0.22, sw, 0.20, sub, size=9.5, color=MUTED, ls=1.0)
        if i % 5 != 4:
            vline(slide, cx + 0.21, y + 0.52, 0.44)
    return slide


# ============================================================
# N1 系统框架：端到端数据流
# ============================================================

FLOW = [
    ("自然语言指令", "中文一句话，可包含形状、纹理、层级与背景描述", TEAL, TEAL_L),
    ("PatternPlan", "LLM 解析为 layout_type / regions / effects / params，"
     "JSON 可读可改", BLUE, BLUE_L),
    ("确定性编译三级", "VisualSpec → RealizationPlan → EventPoint，"
     "几何与效果全部参数化", PURPLE, PURPLE_L),
    ("ActionGrid 40×40", "15 类动作填格，随后规则检查 H1–H8 / C1–C6 与保守修复",
     GREEN, GREEN_L),
    ("双图产物", "色码图 BMP + 结构感知视觉预览图，附计划与统计 JSON",
     ORANGE, ORANGE_L),
]


def make_framework(prs):
    slide = page(prs, "系统框架：端到端数据流",
                 "自然语言 → 中间表示 → 动作网格 → 视觉预览")
    xs = [0.45 + i * 2.51 for i in range(5)]
    for (h, b, ac, li), x in zip(FLOW, xs):
        card(slide, x, 1.35, 2.35, 1.60, h, b, accent=ac, light=li,
             head_h=0.44, head_size=10, body_size=8.4)
    for x in xs[:-1]:
        K.connector(slide, x + 2.35, 2.15, 0.16, color=MUTED, lw=2.0)

    card(slide, 0.45, 3.15, 6.10, 2.00, "回路一 · 规则检查与保守修复（机内自动）",
         ["检查在动作网格上逐格进行，每条违例都记录规则号与坐标。",
          "修复只动违例格及其必要邻格，一律选择更保守的动作；改完之后经过"
          "后处理守卫，再做一次完整复检才允许落盘。",
          "当前版本 max_repair_rounds = 1，即最多一轮 LLM 定点修复。"],
         accent=GREEN, light=GREEN_L, head_h=0.40, head_size=10,
         body_size=8.5)
    card(slide, 6.80, 3.15, 6.10, 2.00, "回路二 · 增量编辑（设计师在环）",
         ["在原指令上改写局部描述，系统比对新旧 PatternPlan 的字段差异。",
          "只有差异涉及的区域会被重新编译，其余区域逐格保持原值，"
          "因此“改一个词”的影响范围是可度量的。",
          "每次生成与编辑都留一条版本记录，可回溯到任意一版。"],
         accent=ORANGE, light=ORANGE_L, head_h=0.40, head_size=10,
         body_size=8.5)

    kx = [0.45, 3.60, 6.75, 9.90]
    kv = [("端到端总耗时", "2.38 s", GREEN), ("其中语义解析（LLM）", "2.17 s", BLUE),
          ("确定性编译四步", "0.17 s", PURPLE),
          ("规则检查 + 修复", "0.015 s", ORANGE)]
    for x, (lb, vl, cl) in zip(kx, kv):
        kpi(slide, x, 5.42, 3.00, 0.65, lb, vl, value_color=cl)

    band(slide, 0.45, 6.25, 12.45, 0.50,
         "不确定性只出现在语义解析与违例修复两处；其余步骤给定输入即给定输出，"
         "同一 PatternPlan 可逐格复现", size=10)
    return slide


# ============================================================
# N2 确定性编译：四步变换
# ============================================================

STEPS = [
    ("① VisualSpec 视觉规格",
     "把 regions 转成归一化几何：每个区域的中心、尺寸与边界，"
     "并确定层级与绘制顺序。", BLUE, BLUE_L, "0.001 s"),
    ("② RealizationPlan 实现方案",
     "为每个区域选定动作组合与工艺参数：床位、移针偏移、纹理周期、"
     "列偏移。", PURPLE, PURPLE_L, "0.005 s"),
    ("③ EventPoint 事件点",
     "把纹理效果离散为稀疏事件，每点携带动作类型与网格坐标；"
     "本例共 40 个事件点。", GREEN, GREEN_L, "0.155 s"),
    ("④ ActionGrid 动作网格",
     "按事件与区域填充 40×40 整型网格，再做精修与美学处理，"
     "交给规则检查。", ORANGE, ORANGE_L, "0.011 s"),
]


def make_compile(prs):
    slide = page(prs, "确定性编译：从计划到动作网格的四步变换",
                 "无随机性 · 逐格可追溯")
    band(slide, 0.45, 1.18, 12.45, 0.42,
         "第 0 步 · 输入 PatternPlan：layout_type / regions / effects / params，"
         "一份可读可改的 JSON", fill=NAVY, size=10)

    xs = [0.45, 3.65, 6.85, 10.05]
    for (h, b, ac, li, _), x in zip(STEPS, xs):
        card(slide, x, 1.80, 3.00, 1.42, h, b, accent=ac, light=li,
             head_h=0.44, head_size=10.5, body_size=8.8)
    for x in xs[:-1]:
        K.connector(slide, x + 3.00, 2.51, 0.17, color=MUTED, lw=2.0)

    for (h, _, ac, _, t), x in zip(STEPS, xs):
        kpi(slide, x, 3.45, 3.00, 0.62, h[2:] + " 耗时", t, value_color=ac)

    band(slide, 0.45, 4.32, 12.45, 0.45,
         "本例：3 个区域 · 40 个事件点 · 1600 格动作 · 四步合计 0.17 s",
         fill=NAVY, size=10.5)

    card(slide, 0.45, 5.05, 4.05, 1.65, "为什么要留这些中间层",
         "每一级都是可读的 JSON 或整型数组：出错时能判断问题出在解析、几何、"
         "事件还是填充环节，改动一处不必重跑整条链路。",
         accent=BLUE, light=BLUE_L, head_h=0.38, head_size=10,
         body_size=8.8)
    card(slide, 4.65, 5.05, 4.05, 1.65, "与端到端生成的差别",
         "端到端模型直接输出像素，错在哪一步无法回答；本系统把不确定性收在"
         "第 1 步，后面四步是确定性变换，可以逐格解释。",
         accent=PURPLE, light=PURPLE_L, head_h=0.38, head_size=10,
         body_size=8.8)
    card(slide, 8.85, 5.05, 4.05, 1.65, "可复现性与可追溯",
         "每一步的输入输出都落盘：pattern_plan.json 记录计划，"
         "pattern_stats.json 记录区域数、事件数、动作分布与违例明细，"
         "出问题可以回到具体环节复算。",
         accent=GREEN, light=GREEN_L, head_h=0.38, head_size=10,
         body_size=8.8)
    return slide


# ============================================================
# N3 动作空间
# ============================================================

def make_action_space(prs):
    slide = page(prs, "动作空间：15 类横机动作与工艺分级",
                 "对照 assets/instructions.json")
    pic_fit(slide, FIG / "fig_action_space.png", 0.45, 1.18, 12.45, 4.42,
            frame=True)
    card(slide, 0.45, 5.72, 4.05, 1.28, "语义层：8 个类别",
         "EMPTY / LOOP / TUCK / TRANSFER / SPLIT / MISS / DROP / CARRIER。"
         "规则判定按语义类别而非颜色进行，新增动作只需登记它的语义。",
         accent=BLUE, light=BLUE_L, head_h=0.36, head_size=10,
         body_size=8.6)
    card(slide, 4.65, 5.72, 4.05, 1.28, "工艺分级：张力 / 弹性 / 稳定",
         "张力 1–7、弹性 1–5、稳定 1–5，用于比较动作之间的相对风险，"
         "是项目内的保守工程分级，不等同于机器实测张力值。",
         accent=PURPLE, light=PURPLE_L, head_h=0.36, head_size=10,
         body_size=8.6)
    card(slide, 8.85, 5.72, 4.05, 1.28, "待校准项",
         "分级来自工艺经验整理，后续需要用工厂样片实测标定；颜色沿用现有"
         "色码图规范，便于与既有产线文件对接。",
         accent=ORANGE, light=ORANGE_L, head_h=0.36, head_size=10,
         body_size=8.6)
    return slide


# ============================================================
# N4 硬约束 H1–H8
# ============================================================

HARD = [
    ("H1", "动作索引不在 0–14 范围内", "该格无法映射到任何机器动作",
     "改为 f_knit（前床成圈）"),
    ("H2", "成圈或劈针的下方是空针 / 落布",
     "新线圈没有旧线圈托住，会直接脱散",
     "下方是空针则本格也降为空针（保留镂空），否则在下方补一格 b_knit"),
    ("H3", "移针的横向偏移超过 ±2 针", "超出机头单次移针能力",
     "改为 f_split（前床劈针）"),
    ("H4", "移针目标列越出针床边界", "目标针位不存在，指令无法执行",
     "改为 f_split"),
    ("H5", "同一行两处移针指向同一根针", "两个线圈叠到一针上，机器报警",
     "冲突的一方改为 f_split"),
    ("H6", "同一列连续集圈超过 3 行", "集圈层层累积，纱线张力过大",
     "在超限处插入 f_knit 打断"),
    ("H7", "落布动作的上方仍未清空", "落布后残留线圈，布片脱离失败",
     "上方改为空针"),
    ("H8", "同一行同时存在前→后与后→前移针", "机头单趟无法完成双向床间移针",
     "该行所有床间移针改为 f_split"),
]


def make_hard_rules(prs):
    slide = page(prs, "规则体系（一）：硬约束 H1–H8",
                 "generator_one/rules/physical_rules.py")
    band(slide, 0.45, 1.16, 12.45, 0.40,
         "硬约束 = 违反即不可织造，必须全部消除；检查在 40×40 动作网格上逐格进行",
         fill=NAVY, size=10)
    rows = [["", "判定条件", "违反后果", "保守修复动作"]] + [list(r) for r in HARD]
    table(slide, 0.45, 1.72, [0.62, 3.55, 4.40, 3.88], rows,
          head_h=0.34, row_h=0.42, size=8.5, head_size=9,
          aligns=["c", "l", "l", "l"],
          col_colors={0: NAVY}, cell_bold={(i, 0): True for i in range(1, 9)})

    card(slide, 0.45, 5.58, 4.05, 1.42, "检查怎么做",
         "遍历整张网格，按语义类别查表判定；每条违例记录规则编号与 (行, 列) "
         "坐标，可以直接定位到色码图上的那一格。",
         accent=BLUE, light=BLUE_L, head_h=0.36, head_size=10, body_size=8.6)
    card(slide, 4.65, 5.58, 4.05, 1.42, "修复的取向",
         "一律选择更保守的动作——宁可把花型做简单，也不输出机器执行不了的"
         "指令；每次修复之后必须重新检查一遍。",
         accent=GREEN, light=GREEN_L, head_h=0.36, head_size=10,
         body_size=8.6)
    card(slide, 8.85, 5.58, 4.05, 1.42, "本次实测",
         "对落盘网格 intermediate/*_action.npy 重跑检查：消融 8 组全部硬违例 0；"
         "压力 24 组中 23 组硬违例 0，余 1 组根因已定位。",
         accent=ORANGE, light=ORANGE_L, head_h=0.36, head_size=10,
         body_size=8.6)
    return slide


# ============================================================
# N5 工艺启发式 C1–C6 与成组完整性
# ============================================================

CRAFT = [
    ("C1", "同一列正反针频繁切换",
     "连续 ≥3 行来回切换会把布面拉皱。三种情形豁免：落在绞花影响区 ±2 列内、"
     "邻域内已有劈针或移针做过渡、属于反针边框像素。",
     "在切换处改为 f_split 过渡"),
    ("C2", "同一行左右移针负载不均",
     "左移与右移的数量相差超过 2 针时，机头受力偏向一侧，长期运行影响精度。",
     "把数量多的一侧改为 f_split"),
    ("C3", "床间移针行缺少缓冲",
     "整行都是床间移针，却没有集圈、空针或反针，纱线张力无处释放。",
     "在行中补入空针作为缓冲"),
    ("C4–C6", "语义层的支撑与一致性",
     "作用在 9 类 StitchType 的语义层上：结构针缺少必要边框支撑、语义类型与"
     "所在区域声明不匹配、语义层与动作层描述冲突。",
     "统一降级为通用结构针 STRUCTURE"),
]


def make_craft_rules(prs):
    slide = page(prs, "规则体系（二）：工艺启发式与成组完整性",
                 "generator_one/rules/craft_rules.py")
    band(slide, 0.45, 1.16, 12.45, 0.40,
         "工艺启发式 = 不违反机器约束，但影响成品品相；按经验阈值提示并做保守调整",
         fill=NAVY, size=10)
    rows = [["", "关注的问题", "为什么要管", "保守调整"]] + [list(r) for r in CRAFT]
    table(slide, 0.45, 1.70, [0.72, 2.68, 4.45, 4.60], rows,
          head_h=0.32, row_h=0.52, size=8.5, head_size=9,
          aligns=["c", "l", "l", "l"],
          col_colors={0: NAVY}, cell_bold={(i, 0): True for i in range(1, 5)})

    card(slide, 0.45, 4.28, 8.60, 1.22, "成组完整性守卫（wave_pair_guard）",
         "波浪吊目由“劈针 + 集圈”成对构成：劈针负责让线圈过渡，集圈负责堆出凸起。"
         "任何一方被规则单独改掉，纹理就会断成孤立的点。守卫在规则修复之后运行，"
         "逐对检查两者是否同时存在，缺失的一侧按 PatternPlan 里的列偏移补回来。",
         accent=PURPLE, light=PURPLE_L, head_h=0.38, head_size=10.5,
         body_size=8.8)

    kv = [("集圈 f_tuck", "20 格", PURPLE), ("劈针 f_split", "20 格", BLUE),
          ("成对率", "20 / 20", GREEN)]
    for i, (lb, vl, cl) in enumerate(kv):
        kpi(slide, 0.45 + i * 2.95, 5.62, 2.70, 0.62, lb, vl, value_color=cl)

    txt(slide, 0.45, 6.32, 8.60, 0.50,
        "右图是本次真实产物的 13×13 局部放大，黑框标出 8 对相邻的劈针（左）"
        "与集圈（右）。", size=8.6, color=MUTED, anchor="top", ls=1.3)

    pic_fit(slide, FIG / "cmap_zoom.png", 9.30, 4.20, 3.60, 2.75, frame=True)
    return slide


# ============================================================
# N6 检查与修复回路
# ============================================================

REPAIR_FLOW = ["生成动作网格", "规则检查 H / C", "LLM 定点修复", "后处理守卫",
               "规则复检", "落盘产物"]


def make_repair(prs):
    slide = page(prs, "检查与修复：保守降级与失败边界",
                 "含一例未收敛的真实记录")
    pic_fit(slide, FIG / "fig_repair_cases.png", 0.45, 1.16, 12.45, 2.60,
            frame=True)

    accents = [MUTED, BLUE, PURPLE, GREEN, BLUE, ORANGE]
    for i, (label, ac) in enumerate(zip(REPAIR_FLOW, accents)):
        x = 0.55 + i * 2.09
        chip(slide, x, 4.08, 1.80, 0.46, label, fill=ac, size=9.5)
        if i < 5:
            K.arrow(slide, x + 1.87, 4.21, 0.24, 0.20, color=MUTED)

    card(slide, 0.45, 4.75, 4.05, 1.95, "修复策略",
         "只改违例格及其必要邻格，不重画整幅图。每一次修改都对应源码里一个明确的"
         "规则分支，因此改动可以逐格解释。当前版本 max_repair_rounds = 1，"
         "即最多进行一轮 LLM 定点修复，之后交给守卫与复检。",
         accent=GREEN, light=GREEN_L, head_h=0.40, head_size=10.5,
         body_size=8.7)
    card(slide, 4.65, 4.75, 4.05, 1.95, "实测结果",
         "对 24 组压力测试真正落盘的网格 intermediate/*_action.npy 重跑检查："
         "工艺违例在全部 24 组中均为 0；硬违例为 0 的有 23 组。数字直接来自网格"
         "重算，不采用运行过程中的中间统计快照。",
         accent=BLUE, light=BLUE_L, head_h=0.40, head_size=10.5,
         body_size=8.7)
    card(slide, 8.85, 4.75, 4.05, 1.95, "失败边界：stress_19",
         "残留 18 处 H2，根因已定位：C3 补入的空针使其上方失去支撑而触发 H2，"
         "而 H2 在下方为空针时会把本格也降为空针，违例沿着列向上移一行——单轮"
         "修复因此无法收敛。实测反复调用需 7 轮才归零，修法明确：把修复循环"
         "改为迭代至不动点。",
         accent=ORANGE, light=ORANGE_L, head_h=0.40, head_size=10.5,
         body_size=8.7)
    return slide


# ============================================================
# N7 流程分步走查
# ============================================================

WALK = [
    ("语义解析（LLM，2.17 s）",
     "layout_type = framed；解析出 3 个区域——frame_1（边框，厚度 2，尺寸 17）、"
     "circle_1（圆形，尺寸 16，效果 wave_tuck）、背景平针。参数全部落到 JSON。",
     BLUE),
    ("几何与实现方案（0.006 s）",
     "VisualSpec 定下各区域的中心与边界；RealizationPlan 为圆形内部选定 wave_tuck："
     "period_x = 5、period_y = 4、split_col_offset = 2、tuck_col_offset = 3。",
     PURPLE),
    ("事件离散与网格填充（0.155 s）",
     "把波浪吊目离散为 40 个事件点，逐点写入 40×40 动作网格；随后的细节精修与"
     "美学处理各约 0.01 s。", GREEN),
    ("规则检查与修复（0.015 s）",
     "H1–H8 与 C1–C6 全部通过，硬违例 0、工艺违例 0，因此本次 LLM 修复轮数为 0；"
     "成组完整性守卫复核后直接放行。", ORANGE),
    ("双图产物输出（合计 2.38 s）",
     "色码图 BMP（40×40，4 种动作共 1600 格）与结构感知视觉预览图同时落盘，"
     "附 pattern_plan.json 与 pattern_stats.json。", TEAL),
]


def make_walkthrough(prs):
    slide = page(prs, "流程分步走查：一句话指令到双图产物",
                 "真实运行 outputs/runs/gui/20260731_155816")
    band(slide, 0.45, 1.14, 12.45, 0.40,
         "输入指令：“矩形边框包裹中心圆形，圆形内部为波浪吊目，背景为平针”",
         fill=NAVY, size=10.5)

    for i, (title, body, ac) in enumerate(WALK):
        numbered(slide, 0.45, 1.70 + i * 1.08, i + 1, title, body,
                 w_title=6.55, w_body=6.55, accent=ac)

    pic_fit(slide, FIG / "cmap_full.png", 8.25, 1.70, 4.65, 3.70, frame=True)
    txt(slide, 8.25, 5.45, 4.65, 0.22, "本次落盘的色码图与动作分布（共 1600 格）",
        size=8.6, color=MUTED, align="c", ls=1.0)

    dist = [("f_knit 前床成圈", "1196", BLUE), ("b_knit 后床成圈", "364", GREEN),
            ("f_tuck 前床集圈", "20", PURPLE),
            ("f_split 前床劈针", "20", ORANGE)]
    for i, (lb, vl, cl) in enumerate(dist):
        kpi(slide, 8.25 + (i % 2) * 2.40, 5.72 + (i // 2) * 0.62, 2.25, 0.55,
            lb, vl, value_color=cl, label_size=8, value_size=10)
    return slide


# ============================================================
# N8 桌面工作台总览
# ============================================================

GUI_NOTES = [
    ("指令输入与预设", "可以直接输入中文指令，也可以从预设模板里挑一条再改写。",
     BLUE),
    ("生成模式选择", "仅色码图 / 色码图 + 视觉预览两种；前者纯 CPU 秒级返回，"
     "后者调用远程 GPU。", PURPLE),
    ("双图对照视图", "色码图与视觉预览图并排显示，支持缩放与像素级查看。", GREEN),
    ("规则与统计面板", "实时展示区域数、事件数、动作分布、违例条目与修复轮数。",
     ORANGE),
    ("一键导出", "把 BMP、预览图、pattern_plan.json、pattern_stats.json 打包"
     "到本地目录。", TEAL),
]


def make_gui(prs):
    slide = page(prs, "桌面工作台：面向设计师的本地界面",
                 "PySide6 + QFluentWidgets")
    pic_fit(slide, GUI / "shot_workspace_dual.png", 0.45, 1.20, 8.05, 5.03,
            frame=True)
    for i, (title, body, ac) in enumerate(GUI_NOTES):
        numbered(slide, 8.70, 1.22 + i * 1.02, i + 1, title, body,
                 w_title=3.52, w_body=3.52, accent=ac, cw=0.42, size=10.5,
                 body_size=8.4)
    band(slide, 0.45, 6.40, 12.45, 0.50,
         "界面只负责交互与产物管理，生成逻辑与命令行用的是同一套代码，"
         "界面里跑出的结果可以直接复现", fill=NAVY, size=10)
    return slide


# ============================================================
# N9 增量编辑、版本谱系与部署
# ============================================================

def make_edit_deploy(prs):
    slide = page(prs, "增量编辑、版本谱系与部署形态",
                 "本地界面 + 远程 GPU 推理")
    pic_fit(slide, GUI / "shot_workspace_edit.png", 0.45, 1.18, 6.05, 3.78,
            frame=True)
    pic_fit(slide, GUI / "shot_history.png", 6.85, 1.18, 6.05, 3.78,
            frame=True)
    txt(slide, 0.45, 5.02, 6.05, 0.30,
        "编辑模式：只改一个词，系统只重编译受影响的区域（图中为仅色码图模式）",
        size=8.8, color=MUTED, align="c", ls=1.0)
    txt(slide, 6.85, 5.02, 6.05, 0.30,
        "版本谱系：每次生成与编辑都留一条记录，可回溯到任意一版",
        size=8.8, color=MUTED, align="c", ls=1.0)

    card(slide, 0.45, 5.40, 6.05, 1.38, "增量编辑是怎么做到的",
         "在原指令上改写局部描述，系统比对新旧 PatternPlan 的字段差异，只对差异"
         "涉及的区域重新编译，其余区域逐格保持原值——因此“改一个词”的影响范围"
         "是可度量的。",
         accent=BLUE, light=BLUE_L, head_h=0.40, head_size=10.5,
         body_size=8.8)
    card(slide, 6.85, 5.40, 6.05, 1.38, "部署形态",
         "界面与轻量编译在本地 CPU 运行；视觉预览的扩散模型推理放在远程 GPU，"
         "产物回传后在本地统一归档；网络或显卡异常时自动退回仅色码图模式，"
         "不阻断设计流程。",
         accent=PURPLE, light=PURPLE_L, head_h=0.40, head_size=10.5,
         body_size=8.8)
    return slide


# ============================================================
# 组装
# ============================================================

def main():
    prs = Presentation(str(SRC))
    old = list(prs.slides)
    print(f"原稿 {len(old)} 页")

    # 原第 5 页：改措辞与零违例口径
    s5 = old[4]
    for a, b in [("Zero-Violation 保证", "规则检查与保守修复"),
                 ("= 0（全部测试集）", "= 0（消融 8/8 · 压力 23/24）"),
                 ("= 0（典型场景）", "= 0（压力 24/24）")]:
        print(f"  retext {a!r} -> {retext(s5, a, b)}")

    toc = make_toc(prs)
    n1 = make_framework(prs)
    n2 = make_compile(prs)
    n3 = make_action_space(prs)
    n4 = make_hard_rules(prs)
    n5 = make_craft_rules(prs)
    n6 = make_repair(prs)
    n7 = make_walkthrough(prs)
    n8 = make_gui(prs)
    n9 = make_edit_deploy(prs)

    # 原目录 old[1] 与原“创新点②” old[6] 不列入，即删除
    set_order(prs, [old[0], toc, old[2], old[3], old[4], n1, old[5], n2,
                    n3, n4, n5, n6, old[7], n7, old[8], old[9], old[10],
                    n8, n9, old[11], old[12]])
    renumber(prs)
    prs.save(str(OUT))
    print(f"已写出 {OUT}，共 {len(prs.slides._sldIdLst)} 页")


if __name__ == "__main__":
    main()
