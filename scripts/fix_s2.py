# -*- coding: utf-8 -*-
"""S2 目录页修复：左右两栏顶部对齐、同节奏排布，整体水平居中。"""
from pptx import Presentation
from pptx.util import Inches as IN

PPTX = r"e:\DeskTop\knitting_ai_competition.pptx"

p = Presentation(PPTX)
s = p.slides[1]
m = {sh.shape_id: sh for sh in s.shapes}

# 四行统一节奏（左栏 4 项 / 右栏 3 项均从同一顶部开始）
rows = [1.75, 3.1, 4.45, 5.8]


def place(ids_rows, conn_ids, chip_l, txt_l, sub_w):
    for k, (chip, chiptxt, title, sub) in enumerate(ids_rows):
        t = rows[k]
        for sid in (chip, chiptxt):
            m[sid].left = IN(chip_l)
            m[sid].top = IN(t)
        m[title].left = IN(txt_l)
        m[title].top = IN(t)
        m[sub].left = IN(txt_l)
        m[sub].top = IN(t + 0.22)
        m[sub].width = IN(sub_w)
    for k, cid in enumerate(conn_ids):
        m[cid].left = IN(chip_l + 0.21)
        m[cid].top = IN(rows[k] + 0.55)
        m[cid].height = IN(0.67)


# 左栏 01-04
place([(6, 7, 8, 9), (11, 12, 13, 14), (16, 17, 18, 19), (21, 22, 23, 24)],
      [10, 15, 20], 1.35, 1.9, 5.0)
# 右栏 05-07（与左栏顶部对齐）
place([(26, 27, 28, 29), (31, 32, 33, 34), (36, 37, 38, 39)],
      [30, 35], 7.85, 8.4, 4.6)

p.save(PPTX)
print("S2 fixed OK")
