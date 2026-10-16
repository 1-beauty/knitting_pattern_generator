# -*- coding: utf-8 -*-
"""第七轮补丁（正文）：张力等级→切换风险等级 全文统一（para79/86 等）。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
m = Document(MAIN)
n = 0
for p in m.paragraphs:
    for t in p._p.findall('.//' + qn('w:t')):
        if t.text and '张力等级' in t.text:
            n += t.text.count('张力等级'); t.text = t.text.replace('张力等级', '切换风险等级')
for tbl in m.tables:
    for row in tbl.rows:
        for c in row.cells:
            for t in c._tc.findall('.//' + qn('w:t')):
                if t.text and '张力等级' in t.text:
                    n += t.text.count('张力等级'); t.text = t.text.replace('张力等级', '切换风险等级')
assert n >= 2, n
txt = '\n'.join(p.text for p in m.paragraphs)
assert '张力等级' not in txt
m.save(MAIN)
print('OK round7c main | renamed:', n)
