# -*- coding: utf-8 -*-
"""Round 8q: Close the evidence loop in discussion + conclusion."""
import sys
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(SRC)

def rewrite(p, new_text):
    """Replace paragraph text, preserving first run's formatting."""
    assert not p._p.findall('.//' + qn('m:oMath')), "Has oMath!"
    runs = p.runs
    if not runs:
        p.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def wt_sub(p, old, new):
    count = 0
    for t in p._p.findall('.//' + qn('w:t')):
        if t.text and old in t.text:
            t.text = t.text.replace(old, new)
            count += 1
    assert count >= 1, f"Not found: '{old[:50]}...'"

# ============================================================
# 1. p128 — Discussion opening: 4-level evidence chain
# ============================================================
p128_new = (
    "上述实验从语义、编译、视觉与物理四个层次构成闭环验证。"
    "语义层：PatternPlan 在 100 条受控输入上达 97.2% 字段级准确率，"
    "三次重复调用共识标准差 2.6 个百分点，证实设计意图可被稳定解析为可编译结构；"
    "编译层：44 条硬嵌套输入的受控消融证实事件级编译在硬约束零违例（构造性保证）"
    "与成组完整性（1.000 vs 0.923, p<0.001）上显著优于密集填充，"
    "且与 CP-SAT 全局约束求解的对比区分了"
    "\u201c合法性可后验闭合\u201d与\u201c成组语义必须编译时注入\u201d两类本质不同的差距；"
    "视觉层：KID 与 BoundaryF1 在 Stage 2\u20134 间呈反向趋势，"
    "定量刻画了全局分布保真与边界可读性的设计权衡，"
    "CLIPScore 四档持平（0.282\u20130.283）表明语义一致性不构成区分维度；"
    "物理层：5 片工厂样片均成功织出且区域结构与设计一致，"
    "从数字生成到真实制造的完整链路得到验证。"
    "各对比与消融的具体结论已在 4.3 节给出，"
    "本节聚焦方法层面的解读与验证边界。"
    "方法层面的详细解读、PatternPlan 表示层作用与后续约束建模方向见补充材料 Section N。"
)
rewrite(doc.paragraphs[128], p128_new)
print("p128 rewritten (4-level evidence chain)")

# ============================================================
# 2. p132 — Conclusion PatternPlan: add stability
# ============================================================
old_frag = "97.2% 的字段级解析准确率与 83.0% 的样本完全正确率（4.2 节）"
new_frag = ("97.2% 的字段级解析准确率与 83.0% 的样本完全正确率、"
            "三次重复调用共识标准差 2.6 个百分点（4.2 节）")
wt_sub(doc.paragraphs[132], old_frag, new_frag)
print("p132 stability added")

# ============================================================
# 3. p134 — Conclusion visual preview: mention KID/CLIP
# ============================================================
p134_new = (
    "作为辅助机制，结构感知视觉预览以条件流匹配模型将动作色码图转换为针织外观预览，"
    "为设计阶段提供边界、纹理与连续性的视觉反馈。"
    "Stage 2\u20133 在像素保真与 KID 分布指标上占优，"
    "Stage 4 在 BoundaryF1 与 SeamErr 上占优——"
    "重建\u2013边界的反向权衡同时由分布指标（KID）与结构指标定量刻画；"
    "CLIPScore 四档持平表明语义对齐不构成 Stage 间的区分维度，"
    "单一图像重建指标不足以完整评价面向设计辅助的针织预览效果（4.4 节）。"
)
rewrite(doc.paragraphs[134], p134_new)
print("p134 rewritten (KID/CLIP in conclusion)")

# ============================================================
# 4. p135 — Overall conclusion: strengthen evidence chain phrase
# ============================================================
# Current: "...在受测输入上实现动作表示的硬约束零违例（例外见 4.3.3 节）、区域结构与可编辑性；"
# Insert before "在此基础上": "从语义解析稳定性（std 2.6 pp）、编译零硬违例与成组完整性（p<0.001）、视觉预览分布—边界权衡量化到 5 片工厂实织验证，四级证据链闭合了从设计意图到物理制品的完整链路。"
old_anchor = "在此基础上，本文系统已通过合作工厂的真实落地检测"
new_anchor = ("从语义解析稳定性（std 2.6 pp）、编译零硬违例与成组完整性（p<0.001）、"
              "视觉预览分布\u2013边界权衡量化到 5 片工厂实织验证，"
              "四级证据链闭合了从设计意图到物理制品的完整链路。"
              "在此基础上，本文系统已通过合作工厂的真实落地检测")
wt_sub(doc.paragraphs[135], old_anchor, new_anchor)
print("p135 evidence chain phrase inserted")

# ============================================================
doc.save(SRC)
print("\nOK round8q")
