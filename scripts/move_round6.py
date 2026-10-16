# -*- coding: utf-8 -*-
"""第六轮结构性调整 B：结构守卫阴性结果块（para150-152）下放 Supp Section K。
保留 para148-149（表5量化讨论 + CSP不能替代事件级编译的核心论点）。
para150 含 oMath×4 → 必须 deepcopy（保留公式）。迁移段交叉引用加"正文"前缀。
所有断言通过后才落盘 main + supp。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'

m = Document(MAIN)
s = Document(SUPP)
ps = m.paragraphs

def cn(t): return sum(1 for c in t if '\u4e00' <= c <= '\u9fff')
def find(pre): return next(p for p in ps if p.text.startswith(pre))

before_cn = sum(cn(p.text) for p in ps)

# ---- 定位 ----
p149 = find('加入 Plan-Grid+CSP 后，递进对比链进一步闭合')
p150 = find('值得注意的是，Plan-Grid 的 2 条残留硬违例')
p151 = find('事件级结构守卫的作用范围。')
p152 = find('该组为阴性结果：')

assert len(p150._p.findall('.//'+qn('m:oMath'))) == 4, 'p150 oMath!=4'
assert len(p151._p.findall('.//'+qn('m:oMath'))) == 0
assert len(p152._p.findall('.//'+qn('m:oMath'))) == 0
assert len(p149._p.findall('.//'+qn('m:oMath'))) == 0

# ---- deepcopy + 交叉引用加"正文"前缀 ----
def fix_copy(el, fixes):
    ts = el.findall('.//'+qn('w:t'))
    for old, new in fixes:
        cnt = 0
        for t in ts:
            if t.text and old in t.text:
                t.text = t.text.replace(old, new); cnt += 1
        assert cnt == 1, ('fix not single-hit', old, cnt)
    return el

cp150 = fix_copy(copy.deepcopy(p150._p), [('4.3.1 节', '正文 4.3.1 节')])
cp151 = fix_copy(copy.deepcopy(p151._p), [('表 4 与', '正文表 4 与')])
cp152 = fix_copy(copy.deepcopy(p152._p), [('（表 5）', '（正文表 5）')])

# ---- Supp Section K ----
sup_body = s.element.body
sect = sup_body.find(qn('w:sectPr'))
h3_style = next(p.style for p in s.paragraphs if p.style.name == 'Heading 3')
body_style = next(p.style for p in s.paragraphs if p.style.name == 'Body Text')
body_sid = body_style.style_id

def sup_new(text, style):
    p = s.add_paragraph(text, style=style)
    sect.addprevious(p._p)
    return p

# 规范化迁移段段落样式为 Supp Body Text（保留 run 级格式与 oMath）
for cp in (cp150, cp151, cp152):
    pPr = cp.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); cp.insert(0, pPr)
    pStyle = pPr.find(qn('w:pStyle'))
    if pStyle is None:
        pStyle = OxmlElement('w:pStyle'); pPr.insert(0, pStyle)
    pStyle.set(qn('w:val'), body_sid)

sup_new('K. 结构守卫消融阴性结果与残留硬违例溯源（Negative Result of '
        'Structural-Guard Ablation and Trace of Residual Hard Violations）', h3_style)
sup_new('本节承接正文 4.3.2 节表 5 的内部基线控变量对比，补充 Plan-Grid 残留硬违例在'
        '确定性规则修复阶段的溯源、事件级结构守卫（Layer Clipping 与 Wave Pair Guard）'
        '的作用范围，及其在本批受控嵌套输入上的阴性结果。', body_style)
sect.addprevious(cp150)
sect.addprevious(cp151)
sect.addprevious(cp152)

# ---- main：删 para150-152，para149 追加指针 ----
for p in (p150, p151, p152):
    p._p.getparent().remove(p._p)

POINTER = ('Plan-Grid 残留硬违例的规则修复溯源、事件级结构守卫的作用范围与其在本批'
           '受控输入上的阴性结果，详见补充材料 Section K。')
last_r = p149.runs[-1]._r
new_r = copy.deepcopy(last_r)
for t in new_r.findall(qn('w:t')):
    new_r.remove(t)
for om in new_r.findall('.//'+qn('m:oMath')):
    om.getparent().remove(om)
nt = OxmlElement('w:t'); nt.set(qn('xml:space'), 'preserve'); nt.text = POINTER
new_r.append(nt)
last_r.addnext(new_r)

# ---- 校验 + 落盘 ----
after_cn = sum(cn(p.text) for p in m.paragraphs)
assert 'Plan-Grid 残留硬违例的规则修复溯源' in p149.text, 'pointer missing'
assert not any(p.text.startswith('该组为阴性结果：') for p in m.paragraphs), 'p152 still in main'
assert not any(p.text.startswith('值得注意的是，Plan-Grid 的 2 条') for p in m.paragraphs), 'p150 still'
assert not any(p.text.startswith('事件级结构守卫的作用范围。') for p in m.paragraphs), 'p151 still'

sup_txt = '\n'.join(p.text for p in s.paragraphs)
assert 'K. 结构守卫消融阴性结果' in sup_txt
assert '正文 4.3.1 节协议 A' in sup_txt
assert '正文表 4 与表 S12' in sup_txt
assert '（正文表 5）' in sup_txt

m.save(MAIN)
s.save(SUPP)
print('OK | main CN: %d -> %d (net %+d)' % (before_cn, after_cn, after_cn - before_cn))
print('Supp paras:', len(s.paragraphs))
