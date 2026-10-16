# -*- coding: utf-8 -*-
"""Batch 3: S1/S13 light restyle + embed real result figures into S6/S8/S11."""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = r"e:\DeskTop\knitting_pattern_generator"
PPTX = r"e:\DeskTop\knitting_ai_competition.pptx"
ASSETS = os.path.join(ROOT, "scripts", "assets_ppt")
os.makedirs(ASSETS, exist_ok=True)

# ---------- 0. preprocess images (nearest-neighbor upscale for pixel maps) ----------
def nn_upscale(src, dst, size):
    im = Image.open(src).convert("RGB")
    im = im.resize((size, size), Image.NEAREST)
    im.save(dst)

SIM = os.path.join(ROOT, r"outputs\runs\simulation\sim_sem01_stage4\simulation")
EDT = os.path.join(ROOT, r"outputs\runs\edit")
nn_upscale(os.path.join(SIM, "id_matrix_preview.png"), os.path.join(ASSETS, "sem01_colormap.png"), 480)
nn_upscale(os.path.join(SIM, "stage4_result.png"), os.path.join(ASSETS, "sem01_stage4.png"), 640)
nn_upscale(os.path.join(EDT, r"edt_00\color_map\edt_00_sim_input.png"), os.path.join(ASSETS, "edt_00.png"), 480)
nn_upscale(os.path.join(EDT, r"edt_05\color_map\edt_05_sim_input.png"), os.path.join(ASSETS, "edt_05.png"), 480)
STAGES_FIG = os.path.join(ROOT, r"docs\figures\fig_sim_sem01_stages.png")

prs = Presentation(PPTX)
IN = Inches

def shp_map(slide):
    return {sh.shape_id: sh for sh in slide.shapes}

def set_pos(sh, L=None, T=None, W=None, H=None):
    if L is not None: sh.left = IN(L)
    if T is not None: sh.top = IN(T)
    if W is not None: sh.width = IN(W)
    if H is not None: sh.height = IN(H)

def set_fill(sh, rgb):
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(rgb)

def set_line(sh, rgb, w=1.0):
    sh.line.color.rgb = RGBColor.from_string(rgb)
    sh.line.width = Pt(w)

def set_text_color(sh, rgb, size=None, bold=None):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(rgb)
            if size is not None: r.font.size = Pt(size)
            if bold is not None: r.font.bold = bold

def add_text(slide, L, T, W, H, text, size, color, bold=False, align=None, font="Microsoft YaHei"):
    tb = slide.shapes.add_textbox(IN(L), IN(T), IN(W), IN(H))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    if align: p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = font
    return tb

def add_pic(slide, path, L, T, W=None, H=None, border="#E2E8F0"):
    kw = {}
    if W is not None: kw["width"] = IN(W)
    if H is not None: kw["height"] = IN(H)
    pic = slide.shapes.add_picture(path, IN(L), IN(T), **kw)
    if border:
        pic.line.color.rgb = RGBColor.from_string(border.lstrip("#"))
        pic.line.width = Pt(0.75)
    return pic

# =============== S1 cover -> light ===============
s1 = prs.slides[0]
m = shp_map(s1)
set_fill(m[2], "F4F6FA"); set_line(m[2], "F4F6FA")
for sid in range(4, 44):          # decorative dot grid
    set_fill(m[sid], "BFDBFE"); set_line(m[sid], "BFDBFE")
set_text_color(m[44], "1E3A8A")   # main title
set_text_color(m[45], "2563EB")   # competition subtitle
for pill, txt in [(46, 47), (48, 49), (50, 51), (52, 53)]:
    set_fill(m[pill], "EFF6FF"); set_line(m[pill], "93C5FD")
    set_text_color(m[txt], "2563EB")
# byline (54) #64748B already fine on light bg

# =============== S13 closing -> light ===============
s13 = prs.slides[12]
m = shp_map(s13)
set_fill(m[2], "F4F6FA"); set_line(m[2], "F4F6FA")
set_text_color(m[4], "1E3A8A", size=20, bold=True)   # 核心成果总结 -> real title style
for card in (5, 8, 11, 14, 17):
    set_fill(m[card], "FFFFFF"); set_line(m[card], "CBD5E1")
for txt in (7, 10, 13, 16, 19):
    set_text_color(m[txt], "2D3561")
set_line(m[20], "2563EB", 1.5)                       # divider above thanks
set_text_color(m[21], "1E3A8A")                      # 感谢聆听
set_text_color(m[22], "64748B")                      # data source footnote

# =============== S6: real evidence pair in bottom band ===============
s6 = prs.slides[5]
add_text(s6, 0.45, 6.4, 3.15, 0.3, "端到端真实样例（SEM-01）", 10.5, "4C1D95", bold=True)
add_pic(s6, os.path.join(ASSETS, "sem01_colormap.png"), 3.75, 5.95, H=1.2, W=1.2)
add_text(s6, 5.0, 6.4, 0.45, 0.3, "→", 14, "7C3AED", bold=True, align=PP_ALIGN.CENTER)
add_pic(s6, os.path.join(ASSETS, "sem01_stage4.png"), 5.5, 5.95, H=1.2, W=1.2)
add_text(s6, 6.95, 6.4, 5.9, 0.3, "左：编译得到的动作色码图（40×40）    右：Stage 4 结构感知视觉预览", 9, "64748B")

# =============== S8: compress cards + four-stage preview strip ===============
s8 = prs.slides[7]
m = shp_map(s8)
CT = [1.42, 2.34, 3.26, 4.18]     # new card tops, H 0.82
left_cards  = [(7, 8, 9, 10, 11, 12), (13, 14, 15, 16, 17, 18),
               (19, 20, 21, 22, 23, 24), (25, 26, 27, 28, 29, 30)]
for k, (bg, chip, chiptxt, desc, mse, ssim) in enumerate(left_cards):
    t = CT[k]
    set_pos(m[bg],   T=t, H=0.82)
    set_pos(m[chip], T=t, H=0.82)
    set_pos(m[chiptxt], T=t + 0.20, H=0.42)
    set_pos(m[desc], T=t + 0.05, H=0.44)
    set_pos(m[mse],  T=t + 0.48, H=0.24)
    set_pos(m[ssim], T=t + 0.48, H=0.24)
set_pos(m[31], T=4.51)            # 最终输出模型 ✓
set_pos(m[32], T=5.10)            # footnote
right_cards = [(34, 35), (36, 37), (38, 39), (40, 41)]
for k, (bg, txt) in enumerate(right_cards):
    t = CT[k]
    set_pos(m[bg],  T=t, H=0.82)
    set_pos(m[txt], T=t + 0.05, H=0.72)
set_pos(m[42], T=5.10)            # training data bar
set_pos(m[43], T=5.14)
# four-stage progression strip (real model outputs)
add_pic(s8, STAGES_FIG, 0.45, 5.62, H=1.42, W=1.42 * 1304 / 360)
add_text(s8, 5.85, 6.05, 7.0, 0.6,
         "Case SEM-01 真实输出：随训练阶段推进，预览由噪声逐步收敛为边界清晰、纹理可读的针织结构",
         10, "78350F", bold=True)

# =============== S11: edit before/after color maps in empty slot ===============
s11 = prs.slides[10]
add_text(s11, 8.85, 3.86, 4.0, 0.26, "色码图对照：基准 → 形状替换", 9.5, "64748B", bold=True, align=PP_ALIGN.CENTER)
add_pic(s11, os.path.join(ASSETS, "edt_00.png"), 9.32, 4.22, H=1.28, W=1.28)
add_text(s11, 10.6, 4.7, 0.5, 0.32, "→", 14, "DC2626", bold=True, align=PP_ALIGN.CENTER)
add_pic(s11, os.path.join(ASSETS, "edt_05.png"), 11.1, 4.22, H=1.28, W=1.28)
add_text(s11, 9.32, 5.54, 1.28, 0.2, "EDT-00 基准", 8, "64748B", align=PP_ALIGN.CENTER)
add_text(s11, 11.1, 5.54, 1.28, 0.2, "EDT-05 菱形", 8, "64748B", align=PP_ALIGN.CENTER)

prs.save(PPTX)
print("SAVED batch3")
