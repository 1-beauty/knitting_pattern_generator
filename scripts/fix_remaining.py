# -*- coding: utf-8 -*-
"""收尾修复：
① para[31]/[175] 从 .bak 动态提取 bold=False 的"图 1"/"图 10"正文句原文恢复（不手工誊写）；
② 参考文献 [1][2][40][49]：删 pandoc 遗留 hyperlink(URL 粘连) + 用 REFS 重写 run + 补悬挂缩进(与其余45条一致)。
"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Twips

MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
BAK = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx.bak'

REFS = {
 1: 'Shima Seiki, SDS-ONE APEX Series: All-in-One Design System for Computerized Flat Knitting, https://www.shimaseiki.com/product/design, 2026 (accessed 29 July 2026).',
 2: 'STOLL by KARL MAYER, M1plus Pattern Software for Computerized Flat Knitting Machines, https://www.stoll.com/en/, 2026 (accessed 29 July 2026).',
 40: 'L. Perron, V. Furnon, OR-Tools: Software suite for combinatorial optimization, https://developers.google.com/optimization, 2026 (accessed 20 September 2026).',
 49: 'A. Kaspar, J. McCann, W. Matusik, Knitout: A machine-independent representation for machine knitting, https://textiles-lab.github.io/knitout/, 2026 (accessed 29 July 2026).',
}
REF_PARA = {206: 1, 207: 2, 245: 40, 254: 49}

# --- ① 从 .bak 提取正文句原文 ---
bak = Document(BAK)
b31 = b175 = None
for p in bak.paragraphs:
    bold = any(r.bold for r in p.runs if r.bold is not None)
    if bold:
        continue
    ts = p.text.strip()
    if b31 is None and re.match(r'^图\s*1\s', ts):
        b31 = ts
    if b175 is None and re.match(r'^图\s*10\s', ts):
        b175 = ts
assert b31 and b175, (b31 is None, b175 is None)
print('BAK 正文句提取成功:')
print('  b31 :', b31[:50], '...')
print('  b175:', b175[:50], '...')

# --- 主文件修复 ---
doc = Document(MAIN)
ps = doc.paragraphs

# ① 恢复 para[31]/[175]
assert ps[31].text.strip().startswith('Fig. 1.'), ps[31].text[:30]
assert ps[175].text.strip().startswith('Fig. 10.'), ps[175].text[:30]
for p, txt in ((ps[31], b31), (ps[175], b175)):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(txt)
print('para[31]/[175] 已恢复为中文正文句')

# ② 参考文献 4 条
for i, k in REF_PARA.items():
    p = ps[i]
    for hl in p._p.findall(qn('w:hyperlink')):
        p._p.remove(hl)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(f'[{k}] ' + REFS[k])
    pf = p.paragraph_format
    pf.left_indent = Twips(480)
    pf.first_line_indent = Twips(-480)
print('参考文献 [1][2][40][49] 已删 hyperlink + 重写 + 悬挂缩进')

doc.save(MAIN)

# --- 验证 ---
d = Document(MAIN)
dps = d.paragraphs
figs = [i for i, p in enumerate(dps) if p.text.strip().startswith('Fig.')]
print('\n=== 验证 ===')
print('Fig. 段落数(应=10):', len(figs), '->', figs)
print('para[31] 开头:', dps[31].text.strip()[:24])
print('para[175] 开头:', dps[175].text.strip()[:24])
refs = [(i, p.text.strip()) for i, p in enumerate(dps) if re.match(r'^\[\d+\]', p.text.strip())]
print('参考文献条数(应=49):', len(refs))
for i, t in refs:
    n = int(re.match(r'^\[(\d+)\]', t).group(1))
    if n in (1, 2, 40, 49):
        print(f'  [{i}] {t[:88]}')
nums = sorted(int(re.match(r'^\[(\d+)\]', t).group(1)) for _, t in refs)
print('编号连续性:', '完整1-49' if nums == list(range(1, 50)) else f'缺{set(range(1,50))-set(nums)}')
