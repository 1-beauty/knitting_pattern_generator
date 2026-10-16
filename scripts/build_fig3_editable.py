"""按图1风格重绘原生可编辑图3（事件级约束编译机制），全英文 TNR，并执行交付审计。"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
import re
from io import BytesIO
from pathlib import Path
import sys
import time
from zipfile import ZipFile

sys.dont_write_bytecode = True
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
import build_fig1_editable as f1

ROOT = f1.ROOT
TEMPLATE = f1.TEMPLATE
OUT = ROOT / 'outputs/fig3_pptx_en_20260922'
STEM = ROOT / 'docs/figures/Fig3_Event_Level_Constraint_Compilation_EN'
PPTX, PNG, PDF = (STEM.with_suffix(ext) for ext in ('.pptx', '.png', '.pdf'))
INK, MUTED, BLUE = f1.INK, f1.MUTED, f1.BLUE
PLAN = ROOT / 'outputs/runs/patternplan/pp_01/color_map/pp_01_pattern_plan.json'
STATS = ROOT / 'outputs/runs/patternplan/pp_01/color_map/pp_01_stats.json'
ACTION_BMP = ROOT / 'outputs/runs/patternplan/pp_01/color_map/pp_01_output.bmp'
ACTION_NPY = ACTION_BMP.parent / 'intermediate/pp_01_action.npy'
COL_W, GAP, X0 = 4.05, .30, .16
COLS = [X0 + i * (COL_W + GAP) for i in range(4)]
PALETTES = [('EAF3FC', '4E86B5'), ('EAF4ED', '487F55'), ('FDF3E7', 'B07A3C'), ('F1EDF8', '8062AA')]
GROUND, RING, DISC, WAVE = 'E8EDF2', '4E86B5', '9FB8CF', '8062AA'
CANVAS_SIZE = 1.78
TEXT_RECORDS = []
VISUAL_RECORDS = []


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(name, data):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def instance():
    # 只读 Generation Case 01（pp_01）的历史运行产物；掩膜公式与图1同源并已对真实色码图核验。
    import numpy as np
    plan = json.loads(PLAN.read_text(encoding='utf-8'))
    stats = json.loads(STATS.read_text(encoding='utf-8'))
    motifs = {m['motif_id']: m for m in plan['motifs']}
    assert motifs['wave_1']['parent_id'] == 'circle_1'
    assert motifs['frame_1']['params']['border_thickness'] == 2
    assert motifs['wave_1']['params']['period_x'] == 5 and motifs['wave_1']['params']['period_y'] == 4
    assert motifs['wave_1']['params']['split_col_offset'] == 2
    assert motifs['wave_1']['params']['tuck_col_offset'] == 3
    actions = np.load(ACTION_NPY, allow_pickle=False)
    assert actions.shape == (40, 40) and set(np.unique(actions).tolist()) == {0, 1, 7, 8}
    grid = Image.open(ACTION_BMP).convert('RGB')
    colors = {0: (255, 0, 16), 1: (43, 206, 72), 7: (0, 164, 255), 8: (255, 230, 6)}
    assert all(grid.getpixel((c, r)) == colors[int(actions[r, c])] for r in range(40) for c in range(40))
    full = {(r, c) for r in range(40) for c in range(40)}
    # 与 generator_one/generation/visual_to_realization.py 的 lowering 几何一致（图1证据同款）。
    masks = {'frame_1': {(r, c) for r, c in full if 15 < max(abs(r - 20), abs(c - 20)) <= 17},
             'circle_1': {(r, c) for r, c in full if (r - 20) ** 2 + (c - 20) ** 2 <= 13 ** 2},
             'wave_1': {(r, c) for r, c in full if (r - 20) ** 2 + (c - 20) ** 2 <= 12 ** 2}}
    assert {k: len(v) for k, v in masks.items()} == {'frame_1': 264, 'circle_1': 529, 'wave_1': 441}
    assert masks['wave_1'] <= masks['circle_1'], '子作用域必须位于父作用域内'
    green = {(r, c) for r, c in full if grid.getpixel((c, r)) == colors[1]}
    assert green == masks['frame_1'] | (masks['circle_1'] - masks['wave_1']), '掩膜必须与真实色码图一致'
    event_cells = {(r, c) for r, c in full if int(actions[r, c]) in (7, 8)}
    assert event_cells <= masks['wave_1'], '事件锚点必须落在纹理作用域内'
    assert len(event_cells) == 32 and stats['total_events'] == 32
    assert all(stats['violations'][k] == 0 for k in ('hard_before', 'craft_before', 'hard_after', 'craft_after'))
    assert stats['violations']['hard'] == [] and stats['violations']['craft'] == []
    ring_img = Image.new('RGB', (40, 40), tuple(bytes.fromhex(GROUND)))
    for r, c in masks['frame_1']:
        ring_img.putpixel((c, r), tuple(bytes.fromhex(RING)))
    disc_img = Image.new('RGB', (40, 40), tuple(bytes.fromhex(GROUND)))
    for r, c in masks['circle_1']:
        disc_img.putpixel((c, r), tuple(bytes.fromhex(DISC)))
    for r, c in masks['wave_1']:
        disc_img.putpixel((c, r), tuple(bytes.fromhex(WAVE)))
    evidence = {
        'sample': 'Generation Case 01 (pp_01), historical run outputs only',
        'plan_path': str(PLAN), 'plan_sha256': sha(PLAN), 'stats_sha256': sha(STATS),
        'action_npy_sha256': sha(ACTION_NPY), 'action_bmp_sha256': sha(ACTION_BMP),
        'geometry_source': 'generator_one/generation/visual_to_realization.py lowering geometry, '
                           'cross-checked against the real pp_01 color map in Figure 1 evidence',
        'mask_formulas': {'frame_1': '15 < chebyshev(r, c; 20, 20) <= 17 (half 17, thickness 2)',
                          'circle_1': '(r-20)^2 + (c-20)^2 <= 13^2',
                          'wave_1': '(r-20)^2 + (c-20)^2 <= 12^2'},
        'mask_areas': {k: len(v) for k, v in masks.items()},
        'green_pixels_equal_frame_union_circle_minus_wave': True,
        'event_cells_inside_wave_scope': True, 'event_count': 32, 'wave_groups': 16,
        'transfer_events': 0, 'violations': stats['violations'],
    }
    write_json('sample_evidence.json', evidence)
    return masks, ring_img, disc_img


def text(sl, name, x, y, w, h, value, size=20, color=INK, bold=False, align='left'):
    s = f1.text(sl, name, x, y, w, h, value, size, color, bold, align)
    TEXT_RECORDS.append({'name': name, 'text': value, 'font_pt': size, 'box_inches': [x, y, w, h]})
    return s


def scope_runs(sl, image, x, y, size, prefix):
    # 与图1/图2同款：同行同色连续格合并为原生矩形，颜色取自掩膜参考图。
    elements = []
    for r in range(40):
        c = 0
        while c < 40:
            rgb = image.getpixel((c, r))
            start = c
            while c < 40 and image.getpixel((c, r)) == rgb:
                c += 1
            hexrgb = ''.join(f'{k:02X}' for k in rgb)
            elements.append(f1.shape(sl, f'{prefix}_run_{r}_{start}_{c}', x + start * size / 40,
                                     y + r * size / 40, (c - start) * size / 40, size / 40,
                                     hexrgb, hexrgb, lw=.18))
    f1.grouped(sl, prefix + '_group', elements)
    VISUAL_RECORDS.append({'prefix': prefix, 'layout': [x, y, size]})


def chip(sl, name, x, y, w, h, fill, accent, line1, line2=None, s1=24, s2=22, m=.06):
    f1.shape(sl, name, x, y, w, h, fill, accent, 'round')
    if line2 is None:
        text(sl, name + '_text', x + m, y + (h - .40) / 2, w - 2 * m, .40, line1, s1, accent, True, 'center')
    else:
        text(sl, name + '_title', x + m, y + .10, w - 2 * m, .40, line1, s1, bold=True, align='center')
        text(sl, name + '_sub', x + m, y + .52, w - 2 * m, .38, line2, s2, MUTED, align='center')


def build_deck():
    OUT.mkdir(parents=True, exist_ok=True)
    masks, ring_img, disc_img = instance()
    prs = Presentation()
    if TEMPLATE.exists():
        tpl = Presentation(TEMPLATE)
        prs.slide_width, prs.slide_height = tpl.slide_width, tpl.slide_height
    else:
        # 模板随2026-09-29 docs整理删除；幻灯片尺寸改读图1证据清单（与图2同款回退）。
        ev = ROOT / 'outputs/fig1_pptx_en_20260921/source_evidence.json'
        if ev.exists():
            w, h = json.loads(ev.read_text(encoding='utf-8'))['template_inches']
        else:
            w, h = 17.416666666666668, 9.802083333333334
        prs.slide_width, prs.slide_height = Inches(w), Inches(h)
    # 导师标准：底部横带删除后回收幻灯片高度至9.00in，上下留白 .40/.40 平衡。
    prs.slide_height = Inches(9.00)
    prs.core_properties.title = 'Event-level Constraint Compilation from PatternPlan to ActionGrid'
    prs.core_properties.subject = 'Figure 3: editable native shapes; Generation Case 01 evidence'
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    TEXT_RECORDS.clear()
    VISUAL_RECORDS.clear()
    # 导师标准：栏题合并为单行；副标题保留（VisualSpec/RealizationPlan/ActionGrid 三个IR名为图内唯一出现处）。
    headers = [('Geometric Scope', 'VisualSpec · binary masks'),
               ('Craft Assignment', 'RealizationPlan · θ'),
               ('Sparse Events', 'masked · all-or-nothing'),
               ('Layered Synthesis', 'ActionGrid · L0 to L3')]
    for i, (x, (bg, accent), (heading, sub)) in enumerate(zip(COLS, PALETTES, headers)):
        tag = 'abcd'[i]
        f1.shape(sl, f'column_{i+1}', x, .40, COL_W, 8.20, bg, kind='round')
        f1.shape(sl, f'panel_disc_{tag}', x + .12, .58, .44, .44, accent, kind='oval')
        text(sl, f'panel_label_{tag}', x + .12, .58, .44, .44, tag, 26, 'FFFFFF', align='center')
        # 审稿意见：字号再升一档对齐图1/2——栏题30pt独立居中行（窄栏容不下圆盘旁排），副题24pt。
        text(sl, f'heading_{i+1}', x + .10, 1.12, COL_W - .20, .54, heading, 30, bold=True, align='center')
        text(sl, f'subheading_{i+1}', x + .10, 1.70, COL_W - .20, .44, sub, 24, accent, align='center')
        f1.shape(sl, f'inner_panel_{i+1}', x + .10, 2.18, COL_W - .20, 6.32, 'FFFFFF', bg, 'round', 1)
    for i in range(3):
        f1.line(sl, f'stage_arrow_{i+1}', COLS[i] + COL_W + .04, 5.30, COLS[i + 1] - .04, 5.30, INK, 2.25, True)

    # ① 几何作用域：两张 40×40 掩膜画布（实心原生色块）。
    # 导师标准：三段裁剪说明与灰色注均与图注(a)逐字重复，删除；画布放大到内栏可容上限1.78in。
    # 目检修正：1.95in双画布会穿出内栏（4.15in>3.85in）；标签改三行居中图例（行色对应画布色）避免横排连读。
    cx, cy = COLS[0] + .14, 3.80
    scope_runs(sl, ring_img, cx, cy, CANVAS_SIZE, 'frame_scope')
    scope_runs(sl, disc_img, cx + CANVAS_SIZE + .22, cy, CANVAS_SIZE, 'disc_scope')
    f1.shape(sl, 'frame_scope_outline', cx - .005, cy - .005, CANVAS_SIZE + .01, CANVAS_SIZE + .01, None, '7C98AE', lw=1)
    f1.shape(sl, 'disc_scope_outline', cx + CANVAS_SIZE + .215, cy - .005, CANVAS_SIZE + .01, CANVAS_SIZE + .01, None, '7C98AE', lw=1)
    text(sl, 'scope_label_frame', COLS[0] + .20, 5.72, COL_W - .40, .34, 'frame_1 · 264 cells', 22, RING, align='center')
    text(sl, 'scope_label_circle', COLS[0] + .20, 6.10, COL_W - .40, .34, 'circle_1 · 529 cells', 22, '7C98AE', align='center')
    text(sl, 'scope_label_wave', COLS[0] + .20, 6.48, COL_W - .40, .34, 'wave_1 · 441 cells', 22, WAVE, align='center')

    # ② 工艺分配：取值全部来自 Generation Case 01 的真实 PatternPlan。
    # 导师标准：三行注复述芯片内容/图注，删除；审稿意见再升档：题24pt、副22pt（与图2正文同档）。
    # 22pt副行需265.7pt盒：边距.02in；“wave_tuck on wave_1 scope”省the后239pt安全。
    x2 = COLS[1] + .16
    cw = COL_W - .32
    chip(sl, 'assign_chip_base', x2, 3.14, cw, .98, 'EAF4ED', '487F55', 'Base fabric', 'f_knit · b_knit', m=.02)
    chip(sl, 'assign_chip_border', x2, 4.28, cw, .98, 'EAF4ED', '487F55', 'Border width', 't = 2 on the frame_1 ring', m=.02)
    chip(sl, 'assign_chip_event', x2, 5.42, cw, .98, 'EAF4ED', '487F55', 'Event pattern', 'wave_tuck on wave_1 scope', m=.02)
    chip(sl, 'assign_chip_theta', x2, 6.56, cw, .98, 'EAF4ED', '487F55', 'Parameters', 'period 5 × 4 · offsets 2, 3', m=.02)

    # ③ 稀疏事件规划：三类事件族＋本例真实事件统计（数据元素保留，22pt底线）。
    # 导师标准：两行“sparse”注与图注③逐字重复，删除；22pt副行边距收0保盒宽。
    x3 = COLS[2] + .16
    chip(sl, 'event_chip_wave', x3, 3.14, cw, 1.06, 'FDF3E7', 'B07A3C', 'Split-tuck wave groups', 'all-or-nothing placement', m=0)
    chip(sl, 'event_chip_transfer', x3, 4.34, cw, 1.06, 'FDF3E7', 'B07A3C', 'Transfer events', 'bed direction · loop support', m=0)
    chip(sl, 'event_chip_miss', x3, 5.54, cw, 1.06, 'FDF3E7', 'B07A3C', 'Miss events', 'hole density · isolation', m=0)
    text(sl, 'event_case_1', COLS[2] + .14, 6.94, COL_W - .28, .34, 'Generation Case 01:', 22, MUTED, align='center')
    text(sl, 'event_case_2', COLS[2] + .14, 7.30, COL_W - .28, .34, '16 groups, 32 events,', 22, MUTED, align='center')
    text(sl, 'event_case_3', COLS[2] + .14, 7.66, COL_W - .28, .34, '0 transfer events.', 22, MUTED, align='center')
    text(sl, 'event_case_4', COLS[2] + .14, 8.02, COL_W - .28, .34, 'Groups pair split + tuck.', 22, MUTED, align='center')

    # ④ 层次合成：L0–L3 写入层次（数据元素保留）＋动作字母表；审稿意见层芯片升24pt。
    # 导师标准：guard两行与determinism一行互相重复且图注④逐字承载，删除。
    x4 = COLS[3] + .16
    layers = [('L0 · base fabric', 'E8EDF2'), ('L1 · border ring', 'CBDAE8'),
              ('L2 · event groups', 'AEC4DA'), ('L3 · child clipping', '8FA9C6')]
    for i, (label, fill) in enumerate(layers):
        y = 3.49 + i * 1.00
        f1.shape(sl, f'layer_chip_{i}', x4, y, cw, .66, fill, '6E87A8', 'round')
        text(sl, f'layer_chip_{i}_text', x4 + .06, y + .13, cw - .12, .40, label, 24, '24364E', True, 'center')
        if i < 3:
            f1.line(sl, f'layer_arrow_{i}', x4 + cw / 2, y + .66, x4 + cw / 2, y + 1.00, '6E87A8', 1.6, True)
    text(sl, 'alphabet_line', COLS[3] + .14, 7.50, COL_W - .28, .42, '15 action classes, 40 × 40', 24, align='center')

    # 导师标准：底部横带删除（主线句复述栏题序列，图4衔接句图注末逐字承载）。
    notes = (
        'Figure 3 is rebuilt as native editable PowerPoint objects in the Figure 1 visual language; all figure text is English Times New Roman.\n'
        'Scope masks are rasterized from the Generation Case 01 (pp_01) run plan with the compiler lowering geometry '
        '(frame_1 ring 264 cells, circle_1 disc 529 cells, wave_1 scope 441 cells) and are cross-checked against the real pp_01 color map: '
        'green pixels equal frame_1 union (circle_1 minus wave_1).\n'
        'Event statistics are the real run values: 16 split-tuck groups, 32 events, 0 transfer events, 0 hard and 0 craft violations; '
        'all 32 event cells lie inside the wave_1 scope.\n'
        'Stage 2 parameter values (period 5 x 4, border thickness 2, column offsets 2 and 3) are read from pp_01_pattern_plan.json.\n'
        'The figure carries no axes, ticks, monospace code blocks or bitmap overlays; scope canvases are merged native rectangles.\n'
        '2026-09-29 advisor-style pass: the bottom lowering band and the restatement notes (scope clipping, theta contract, sparse anchors, write guard, determinism) were removed as verbatim duplicates of the caption; number discs relabeled a-d per Elsevier style and the caption cites (a)-(d); remaining body fonts raised to 20-28pt and the slide height trimmed to 9.00in.\n'
        '2026-09-29 reviewer pass: all body text raised one more step to a uniform 22pt floor (24pt chip and layer titles, 24pt subheadings, 30pt column headings) so the printed figure matches Figs. 1-2; headings moved to their own centered row because the narrow columns cannot fit 30pt beside the letter disc.')
    tf = sl.notes_slide.notes_text_frame
    tf.text = notes
    for para in tf.paragraphs:
        for run in para.runs:
            run.font.name, run.font.size = 'Times New Roman', Pt(18)
    prs.save(PPTX)
    normalize_package()
    write_json('text_manifest.json', TEXT_RECORDS)
    write_json('native_visual_manifest.json', VISUAL_RECORDS)
    return audit_deck()


def normalize_package():
    # 与图1/图2同一规范化：包内所有typeface统一Times New Roman并补latin/ea/cs。
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    target = BytesIO()
    with ZipFile(PPTX) as src, ZipFile(target, 'w') as dest:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename.endswith('.xml'):
                root = etree.fromstring(blob)
                for node in root.iter():
                    if 'typeface' in node.attrib:
                        node.set('typeface', 'Times New Roman')
                    if node.tag in {qn('a:rPr'), qn('a:defRPr'), qn('a:endParaRPr')}:
                        node.set('lang', 'en-US')
                        if 'sz' in node.attrib:
                            node.set('sz', str(max(2200, int(node.get('sz')))))
                        for tag in ('a:latin', 'a:ea', 'a:cs'):
                            face = node.find(qn(tag))
                            if face is None:
                                face = OxmlElement(tag)
                                node.append(face)
                            face.set('typeface', 'Times New Roman')
                blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            dest.writestr(item, blob)
    PPTX.write_bytes(target.getvalue())


def flatten(shapes):
    for s in shapes:
        yield s
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from flatten(s.shapes)


def reference_images():
    _, ring_img, disc_img = instance()
    return {'frame_scope': ring_img, 'disc_scope': disc_img}


def audit_deck():
    assert PPTX.exists()
    p = Presentation(PPTX)
    assert len(p.slides) == 1
    shapes = list(flatten(p.slides[0].shapes))
    refs = reference_images()
    covered = {}
    for prefix, image in refs.items():
        runs = [s for s in shapes if s.name.startswith(prefix + '_run_')]
        seen = set()
        for s in runs:
            r, start, end = (int(v) for v in s.name.rsplit('_', 3)[1:])
            rgb = image.getpixel((start, r))
            assert tuple(s.fill.fore_color.rgb) == rgb, s.name
            assert all(image.getpixel((c, r)) == rgb for c in range(start, end)), s.name
            assert start == 0 or image.getpixel((start - 1, r)) != rgb, s.name
            assert end == 40 or image.getpixel((end, r)) != rgb, s.name
            x, y, size = layout_of(prefix)
            for actual, want in zip((s.left, s.top, s.width, s.height),
                                    (x + start * size / 40, y + r * size / 40, (end - start) * size / 40, size / 40)):
                assert abs(actual - Inches(want)) <= 3, s.name
            cells = {(r, c) for c in range(start, end)}
            assert not seen & cells
            seen |= cells
        assert seen == {(r, c) for r in range(40) for c in range(40)}, prefix
        covered[prefix] = len(seen)
    pictures = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert not pictures, '图3不使用任何位图贴图'
    fonts, bounds = [], []
    for s in shapes:
        if s.left < -12700 or s.top < -12700 or s.left + s.width > p.slide_width + 12700 or s.top + s.height > p.slide_height + 12700:
            bounds.append(s.name)
        if s.has_text_frame and s.text.strip():
            for para in s.text_frame.paragraphs:
                for run in para.runs:
                    assert run.font.size is not None, s.name
                    fonts.append(run.font.size.pt)
    assert not bounds, bounds
    assert min(fonts) >= 22, min(fonts)
    ratio = 160 / (p.slide_width / 914400 * 25.4)
    assert min(fonts) * ratio >= 6
    with ZipFile(PPTX) as z:
        forbidden = [n for n in z.namelist() if n.lower().endswith('.svg')]
        assert not forbidden
        assert all(b'image/svg+xml' not in z.read(n) and b'<asvg:' not in z.read(n)
                   for n in z.namelist() if n.endswith(('.xml', '.rels')))
    from lxml import etree
    from pptx.oxml.ns import qn
    faces = set()
    with ZipFile(PPTX) as z:
        for name in z.namelist():
            if name.endswith('.xml'):
                for node in etree.fromstring(z.read(name)).iter():
                    if 'typeface' in node.attrib:
                        faces.add(node.get('typeface'))
                    if node.tag == qn('a:t'):
                        assert not re.search(r'[\u3400-\u9fff]', node.text or ''), name
    assert faces == {'Times New Roman'}, faces
    alltext = '\n'.join(s.text for s in shapes if s.has_text_frame)
    for phrase in ('VisualSpec', 'RealizationPlan', 'ActionGrid', 'wave_tuck', 'period 5 × 4',
                   '16 groups', 'L3 · child clipping', 'all-or-nothing', 'Miss events', 'Generation Case 01'):
        assert phrase in alltext, phrase
    result = {
        'passed': True, 'slide_count': 1, 'total_objects_recursive': len(shapes),
        'object_types': dict(Counter(str(s.shape_type) for s in shapes)),
        'scope_runs': sum(len([s for s in shapes if s.name.startswith(pfx + '_run_')]) for pfx in refs),
        'scope_cells_covered': covered, 'scope_pixel_exact': True, 'pictures': len(pictures),
        'text_objects': sum(s.has_text_frame and bool(s.text.strip()) for s in shapes),
        'min_font_pt': min(fonts), 'max_font_pt': max(fonts),
        'figure_width_mm': 160, 'min_reduced_font_pt': min(fonts) * ratio,
        'shape_out_of_slide': bounds, 'svg_files': forbidden,
        'package_typefaces': sorted(faces), 'pptx_sha256': sha(PPTX),
    }
    write_json('object_audit.json', result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def layout_of(prefix):
    cx, cy = COLS[0] + .14, 3.80
    return {'frame_scope': [cx, cy, CANVAS_SIZE],
            'disc_scope': [cx + CANVAS_SIZE + .22, cy, CANVAS_SIZE]}[prefix]


def render_deck(export=True):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app, pres = None, None
    try:
        if not export:
            prior = json.loads((OUT / 'export_audit.json').read_text(encoding='utf-8'))
            for path in (PPTX, PNG, PDF):
                assert sha(path) == prior['exports'][path.suffix]['sha256'], f'交付物与原导出记录不符: {path.name}'
        try:
            app = win32com.client.GetActiveObject('PowerPoint.Application')
        except Exception:
            app = win32com.client.Dispatch('PowerPoint.Application')
        for attempt in range(12):
            try:
                pres = app.Presentations.Open(str(PPTX), True, False, False)
                assert pres.Slides.Count == 1
                break
            except Exception:
                if attempt == 11:
                    raise
                time.sleep(1)
        slide = pres.Slides.Item(1)
        if export:
            slide.Export(str(PNG), 'PNG', 4000, round(4000 * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth))
            pres.SaveAs(str(PDF), 32)
        records, overflow = [], []

        def inspect(shapes):
            for i in range(1, shapes.Count + 1):
                s = shapes.Item(i)
                if s.Type == 6:
                    inspect(s.GroupItems)
                if s.HasTextFrame and s.TextFrame2.HasText:
                    tf, tr = s.TextFrame2, s.TextFrame2.TextRange
                    w, h = s.Width - tf.MarginLeft - tf.MarginRight, s.Height - tf.MarginTop - tf.MarginBottom
                    row = {'name': s.Name, 'text': tr.Text, 'available_pt': [w, h],
                           'bound_pt': [tr.BoundWidth, tr.BoundHeight],
                           'glyph_rect_pt': [tr.BoundLeft, tr.BoundTop, tr.BoundWidth, tr.BoundHeight],
                           'com_font': tr.Font.Name, 'com_font_size': tr.Font.Size}
                    assert tr.Font.Name == 'Times New Roman' and tr.Font.Size >= 22, row
                    records.append(row)
                    if tr.BoundWidth > w + 1.2 or tr.BoundHeight > h + 1.2:
                        overflow.append(row)
        # 僵尸/忙 COM 会拒绝呼叫（图1/图2同款）：导出后静置，幂等重取slide整体重试；真实字体违规立即上抛不掩盖。
        time.sleep(6)
        for walk in range(14):
            records.clear()
            overflow.clear()
            try:
                inspect(pres.Slides.Item(1).Shapes)
                break
            except AssertionError:
                raise
            except Exception:
                if walk == 13:
                    raise
                time.sleep(3)
        overlaps = []
        for i, a in enumerate(records):
            ax, ay, aw, ah = a['glyph_rect_pt']
            for b in records[i + 1:]:
                bx, by, bw, bh = b['glyph_rect_pt']
                ix = min(ax + aw, bx + bw) - max(ax, bx)
                iy = min(ay + ah, by + bh) - max(ay, by)
                if ix > 2 and iy > 1.2:
                    overlaps.append({'names': [a['name'], b['name']], 'intersection_pt': [ix, iy]})
        report = {'png': str(PNG), 'png_size': Image.open(PNG).size if PNG.exists() else None,
                  'pdf': str(PDF), 'pptx_sha256': sha(PPTX), 'actual_powerpoint_com': True,
                  'text_objects_checked': len(records), 'text_overflow': overflow,
                  'text_overlap': overlaps, 'overlap_tolerance_pt': [2, 1.2],
                  'text_measurements': records, 'export_performed_this_call': export,
                  'ppt_version': app.Version}
        write_json('render_audit.json', report)
        if overflow or overlaps:
            write_json('render_failure_' + str(time.time_ns()) + '.json', report)
        im = Image.open(PNG).convert('RGB')
        im.resize((1600, round(im.height * 1600 / im.width)), Image.Resampling.LANCZOS).save(OUT / 'overview.png')
        im.resize((1890, round(im.height * 1890 / im.width)), Image.Resampling.LANCZOS).save(OUT / 'print_160mm_300dpi.png', dpi=(300, 300))
        k = im.width / (pres.PageSetup.SlideWidth / 72)
        for i, key in enumerate(('scope', 'assign', 'events', 'synthesis')):
            im.crop((round(COLS[i] * k), round(.40 * k), round((COLS[i] + COL_W) * k), round(8.60 * k))).save(OUT / f'detail_{key}.png')
        print(json.dumps({'rendered': True, 'text_overflow_count': len(overflow), 'overflow': overflow,
                          'text_overlap_count': len(overlaps), 'overlaps': overlaps}, ensure_ascii=False, indent=2))
        assert not overflow, 'PowerPoint 实测存在文字溢出'
        assert not overlaps, 'PowerPoint 实测存在文字重叠'
        audit_exports()
    except Exception as exc:
        write_json('render_error_' + str(time.time_ns()) + '.json', {'error': repr(exc), 'pptx': str(PPTX)})
        raise
    finally:
        # 僵尸COM下Close也会拒绝呼叫，不能让它掩盖真实结果。
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def audit_exports():
    import pymupdf as fitz
    with fitz.open(PDF) as doc:
        assert len(doc) == 1
        page = doc[0]
        spans = [s for b in page.get_text('dict')['blocks'] if 'lines' in b
                 for line_ in b['lines'] for s in line_['spans'] if s['text'].strip()]
        # 删冗+升档后文本对象38个（与 object_audit 一致），下限钉36防静默丢字。
        assert len(spans) >= 36, len(spans)
        assert not page.get_images(full=True), '图3 PDF不应含位图'
        # PDF导出字号有-0.04pt舍入（图1的19pt实测18.96同款），阈值取标称-0.05。
        assert min(s['size'] for s in spans) >= 21.95
        assert Image.open(PNG).width == 4000
        pdf_fonts = sorted({s['font'] for s in spans})
        assert all('TimesNewRoman' in f.replace('-', '') for f in pdf_fonts), pdf_fonts
        assert not re.search(r'[\u3400-\u9fff]', page.get_text())
        assert all(f[1] != 'n/a' for f in page.get_fonts(full=True)), 'PDF字体未嵌入'
        image = Image.open(PNG).convert('RGB')
        refs = reference_images()
        scale = image.width / (page.rect.width / 72)
        checks, bad_total = {}, 0
        for prefix, source in refs.items():
            x, y, width = layout_of(prefix)
            pitch = width / 40
            bad = [(r, c) for r in range(40) for c in range(40)
                   if image.getpixel((round((x + (c + .5) * pitch) * scale), round((y + (r + .5) * pitch) * scale))) != source.getpixel((c, r))]
            checks[prefix] = {'cell_centers_checked': 1600, 'rgb_mismatches': len(bad)}
            bad_total += len(bad)
        result = {'pdf_pages': 1, 'pdf_text_spans': len(spans),
                  'pdf_min_font_pt': min(s['size'] for s in spans), 'pdf_fonts': pdf_fonts,
                  'pdf_image_objects': len(page.get_images(full=True)),
                  'pdf_vector_drawings': len(page.get_drawings()),
                  'scope_pixel_checks': checks, 'scope_rgb_mismatches': bad_total,
                  'png_pixels': Image.open(PNG).size,
                  'exports': {p.suffix: {'path': str(p), 'sha256': sha(p), 'bytes': p.stat().st_size}
                              for p in (PPTX, PNG, PDF)}}
        assert not bad_total, result
        write_json('export_audit.json', result)
        print(json.dumps(result, ensure_ascii=False, indent=2))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--verify', action='store_true')
    args = ap.parse_args()
    f1.FONT = 'Times New Roman'
    if args.verify:
        audit_deck()
        render_deck(export=False)
    else:
        build_deck()
        render_deck()


if __name__ == '__main__':
    main()
