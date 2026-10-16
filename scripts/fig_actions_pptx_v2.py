# -*- coding: utf-8 -*-
"""Rebuild Fig_Actions_Primitives_EN.pptx as a native-vector editable twin of
scripts/fig_actions_v3.py (slender-loop geometry). Logical canvas 1800x1020 px
mapped onto a 13.333 x 7.556 in slide; every loop / lattice node / arrow is a
real PowerPoint shape (oval, rounded rectangle, connector) so it stays editable.
"""
import os, math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

ROOT = r"e:\DeskTop\knitting_pattern_generator"
OUTP = os.path.join(ROOT, "docs", "figures", "Fig_Actions_Primitives_EN.pptx")

# ---------- canvas mapping (logical px -> inches) ----------
LW, LH = 1800.0, 1020.0
SLIDE_W_IN = 13.333
K = SLIDE_W_IN / LW                      # inches per logical px
SLIDE_H_IN = LH * K
PT_PER_PX = (SLIDE_W_IN * 72.0) / LW     # pt per logical px (font sizes)

def IN(v):  return Inches(v * K)
def PT(sz): return Pt(sz * PT_PER_PX)

PAGE = RGBColor(233, 242, 250); CARD = RGBColor(255, 255, 255)
BORDER = RGBColor(214, 228, 242); NAVY = RGBColor(31, 56, 100)
TEAL = RGBColor(52, 168, 166); MAG = RGBColor(230, 0, 126)
GRAY = RGBColor(150, 158, 168); WHITE = RGBColor(255, 255, 255)
DASH = RGBColor(50, 50, 50)

prs = Presentation()
prs.slide_width = IN(LW); prs.slide_height = IN(LH)
slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAGE
SH = slide.shapes

def _set_line(sp, color, w, dash=None, tail=False):
    ln = sp.line._get_or_add_ln()
    ln.set("w", str(Emu(IN(w)).emu if False else int(w * K * 914400)))
    for tag in ("a:prstDash", "a:tailEnd"):
        for e in ln.findall(qn(tag)):
            ln.remove(e)
    if dash:
        pd = ln.makeelement(qn("a:prstDash"), {"val": "dash"}); ln.append(pd)
    if tail:
        te = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(te)
    sp.line.color.rgb = color

def _nofill(sp):
    sp.fill.background()

def ellipse(box, fill=None, outline=None, width=1):
    x0, y0, x1, y1 = box
    sp = SH.add_shape(MSO_SHAPE.OVAL, IN(x0), IN(y0), IN(x1 - x0), IN(y1 - y0))
    if fill is None: _nofill(sp)
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if outline is None: sp.line.fill.background()
    else: _set_line(sp, outline, width)
    sp.shadow.inherit = False
    return sp

def rrect(box, radius, fill=None, outline=None, width=1):
    x0, y0, x1, y1 = box
    w = x1 - x0; h = y1 - y0
    sp = SH.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x0), IN(y0), IN(w), IN(h))
    try: sp.adjustments[0] = max(0.0, min(0.5, radius / min(w, h)))
    except Exception: pass
    if fill is None: _nofill(sp)
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if outline is None: sp.line.fill.background()
    else: _set_line(sp, outline, width)
    sp.shadow.inherit = False
    return sp

def line(x1, y1, x2, y2, color, w, dash=False, tail=False):
    cn = SH.add_connector(MSO_CONNECTOR.STRAIGHT, IN(x1), IN(y1), IN(x2), IN(y2))
    _set_line(cn, color, w, dash=dash, tail=tail)
    return cn

def arrow(x1, y1, x2, y2, color, w, hs):
    return line(x1, y1, x2, y2, color, w, tail=True)

def dashed_box(box, color, w):
    x0, y0, x1, y1 = box
    sp = SH.add_shape(MSO_SHAPE.RECTANGLE, IN(x0), IN(y0), IN(x1 - x0), IN(y1 - y0))
    _nofill(sp); _set_line(sp, color, w, dash=True)
    sp.shadow.inherit = False
    return sp

def text(x, y, s, sz, color, bold=False, ital=False, align=PP_ALIGN.LEFT, w_px=None):
    tb = SH.add_textbox(IN(x), IN(y), IN(w_px or 400), IN(sz * 2.2))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    f = r.font; f.name = "Times New Roman"; f.size = PT(sz)
    f.bold = bold; f.italic = ital; f.color.rgb = color
    return tb

# ---------- geometry ported from fig_actions_v3.py (SS = 1) ----------
def s(v): return v
EMU_PX = K * 914400
def arch(cx, ytop, w, h, color, t, dot=True):
    pts = [(cx - w * .68, ytop + h), (cx - w * .50, ytop + h * .80), (cx - w / 2, ytop + h * .52)]
    hc = ytop + w / 2
    for i in range(13):
        th = math.pi * (1 - i / 12.0)
        pts.append((cx + (w / 2) * math.cos(th), hc - (w / 2) * math.sin(th)))
    pts += [(cx + w / 2, ytop + h * .52), (cx + w * .50, ytop + h * .80), (cx + w * .68, ytop + h)]
    fb = SH.build_freeform(pts[0][0], pts[0][1], scale=EMU_PX)
    fb.add_line_segments(pts[1:], close=False)
    sp = fb.convert_to_shape()
    _nofill(sp); _set_line(sp, color, t)
    sp.shadow.inherit = False
    if dot:
        r1 = t * 1.5; r2 = t * .8
        ellipse([cx - r1, ytop - r1, cx + r1, ytop + r1], fill=color)
        ellipse([cx - r2, ytop - r2, cx + r2, ytop + r2], fill=WHITE)

def row_loops(cxc, ytop, n, u, color, wall, ov=.30):
    w = u * .48; step = w * (1 - ov); total = step * (n - 1) + w
    for i in range(n):
        arch(cxc - total / 2 + w / 2 + i * step, ytop, w, u, color, wall)

FAM = [("knit", "Knit", "f_knit / b_knit", "New loop is drawn through the loop below."),
       ("tuck", "Tuck", "f_tuck / b_tuck", "Old loop retained; a new yarn is tucked."),
       ("miss", "Miss", "miss", "Needle not selected; no new loop formed."),
       ("xfer", "Transfer", "xfer_l1 / xfer_f_b", "A loop moves to an adjacent needle or bed."),
       ("split", "Split", "f_split / b_split", "A loop is divided for support / transition."),
       ("carrier", "Drop-off", "carrier", "Loop support removed at this wale (bind-off).")]

SPEC = {"knit": {"mag": {(0, 2)}, "boxes": [(0, 2, 1, 2)]},
        "tuck": {"mag": {(0, 1), (0, 2)}, "boxes": [(0, 1, 0, 2), (2, 1, 2, 2)]},
        "miss": {"mag": {(1, 2)}, "boxes": [(0, 2, 0, 2), (2, 2, 2, 2)]},
        "xfer": {"mag": {(0, 1), (1, 3)}, "boxes": [(0, 0, 0, 1), (1, 2, 1, 3)], "diag": True},
        "split": {"mag": {(0, 2), (1, 2)}, "boxes": [(0, 1, 1, 2)]},
        "carrier": {"mag": {(2, 2)}, "boxes": [(2, 1, 2, 3)]}}

def lattice(x, y, w, h, spec, u):
    rows, cols = 3, 5
    xs = [x + w * (c + .5) / cols for c in range(cols)]
    ys = [y + h * (r + .5) / rows for r in range(rows)]
    nr = max(2, u * .10); lw = max(1, u * .030); alw = lw
    for yy in ys: line(xs[0], yy, xs[-1], yy, TEAL, lw)
    for xx in xs:
        line(xx, ys[0] - h * .09, xx, ys[-1] + h * .05, MAG, alw, tail=True)
    mag = spec.get("mag", set()); boxes = spec.get("boxes", [])
    for r in range(rows):
        for c in range(cols):
            cx, cy = xs[c], ys[r]
            if (r, c) in mag: ellipse([cx - nr, cy - nr, cx + nr, cy + nr], fill=MAG)
            elif r == 0: ellipse([cx - nr, cy - nr, cx + nr, cy + nr], fill=WHITE, outline=GRAY, width=lw)
            else: ellipse([cx - nr, cy - nr, cx + nr, cy + nr], fill=TEAL)
    cw_ = w / cols; chh = h / rows
    for (r0, c0, r1, c1) in boxes:
        dashed_box([x + c0 * cw_ + cw_ * .10, y + r0 * chh + chh * .10,
                    x + (c1 + 1) * cw_ - cw_ * .10, y + (r1 + 1) * chh - chh * .10], DASH, max(1, u * .040))
    if spec.get("diag"): arrow(xs[1], ys[0], xs[3], ys[1], MAG, alw, u * .12)

def panel_b(cxc, by, bw, bh, kind, u, wall):
    w = u * .48; ymid = by + bh / 2; mw = u * .085
    if kind == "knit":
        row_loops(cxc, ymid + u * .05, 3, u, TEAL, wall)
        row_loops(cxc, ymid - u * .80, 3, u, TEAL, wall)
        arrow(cxc, ymid - u * .70, cxc, ymid + u * .45, MAG, mw, u * .14)
    elif kind == "tuck":
        row_loops(cxc, ymid - u * .28, 3, u, TEAL, wall)
        arch(cxc, ymid - u * 1.02, w * .62, u * .72, MAG, wall, dot=False)
    elif kind == "miss":
        row_loops(cxc, ymid - u * .38, 3, u, TEAL, wall)
        line(cxc - w * .95, ymid + u * .05, cxc + w * .95, ymid + u * .05, MAG, mw * 1.2)
    elif kind == "xfer":
        row_loops(cxc - w * .60, ymid - u * .38, 1, u, TEAL, wall)
        row_loops(cxc + w * .95, ymid - u * .38, 1, u, TEAL, wall)
        arch(cxc - w * .05, ymid - u * .62, w * .8, u * .95, MAG, wall, dot=False)
        arrow(cxc - w * .05, ymid - u * .85, cxc + w * .75, ymid - u * .42, MAG, mw, u * .14)
    elif kind == "split":
        row_loops(cxc, ymid - u * .42, 3, u, TEAL, wall)
        line(cxc, ymid - u * .30, cxc + w * .62, ymid + u * .30, MAG, mw)
        line(cxc, ymid - u * .30, cxc - w * .62, ymid + u * .30, MAG, mw)
    elif kind == "carrier":
        row_loops(cxc - w * .95, ymid - u * .38, 1, u, TEAL, wall)
        row_loops(cxc + w * .95, ymid - u * .38, 1, u, TEAL, wall)
        line(cxc - w * .30, ymid - u * .80, cxc - w * .10, ymid + u * .20, MAG, mw)
        line(cxc + w * .30, ymid - u * .80, cxc + w * .10, ymid + u * .20, MAG, mw)
        arrow(cxc, ymid + u * .15, cxc, ymid + u * .62, MAG, mw, u * .14)

m = s(26); gap = s(22); top0 = s(18)
cw = (LW - 2 * m - 2 * gap) / 3; ch = (LH - top0 - m - gap) / 2
for idx, (key, title, abbr, cap) in enumerate(FAM):
    r, c = divmod(idx, 3); x = m + c * (cw + gap); y = top0 + r * (ch + gap)
    rrect([x, y, x + cw, y + ch], 16, fill=CARD, outline=BORDER, width=2)
    text(x + 26, y + 16, title, 33, NAVY, bold=True)
    text(x + 26 + len(title) * 19 + 12, y + 30, abbr, 16, GRAY, ital=True)
    by = y + s(80); bh = ch - s(80) - s(66)
    pad = s(12); aw = cw * .30; bw = cw * .34; cwid = cw * .30
    axc = x + pad + aw / 2; bxc = x + pad + aw + bw / 2; cxc = x + pad + aw + bw + cwid / 2
    u = aw * .52; wall = u * .075
    row_loops(axc, by + bh / 2 - u / 2, 4, u, TEAL, wall)
    text(axc - 7, by + bh * .82, "(a)", 15, NAVY)
    panel_b(bxc, by, bw, bh, key, u, wall)
    text(bxc - 7, by + bh * .82, "(b)", 15, NAVY)
    lb = cwid * .94
    lattice(cxc - lb / 2, by + bh / 2 - lb / 2, lb, lb, SPEC[key], u)
    text(cxc - 7, by + bh * .82, "(c)", 15, NAVY)
    text(x, y + ch - s(42), cap, 17, NAVY, align=PP_ALIGN.CENTER, w_px=cw)

prs.save(OUTP)
print("saved", OUTP, "shapes:", len(SH))
FAM = [("knit", "Knit", "f_knit / b_knit", "New loop is drawn through the loop below."),
       ("tuck", "Tuck", "f_tuck / b_tuck", "Old loop retained; a new yarn is tucked."),
       ("miss", "Miss", "miss", "Needle not selected; no new loop formed."),
       ("xfer", "Transfer", "xfer_l1 / xfer_f_b", "A loop moves to an adjacent needle or bed."),
       ("split", "Split", "f_split / b_split", "A loop is divided for support / transition."),
       ("carrier", "Drop-off", "carrier", "Loop support removed at this wale (bind-off).")]
SPEC = {"knit": {"mag": {(0, 2)}, "boxes": [(0, 2, 1, 2)]},
        "tuck": {"mag": {(0, 1), (0, 2)}, "boxes": [(0, 1, 0, 2), (2, 1, 2, 2)]},
        "miss": {"mag": {(1, 2)}, "boxes": [(0, 2, 0, 2), (2, 2, 2, 2)]},
        "xfer": {"mag": {(0, 1), (1, 3)}, "boxes": [(0, 0, 0, 1), (1, 2, 1, 3)], "diag": True},
        "split": {"mag": {(0, 2), (1, 2)}, "boxes": [(0, 1, 1, 2)]},
        "carrier": {"mag": {(2, 2)}, "boxes": [(2, 1, 2, 3)]}}
def lattice(x, y, w, h, spec, u):
    rows, cols = 3, 5
    xs = [x + w * (c + .5) / cols for c in range(cols)]
    ys = [y + h * (r + .5) / rows for r in range(rows)]
    nr = max(2, u * .10); lw = max(1, u * .030); alw = lw
    for yy in ys: line(xs[0], yy, xs[-1], yy, TEAL, lw)
    for xx in xs:
        line(xx, ys[0] - h * .09, xx, ys[-1] + h * .05, MAG, alw, tail=True)
    mag = spec.get("mag", set()); boxes = spec.get("boxes", [])
    for r in range(rows):
        for c in range(cols):
            cx, cy = xs[c], ys[r]
            if (r, c) in mag: ellipse([cx - nr, cy - nr, cx + nr, cy + nr], fill=MAG)
            elif r == 0: ellipse([cx - nr, cy - nr, cx + nr, cy + nr], fill=WHITE, outline=GRAY, width=lw)
            else: ellipse([cx - nr, cy - nr, cx + nr, cy + nr], fill=TEAL)
    cw_ = w / cols; chh = h / rows
    for (r0, c0, r1, c1) in boxes:
        dashed_box([x + c0 * cw_ + cw_ * .10, y + r0 * chh + chh * .10,
                    x + (c1 + 1) * cw_ - cw_ * .10, y + (r1 + 1) * chh - chh * .10], DASH, max(1, u * .040))
    if spec.get("diag"): arrow(xs[1], ys[0], xs[3], ys[1], MAG, alw, u * .12)
