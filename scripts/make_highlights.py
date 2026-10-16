# -*- coding: utf-8 -*-
"""独立 Highlights.docx（5 条，均 <=85 字符）。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
HL = [
    'Formulates flat-knitting action-graph generation as semantics-to-action compilation.',
    'PatternPlan IR captures geometry, texture and parent-child hierarchy for knitting.',
    'Sparse event planning plus constraint repair yields zero-violation action graphs.',
    '97.2% field-level parsing accuracy on 100 in-coverage natural-language inputs.',
    'Five designs knitted at a partner factory confirm the digital-to-physical pipeline.',
]
d = Document()
d.add_heading('Highlights', level=1)
for h in HL:
    assert len(h) <= 85, (len(h), h)
    d.add_paragraph(h, style='List Bullet')
p = r'e:\DeskTop\knitting_pattern_generator\docs\Highlights.docx'
d.save(p)
print('Highlights saved, chars:', [len(h) for h in HL])
