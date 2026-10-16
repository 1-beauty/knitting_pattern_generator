# -*- coding: utf-8 -*-
"""第四轮（F+G+H）：
F) 3.1–3.5 实现细节：para33 清理细节/para68 防御步骤全文/para58 规则溯源 移 Supp 新 Section I；
   para41/50/58/68/71/86 正文凝缩（全部 oMath=0）。
G) 4.1：para105/116 整段移 Supp 新 Section J；para117 前半（含 9 oMath）元素级迁移；
   para114 API 细节 span 删除并在 J 补记；para119 基线段在 oMath 后整尾重写（编号用旧值，交给统一重编号）。
H) 表 2（视觉语义→工艺实现映射，8x4）连题注移 Supp 成表 S20；para56 引用改泛指；
   全文（含表格单元格与 Supp 中正文表引用）表 3–7 → 表 2–6 一次性同步重编号。
纪律：para74（H1–H8 定义）、para120（基线覆盖辩护）、para106（数据充分性辩护）、4.5 工厂段不动。"""
import copy
import re
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MAIN = r'docs\论文to_cii.docx'
SUPP = r'docs\Supplementary_Material.docx'
W = qn('w:t')

# ---------- 工具 ----------
def split_runs(p, positions):
    cuts = sorted(set(positions))
    pos = 0
    for ch in list(p._p):
        if ch.tag != qn('w:r'):
            continue
        txt = ''.join(x.text or '' for x in ch.iter(W))
        s, e = pos, pos + len(txt)
        pos = e
        inner = [c for c in cuts if s < c < e]
        if not inner:
            continue
        frags, last = [], s
        for c in inner:
            frags.append(txt[last - s:c - s]); last = c
        frags.append(txt[last - s:])
        ch.find(W).text = frags[0]
        prev = ch
        for fr in frags[1:]:
            nr = copy.deepcopy(ch)
            nr.find(W).text = fr
            prev.addnext(nr)
            prev = nr

def segs(p, bounds):
    groups = [[] for _ in range(len(bounds) + 1)]
    pos, gi = 0, 0
    for ch in list(p._p):
        if ch.tag == qn('w:pPr'):
            continue
        txt = ''.join(x.text or '' for x in ch.iter(W))
        while gi < len(bounds) and pos >= bounds[gi]:
            gi += 1
        groups[gi].append(ch)
        pos += len(txt)
    return groups

def make_run(text, template_el=None):
    nr = copy.deepcopy(template_el) if template_el is not None else OxmlElement('w:r')
    tnode = nr.find(W)
    if tnode is None:
        tnode = OxmlElement('w:t')
        nr.append(tnode)
    tnode.set(qn('xml:space'), 'preserve')
    tnode.text = text
    return nr

def rewrite(p, new, keep_cites=True):
    assert not p._p.findall('.//' + qn('m:oMath')), '含内联公式禁止整段重建'
    if keep_cites:
        old = set(re.findall(r'\[\d+[^\]]*\]', p.text))
        miss = [c for c in old if c not in new]
        assert not miss, (p.text[:20], miss)
    rr = p.runs
    rr[0].text = new
    for r in rr[1:]:
        r._element.getparent().remove(r._element)

def strip_omath_and_rebuild(p, new):
    for om in p._p.findall('.//' + qn('m:oMath')):
        om.getparent().remove(om)
    rewrite(p, new, keep_cites=False)

doc = docx.Document(MAIN)
ps = doc.paragraphs
def find(prefix):
    return next(q for q in doc.paragraphs if q.text.startswith(prefix))

sup = docx.Document(SUPP)
head_style = next(q.style for q in sup.paragraphs if q.style.name == 'Heading 3')
body_style = next(q.style for q in sup.paragraphs if q.style.name == 'Body Text')
sect = sup.element.body.find(qn('w:sectPr'))

def sup_add(text):
    return sup.add_paragraph(text, style=body_style)

def sup_add_copy(src_p, fixes=(), drop_tail_from=None, drop_head_to=None):
    np = copy.deepcopy(src_p._p)
    sect.addprevious(np)
    if drop_tail_from is not None or drop_head_to is not None:
        pos = 0
        for ch in list(np):
            if ch.tag == qn('w:pPr'):
                continue
            txt = ''.join(x.text or '' for x in ch.iter(W))
            s = pos
            pos += len(txt)
            if drop_tail_from is not None and s >= drop_tail_from:
                np.remove(ch)
            elif drop_head_to is not None and s < drop_head_to:
                np.remove(ch)
    para = next(q for q in sup.paragraphs if q._p is np)
    para.style = body_style
    for r in para.runs:
        for old, new in fixes:
            if old in r.text:
                r.text = r.text.replace(old, new)
    return para

# ================= F：第 3 章 =================
p33, p41, p50 = find('本文使用的动作标签来源于'), find('从编译器视角看'), find('该分解并非实现偏好')
p58, p68, p71, p86 = find('规则库按（视觉语义，形状）二元组组织'), find('初始网格生成后'), find('条件化生成保证'), find('动作网格是面向机器与规则的离散表示')

# para68 原文整段迁 Supp I（先拷贝再改写正文）
sup.add_paragraph('I. 方法实现细节（Method Implementation Details：Data Cleaning, Defensive Steps, Rule Provenance）', style=head_style)
sup_add('样本级数据清理细节：原始数据集包含 17 类规则化针织指令；移针并左移 2 针（xfer_l2）与脱圈（drop）两类低频动作在原始数据中出现频率极低，本文删除包含任一目标动作的完整配对样本（同步删除对应的动作图和针织视觉预览图），不对保留样本中的动作进行替换，过滤后保留 4028 对配对样本。事后逐像素核验发现，4028 对中仍有 5 个平移增广变体残留极少量被删动作像素（每张不超过 9/400 格），已在预处理阶段经最近邻颜色量化映射到保留标签，不影响 15 类标签空间的定义（表 S1a）。')
sup_add_copy(p68, fixes=(('3.3.2 节', '正文 3.3.2 节'), ('4.3.2 节', '正文 4.3.2 节')))
sup_add('规则溯源与回退机制：规则库的 114 条映射项并非任意设定，而是依据针织工艺手册记载的组织结构—动作对应关系，并参考自动机器编织与针织纹理程序合成研究中的动作语义整理而成，每条规则均可溯源到明确的工艺原理或文献依据；规则库外的输入经形状别名归一化与优先级回退后仍无法匹配的，统一降级为素色底布。典型语义—工艺实现映射见表 S20，完整规模统计见表 S3a。')

rewrite(p33, '本文使用的动作标签来源于 Neural Inverse Knitting 数据集，原始数据集包含 17 类规则化针织指令；'
    '本文对两类低频动作（xfer_l2 与 drop）进行样本级清理，删除包含任一目标动作的完整配对样本，最终保留 4028 对配对样本，'
    '形成 15 类离散动作标签空间。模型训练、动作生成、约束检查和视觉预览均基于该过滤后的数据和固定标签定义完成'
    '（清理细则与逐像素核验见补充材料 Section I）。')
rewrite(p41, '从编译器视角看，本文生成链路对应“语义前端—制造中间表示—约束感知后端”三段结构：语义前端将开放式自然语言设计意图'
    '解析为显式结构化字段；PatternPlan 即制造中间表示（manufacturing IR），是前后端之间唯一的交换界面，显式携带区域拓扑、'
    '几何作用域、纹理效果、工艺参数与父子层级关系，使解析结果在进入动作生成之前即可被校验、编辑与复放；约束感知后端将 IR '
    '经事件级约束编译转换为离散动作网格（3.3–3.4 节）。与经典编译器以 IR 隔离源语言与目标机细节同理，PatternPlan 隔离了'
    '自然语言的模糊性与针迹级动作的强约束性，前端解析与后端编译因而可独立替换、独立评估——4.2 节与 4.3 节即分别对应'
    '前端保真度与后端约束满足的实验验证。')
rewrite(p50, '该分解并非实现偏好，而是对两类失败模式的直接回应：其一，端到端预测完整动作矩阵无法保证针迹级合法性——4.3.1 节中 '
    'LLM-Direct 平均每条输入产生 60.3 条硬约束违例，文本到图像（T2I）与颜色量化基线高达数百条；其二，仅有区域 mask 的密集填充'
    '并不自动产生区域结构——4.3.2 节的控变量对比显示，以区域内密集周期填充替换事件级编译后，纹理泄漏率由 0.240 升至 0.311、'
    '成组完整率由 1.000 降至 0.953。据此，本文将“动作编译”界定为将结构计划与纹理事件转换为规则化动作网格的受约束过程，'
    '而非一次性预测完整动作矩阵：三步降低中前两步分别固化显式几何作用域与成组区域归属，第三步与 3.4 节的约束修复在合成阶段'
    '保证合法性，分别阻断上述两类失败模式（定量验证见 4.3.1–4.3.2 节）。')
rewrite(p58, '规则库按（视觉语义，形状）二元组组织，当前包含 114 条显式映射项，覆盖 10 种视觉语义与 18 种形状取值，依据针织工艺'
    '手册 [12] 记载的组织结构—动作对应关系、并参考自动机器编织 [26] 与针织纹理程序合成 [35] 研究中的动作语义整理而成；'
    '规则库外的输入经形状别名归一化与优先级回退后仍无法匹配的，统一降级为素色底布，不生成规则外的特殊动作——语义覆盖范围'
    '因此是可枚举、可审计的，这也是 4.2 节解析测试集能够事先界定“规则库覆盖范围内输入”的前提（完整规模统计、别名回退机制'
    '与典型映射表见补充材料 Section I）。')
rewrite(p68, '初始网格生成后，还需 refinement 与 aesthetic pass 处理区域交界处的动作过渡与动作分布的视觉规则性。编译链路同时保留 '
    'layer clipping、compiler guard、规则修复与成组配对守卫（Wave Pair Guard）等防御性步骤，作为编译之后的兜底安全网：规则修复'
    '仅对支撑缺失、移针目标冲突、连续集圈堆叠和成组组织残缺做保守删除或降级，不新增动作；这层冗余与 3.3.2 节的事件级约束存在'
    '职能重叠，本文有意保留（各步骤职能细则见补充材料 Section I）。Wave Pair Guard 的实际触发范围在 4.3.2 节的受控消融中单独考察。')
rewrite(p71, '条件化生成保证。本文给出一个条件化的经验命题：对于位于规则库覆盖范围内的 PatternPlan，若（i）每个承载纹理的 motif 的'
    '有效作用域满足该纹理事件的最小放置尺寸，（ii）纹理事件周期与区域边界条件相容（成组事件的完整配对可放入作用域内），'
    '且（iii）所声明的底布与边界实现存在于规则库中，则编译流程输出满足 H1–H8 硬约束的动作网格；任一条件不满足时，系统执行'
    '显式降级（跳过无法完整放置的事件组或回退为稳定底布）并在统计报告中记录，而非静默输出非法动作。本文未给出该命题的完整'
    '形式化证明，而以第 4 章的批量运行作为经验验证：残留硬违例的案例（Stress Case 19 与 Field-Edit Case 04，见 4.3.3 节）均对应'
    '稀疏镂空组织下局部支撑条件难以在单点替换范围内满足的情形，其残余违例均被约束检查器显式检测并报告，而非被掩盖；'
    '其余输入的最终输出均满足硬约束零违例。命题的形式化证明与规则库外输入的降级行为刻画留作后续工作。')
rewrite(p86, '动作网格是面向机器与规则的离散表示，设计者难以直接从 15 类色码中检查区域边界与纹理连续性——这是引入视觉预览模块的'
    '直接动因：将内部动作图转换为针织外观预览图，使第一阶段的动作结构在视觉层面可被检视。该阶段被定义为条件图像生成任务，'
    '基于公开数据集的配对样本学习动作图到织物外观的映射。其角色有明确边界：一方面降低非专业设计者理解动作网格的门槛，'
    '辅助暴露边界断裂、纹理不连续等视觉风险；另一方面，它不显式模拟纱线材料、机台参数和织造过程，既不是线圈级物理仿真'
    '（如 Yarn-Level Simulation [41]），其输出也不能作为实际可织造性的证明。')

# ================= G：4.1 =================
p105, p114, p115 = find('本文第二阶段采用 patch 化和多阶段渐进式训练'), find('第一阶段语义解析由通用大语言模型接口完成'), find('第二阶段视觉预览模型采用基于 UNet 的条件流匹配框架')
p116, p117, p119 = find('patch 训练采用固定的空间尺寸'), find('全图级评估为控制推理开销'), find('本文设置三类诊断性外部基线')

# --- Supp J 节 ---
j_head = sup.add_paragraph('J. 训练与评估细节（Training and Evaluation Details：Patch、Resolution and Protocols）', style=head_style)
sup_add('语义解析通过 deepseek-chat 接口调用，解析阶段温度等采样参数保持默认设置。')
sup_add_copy(p105)
sup_add_copy(p116)
bnd117 = p117.text.index('这一分辨率适配导致')
sup_add_copy(p117, drop_tail_from=bnd117)

# para105：迁移后正文换凝缩段（含 3 oMath，先删再重建）
strip_omath_and_rebuild(p105, '第二阶段采用 patch 化与多阶段渐进式训练：每对动作图与对应针织视觉预览图按网格逐单元裁剪，'
    '每对原始样片产生 400 个训练 patch；Stage 1–3 与 Stage 4 采用不同的 patch 像素分辨率，四阶段渐进式训练中后续阶段在前一'
    '阶段权重基础上继续学习动作条件、邻域结构和动作属性（patch 规模、分辨率设置与样本数详见补充材料 Section J）。')
# para116：删除（全文已迁 Supp，模型选择句并入 para115）
p116._p.getparent().remove(p116._p)
p115.runs[-1].text = p115.runs[-1].text.rstrip() + '模型选择依据验证集上的全图级指标和结构感知指标，而非训练集损失单独决定。'
# para114：删除 API 细节 span
t114 = p114.text
s114, e114 = t114.index('，通过 deepseek-chat'), t114.index('），模型仅负责')
split_runs(p114, [s114, e114])
for e in segs(p114, [s114, e114])[1]:
    p114._p.remove(e)
# para117：删除前半，段首插指针
split_runs(p117, [bnd117])
for e in segs(p117, [bnd117])[0]:
    p117._p.remove(e)
ppr = p117._p.find(qn('w:pPr'))
lead = make_run('全图级与窗口级评估的下采样与分辨率适配细节见补充材料 Section J。', template_el=(p117.runs[0]._r if p117.runs else None))
if ppr is not None:
    ppr.addnext(lead)
else:
    p117._p.insert(0, lead)
# para119：oMath 之后整尾重写（表编号写旧值，交由 H 统一重编号）
om119 = p119._p.find('.//' + qn('m:oMath'))
assert om119 is not None
el = om119.getnext()
while el is not None:
    nxt = el.getnext()
    p119._p.remove(el)
    el = nxt
tail119 = (' 动作矩阵，检验直接从自然语言生成能否同时满足语义结构与针迹级离散约束；T2I+颜色量化先以通用文本生成图像模型生成视觉图像，'
    '再将 RGB 像素量化到 15 类动作标签空间，检验通用视觉生成结果能否转换为动作语义；Random Valid Map 随机采样满足局部合法性的动作图，'
    '用于区分“动作合法”与“语义一致”。除上述外部基线外，另设三个内部强基线隔离事件级编译的贡献：Plan-Grid 复用与完整系统相同的 '
    'PatternPlan 与区域实现，但跳过稀疏事件规划，直接在各区域掩码内按相同纹理周期密集写入动作，不维护成组配对、父级作用域交集与'
    '边框环排除，也不启用结构级守卫；Plan-Grid+Repair 在此基础上启用与完整系统相同的保守自修复；Plan-Grid+CSP 与 Plan-Grid 共享 '
    'PatternPlan 与密集周期目标，将 H1–H8 硬约束完整编码为约束满足问题（CSP，以 OR-Tools 的 CP-SAT 求解器 [40] 求解），每格动作域取'
    '“目标动作 ∪ 同语义族动作 ∪ 安全底布动作”，以与密集目标网格汉明距离最小为目标做全局最优投影。四者构成“无约束处理—局部贪心'
    '修复—全局约束投影—事件级编译”的递进对比链。内部基线在两种协议下评估：协议 A 在基线对比测试集（Baseline Case 01–Baseline '
    'Case 10）上各自独立解析 PatternPlan，便于与外部基线同表比较（表 4）；协议 B 在消融测试集（Ablation Case 01–Ablation Case 08）上'
    '强制复用同一份 PatternPlan，以排除语义解析方差、把差异限定在编译阶段（表 6）。此外，Event-NoScope（关闭 Layer Clipping）与 '
    'Event-NoGroup（关闭 Wave Pair Guard）两项针对性消融在 6 条含明确父子嵌套结构的输入上按协议 B 运行，结果与输入集见补充材料。')
tmpl = p119.runs[-1]._r if p119.runs else None
p119._p.append(make_run(tail119, template_el=tmpl))

# ================= H：表 2 迁移与重编号 =================
p56 = find('式（3）中，')
t56 = p56.text
s56, e56 = t56.index('表 2 给出典型'), t56.index('该映射有意保')
split_runs(p56, [s56, e56])
g56 = segs(p56, [s56, e56])
for e in g56[1]:
    p56._p.remove(e)
g56[0][-1].addnext(make_run('典型的语义—工艺实现映射见补充材料。', template_el=(g56[0][-1] if g56[0][-1].tag == qn('w:r') else None)))

cap2 = next(q for q in doc.paragraphs if q.text.strip().startswith('Table 2 Mapping from visual semantics'))
tbl2 = cap2._p.getnext()
while tbl2 is not None and tbl2.tag != qn('w:tbl'):
    tbl2 = tbl2.getnext()
assert tbl2 is not None and tbl2.tag == qn('w:tbl')
cap_cp, tbl_cp = copy.deepcopy(cap2._p), copy.deepcopy(tbl2)
j_head._p.addprevious(cap_cp)
j_head._p.addprevious(tbl_cp)
cap_para = next(q for q in sup.paragraphs if q._p is cap_cp)
for r in cap_para.runs:
    if 'Table 2' in r.text:
        r.text = r.text.replace('Table 2', 'Table S20')
cap2._p.getparent().remove(cap2._p)
tbl2.getparent().remove(tbl2)

def renumber_paras(paras):
    n = 0
    for p in paras:
        if not re.search(r'表\s*[3-7]', p.text):
            continue
        ms = list(re.finditer(r'表\s*[3-7]', p.text))
        bounds = []
        for m in ms:
            bounds += [m.start(), m.end()]
        if p._p.findall('.//' + qn('m:oMath')):
            split_runs(p, bounds)
        else:
            split_runs(p, bounds)
        for r in p.runs:
            if re.search(r'表\s*[3-7]', r.text):
                r.text = re.sub(r'表\s*([3-7])', lambda m: '表 ' + str(int(m.group(1)) - 1), r.text)
        n += 1
    return n

def cell_paras(d):
    out = []
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                out.extend(c.paragraphs)
    return out

# Supp 中对正文表的裸引用同步重编号（Supp 自身表均为 S 编号，不受影响）
sup_hits = [q.text[:60] for q in sup.paragraphs if re.search(r'表\s*[2-7](?!\d)', q.text) and '表 S' not in q.text[:6]]
print('supp bare 表N refs before renumber:', sup_hits)
n_main = renumber_paras(doc.paragraphs)
n_cell = renumber_paras(cell_paras(doc))
n_sup = renumber_paras(sup.paragraphs) + renumber_paras(cell_paras(sup))
print('renumbered paras: main', n_main, '| cells', n_cell, '| supp', n_sup)

sup.save(SUPP)
doc.save(MAIN)

# ================= 校验 =================
chk = docx.Document(MAIN)
cts = [p.text for p in chk.paragraphs]
ref = next(i for i, t in enumerate(cts) if t.strip() == 'References')
body = '\n'.join(cts[:ref])
tot = sum(len(re.findall(r'[\u4e00-\u9fff]', t)) for t in cts[:ref])
for probe in ['补充材料 Section I', '补充材料 Section J', '典型的语义—工艺实现映射见补充材料',
              '（清理细则与逐像素核验见补充材料 Section I）', '[12]', '[26]', '[35]', '[40]', '[41]', '[47]']:
    assert probe in body, probe
assert not re.search(r'Table 2 Mapping', body)
assert re.search(r'表\s*7', body) is None, '仍存在表 7 引用'
s2 = docx.Document(SUPP)
st = [p.text for p in s2.paragraphs]
assert any(t.startswith('I. 方法实现细节') for t in st)
assert any(t.startswith('J. 训练与评估细节') for t in st)
assert any('Table S20' in t for t in st)
print('main CN chars now:', tot, '| supp paragraphs:', len(s2.paragraphs), '| supp tables:', len(s2.tables))
