# -*- coding: utf-8 -*-
"""公式手动换行：在目标 oMath 顶层子元素前插入 <m:r><m:brk/></m:r>，使续行在逻辑运算符处干净起始。
(1) before child7(→EventPoint)；(3) before child3(→base)；(6) before child12(∀G)+child30(Vh)；(11) before child6(Leak)。
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
M = qn('m:oMath')
BREAKS = {0: [7], 3: [3], 7: [12, 30], 12: [6]}

doc = Document(MAIN)
for ti, idxs in BREAKS.items():
    t = doc.tables[ti]
    om = t.cell(0, 1)._tc.find('.//' + M)
    children = list(om)
    for i in idxs:
        target = children[i]
        brk_run = OxmlElement('m:r')
        brk = OxmlElement('m:brk')
        brk_run.append(brk)
        target.addprevious(brk_run)
    num = t.cell(0, 2).text.strip()
    print(f'table[{ti}] {num}: inserted {len(idxs)} brk')
doc.save(MAIN)
print('saved')
