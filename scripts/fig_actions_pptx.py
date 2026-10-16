import os, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
ROOT=r"e:\DeskTop\knitting_pattern_generator"
OUT=os.path.join(ROOT,"outputs","fig_actions_proto"); os.makedirs(OUT,exist_ok=True)
STEM=os.path.join(ROOT,"docs","figures","Fig_Actions_Primitives_EN")
os.makedirs(os.path.dirname(STEM),exist_ok=True)
PXW,PXH=1680,1240; SW=13.333; SH=SW*PXH/PXW; k=SW/PXW
def IN(px): return Inches(px*k)
INK="172D42"; MUT="52677A"; EDGE="CBD5E1"; GBG="E9EDF2"; BG="FFFFFF"
C_KNIT="FF0010"; C_TUCK="00A4FF"; C_MISS="B3B3B3"; MISSED="69707A"; C_XFER="7E0095"; C_SPLIT="D6B400"; C_CAR="753B3B"
FAM=[("knit","Knit","f_knit / b_knit",C_KNIT,"New loop is drawn through the loop below."),
     ("tuck","Tuck","f_tuck / b_tuck",C_TUCK,"Old loop retained; a new yarn is tucked."),
     ("miss","Miss","miss",C_MISS,"Needle not selected; no new loop formed."),
     ("xfer","Transfer","xfer_l1 / xfer_f_b",C_XFER,"A loop moves to an adjacent needle or bed."),
     ("split","Split","f_split / b_split",C_SPLIT,"A loop is divided for support / transition."),
     ("carrier","Drop-off","carrier",C_CAR,"Loop support removed at this wale (bind-off).")]
def cr(pts,seg=10):
    P=list(pts); out=[]
    for i in range(len(P)-1):
        p0=P[max(i-1,0)];p1=P[i];p2=P[i+1];p3=P[min(i+2,len(P)-1)]
        for j in range(seg):
            t=j/seg;t2=t*t;t3=t2*t
            x=0.5*(2*p1[0]+(-p0[0]+p2[0])*t+(2*p0[0]-5*p1[0]+4*p2[0]-p3[0])*t2+(-p0[0]+3*p1[0]-3*p2[0]+p3[0])*t3)
            y=0.5*(2*p1[1]+(-p0[1]+p2[1])*t+(2*p0[1]-5*p1[1]+4*p2[1]-p3[1])*t2+(-p0[1]+3*p1[1]-3*p2[1]+p3[1])*t3)
            out.append((x,y))
    out.append(P[-1]);return out
def freeform(sl,pts_px,color,wpt,closed=False):
    pts=[(IN(x),IN(y)) for x,y in pts_px]
    fb=sl.shapes.build_freeform(pts[0][0],pts[0][1],scale=1)
    fb.add_line_segments(pts[1:],close=closed)
    sh=fb.convert_to_shape(); sh.fill.background()
    sh.line.color.rgb=RGBColor.from_string(color); sh.line.width=Pt(wpt)
    return sh
def seg(sl,x1,y1,x2,y2,color,wpt): return freeform(sl,[(x1,y1),(x2,y2)],color,wpt)
def arrow(sl,x1,y1,x2,y2,color,wpt=2.2,hs=14):
    seg(sl,x1,y1,x2,y2,color,wpt); a=math.atan2(y2-y1,x2-x1)
    for s in (a+math.radians(150),a-math.radians(150)):
        seg(sl,x2,y2,x2+hs*math.cos(s),y2+hs*math.sin(s),color,wpt)
def loop_pts(cx,ytop,depth,hw):
    return [(cx-hw,ytop-depth*0.28),(cx-hw,ytop+depth*0.7),(cx-hw*0.5,ytop+depth),
            (cx+hw*0.5,ytop+depth),(cx+hw,ytop+depth*0.7),(cx+hw,ytop-depth*0.28)]
def rect(sl,x,y,w,h,fill,line=None,lw=0.75,shape=MSO_SHAPE.RECTANGLE):
    s=sl.shapes.add_shape(shape,IN(x),IN(y),IN(w),IN(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb=RGBColor.from_string(fill)
    else: s.fill.background()
    if line: s.line.color.rgb=RGBColor.from_string(line); s.line.width=Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit=False
    return s
def text(sl,x,y,w,h,val,size,color=INK,bold=False,italic=False,align=PP_ALIGN.LEFT):
    tb=sl.shapes.add_textbox(IN(x),IN(y),IN(w),IN(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=val; f=r.font; f.name="Times New Roman"; f.size=Pt(size)
    f.bold=bold; f.italic=italic; f.color.rgb=RGBColor.from_string(color)
    return tb
def panel_loops(sl,lx,ly,lw,lh,kind):
    cols=5;rows=3;cellw=lw/cols;depth=lh*0.30;rowgap=lh*0.255;wb=2.0;wh=4.2
    ytop0=ly+lh*0.20
    for r in range(rows):
        ytop=ytop0+r*rowgap
        for c in range(cols):
            if kind=="miss" and r==0 and c==2: continue
            freeform(sl,cr(loop_pts(lx+(c+0.5)*cellw,ytop,depth,cellw*0.40)),INK,wb)
        for c in range(cols-1):
            xa=lx+(c+0.5)*cellw+cellw*0.40; xb=lx+(c+1.5)*cellw-cellw*0.40
            seg(sl,xa,ytop,xb,ytop,INK,1.4)
    cx=lx+2.5*cellw;ytop=ytop0;hw=cellw*0.40
    if kind=="knit": freeform(sl,cr(loop_pts(cx,ytop,depth,hw)),C_KNIT,wh)
    elif kind=="tuck":
        freeform(sl,cr(loop_pts(cx,ytop,depth,hw)),C_TUCK,wh)
        freeform(sl,cr(loop_pts(cx,ytop-depth*0.22,depth*0.60,hw*0.64)),C_TUCK,wh-0.6)
    elif kind=="miss":
        seg(sl,lx+2.0*cellw,ytop,lx+3.0*cellw,ytop,MISSED,6.0)
        seg(sl,cx,ytop+2,cx,ytop+depth-2,MISSED,1.6)
    elif kind=="xfer":
        pts=[(cx-hw,ytop-depth*0.2),(cx-hw,ytop+depth*0.5),(cx-hw*0.1,ytop+depth),(cx+hw*1.5,ytop+depth),(cx+hw*1.9,ytop+depth*0.5),(cx+hw*1.9,ytop-depth*0.2)]
        freeform(sl,cr(pts),C_XFER,wh); arrow(sl,cx+hw*0.2,ytop-depth*0.5,cx+hw*1.9,ytop-depth*0.5,C_XFER,2.4,16)
    elif kind=="split":
        freeform(sl,cr(loop_pts(cx,ytop,depth,hw)),C_SPLIT,wh)
        freeform(sl,cr([(cx,ytop+depth),(cx+hw*0.9,ytop+depth*1.05),(cx+hw*1.6,ytop+depth*0.7)]),C_SPLIT,wh-0.6)
    elif kind=="carrier":
        freeform(sl,cr([(cx-hw,ytop-depth*0.2),(cx-hw,ytop+depth*0.6),(cx-hw*0.3,ytop+depth*0.95)]),C_CAR,wh)
        freeform(sl,cr([(cx+hw,ytop-depth*0.2),(cx+hw,ytop+depth*0.6),(cx+hw*0.3,ytop+depth*0.95)]),C_CAR,wh)
        arrow(sl,cx,ytop+depth*0.7,cx,ytop+depth*1.35,C_CAR,2.4,14)
def panel_grid(sl,gx,gy,gs,kind):
    col=dict(knit=C_KNIT,tuck=C_TUCK,miss=C_MISS,xfer=C_XFER,split=C_SPLIT,carrier=C_CAR)[kind]
    n=5;cs=gs/n
    motif=dict(knit=[(r,2) for r in range(5)],tuck=[(1,1),(2,1),(1,3),(2,3)],
               miss=[(1,1),(1,2),(1,3),(2,1),(2,3),(3,1),(3,2),(3,3)],
               xfer=[(1,1),(2,2),(3,3)],split=[(1,2),(2,2),(3,2),(1,1)],carrier=[(4,c) for c in range(5)])
    for r in range(n):
        for c in range(n): rect(sl,gx+c*cs,gy+r*cs,cs,cs,GBG,BG,0.75)
    for (r,c) in motif[kind]: rect(sl,gx+c*cs,gy+r*cs,cs,cs,col,BG,0.75)
    if kind=="miss": rect(sl,gx+2*cs+cs*0.28,gy+2*cs+cs*0.28,cs*0.44,cs*0.44,BG,C_MISS,1.0)
    if kind=="xfer": arrow(sl,gx+1.5*cs,gy+1.5*cs,gx+3.5*cs,gy+3.5*cs,C_XFER,2.2,10)
    rect(sl,gx,gy,gs,gs,None,EDGE,1.0)
def build():
    prs=Presentation(); prs.slide_width=Inches(SW); prs.slide_height=Inches(SH)
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    rect(sl,0,0,PXW,PXH,BG,None)
    text(sl,26,8,1200,40,"Action primitives of the 15-class flat-knitting label space",26,INK,True)
    m=24;gap=20;top0=64;cw=(PXW-2*m-2*gap)/3;ch=(PXH-top0-m-gap)/2
    for idx,(key,title,abbr,col,cap) in enumerate(FAM):
        r,c=divmod(idx,3);x=m+c*(cw+gap);y=top0+r*(ch+gap)
        rect(sl,x,y,cw,ch,BG,EDGE,1.2,MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl,x+16,y+12,cw-90,42,title,30,INK,True)
        text(sl,x+18,y+52,cw-40,26,abbr,18,MUT,False,True)
        rect(sl,x+cw-58,y+18,42,24,col,EDGE,0.75,MSO_SHAPE.ROUNDED_RECTANGLE)
        cy0=y+70;cy1=y+ch-46;lw=cw*0.57
        panel_loops(sl,x+16,cy0,lw,cy1-cy0,key)
        gs=cw*0.33;panel_grid(sl,x+cw-14-gs,cy0+((cy1-cy0)-gs)/2,gs,key)
        text(sl,x+16,y+ch-40,cw-32,40,cap,18,MUT)
    p=STEM+".pptx"; prs.save(p); print("saved",p)
    return p
build()
