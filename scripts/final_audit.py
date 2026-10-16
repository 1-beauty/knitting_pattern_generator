# -*- coding: utf-8 -*-
"""Final comprehensive audit before sending to advisor."""
import sys, re, os
sys.stdout.reconfigure(encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

SRC = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'

doc = Document(SRC)
body = doc.element.body

# ============================================================
print("=" * 60)
print("1. 表格完整性检查")
print("=" * 60)
content_tables = []
for ti, tbl in enumerate(doc.tables):
    nr = len(tbl.rows); nc = len(tbl.columns)
    if nr <= 1 and nc <= 3:
        continue  # skip formula/layout tables
    hdr = [tbl.cell(0, c).text.strip()[:25] for c in range(min(nc, 5))]
    content_tables.append((ti, nr, nc, hdr))
    print(f"  tbl[{ti:02d}] {nr}r x {nc}c | {hdr}")

# Check table numbering in captions
print("\n  题注检查:")
for i, p in enumerate(doc.paragraphs):
    if p.style and 'Caption' in (p.style.name or ''):
        txt = p.text.strip()
        # check for duplicate prefix
        dup = re.search(r'^(Table|Figure)\s+\d+\s+\1', txt)
        if dup:
            print(f"    ⚠️ p{i:03d} DUPLICATE PREFIX: {txt[:60]}")
        else:
            print(f"    ✓ p{i:03d}: {txt[:80]}")

# ============================================================
print("\n" + "=" * 60)
print("2. 图片引用与文件检查")
print("=" * 60)
# Find figure captions
fig_captions = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^(?:图|Fig\.?)\s*(\d+)', txt)
    if m:
        fig_captions.append((i, int(m.group(1)), txt[:80]))
    # Also check English "Fig. N." in Table Caption style
    if p.style and 'Caption' in (p.style.name or ''):
        m2 = re.match(r'^Fig\.\s*(\d+)', txt)
        if m2:
            fig_captions.append((i, int(m2.group(1)), txt[:80]))

fig_captions.sort(key=lambda x: x[1])
print(f"  Found {len(fig_captions)} figure captions:")
for idx, num, txt in fig_captions:
    print(f"    Fig {num} (p{idx:03d}): {txt[:70]}")

# Check for embedded images
img_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        img_count += 1
print(f"\n  Embedded images in docx: {img_count}")

# ============================================================
print("\n" + "=" * 60)
print("3. 摘要词数 + 关键词")
print("=" * 60)
p002 = doc.paragraphs[2]  # Abstract
words = [w for w in p002.text.split() if re.search(r'[a-zA-Z0-9]', w)]
print(f"  Abstract word count: {len(words)} (limit: 250)")
if len(words) > 250:
    print("  ⚠️ EXCEEDS 250!")
else:
    print("  ✓ Within limit")

# Keywords
p003 = doc.paragraphs[3]
print(f"  Keywords: {p003.text[:120]}")

# ============================================================
print("\n" + "=" * 60)
print("4. 引用闭合 (via verify script)")
print("=" * 60)
import subprocess
result = subprocess.run(
    [sys.executable, '-X', 'utf8', 'scripts/_r8_verify.py'],
    capture_output=True, text=True, encoding='utf-8'
)
for line in result.stdout.splitlines():
    if any(k in line for k in ['PASS', 'FAIL', '孤儿', '缺失', 'range']):
        print(f"  {line.strip()}")

# ============================================================
print("\n" + "=" * 60)
print("5. 补充材料 + 过期文件")
print("=" * 60)
# List docs/ contents
docs_dir = r'e:\DeskTop\knitting_pattern_generator\docs'
for f in sorted(os.listdir(docs_dir)):
    if f.endswith('.docx') or f.endswith('.png'):
        fp = os.path.join(docs_dir, f)
        sz = os.path.getsize(fp) // 1024
        print(f"  {f:50s} {sz:>5} KB")

# ============================================================
print("\n" + "=" * 60)
print("6. 正文中遗留占位符/方括号检查")
print("=" * 60)
placeholders = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text
    if '[Author' in txt or '[Replace' in txt or '[grant' in txt.lower():
        placeholders.append((i, txt[:80]))
    if '<DOI-to-be-assigned>' in txt:
        placeholders.append((i, txt[:80]))
if placeholders:
    for idx, t in placeholders:
        print(f"  ⚠️ p{idx:03d}: {t}")
else:
    print("  ✓ No remaining bracket placeholders")

print("\n" + "=" * 60)
print("AUDIT COMPLETE")
print("=" * 60)
