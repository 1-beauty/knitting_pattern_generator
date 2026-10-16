# -*- coding: utf-8 -*-
"""第三层精简 · 第二轮（结构性移动）：
A) 3.6 GUI 块：para98 截图溯源 + 远程 GPU 工程细节 -> 补充材料 Section F；
   正文 para92/93 合并凝缩为指针段（Fig 6/7 与"工程封装"免责声明保留）。
B) para64 贪心策略 (i)(ii)(iii) 细节（含 6 个内联 oMath）-> 补充材料新 Section G（deepcopy 整段）；
   正文 para64 替换为凝缩指针段并删除内联公式元素。
纪律：工厂口径不动；Fig/表/节交叉引用闭合；补充材料 F 内自指句删除。"""
import copy
import re
import docx
from docx.oxml.ns import qn

MAIN = r'docs\论文to_cii.docx'
SUPP = r'docs\Supplementary_Material.docx'

COND_92 = ('为便于设计者使用，本文将完整管线封装为桌面原型系统（PySide6）：工作台按“语义解析—PatternPlan—色码图—视觉预览”'
    '四阶段展示流水线进度，生成完成后并排呈现动作色码图与视觉预览图以便对照检查（图 6）；设计者可在任一历史版本上发起增量编辑，'
    '系统以该版本的 PatternPlan 为基础对目标字段做增量修改后重新编译，新版本与其基础版本以链式关系记录，支持设计演化路径的追踪与回滚'
    '（图 7）。需要说明的是，该原型仅是本文方法的工程封装，不构成独立的方法贡献；界面细节、截图溯源与案例映射见补充材料 Section F。')

COND_64 = ('本文不通过全局优化求解该问题，而是采用“生成—作用域裁剪—成组完整性检查—保守修复”的确定性单遍贪心编译策略：'
    '区域按版面角色的固定优先级稳定排序，事件按行优先周期顺序发射，写入依次经过画布越界、区域归属与边框轮廓保护三重过滤，'
    '降级次序固定为“整组取消 → 回退底布动作”。该策略的确定性、终止性与时间复杂度分析见补充材料 Section G。贪心策略的固有代价是'
    '先写入的低风险结构可能挤占后续事件的可行位置，4.3.2 节的事件保留率实验对此代价给出了量化；其与像素级后处理的本质区别在于：'
    '约束在事件与事件组的粒度上执行，删除与降级以组为单位进行。')

SUPP_F1 = ('正文图 7 的 GUI 截图均为未经修改的原始采集，保留原始运行 ID、时间戳与文件名，界面字段与英文对照键的对应关系如下。'
    '历史面板中，字段编辑（native-edit）条目依次对应 Field-Edit Case 00（基线）、Field-Edit Case 01（蜂窝纹理）、'
    'Field-Edit Case 02（蕾丝纹理）、Field-Edit Case 03（加宽边框）、Field-Edit Case 04（稀疏网眼背景）与 '
    'Field-Edit Case 05（菱形形状）；嵌套条目对应 Nested Case 04（仅见于 GUI 历史，不属于三案例嵌套测试集）。'
    '以时间戳命名的 GUI 会话仅为界面历史记录，不构成额外编号的实验案例。')
SUPP_F2 = ('考虑到视觉预览模块的推理开销，原型系统提供仅生成动作色码图的轻量模式，并支持将预览推理切换至远程 GPU 执行、'
    '产物自动回传。')
G_HEAD = 'G. 贪心编译策略：确定性、终止性与复杂度分析（Greedy Compilation Strategy: Determinism, Termination and Complexity）'

doc = docx.Document(MAIN)
ps = doc.paragraphs
p64, p92, p93, p98 = ps[64], ps[92], ps[93], ps[98]
assert p98.text.startswith('图 7 的 GUI 截图'), p98.text[:20]
assert p64.text.startswith('式（6）中，'), p64.text[:10]
assert not p92._p.findall('.//' + qn('m:oMath')) and not p98._p.findall('.//' + qn('m:oMath'))
n_om = len(p64._p.findall('.//' + qn('m:oMath')))
assert n_om == 6, n_om

sup = docx.Document(SUPP)
head_style = next(p.style for p in sup.paragraphs if p.style.name == 'Heading 3')
body_style = next(p.style for p in sup.paragraphs if p.style.name == 'Body Text')

# --- 补充材料：F 追加两段 ---
for txt in (SUPP_F1, SUPP_F2):
    np = sup.add_paragraph(txt, style=body_style)
# --- 补充材料：新 Section G = 标题 + deepcopy(para64) ---
sup.add_paragraph(G_HEAD, style=head_style)
moved64 = copy.deepcopy(p64._p)
sup.element.body.find(qn('w:sectPr')).addprevious(moved64)
# 移动段样式对齐补充材料正文
mp = sup.paragraphs[-1]
for p in sup.paragraphs:
    if p._p is moved64:
        mp = p
        break
mp.style = body_style
# 跨文档引用措辞：移动段内的正文节号/公式号加“正文”前缀
for r in mp.runs:
    for old, new in (('式（6）中，', '正文式（6）中，'), ('见 3.4 节', '见正文 3.4 节'), ('4.3.2 节的事件保留率', '正文 4.3.2 节的事件保留率')):
        if old in r.text:
            r.text = r.text.replace(old, new)
sup.save(SUPP)
print('supp: F += 2 paras, G created with', n_om, 'inline oMath')

# --- 正文：para92 凝缩 ---
runs = p92.runs
runs[0].text = COND_92
for r in runs[1:]:
    r._element.getparent().remove(r._element)
# --- 正文：删除 para93 / para98 ---
p93._p.getparent().remove(p93._p)
p98._p.getparent().remove(p98._p)
# --- 正文：para64 替换（删内联 oMath + 凝缩） ---
for om in p64._p.findall('.//' + qn('m:oMath')):
    om.getparent().remove(om)
runs = p64.runs
runs[0].text = COND_64
for r in runs[1:]:
    r._element.getparent().remove(r._element)
doc.save(MAIN)

# --- 校验 ---
chk = docx.Document(MAIN)
ts = [p.text for p in chk.paragraphs]
ref = next(i for i, t in enumerate(ts) if t.strip() == 'References')
body = '\n'.join(ts[:ref])
tot = sum(len(re.findall(r'[\u4e00-\u9fff]', t)) for t in ts[:ref])
for probe in ['图 6', '图 7', '补充材料 Section F', '补充材料 Section G', '4.3.2 节的事件保留率']:
    assert probe in body, probe
print('main CN chars now:', tot)
s2 = docx.Document(SUPP)
st = [p.text for p in s2.paragraphs]
assert any(t.startswith('G. 贪心编译策略') for t in st)
assert any('式（6）中，' in t for t in st)
print('supp paragraphs:', len(s2.paragraphs), '| G present, moved para64 present')
