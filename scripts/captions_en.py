# -*- coding: utf-8 -*-
"""P1：图1-10 图注英文化（Fig. N 前缀加粗，其余常规）。"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'

CAPS = {
1: ('Hierarchical semantic-action constraint compilation framework for flat knitting. Using Factory '
 'Sample 1 (Generation Case 01), five modules are shown left to right: (a) natural-language design '
 'intent (the semantic front end parses only constrained fields and does not emit action grids '
 'directly); (b) PatternPlan, the hierarchical manufacturing IR, explicitly encoding the hierarchy '
 'and parameters of the background, frame, circular region and its inner wave tuck; (c) event-level '
 'constraint compilation organized (Sec. 3.3) into visual-target generation, sparse event planning '
 'and action-grid construction, completed in a single lowering pass with all-or-nothing grouped '
 'motifs; the three insets are the true layered views of the same Generation Case 01 color-code map: '
 'region process layer (L0+L1, 1600 cells), event layer (L2, 32 split-tuck cells) and composed action '
 'graph (L0-L3, 1600 cells); constraint checking and conservative repair (Sec. 3.4) comprise hard '
 'constraints H1-H8, process rules and grouping completeness (deletion or downgrade only); no transfer '
 'or repair occurs in this example; (d) the 40x40 editable action color-code map (solid color blocks, '
 'sharing the paper color table with the factory-delivered maps in Fig. 10), where "0 violations" is '
 'the result of the constraint checker; (e) the factory-knitted physical sample, whose color-code map '
 'was converted to machine programs via the factory CAD patterning workflow, with geometry and texture '
 'consistent with the design (Sec. 4.6). The bottom loop denotes returning to the IR for local '
 'recompilation after the user updates target fields. The figure connects the evidence chain from '
 'design semantics and constrained action representation to factory validation.'),
2: ('Instance of the PatternPlan hierarchical manufacturing IR. For the design "rectangular frame '
 'enclosing a wave-tuck circular body", three columns: (a) hierarchy tree - root PatternPlan '
 '(layout_type = framed) with m0 (rectangle, role = border, effect = rib, scope center (20,20), size '
 '32x32, frame thickness 2, parent_id = null) and m1 (circle, role = center, effect = wave_tuck, center '
 '(20,20), size 18x18, period 4x4, parent_id = m0); parent_id edges are the structural premise for '
 'compile-time scope clipping; (b) spatial scopes - the 40x40 canvas drawn cell-by-cell with solid '
 'native color blocks, geometries rasterized by the system mask functions from instance parameters: the '
 'm0 frame ring (cheb <= 16 and cheb > 14, 248 cells) and the m1 disk (radius 9, 253 cells) are mutually '
 'exclusive with the disk inside the frame, and child-region event writing is confined to the parent '
 'scope at compile time; (c) seven-tuple field contract - shape fixes the writing scope, texture fixes '
 'the action pattern, parent fixes the topology. No axes or monospace code blocks; all text in Times New '
 'Roman; the full JSON schema and this instance are in Supplementary Material E.'),
3: ('Event-level constraint compilation from PatternPlan to ActionGrid. The mechanism is a four-column '
 'lowering pipeline left to right, using the true run output of Generation Case 01: (a) geometric scopes '
 '(VisualSpec, binary masks) - two 40x40 canvases drawn cell-by-cell with solid native color blocks for '
 'the frame ring frame_1 (264 cells) and the body disk circle_1 (529 cells, containing the nested texture '
 'scope wave_1, 441 cells); mask geometries match the system lowering formulas cell-by-cell and are '
 'cross-checked against the true color-code map and event anchors; parent-child clipping and boundary '
 'margins complete before any action planning; (b) process assignment (RealizationPlan, theta) - one '
 'contract per region: base fabric f_knit/b_knit, frame width t = 2 (on frame_1), event pattern wave_tuck '
 '(on wave_1), period 5x4 and offsets 2,3; (c) sparse events - anchors outside masks are never '
 'instantiated: split/tuck groups follow all-or-nothing placement, transfer events obey bed-direction and '
 'loop-support constraints, miss events obey hole-density and isolation constraints; Generation Case 01 '
 'has 16 groups, 32 events, 0 transfers, each group a split+tuck pair; (d) hierarchical composition '
 '(ActionGrid, L0-L3) - base, frame ring, event groups and parent-child clipping are written '
 'deterministically in fixed order, with structural guards ensuring conflict-free writing; 15 action '
 'labels on the 40x40 grid. The four columns lower stepwise along "scope -> assignment -> event -> '
 'composition"; the composed result enters the constraint checking and conservative repair of Fig. 4; '
 'each stage narrows the search space with explicit process knowledge. All figure text in English, Times '
 'New Roman, delivered as native editable PPTX vector objects (no SVG, no bitmap pasting).'),
4: ('Constraint checking and conservative repair loop. Four columns left to right, run instance '
 'Generation Case 01: (a) candidate ActionGrid - the true color-code map of this case (40x40, 4 action '
 'classes, 1600 cells drawn cell-by-cell with solid native color blocks); this example passes all checks '
 'in the first round (0 hard-constraint and 0 process violations, 0 repair rounds and 0 edits, 0 '
 'conservative-guard deletions), so the repair loop is idle for it; (b) four deterministic check groups - '
 'write-time structural guards (risk ascending, outline before texture), hard constraints H1-H8, process '
 'heuristics C1-C6 and grouping structural completeness; each violation v is localized to a small support '
 'domain support(v) and checks run in fixed order; (c) conservative repair R(A, v) - four operator '
 'contracts: roll back only stable base actions, cancel incomplete event groups, delete or downgrade '
 'high-risk actions, never create new events; modifications are strictly confined to the support domain; '
 'the true trigger frequency of repair paths is given by archived run statistics (of 263 runs, 33 found '
 'process violations at first check, 17 entered the repair loop, 100 produced deletions by conservative '
 'guards, and hard violations during detection were 0); (d) re-check and release - the dashed loop returns '
 'along orthogonal paths to the check column; the violation count is monotonically non-increasing along '
 'the loop (V_hard converges to 0); region outlines and base fabric are not reshaped; on passing, the '
 'action color-code map and machine export are output. The pipeline does not export pre-repair grids, so '
 'the figure shows repair paths only as the four operator contracts (aligned vertically with the four '
 'chips of columns (b) and (d)); their true trigger frequency is given by this caption and the body text '
 'rather than pre/post-repair grid pairs; column (a) likewise shows only the true color-code grid and its '
 'label, and the first-round zero violations are also given by this caption. All figure text in English, '
 'Times New Roman, delivered as native editable PPTX vector objects (no SVG, no bitmap pasting).'),
5: ('Structure-aware visual preview module. A single-row four-column chain left to right, run instance '
 'Generation Case 01: (a) conditional input - the true action color-code map A of this case (40x40, '
 '15-class discrete action space; this example uses the four classes f_knit/b_knit/f_tuck/f_split, 1600 '
 'cells drawn cell-by-cell with solid native color blocks) and the neighborhood context window N8 (center '
 'action + 8 neighbors, mechanism schematic); (b) condition encoder E - a neighbor-aware encoder fused '
 'with physical-property priors (tension, elasticity, stability levels per action class) into a single '
 'condition vector c injected into the UNet; (c) generation core - conditional flow matching learns the '
 'velocity field v(x,t,c); ODE midpoint integration (300 steps, globally correlated noise, seed 42, patch '
 '16) generates preview x1 from noise x0 on a 20x20 patch grid; (d) appearance preview - embeds the true '
 'Stage 4 output of Generation Case 01 (320x320, same verified asset as Fig. 1, not re-inferred), '
 'providing the designer visual feedback on region boundaries, texture and continuity. The module reads '
 'only the action graph A and introduces no yarn material or machine state; the preview is an appearance '
 'rendering, not yarn-level physical simulation, and does not constitute a knitability proof. The '
 'four-stage progressive training curriculum, network architecture and loss configuration are in '
 'Supplementary Material D; evaluation protocol and results in Sec. 4.4. All figure text in English, '
 'Times New Roman, delivered as native editable PPTX vector objects (no SVG; the only bitmap is the true '
 'Stage 4 preview output).'),
6: ('Workbench interface of the desktop prototype system. The left panel shows design-command input, '
 'generation parameters and pipeline-stage progress; the top shows chained version records; the right '
 'displays the generated action color-code map and structure-aware visual preview side by side. The '
 'interface shows that the method can be packaged as an interactive desktop prototype; the prototype is '
 'only an engineering wrapper of the method and not an independent methodological contribution.'),
7: ('Interactive incremental editing and chained version rollback interface. The upper and lower '
 'screenshot groups are the interactive incremental-editing interface and the historical version-chain '
 'interface, respectively: the designer can initiate incremental editing on target fields from any '
 'historical version; the system recompiles only the modified fields rather than re-parsing the whole '
 'design; new versions produced by editing are recorded in a chained relation with their base versions, '
 'supporting tracking and rollback of the design evolution path.'),
8: ('Editable generation cases. (a) Incremental editing mechanism on Description-Edit Case 01: starting '
 'from the base Description-Edit Case 00 (rectangular frame + circular body, wave tuck inside the circle, '
 'plain background), the edit command modifies only target fields (center.effect: wave -> honeycomb); the '
 'target region R is automatically derived from the PatternPlan difference (797 cells, white '
 'semi-transparent mask overlay); local recompilation replaces the target-field event group and rewrites '
 'the action grid within R; the resulting After changes only 76 cells, all inside R. (b) After images of '
 'the five Group-B (text re-parsing) editing variants overlaid with changed-cell masks; edit types are '
 'texture to honeycomb (Description-Edit Case 01), texture to open mesh (Description-Edit Case 02), frame '
 'thickening 2 to 3 (Description-Edit Case 03), background to sparse mesh (Description-Edit Case 04) and '
 'body shape to diamond (Description-Edit Case 05); below each variant the number of changed cells and '
 'the changes outside R are annotated (0 outside R, i.e. non-target preservation Preserve = 1.000); '
 'Description-Edit Case 04 is a background-level edit whose R covers the whole canvas and whose non-target '
 'set is empty. All color-code maps and change masks are true run outputs; locality metric definitions in '
 'Sec. 4.3.3. The results show that the hierarchical representation supports localized, leakage-free '
 'incremental editing.'),
9: ('Global-consistency versus boundary-readability trade-off across visual preview training stages. The '
 'three sample rows from top to bottom are Semantic Case 01, Semantic Case 03 and Semantic Case 06. '
 'Stages 1-3 favor global pixel consistency, while Stage 4 emphasizes boundary response and local texture '
 'readability. Window-level evaluation, complete whole-image metrics and training configurations are in '
 'Supplementary Material Tables S9, S10 and S19, respectively.'),
10: ('End-to-end cases and factory production test comparison. Factory Samples 1-5 are the end-to-end '
 'designs delivered to the factory (Generation Case 01, Description-Edit Case 05, Generation Case 04, '
 'Field-Edit Case 01 and Random-Generation Case 03, respectively); each row gives the PatternPlan summary '
 '(two-line wrap in the cell: region shape, texture and position/density modifiers; all five plans use '
 'framed layout), the 40x40 action color-code map, the structure-aware visual preview and the '
 'factory-knitted physical sample. The region geometry and texture of all five samples are consistent '
 'with the designs; all plans and color-code maps in the table are reproducible run outputs; visual '
 'preview quality and training-stage trade-offs in Sec. 4.4 Fig. 9.'),
}

doc = Document(MAIN)
done = 0
for p in doc.paragraphs:
    m = re.match(r'^图\s*(\d+)\s', p.text.strip())
    if not m:
        continue
    k = int(m.group(1))
    if k not in CAPS:
        continue
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r1 = p.add_run(f'Fig. {k}. '); r1.bold = True
    p.add_run(CAPS[k])
    done += 1
print('captions translated:', done)
doc.save(MAIN)
print('saved')
