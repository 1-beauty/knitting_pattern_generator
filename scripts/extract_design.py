import json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

path = r"e:\DeskTop\knitting_ai_competition.pptx"
prs = Presentation(path)

def color_of(fmt):
    try:
        if fmt.type is None:
            return None
        from pptx.enum.dml import MSO_FILL
        if fmt.type == MSO_FILL.SOLID:
            c = fmt.fore_color
            try:
                return "#" + str(c.rgb)
            except Exception:
                try:
                    return f"theme:{c.theme_color}"
                except Exception:
                    return "solid?"
        return str(fmt.type)
    except Exception as e:
        return None

def line_of(sh):
    try:
        ln = sh.line
        w = None
        try:
            w = round(Pt(ln.width).pt,2) if ln.width else None
        except Exception:
            w = None
        col = None
        try:
            col = "#" + str(ln.color.rgb)
        except Exception:
            col = None
        return {"w": w, "color": col}
    except Exception:
        return None

def emu_in(v):
    try:
        return round(Emu(v).inches,3)
    except Exception:
        return None

out = {"size_in": [emu_in(prs.slide_width), emu_in(prs.slide_height)], "slides": []}
for i, slide in enumerate(prs.slides):
    sl = {"idx": i+1, "shapes": []}
    for sh in slide.shapes:
        d = {
            "id": sh.shape_id,
            "name": sh.name,
            "geom": [emu_in(sh.left), emu_in(sh.top), emu_in(sh.width), emu_in(sh.height)],
        }
        # autoshape type + adjustments
        try:
            d["auto"] = str(sh.adjustments._prstGeom.prst) if hasattr(sh,'adjustments') else None
        except Exception:
            d["auto"] = None
        try:
            d["fill"] = color_of(sh.fill)
        except Exception:
            d["fill"] = None
        d["line"] = line_of(sh)
        if sh.has_text_frame:
            tf = sh.text_frame
            d["wrap"] = tf.word_wrap
            try:
                d["anchor"] = str(tf.vertical_anchor)
            except Exception:
                d["anchor"] = None
            try:
                d["margins"] = [emu_in(tf.margin_left), emu_in(tf.margin_top), emu_in(tf.margin_right), emu_in(tf.margin_bottom)]
            except Exception:
                d["margins"] = None
            paras = []
            for p in tf.paragraphs:
                runs = []
                for r in p.runs:
                    rf = r.font
                    col = None
                    try:
                        col = "#" + str(rf.color.rgb)
                    except Exception:
                        col = None
                    runs.append({
                        "t": r.text,
                        "sz": round(rf.size.pt,1) if rf.size else None,
                        "b": rf.bold,
                        "i": rf.italic,
                        "name": rf.name,
                        "color": col,
                    })
                paras.append({
                    "align": str(p.alignment) if p.alignment else None,
                    "lvl": p.level,
                    "runs": runs,
                })
            d["paras"] = paras
        sl["shapes"].append(d)
    out["slides"].append(sl)

with open(r"e:\DeskTop\knitting_pattern_generator\scripts\pptx_design.json","w",encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print("written", len(out["slides"]), "slides")
