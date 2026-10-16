# -*- coding: utf-8 -*-
"""第八轮 h（补齐引用闭合）：
1) 式(6)（动作网格合成方程，p061）此前未被任何散文引用 → 在 p063 经验命题处补引“按式（6）合成”。
2) 补充图 S6（工作台界面）此前在 main/supp 均无引用 → 在 Supp Section L 引言(s167)补引 图 S6/图 S7。
3) 全面审计未引用对象：公式 1–8、主图/主表、补充图 S1–S7、补充表 Sx；断言公式与补充图全部被引。"""
import sys; sys.stdout.reconfigure(encoding='utf-8')
import re
from docx import Document
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
m, s = Document(MAIN), Document(SUPP)
ps, sp = m.paragraphs, s.paragraphs


def rewrite(paras, pre, new):
    p = next(p for p in paras if p.text.startswith(pre))
    from docx.oxml.ns import qn
    assert not p._p.findall('.//' + qn('m:oMath')), ('oMath', pre)
    rr = p.runs; rr[0].text = new
    for r in rr[1:]: r._element.getparent().remove(r._element)
    return p


# ---------- 1) 式(6) 补引 ----------
rewrite(ps, '条件化生成的经验命题。',
        '条件化生成的经验命题。对规则库覆盖范围内的 PatternPlan，若（i）每个承载纹理的 motif 的有效作用域满足该纹理事件的最小放置尺寸，'
        '（ii）纹理事件周期与区域边界条件相容，且（iii）所声明的底布与边界实现存在于规则库中，则编译流程按式（6）合成满足 H1–H8 硬约束的动作网格；'
        '任一条件不满足时系统执行显式降级并在统计报告中记录，而非静默输出非法动作。本文未给出完整形式化证明，而以第 4 章批量运行作为经验验证：'
        '残留硬违例案例（Stress Case 19 与 Field-Edit Case 04，见 4.3.3 节）均对应稀疏镂空组织下局部支撑条件难以在单点替换范围内满足的情形，'
        '且均被约束检查器显式检测并报告；其余输入的最终输出均满足硬约束零违例。形式化证明与规则库外输入的降级刻画留作后续工作。')

# ---------- 2) 图 S6/S7 补引（Supp Section L 引言） ----------
rewrite(sp, '本节给出正文第 3 章管线的桌面原型实现',
        '本节给出正文第 3 章管线的桌面原型实现与交互编辑界面截图，供复现与试用参考：工作台界面见图 S6，交互增量编辑与链式版本回溯界面见图 S7。')

m.save(MAIN); s.save(SUPP)

# ---------- 3) 审计 ----------
mtxt = '\n'.join(p.text for p in ps)
stxt = '\n'.join(p.text for p in sp)
eqs = sorted(set(int(x) for x in re.findall(r'式\s*[（(]\s*(\d+)\s*[)）]', mtxt)))
print('公式被引编号:', eqs, '| 缺:', [n for n in range(1, 9) if n not in eqs])
assert eqs == list(range(1, 9)), eqs

# 补充图：引用（排除题注段本身）
def fig_cited(n):
    pat = re.compile(r'(?:图|Fig\.)\s*S%d\b' % n)
    for p in ps:
        if pat.search(p.text): return 'main'
    for p in sp:
        t = p.text.strip()
        if t.startswith(('图 S%d ' % n, 'Fig. S%d.' % n, 'Fig. S%d ' % n)): continue  # 题注
        if 'GUI 截图均为未经修改' in t and n == 7: pass
        if pat.search(p.text): return 'supp'
    return None
uncit_fig = [n for n in range(1, 8) if fig_cited(n) is None]
print('补充图被引:', {n: fig_cited(n) for n in range(1, 8)}, '| 未引:', uncit_fig)
assert not uncit_fig, uncit_fig

# 主图/主表被引
mf = sorted(set(int(x) for x in re.findall(r'图\s*(\d+)', mtxt)))
mt = sorted(set(int(x) for x in re.findall(r'表\s*(\d+)', mtxt)))
print('主图被引:', mf, '| 主表被引:', mt)

# 补充表被引（排除题注）
uncit_tab = []
for n in range(1, 22):
    pat = re.compile(r'(?:表|Table)\s*S%d\b' % n)
    hit = False
    for p in sp:
        t = p.text.strip()
        if t.startswith(('表 S%d ' % n, 'Table S%d ' % n)): continue
        if pat.search(p.text): hit = True; break
    if not hit:
        for p in ps:
            if pat.search(p.text): hit = True; break
    if not hit: uncit_tab.append(n)
print('补充表未引用(仅题注出现):', uncit_tab if uncit_tab else 'none')
print('OK round8h | 式1-8 全引 | 补充图 S1-S7 全引')
