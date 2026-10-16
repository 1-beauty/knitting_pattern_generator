# -*- coding: utf-8 -*-
"""主文件侧 P0-1：删除 '补充材料' H2 及之后全部 body 子元素(保留末尾 sectPr)，
并在 参考文献 前插入 Supplementary data online 句。不触碰 Supplementary_Material.docx。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
MAIN = r'e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx'
doc = Document(MAIN)
h = None
for child in doc.element.body.iterchildren():
    if child.tag.endswith('}p'):
        p = Paragraph(child, doc)
        if p.style.name == 'Heading 2' and p.text.strip().startswith('补充材料'):
            h = child; break
assert h is not None
removing = False; removed = 0
for child in list(doc.element.body.iterchildren()):
    if child is h:
        removing = True
    if removing and child.tag != qn('w:sectPr'):
        doc.element.body.remove(child); removed += 1
print('removed:', removed)
ref_h = None
for child in doc.element.body.iterchildren():
    if child.tag.endswith('}p'):
        p = Paragraph(child, doc)
        if p.style.name == 'Heading 2' and p.text.strip() == '参考文献':
            ref_h = child; break
newp = OxmlElement('w:p')
ppr = OxmlElement('w:pPr'); ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), 'BodyText'); ppr.append(ps); newp.append(ppr)
r = OxmlElement('w:r'); t = OxmlElement('w:t')
t.text = ('Supplementary data to this article can be found online at https://doi.org/<DOI-to-be-assigned> '
          '(Supplementary Material, Tables S1-S19 and Figures S1-S3).')
t.set(qn('xml:space'), 'preserve'); r.append(t); newp.append(r)
ref_h.addprevious(newp)
doc.save(MAIN)
d = Document(MAIN)
print('main tables now:', len(d.tables), 'paras:', len(d.paragraphs))
