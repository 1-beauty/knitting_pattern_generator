# -*- coding: utf-8 -*-
"""Algorithm 1 框修补：①题注 keep_with_next 与框同页；②步骤编号改为 "N.<tab>文本"
（显式 360twip 制表位对齐折行）；③行 cantSplit 防框体跨页断裂。"""
import re
import docx
from docx.shared import Twips
from docx.enum.text import WD_TAB_ALIGNMENT

SRC = r'docs\论文to_cii.docx'
doc = docx.Document(SRC)

# ① 题注与表格同页
for p in doc.paragraphs:
    if p.text.startswith('Algorithm 1'):
        p.paragraph_format.keep_with_next = True
        break

# 定位算法表
tbl = None
for t in doc.tables:
    if 'Input' in t.rows[0].cells[0].text and 'Output' in t.rows[0].cells[0].text:
        tbl = t
        break
assert tbl is not None

# ③ 禁止行跨页拆分
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
for row in tbl.rows:
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))

# ② 步骤编号格式：原 runs = [num+' .', rest]  ->  [num+'.', '\t'+rest]
n = 0
for p in tbl.rows[0].cells[0].paragraphs:
    txt = p.text
    m = re.match(r'^(\d+)\s*\.\s*(.*)$', txt)
    if m:
        num, rest = m.group(1), m.group(2)
        runs = p.runs
        assert len(runs) == 2, (txt, len(runs))
        runs[0].text = num + '.'
        runs[1].text = '\t' + rest
        pf = p.paragraph_format
        pf.left_indent = Twips(360)
        pf.first_line_indent = Twips(-360)
        pf.tab_stops.add_tab_stop(Twips(360), WD_TAB_ALIGNMENT.LEFT)
        n += 1
assert n == 11, n

doc.save(SRC)
print('steps fixed:', n)
