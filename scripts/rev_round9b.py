"""rev_round9b: length compression of figure-narration / related-work / limitation paragraphs.
All targets verified oMath==0 -> safe whole-paragraph rewrite via runs[0]=new + clear rest.
Every rewrite keeps ALL numbers, citation indices, terms and math symbols; only removes
narration that duplicates the figure captions and wordiness. Each segment carries a
must-keep token list asserted both before write and after reopen.
"""
from docx import Document
from docx.oxml.ns import qn
import re

PATH = r"e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx"
d = Document(PATH)
WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'\u2019\.\-]*")

def find(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None

def omath(p):
    return p._p.findall('.//' + qn('m:oMath'))

SEG = []

# ---- p008 intro background ----
SEG.append(("p008", "Flat knitting, as a core technology of flexible manufacturing",
"""Flat knitting, a core technology of flexible manufacturing and mass personalization, depends on precisely transforming design intent into machine instructions. In industrial practice the action color map is the dominant pattern representation and machine-input format: pattern engineers must specify the action class of every needle cell by cell in CAD systems such as Stoll M1 Plus and Shima Seiki SDS-ONE APEX [1, 2]. This tabular representation is also the mainstream paradigm in knitting-machine programming research [3, 4], yet its automated generation remains an unmet industrial need. The transformation is still performed manually by professional pattern engineers; as smart manufacturing demands rapid iteration, this experience-dependent mode has become a production bottleneck, because the professional compilation barrier between design intent and manufacturable actions prevents industrial designers without a knitting background from producing production-ready patterns. Removing this barrier so that non-expert designs can go directly into production (democratized design-for-manufacturing) is therefore a core problem in knitting industrial informatics. Designers describe textures and layouts in unstructured natural language, whereas computerized flat-knitting machines require stitch-level actions defined precisely on a discrete grid; this cross-modal, cross-level transformation exhibits a pronounced representation gap\u2014the ambiguity of high-level semantics and the strict constraints of the low-level action space lack an effective mapping mechanism, leading to long and uncontrollable idea-to-product cycles.""",
["[1, 2]", "[3, 4]", "Stoll M1 Plus", "SDS-ONE APEX", "democratized design-for-manufacturing", "representation gap"]))

# ---- p025 related work 2.2 ----
SEG.append(("p025", "Existing research on digital design for knitting can be divided",
"""Existing research on digital knitting design falls into three categories. (1) Visual texture generation: ColorwAI [5] and FabricDiffusion [6] use GANs or diffusion models [21, 22] for high-fidelity texture transfer and color matching, with conditional control [23] and Transformer-based architectures [24] improving controllability and scale; Knit-FLUX [13] enhances photorealism via low-rank adaptation but outputs only RGB pixels without action semantics. (2) Instruction inverse recovery and synthesis: Neural Inverse Knitting [15] derives machine instructions from sample images, while deepKnit [25], AutoKnit [26], and Knit Sketching [27] map 3D meshes, cut-and-sew patterns, or text to instruction sequences; these focus on linear instruction streams and struggle with the topological relations of two-dimensional action grids. (3) New representations and compilers: knitting skeletons [28], TopoKnit [29], visual flat-knitting programming [4], KNITIT [30], 3D machine-knitting compilers [31] with their formal semantics and scheduling framework [32], skCAD [33], and KnitGIST [34] simplify the process through high-level primitives but lack flexibility for open-ended natural-language semantics. Overall, these methods target visual texture, inverse recovery, or constraint optimization, and none unifies generation from natural-language intent to two-dimensional action structures. From the perspective of representation and constraint-injection timing, existing pipelines split into linear instruction streams (Neural Inverse Knitting [15], knitout [48], 3D knitting compilers [31, 32]), which treat one-dimensional sequences as objects and handle constraints at late-compilation or post-hoc stages, and program-synthesis texture generation (KnitGIST [34]); this paper instead takes the two-dimensional stitch action grid as the direct generation target, injects region attribution and grouped-pairing semantics at the event-planning stage, and supports field-level region editing.""",
["[5]", "[6]", "[13]", "[15]", "[25]", "[26]", "[27]", "[28]", "[29]", "[4]", "[30]", "[31]", "[32]", "[33]", "[34]", "[48]", "[21, 22]", "[23]", "[24]"]))

# ---- p029 related work 2.4 ----
SEG.append(("p029", "Two lines of research are directly related",
"""Two research lines relate directly to the strong-constraint generation paradigms compared in Section 4. The first is constrained decoding: PICARD [11] checks formal grammar constraints during autoregressive decoding, Grammar Prompting [10] guides LLMs to generate legal programs using a domain-specific-language grammar, and automaton-based decoding [37] generalizes this to regular and context-free constraints; these guarantee formal legality at the symbol level but do not address structural semantics. The second is constraint-solving-driven synthesis: constraint programming [38] declares and solves constraint satisfaction problems over discrete domains, with solvers such as OR-Tools CP-SAT [39] widely verified in scheduling and combinatorial optimization; in knitting, KnitGIST [34] synthesizes textures and the formal semantics and scheduling framework of knitting compilers [32] incorporates constraint solving into compilation. Both guarantee local constraint satisfiability at the symbolic level, but structural semantics such as region attribution and grouped pairing cannot be derived from the constraints themselves and must be carried explicitly by the generation process. Accordingly, this paper encodes the H1\u2013H8 hard constraints as a CP-SAT model, constructs a global constraint-projection baseline (definitions in Supplementary Section M), and quantitatively examines in Section 4.3.2 the boundary of replacing event-level compilation with constraint solving.""",
["[11]", "[10]", "[37]", "[38]", "[39]", "[34]", "[32]", "CP-SAT", "H1\u2013H8", "Supplementary Section M", "Section 4.3.2"]))

# ---- p033 Fig.1 narration ----
SEG.append(("p033", "Fig. 1 summarizes the overall framework",
"""Fig. 1 summarizes the overall framework from design intent, through the PatternPlan hierarchical manufacturing IR and event-level constraint compilation, to the editable action map and factory physical verification; the visual preview serves as auxiliary inspection, and users can trigger local recompilation by updating target fields of the IR. Taking Factory Sample 1 (Generation Case 01) as an example, the five modules (a)\u2013(e) are: the natural-language design intent, whose semantic front end parses only constrained fields rather than outputting an action grid directly; the PatternPlan IR encoding the hierarchy and parameters of the background, border, circular region, and the wave tuck inside it; event-level constraint compilation (Section 3.3), a single lowering through visual target generation, sparse event planning, and action grid construction with grouped motifs handled all-or-nothing\u2014the three insets are true layered views of the same case (region process layers L0+L1, 1600 cells; event layer L2, 32 split-tuck cells; composite action map L0\u2013L3, 1600 cells)\u2014followed by constraint checking with conservative repair (Section 3.4) covering H1\u2013H8, process rules, and grouping integrity (deletion or downgrade only), with no transfers or repairs in this case; the 40\u00d740 editable action color map (sharing the paper color table with the factory-delivered map in Fig. 7), where "0 violations" is the checker output; and the factory-knitted sample, whose color map was converted into a machine program through the factory's pattern-making workflow with geometry and texture consistent with the design (Section 4.5). The logical chain of action generation and auxiliary preview is:""",
["1600 cells", "32 split-tuck cells", "L0\u2013L3", "Factory Sample 1", "Generation Case 01", "Fig. 7", "Section 3.3", "Section 3.4", "Section 4.5", "40\u00d740"]))

# ---- p035 Fig.2 narration ----
SEG.append(("p035", "Table 1 lists the 15 discrete action labels",
"""Table 1 lists the 15 discrete action labels retained after sample-level cleaning; each corresponds to a local knitting operation or composite structural event, and the colors serve only as a visual encoding of categories (label sources and the two-stage cleaning process in Supplementary Section A). Fig. 2 groups the 15 labels into six families of action primitives\u2014knit, tuck, miss, transfer, split, and take-down\u2014each card showing (a) a plain-knit single-course loop row, (b) a two-course fabric close-up with the family's primary action loop in magenta, and (c) a local courses \u00d7 needles action grid where magenta nodes and dashed boxes mark the written cells. The six families are: knit (f_knit/b_knit), drawing the new loop through the old loop below to form the ground structure; tuck (f_tuck/b_tuck), retaining the old loop and hanging in the new yarn to form the loop accumulation behind raised textures; miss, selecting no needle and forming no new loop (the empty-needle background); transfer (xfer_l1/xfer_f_b), moving an existing loop to a neighboring needle or the opposite bed and leaving the source needle empty; split (f_split/b_split), dividing one loop into two branches for local support and structural transition; and take-down (carrier), removing loop support in the corresponding wale to finish the edge, causing the affected columns to drop. In the loop diagrams, cyan denotes retained loops and inter-course connections, magenta the yarn path of the primary action and written cells, gray dots the interlocking points of the two-course fabric, and gold boxes needles holding both the old and tuck loops; the diagrams are schematics of stitch mechanisms rather than photorealistic renderings.""",
["15 discrete action labels", "f_knit/b_knit", "xfer_l1/xfer_f_b", "f_split/b_split", "Supplementary Section A", "courses \u00d7 needles", "take-down", "carrier"]))

# ---- p051 Fig.4 narration ----
SEG.append(("p051", "Given a PatternPlan, the task of the back-end compiler",
"""Given a PatternPlan, the back-end compiler lowers its hierarchical regions, texture semantics, and process parameters step by step into a stitch-level discrete action representation. This paper adopts layer-wise semantic lowering: the IR is first converted into explicit geometric scopes (visual target generation), then sparse events with region attribution are generated (sparse event planning), and finally an action grid over the 15 discrete labels is synthesized within the constraint-feasible region (action grid construction). Fig. 4 shows the complete mechanism: (a) geometric scopes as binary masks\u2014264 cells in the frame_1 ring and 529 in the circle_1 body disk (including the 441-cell nested texture scope wave_1)\u2014cell-wise consistent with the lowering formulae and cross-verified against the real color map and event anchors, with parent\u2013child clipping and boundary margins completed before any action planning; (b) one process contract per region (base fabric f_knit/b_knit, upper border width t=2 for frame_1, event pattern wave_tuck for wave_1, period 5\u00d74 with offset 2,3); (c) anchors outside the masks are never instantiated, split/tuck groups are handled all-or-nothing, transfers obey bed-direction and loop-support constraints, and misses obey hole-density and isolation constraints (16 groups, 32 events, 0 transfers in this case); and (d) the base fabric, frame ring, event groups, and parent\u2013child clipping are written deterministically in fixed order, structural guards guarantee conflict-freedom, and the final 40\u00d740 grid carries the 15 labels, passed to the checking and repair of Fig. 5.""",
["264 cells", "529", "441-cell", "t=2", "5\u00d74", "offset 2,3", "16 groups", "32 events", "0 transfers", "40\u00d740", "Fig. 5", "wave_1", "frame_1", "circle_1"]))

# ---- p070 Fig.5 narration ----
SEG.append(("p070", "This section integrates action-legality constraints",
"""This section integrates action-legality constraints, process heuristic constraints, and grouped-structure integrity checks into the constraint-aware refinement stage; the checking and repair loop is shown in Fig. 5. Their decidability follows from discrete definitions: every rule is checked by a deterministic program on a finite grid, without physical simulation or human judgment. Taking Generation Case 01 as an example, Fig. 5 shows four columns: (a) the candidate ActionGrid (this case's real 40\u00d740 color map, 4 action classes), which passes all checks in the first round (0 hard-constraint and 0 process violations, 0 repair rounds, 0 guard deletions), so the loop idles; (b) four deterministic check groups\u2014write-time structural guards (ascending risk, outlines before textures), hard constraints H1\u2013H8, process heuristics C1\u2013C6, and grouped-structure integrity\u2014each violation v localized to a small support domain support(v) and checked in fixed order; (c) the four operator contracts of conservative repair R(A,v): roll back only stable base-fabric actions, cancel incomplete event groups, delete or downgrade high-risk actions, and never create new events, with modifications confined to support(v); and (d) re-check and release, where the dashed loop returns along an orthogonal path, the violation count is monotonically non-increasing (V_hard converges to 0), and region outlines and the base fabric are not reshaped; on passing, the action color map and machine export are output. The archived statistics give the repair path's true triggering frequency: among 263 runs, 33 found process violations at first check, 17 entered the repair loop, 100 produced conservative guard deletions, and 0 hard violations occurred during detection; because pre-repair grids are not exported, the repair path is shown only through the four operator contracts (aligned with the four chips in columns (b) and (d)).""",
["Generation Case 01", "40\u00d740", "H1\u2013H8", "C1\u2013C6", "support(v)", "R(A,v)", "263 runs", "V_hard"]))

# ---- p084 Fig.6 narration ----
SEG.append(("p084", "As shown in Fig. 6, the model adopts",
"""As shown in Fig. 6, the model adopts a UNet-based [41] conditional backbone with a conditional-flow-matching framework [42] (the same continuous-time paradigm as Rectified Flow [43]), fusing the center and 8-neighborhood action conditions through a neighbor-aware encoder and introducing action-semantics-based process-attribute encoding (encoding, attribute definitions, ODE inference, and sampling in Supplementary Section D; the pipeline is further encapsulated as a desktop prototype supporting interactive incremental editing, Supplementary Section L). Taking Generation Case 01 as an example, Fig. 6 shows four columns: (a) the conditional inputs are this case's real action color map A (40\u00d740, classes f_knit/b_knit/f_tuck/f_split) and the neighborhood window N8 (center + 8 neighbors); (b) the neighbor-aware encoder fuses per-class physical priors (tension, elasticity, stability) into a condition vector c injected into the UNet; (c) conditional flow matching learns the velocity field v(x,t,c), and ODE midpoint integration (300 steps, globally correlated noise, seed 42, patch 16) generates the preview x1 from noise x0 on a 20\u00d720 patch grid; and (d) the appearance preview embeds this case's real Stage 4 output (320\u00d7320, the same verified asset as Fig. 1), giving feedback on region boundaries, texture, and continuity. The module reads only the action map A and introduces no yarn material or machine state; the preview is an appearance rendering rather than a yarn-level physical simulation and does not constitute a proof of knittability.""",
["[41]", "[42]", "[43]", "Supplementary Section D", "Supplementary Section L", "Generation Case 01", "40\u00d740", "f_knit/b_knit/f_tuck/f_split", "N8", "v(x,t,c)", "300 steps", "seed 42", "patch 16", "20\u00d720", "320\u00d7320", "Fig. 1"]))

# ---- p137 limitations ----
SEG.append(("p137", "The current system still has six limitations",
"""The current system has six limitations (detailed in Supplementary Section N): first, the semantics-to-process mapping relies on a predefined rule set, so complex structures outside the library adopt a conservative downgrade that may sacrifice texture richness; second, self-repair is mainly local action replacement, and region-level conflicts (large-area dense openwork, continuous front\u2013back bed switching, high-density local textures) still require a region-level reconstruction mechanism; third, the action transition intensity score and process heuristics are proxy metrics that cannot replace measuring yarn tension, yarn-breakage risk, and fabric flatness in real knitting; fourth, the visual preview is a design-stage structural preview rather than a physical knitting simulation, and future work should introduce yarn-level simulation and computational design methods closer to knit structures [47] with real-sample verification; fifth, sample-level cleaning removed the paired samples containing two classes of low-frequency actions, so the complete action space is not yet covered; and sixth, the action color map is a machine-specific-code-agnostic abstract representation [48] whose colors only distinguish internal action classes, while the region geometry and texture of the 5 physical samples in Section 4.5 are consistent with the designs, confirming the accuracy of the approach under real manufacturing conditions. System positioning in the industrial-deployment context, domain transferability, and the boundaries of the human role are also given in Supplementary Section N.""",
["six limitations", "Supplementary Section N", "[47]", "[48]", "5 physical samples", "Section 4.5"]))

# ---- apply ----
print("=== APPLY COMPRESSION ===")
before_after = []
for key, prefix, new, tokens in SEG:
    p = find(prefix)
    assert p is not None, f"{key}: locate fail"
    assert len(omath(p)) == 0, f"{key}: oMath!=0, refuse whole-paragraph rewrite"
    # token safety: every must-keep token present in the NEW text
    missing = [t for t in tokens if t not in new]
    assert not missing, f"{key}: new text missing tokens {missing}"
    w_before = len(WORD.findall(p.text))
    runs = p.runs
    assert runs, f"{key}: no runs"
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""
    w_after = len(WORD.findall(new))
    before_after.append((key, w_before, w_after))
    print(f"  {key:6s} {w_before:4d} -> {w_after:4d}  (-{w_before-w_after})")

# ---- save ----
try:
    d.save(PATH); saved = PATH
except PermissionError:
    saved = r"e:\DeskTop\knitting_pattern_generator\docs\论文to_cii_fix9b.docx"
    d.save(saved); print("WARNING locked, bypass:", saved)

# ---- verify reopen ----
d2 = Document(saved)
j = "\n".join(p.text for p in d2.paragraphs)
print("\n=== VERIFY tokens landed ===")
allok = True
for key, prefix, new, tokens in SEG:
    miss = [t for t in tokens if t not in j]
    if miss:
        allok = False
        print(f"  {key}: MISSING {miss}")
print("all tokens present:", allok)

# recompute body word count (paragraphs before References)
ps2 = d2.paragraphs
rs = next(i for i, p in enumerate(ps2) if p.text.strip() == "References")
body_words = sum(len(WORD.findall(p.text)) for p in ps2[:rs])
print("body words before References now:", body_words)
print("total compression this pass:", sum(b - a for _, b, a in before_after))
