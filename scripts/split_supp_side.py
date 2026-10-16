# -*- coding: utf-8 -*-
"""仅处理补充材料副本：删除 '补充材料' 之前全部 body 子元素，标题改英文。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
SUPP = r'e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx'
s = Document(SUPP)
hs = None
for child in s.element.body.iterchildren():
    if child.tag.endswith('}p'):
        p = Paragraph(child, s)
        if p.style.name == 'Heading 2' and p.text.strip().startswith('补充材料'):
            hs = child; break
assert hs is not None
for child in list(s.element.body.iterchildren()):
    if child is hs:
        break
    if child.tag != qn('w:sectPr'):
        s.element.body.remove(child)
hp = Paragraph(hs, s)
for rr in list(hp.runs):
    rr._r.getparent().remove(rr._r)
hp.add_run('Supplementary Material')
s.save(SUPP)
d = Document(SUPP)
print('supp tables:', len(d.tables), 'paras:', len(d.paragraphs), 'first:', d.paragraphs[0].text.strip()[:40])
