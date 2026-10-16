# -*- coding: utf-8 -*-
"""rev_round10b: global sequential-coding renumber + reference-list rebuild + supplementary sync.

Facts established by the pre-scans (scripts/_scan_tokens2.py, _scan_runsplit.py):
  * main body: 89 numeric citation tokens, 52 distinct ids (1..52), no orphans, no id > 52;
    16 non-numeric placeholders [N1]..[N16] inserted by rev_round10a.
  * every citation token in the MAIN document lives inside a single w:t node (0 spanning),
    so per-node substitution is lossless; the script asserts this before writing.
  * reference list: 52 paragraphs, each 'Body Text'/'First Paragraph' style with exactly 1 run.
  * supplementary: 17 real citation tokens (single w:t node) and 6 multi-node bracket groups
    ([20, 20], [32, 32], [18, 18]) which are COORDINATES, not citations - they are skipped by
    requiring single-node AND non-repeated member list.

Method: one pass in true document order (paragraphs + table cells) assigns new numbers by first
appearance; the rewrite is a SINGLE-PASS dict lookup over the union pattern, never sequential
string replacement, so no chained contamination ([20] -> [35] -> ...) can occur.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
import copy
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

MAIN = r"e:\DeskTop\knitting_pattern_generator\docs\论文to_cii.docx"
SUPP = r"e:\DeskTop\knitting_pattern_generator\docs\Supplementary_Material.docx"

W_P = qn('w:p'); W_TBL = qn('w:tbl'); W_T = qn('w:t'); XML_SPACE = qn('xml:space')
TOK = re.compile(r"\[(N\d+|\d+(?:\s*,\s*\d+)*)\]")

# ---- 16 new Computers in Industry entries, house style (metadata from Crossref, DOIs resolved) ----
NEWREFS = [
 ("N1",  "C. Sassanelli, A. Urbinati, P. Rosa, et al., Addressing circular economy through design for X approaches: A systematic literature review, Comput. Ind. 120 (2020) 103245. doi:10.1016/j.compind.2020.103245."),
 ("N2",  "J. Friederich, D.P. Francis, S. Lazarova-Molnar, et al., A framework for data-driven digital twins of smart manufacturing systems, Comput. Ind. 136 (2022) 103586. doi:10.1016/j.compind.2021.103586."),
 ("N3",  u"S. Jask\u00f3, A. Skrop, T. Holczinger, et al., Development of manufacturing execution systems in accordance with Industry 4.0 requirements: A review of standard- and ontology-based methodologies and tools, Comput. Ind. 123 (2020) 103300. doi:10.1016/j.compind.2020.103300."),
 ("N4",  "A. Majumdar, H. Garg, R. Jain, Managing the barriers of Industry 4.0 adoption and implementation in textile and clothing industry: Interpretive structural model and triple helix framework, Comput. Ind. 125 (2021) 103372. doi:10.1016/j.compind.2020.103372."),
 ("N5",  "F. Chiarello, P. Belingheri, G. Fantoni, Data science for engineering design: State of the art and future directions, Comput. Ind. 129 (2021) 103447. doi:10.1016/j.compind.2021.103447."),
 ("N6",  "S. Colabianchi, F. Costantino, N. Sabetta, Assessment of a large language model based digital intelligent assistant in assembly manufacturing, Comput. Ind. 162 (2024) 104129. doi:10.1016/j.compind.2024.104129."),
 ("N7",  "B. Kwon, J. Kim, H. Lee, et al., Construction of design requirements knowledgebase from unstructured design guidelines using natural language processing, Comput. Ind. 159-160 (2024) 104100. doi:10.1016/j.compind.2024.104100."),
 ("N8",  "Y.J. Liu, D.L. Zhang, M.M. Yuen, A survey on CAD methods in 3D garment design, Comput. Ind. 61 (2010) 576-593. doi:10.1016/j.compind.2010.03.007."),
 ("N9",  "C. Au, Y.S. Ma, Garment pattern definition, development and application with associative feature approach, Comput. Ind. 61 (2010) 524-531. doi:10.1016/j.compind.2010.03.002."),
 ("N10", "B. Andres, R. Poler, R. Sanchis, A data model for collaborative manufacturing environments, Comput. Ind. 126 (2021) 103398. doi:10.1016/j.compind.2021.103398."),
 ("N11", "E.M. Sanfilippo, F. Belkadi, A. Bernard, Ontology-based knowledge representation for additive manufacturing, Comput. Ind. 109 (2019) 182-194. doi:10.1016/j.compind.2019.03.006."),
 ("N12", "Z. Soufi, P. David, Z. Yahouni, Generation of material handling system alternatives: A constraints satisfaction problem approach, Comput. Ind. 155 (2024) 104045. doi:10.1016/j.compind.2023.104045."),
 ("N13", u"C. Gonz\u00e1lez-Lluch, R. Plumed, D. P\u00e9rez-L\u00f3pez, et al., A constraint redundancy elimination strategy to improve design reuse in parametric modeling, Comput. Ind. 129 (2021) 103460. doi:10.1016/j.compind.2021.103460."),
 ("N14", "Y. Zhang, X. Ren, J. Zhang, A semi-automated compliance checking framework for shield tunnel design integrating ontology and natural language processing, Comput. Ind. 173 (2025) 104362. doi:10.1016/j.compind.2025.104362."),
 ("N15", "Y. Gu, S. Qin, W. Liao, et al., Intelligent design of dimensions of reinforced concrete frame structure components using diffusion models, Comput. Ind. 175 (2026) 104428. doi:10.1016/j.compind.2025.104428."),
 ("N16", "V. Haseltalab, A. Dutta, S. Yang, Data-driven inverse parameter design for stable dimensional control in extrusion-based additive manufacturing, Comput. Ind. 180 (2026) 104520. doi:10.1016/j.compind.2026.104520."),
]
NEWREFS = dict(NEWREFS)


def iter_paras(doc):
    """true document order: top-level paragraphs interleaved with table-cell paragraphs"""
    pi = 0
    for child in doc.element.body.iterchildren():
        if child.tag == W_P:
            yield ("p%03d" % pi, Paragraph(child, doc)); pi += 1
        elif child.tag == W_TBL:
            tb = Table(child, doc)
            for ri, row in enumerate(tb.rows):
                for ci, cell in enumerate(row.cells):
                    for cp in cell.paragraphs:
                        yield ("tbl r%d c%d" % (ri, ci), cp)


def tnodes(p):
    ns = [n for n in p._p.iter() if n.tag == W_T]
    text = "".join(n.text or "" for n in ns)
    offs = []; c = 0
    for n in ns:
        L = len(n.text or ""); offs.append((c, c + L)); c += L
    return ns, text, offs


def members(tok):
    """'[3, 4]' -> ['3','4'] ; '[N12]' -> ['N12']"""
    inner = tok[1:-1]
    return [x.strip() for x in inner.split(",")]


def is_citation(tok, nnodes):
    """a bracket group counts as a citation only if it is inside one w:t node and is not a
    repeated-number coordinate pair such as [20, 20]"""
    if nnodes != 1:
        return False
    ms = members(tok)
    if len(ms) > 1 and len(set(ms)) == 1:
        return False
    return True


def rewrite_nodes(p, fn):
    """Rewrite citation tokens node-locally.

    Tokens that span several w:t nodes cannot be matched inside a single node's text, so they are
    left untouched by construction (that is what protects the [20, 20]-style coordinate groups in
    the supplementary). Returns (tokens encountered, tokens whose text actually changed).
    """
    hit = changed = 0
    for n in p._p.iter():
        if n.tag != W_T:
            continue
        cur = n.text or ""
        if "[" not in cur:
            continue
        k = len(TOK.findall(cur))
        if not k:
            continue
        hit += k
        new = TOK.sub(fn, cur)
        if new != cur:
            changed += k
            n.text = new
            n.set(XML_SPACE, 'preserve')
    return hit, changed


# =========================================================================== PASS 1
print("=== PASS 1: assign new numbers by first appearance ===")
d = Document(MAIN)
seq = list(iter_paras(d))
ref_idx = next(i for i, (tag, p) in enumerate(seq) if p.text.strip() == "References")
print("References heading at seq index %d (%s); body nodes before it: %d"
      % (ref_idx, seq[ref_idx][0], ref_idx))

order = []          # identity strings in first-appearance order
seen = set()
first_at = {}
for i in range(ref_idx):
    tag, p = seq[i]
    ns, text, offs = tnodes(p)
    for m in TOK.finditer(text):
        s, e = m.span()
        nn = sum(1 for a, b in offs if a < e and s < b)
        if not is_citation(m.group(0), nn):
            print("   SKIP non-citation in body:", tag, m.group(0)); continue
        for mem in members(m.group(0)):
            if mem not in seen:
                seen.add(mem); order.append(mem); first_at[mem] = tag

NEWNUM = {ident: k + 1 for k, ident in enumerate(order)}
N_NEW = len(order)
print("distinct citation identities:", N_NEW)
assert N_NEW == 68, "expected 52 old + 16 new = 68, got %d" % N_NEW
ph = [x for x in order if x.startswith("N")]
old = [x for x in order if not x.startswith("N")]
assert len(ph) == 16 and len(old) == 52, (len(ph), len(old))
assert sorted(int(x) for x in old) == list(range(1, 53)), "old ids are not exactly 1..52"

print("\nold -> new mapping (old id, new id, first appearance):")
for x in old:
    print("   [%s] -> [%d]   @%s" % (x, NEWNUM[x], first_at[x]))
print("\nplaceholder -> new number:")
for x in ph:
    print("   [%s] -> [%d]   @%s" % (x, NEWNUM[x], first_at[x]))

# =========================================================================== PASS 2
print("\n=== PASS 2: rewrite body tokens (single-pass dict lookup) ===")


def repl(m):
    ms = members(m.group(0))
    if len(ms) > 1 and len(set(ms)) == 1:
        return m.group(0)                      # coordinate-style group, not a citation
    return "[" + ", ".join(str(i) for i in sorted(NEWNUM[x] for x in ms)) + "]"


# safety: the main document must have no token spanning several w:t nodes
spanning = 0
for i in range(ref_idx):
    tag, p = seq[i]
    ns, text, offs = tnodes(p)
    for m in TOK.finditer(text):
        s, e = m.span()
        if sum(1 for a, b in offs if a < e and s < b) > 1:
            spanning += 1
            print("   SPANNING token in main body:", tag, m.group(0))
assert spanning == 0, "%d tokens span several w:t nodes; node-local rewrite would be lossy" % spanning

n_hit = n_chg = 0
for i in range(ref_idx):
    tag, p = seq[i]
    a, b = rewrite_nodes(p, repl)
    n_hit += a; n_chg += b
print("body citation tokens encountered: %d, actually changed: %d" % (n_hit, n_chg))
assert n_hit == 105, "expected 89 numeric + 16 placeholders = 105, got %d" % n_hit
allt = "".join(n.text or "" for n in d.element.body.iter() if n.tag == W_T)
assert not re.search(r"\[N\d+\]", allt), "a placeholder survived the rewrite"
print("no placeholder survived: OK")

# =========================================================================== PASS 3
print("\n=== PASS 3: rebuild the reference list ===")
ref_paras = [p for tag, p in seq[ref_idx + 1:]]
assert len(ref_paras) == 52, "expected 52 reference paragraphs, got %d" % len(ref_paras)

OLD_TEXT = {}
for k, p in enumerate(ref_paras):
    m = re.match(r"^\[(\d+)\]\s*(.*)$", p.text.strip(), re.S)
    assert m, "reference paragraph %d not in '[n] text' form: %r" % (k, p.text[:60])
    assert int(m.group(1)) == k + 1, "reference list not ascending at %d: %s" % (k, m.group(1))
    OLD_TEXT[str(k + 1)] = m.group(2).strip()
print("parsed 52 existing entries; all ascending")

INV = {v: k for k, v in NEWNUM.items()}          # new number -> identity
lines = []
for i in range(1, N_NEW + 1):
    ident = INV[i]
    body = OLD_TEXT[ident] if not ident.startswith("N") else NEWREFS[ident]
    assert body, "empty entry for %s" % ident
    lines.append("[%d] %s" % (i, body))

for i, p in enumerate(ref_paras):
    r = p.runs
    assert len(r) == 1, "reference paragraph %d has %d runs" % (i, len(r))
    r[0].text = lines[i]

# append the 16 extra entries by cloning the last reference paragraph (keeps style + indent)
anchor = ref_paras[-1]._p
parent = ref_paras[-1]._parent
for i in range(52, N_NEW):
    el = copy.deepcopy(ref_paras[-1]._p)
    anchor.addnext(el)
    anchor = el
    np = Paragraph(el, parent)
    np.runs[0].text = lines[i]
print("reference list rebuilt with %d entries (%d rewritten in place, %d appended)"
      % (N_NEW, 52, N_NEW - 52))

# =========================================================================== SUPP
print("\n=== supplementary sync ===")
ds = Document(SUPP)
NUMMAP = {int(k): v for k, v in NEWNUM.items() if not k.startswith("N")}


def srepl(m):
    ms = members(m.group(0))
    assert not any(x.startswith("N") for x in ms), "placeholder found in supplementary!"
    if len(ms) > 1 and len(set(ms)) == 1:
        return m.group(0)
    return "[" + ", ".join(str(i) for i in sorted(NUMMAP[int(x)] for x in ms)) + "]"


# count real (single-node, non-coordinate) citation tokens before the rewrite
expect = 0
for tag, p in iter_paras(ds):
    ns, text, offs = tnodes(p)
    for m in TOK.finditer(text):
        s, e = m.span()
        nn = sum(1 for a, b in offs if a < e and s < b)
        if is_citation(m.group(0), nn):
            expect += 1
print("supplementary citation tokens to remap:", expect)
assert expect == 17, "expected 17 supplementary citations, found %d" % expect

s_hit = s_chg = 0
for tag, p in iter_paras(ds):
    a, b = rewrite_nodes(p, srepl)
    s_hit += a; s_chg += b
print("supplementary citation tokens encountered: %d, changed: %d" % (s_hit, s_chg))
assert s_hit == expect, "encountered %d but expected %d" % (s_hit, expect)

# =========================================================================== SAVE BOTH
# both documents are only written after every assertion has passed, so a failure anywhere above
# leaves the two files mutually consistent.
print("\n=== SAVE ===")
try:
    d.save(MAIN); saved = MAIN
except PermissionError:
    saved = MAIN.replace(".docx", "_fix10b.docx")
    d.save(saved); print("WARNING main locked, bypass:", saved)
print("saved main:", saved)
try:
    ds.save(SUPP); ssaved = SUPP
except PermissionError:
    ssaved = SUPP.replace(".docx", "_fix10b.docx")
    ds.save(ssaved); print("WARNING supp locked, bypass:", ssaved)
print("saved supp:", ssaved)
print("\nDONE round10b")
